from __future__ import annotations

from pathlib import Path
from datetime import datetime
import hashlib
import os
import re
import sys

import pandas as pd


SCRIPT_VERSION = "1.0.0"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
POWERBI_ROOT = PROJECT_ROOT / "powerbi"
CANONICAL_DIR = PROJECT_ROOT / "data" / "canonical"


def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg
            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, key_path) as key:
                raw_value, _ = winreg.QueryValueEx(key, "Desktop")
            return Path(os.path.expandvars(raw_value)).resolve()
        except Exception:
            pass
    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = Path(os.environ.get("TRANSPARENCY_PUF_REVIEW_DIR", PROJECT_ROOT / ".local-review")).expanduser().resolve()
OUTPUT_FILE = REVIEW_DIR / "07_POWER_BI_SEMANTIC_MODEL_VALIDATION.xlsx"

EXPECTED_TABLES = {
    "_Measures",
    "DimAvailabilityStatus",
    "DimIssuer",
    "DimMetric",
    "DimPlan",
    "DimReportingPeriod",
    "DimState",
    "DQExceptionRegister",
    "FactIssuerMetricAvailability",
    "FactIssuerTransparency",
    "FactPlanMetricAvailability",
    "FactPlanTransparency",
}

EXPECTED_RELATIONSHIPS = {
    ("DimIssuer.StateKey", "DimState.StateKey"),
    ("DimPlan.IssuerKey", "DimIssuer.IssuerKey"),
    ("FactIssuerTransparency.IssuerKey", "DimIssuer.IssuerKey"),
    ("FactIssuerMetricAvailability.IssuerKey", "DimIssuer.IssuerKey"),
    ("FactPlanTransparency.PlanKey", "DimPlan.PlanKey"),
    ("FactPlanMetricAvailability.PlanKey", "DimPlan.PlanKey"),
    ("FactIssuerTransparency.ReportingPeriodKey", "DimReportingPeriod.ReportingPeriodKey"),
    ("FactPlanTransparency.ReportingPeriodKey", "DimReportingPeriod.ReportingPeriodKey"),
    ("FactIssuerMetricAvailability.ReportingPeriodKey", "DimReportingPeriod.ReportingPeriodKey"),
    ("FactPlanMetricAvailability.ReportingPeriodKey", "DimReportingPeriod.ReportingPeriodKey"),
    ("FactIssuerMetricAvailability.MetricKey", "DimMetric.MetricKey"),
    ("FactPlanMetricAvailability.MetricKey", "DimMetric.MetricKey"),
    ("FactIssuerMetricAvailability.AvailabilityStatusKey", "DimAvailabilityStatus.StatusKey"),
    ("FactPlanMetricAvailability.AvailabilityStatusKey", "DimAvailabilityStatus.StatusKey"),
}

EXPECTED_ROW_COUNTS = {
    "dim_reporting_period.csv": 1,
    "dim_state.csv": 30,
    "dim_issuer.csv": 348,
    "dim_plan.csv": 4956,
    "dim_metric.csv": 30,
    "dim_availability_status.csv": 8,
    "fact_issuer_transparency.csv": 348,
    "fact_plan_transparency.csv": 4956,
    "fact_issuer_metric_availability.csv": 4176,
    "fact_plan_metric_availability.csv": 89208,
    "dq_exception_register.csv": 38,
}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def find_semantic_model() -> Path:
    candidates = sorted(
        p for p in POWERBI_ROOT.glob("*.SemanticModel") if p.is_dir()
    )
    if len(candidates) != 1:
        raise RuntimeError(
            f"Expected exactly one *.SemanticModel folder; found {len(candidates)}."
        )
    return candidates[0]


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig")


def parse_relationships(text: str) -> pd.DataFrame:
    rows = []
    blocks = re.split(r"(?m)(?=^relationship\s+)", text.strip())

    for block in blocks:
        block = block.strip()
        if not block.startswith("relationship "):
            continue

        name = re.match(r"relationship\s+(.+)", block)
        from_col = re.search(r"(?m)^\s*fromColumn:\s*(.+?)\s*$", block)
        to_col = re.search(r"(?m)^\s*toColumn:\s*(.+?)\s*$", block)
        active = re.search(r"(?m)^\s*isActive:\s*(.+?)\s*$", block)
        cross = re.search(
            r"(?m)^\s*crossFilteringBehavior:\s*(.+?)\s*$",
            block,
        )
        from_card = re.search(
            r"(?m)^\s*fromCardinality:\s*(.+?)\s*$",
            block,
        )
        to_card = re.search(
            r"(?m)^\s*toCardinality:\s*(.+?)\s*$",
            block,
        )

        rows.append(
            {
                "RelationshipName": name.group(1).strip() if name else "",
                "FromColumn": from_col.group(1).strip() if from_col else "",
                "ToColumn": to_col.group(1).strip() if to_col else "",
                "IsActive": active.group(1).strip() if active else "true (default)",
                "CrossFilteringBehavior": (
                    cross.group(1).strip()
                    if cross
                    else "oneDirection (default)"
                ),
                "FromCardinality": (
                    from_card.group(1).strip()
                    if from_card
                    else "many (default)"
                ),
                "ToCardinality": (
                    to_card.group(1).strip()
                    if to_card
                    else "one (default)"
                ),
                "RawDefinition": block,
            }
        )

    return pd.DataFrame(rows)


def model_table_refs(model_text: str) -> list[str]:
    return re.findall(r"(?m)^ref table (.+?)\s*$", model_text)


def parse_table_columns(tables_dir: Path) -> pd.DataFrame:
    rows = []
    for path in sorted(tables_dir.glob("*.tmdl")):
        text = read_text(path)
        table_match = re.search(r"(?m)^table\s+(.+?)\s*$", text)
        if not table_match:
            continue
        table = table_match.group(1).strip().strip("'")
        for col in re.findall(r"(?m)^\tcolumn\s+(.+?)\s*$", text):
            rows.append(
                {
                    "TableName": table,
                    "ColumnName": col.strip().strip("'"),
                    "DefinitionFile": str(path.relative_to(PROJECT_ROOT)),
                }
            )
    return pd.DataFrame(rows)


def canonical_integrity_checks() -> tuple[pd.DataFrame, pd.DataFrame]:
    data = {}
    inventory_rows = []

    for filename, expected_rows in EXPECTED_ROW_COUNTS.items():
        path = CANONICAL_DIR / filename
        if not path.exists():
            raise FileNotFoundError(f"Missing canonical file: {path}")
        df = pd.read_csv(path, low_memory=False)
        data[filename] = df
        inventory_rows.append(
            {
                "File": filename,
                "ExpectedRows": expected_rows,
                "ActualRows": len(df),
                "RowCountStatus": "PASS" if len(df) == expected_rows else "FAIL",
                "SHA256": sha256(path),
            }
        )

    checks = []

    def add(check_id, test, expected, actual, passed):
        checks.append(
            {
                "CheckID": check_id,
                "TestName": test,
                "Expected": expected,
                "Actual": actual,
                "Status": "PASS" if passed else "FAIL",
            }
        )

    dim_state = data["dim_state.csv"]
    dim_issuer = data["dim_issuer.csv"]
    dim_plan = data["dim_plan.csv"]
    dim_metric = data["dim_metric.csv"]
    dim_status = data["dim_availability_status.csv"]
    f_issuer = data["fact_issuer_transparency.csv"]
    f_plan = data["fact_plan_transparency.csv"]
    a_issuer = data["fact_issuer_metric_availability.csv"]
    a_plan = data["fact_plan_metric_availability.csv"]

    add(
        "CAN-001",
        "DimState key uniqueness",
        0,
        int(dim_state["StateKey"].duplicated().sum()),
        int(dim_state["StateKey"].duplicated().sum()) == 0,
    )
    add(
        "CAN-002",
        "DimIssuer key uniqueness",
        0,
        int(dim_issuer["IssuerKey"].duplicated().sum()),
        int(dim_issuer["IssuerKey"].duplicated().sum()) == 0,
    )
    add(
        "CAN-003",
        "DimPlan key uniqueness",
        0,
        int(dim_plan["PlanKey"].duplicated().sum()),
        int(dim_plan["PlanKey"].duplicated().sum()) == 0,
    )
    add(
        "CAN-004",
        "DimMetric key uniqueness",
        0,
        int(dim_metric["MetricKey"].duplicated().sum()),
        int(dim_metric["MetricKey"].duplicated().sum()) == 0,
    )
    add(
        "CAN-005",
        "DimAvailabilityStatus key uniqueness",
        0,
        int(dim_status["StatusKey"].duplicated().sum()),
        int(dim_status["StatusKey"].duplicated().sum()) == 0,
    )

    fk_tests = [
        ("CAN-006", "DimIssuer -> DimState", dim_issuer["StateKey"], dim_state["StateKey"]),
        ("CAN-007", "DimPlan -> DimIssuer", dim_plan["IssuerKey"], dim_issuer["IssuerKey"]),
        ("CAN-008", "FactIssuer -> DimIssuer", f_issuer["IssuerKey"], dim_issuer["IssuerKey"]),
        ("CAN-009", "FactPlan -> DimPlan", f_plan["PlanKey"], dim_plan["PlanKey"]),
        ("CAN-010", "IssuerAvailability -> DimIssuer", a_issuer["IssuerKey"], dim_issuer["IssuerKey"]),
        ("CAN-011", "PlanAvailability -> DimPlan", a_plan["PlanKey"], dim_plan["PlanKey"]),
        ("CAN-012", "IssuerAvailability -> DimMetric", a_issuer["MetricKey"], dim_metric["MetricKey"]),
        ("CAN-013", "PlanAvailability -> DimMetric", a_plan["MetricKey"], dim_metric["MetricKey"]),
        ("CAN-014", "IssuerAvailability -> Status", a_issuer["AvailabilityStatusKey"], dim_status["StatusKey"]),
        ("CAN-015", "PlanAvailability -> Status", a_plan["AvailabilityStatusKey"], dim_status["StatusKey"]),
    ]

    for check_id, test_name, child, parent in fk_tests:
        broken = int((~child.isin(parent)).sum())
        add(check_id, test_name, 0, broken, broken == 0)

    return pd.DataFrame(checks), pd.DataFrame(inventory_rows)


def style_report(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row and ws.max_column:
            ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(1, ws.max_column + 1):
            max_len = 0
            for row_idx in range(1, min(ws.max_row, 200) + 1):
                value = ws.cell(row_idx, col_idx).value
                if value is not None:
                    max_len = max(max_len, len(str(value)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(
                max(max_len + 2, 10),
                60,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(path)


def main() -> int:
    print("=" * 78)
    print("Transparency in Coverage PUF — Power BI Semantic Model Validation")
    print("=" * 78)
    print(f"Script version : {SCRIPT_VERSION}")
    print(f"Power BI root  : {POWERBI_ROOT}")
    print(f"Canonical dir  : {CANONICAL_DIR}")
    print(f"Review file    : {OUTPUT_FILE}")

    semantic_model = find_semantic_model()
    definition_dir = semantic_model / "definition"
    model_file = definition_dir / "model.tmdl"
    relationships_file = definition_dir / "relationships.tmdl"
    expressions_file = definition_dir / "expressions.tmdl"
    tables_dir = definition_dir / "tables"

    for path in [model_file, relationships_file, expressions_file]:
        if not path.exists():
            raise FileNotFoundError(f"Missing required PBIP file: {path}")

    model_text = read_text(model_file)
    relationships_text = read_text(relationships_file)
    expressions_text = read_text(expressions_file)

    refs = model_table_refs(model_text)
    relationships = parse_relationships(relationships_text)
    columns = parse_table_columns(tables_dir)

    validations = []

    def add(check_id, test, expected, actual, passed):
        validations.append(
            {
                "CheckID": check_id,
                "TestName": test,
                "Expected": expected,
                "Actual": actual,
                "Status": "PASS" if passed else "FAIL",
            }
        )

    actual_tables = set(refs)
    add(
        "PBI-001",
        "Business table inventory",
        len(EXPECTED_TABLES),
        len(actual_tables),
        actual_tables == EXPECTED_TABLES,
    )

    auto_date_refs = [
        r for r in refs
        if r.startswith("DateTableTemplate_")
        or r.startswith("LocalDateTable_")
    ]
    add("PBI-002", "Auto-date table refs", 0, len(auto_date_refs), len(auto_date_refs) == 0)

    auto_date_files = [
        p.name
        for p in tables_dir.glob("*.tmdl")
        if p.stem.startswith("DateTableTemplate_")
        or p.stem.startswith("LocalDateTable_")
    ]
    add("PBI-003", "Auto-date table files", 0, len(auto_date_files), len(auto_date_files) == 0)

    auto_date_enabled = "__PBI_TimeIntelligenceEnabled = 1" in model_text
    add("PBI-004", "Auto Date/Time enabled", False, auto_date_enabled, not auto_date_enabled)

    actual_pairs = set(
        zip(relationships["FromColumn"], relationships["ToColumn"])
    )
    add(
        "PBI-005",
        "Governed relationship count",
        14,
        len(relationships),
        len(relationships) == 14,
    )
    add(
        "PBI-006",
        "Exact governed relationship endpoint set",
        "Exact",
        "Exact" if actual_pairs == EXPECTED_RELATIONSHIPS else "Mismatch",
        actual_pairs == EXPECTED_RELATIONSHIPS,
    )

    inactive = int(
        relationships["IsActive"].astype(str).str.lower().eq("false").sum()
    )
    add("PBI-007", "Inactive relationships", 0, inactive, inactive == 0)

    bidirectional = int(
        relationships["CrossFilteringBehavior"]
        .astype(str)
        .str.lower()
        .str.contains("both")
        .sum()
    )
    add("PBI-008", "Bidirectional relationships", 0, bidirectional, bidirectional == 0)

    many_to_many = int(
        (
            relationships["FromCardinality"].astype(str).str.lower().str.contains("many")
            & relationships["ToCardinality"].astype(str).str.lower().str.contains("many")
        ).sum()
    )
    add("PBI-009", "Many-to-many relationships", 0, many_to_many, many_to_many == 0)

    dq_connected = int(
        (
            relationships["FromColumn"].astype(str).str.contains("DQExceptionRegister")
            | relationships["ToColumn"].astype(str).str.contains("DQExceptionRegister")
        ).sum()
    )
    add("PBI-010", "DQExceptionRegister relationships", 0, dq_connected, dq_connected == 0)

    endpoint_strings = set(
        columns["TableName"].astype(str)
        + "."
        + columns["ColumnName"].astype(str)
    )
    broken_endpoints = sorted(
        {
            endpoint
            for pair in actual_pairs
            for endpoint in pair
            if endpoint not in endpoint_strings
        }
    )
    add(
        "PBI-011",
        "Relationship endpoints exist in model columns",
        0,
        len(broken_endpoints),
        len(broken_endpoints) == 0,
    )

    parameter_ok = (
        "expression DataFolderPath" in expressions_text
        and "IsParameterQuery=true" in expressions_text.replace(" ", "")
    )
    add("PBI-012", "DataFolderPath parameter exists", True, parameter_ok, parameter_ok)

    function_ok = "expression fxLoadCanonicalCsv" in expressions_text
    add("PBI-013", "Canonical CSV loader function exists", True, function_ok, function_ok)

    measure_matches = []
    for path in tables_dir.glob("*.tmdl"):
        text = read_text(path)
        measure_matches.extend(
            re.findall(r"(?m)^\tmeasure\s+(.+?)\s*=", text)
        )
    add(
        "PBI-014",
        "Governed measures in the saved portfolio release",
        43,
        len(measure_matches),
        len(measure_matches) == 43 and len(set(measure_matches)) == 43,
    )

    canonical_checks, canonical_inventory = canonical_integrity_checks()

    validations_df = pd.DataFrame(validations)
    all_checks = pd.concat(
        [
            validations_df,
            canonical_checks.rename(
                columns={
                    "CheckID": "CheckID",
                    "TestName": "TestName",
                    "Expected": "Expected",
                    "Actual": "Actual",
                    "Status": "Status",
                }
            ),
        ],
        ignore_index=True,
    )

    failures = int((all_checks["Status"] == "FAIL").sum())

    relationship_review = relationships.copy()
    relationship_review["CardinalityResolved"] = (
        relationship_review["FromCardinality"].astype(str)
        + " -> "
        + relationship_review["ToCardinality"].astype(str)
    )

    table_inventory = pd.DataFrame(
        {
            "ModelTable": sorted(actual_tables),
            "IsExpectedBusinessTable": [
                t in EXPECTED_TABLES for t in sorted(actual_tables)
            ],
        }
    )

    file_inventory = []
    for path in sorted(
        p for p in semantic_model.rglob("*") if p.is_file()
    ):
        file_inventory.append(
            {
                "RelativePath": str(path.relative_to(PROJECT_ROOT)),
                "Bytes": path.stat().st_size,
                "SHA256": sha256(path),
            }
        )
    file_inventory_df = pd.DataFrame(file_inventory)

    summary = pd.DataFrame(
        [
            ("ValidationStatus", "PASS" if failures == 0 else "FAIL"),
            ("ScriptVersion", SCRIPT_VERSION),
            ("ValidatedAtLocal", datetime.now().isoformat(timespec="seconds")),
            ("SemanticModel", str(semantic_model)),
            ("BusinessTables", len(actual_tables)),
            ("Relationships", len(relationships)),
            ("InactiveRelationships", inactive),
            ("BidirectionalRelationships", bidirectional),
            ("ManyToManyRelationships", many_to_many),
            ("AutoDateTableRefs", len(auto_date_refs)),
            ("AutoDateTableFiles", len(auto_date_files)),
            ("Measures", len(measure_matches)),
            ("ValidationChecks", len(all_checks)),
            ("ValidationFailures", failures),
            ("EvidenceState", "POST-DESKTOP-SAVE PBIP SOURCE VALIDATION"),
        ],
        columns=["Item", "Value"],
    )

    REVIEW_DIR.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="00_Summary", index=False)
        all_checks.to_excel(writer, sheet_name="01_Validation", index=False)
        relationship_review.to_excel(
            writer, sheet_name="02_Relationships", index=False
        )
        table_inventory.to_excel(
            writer, sheet_name="03_Table_Inventory", index=False
        )
        columns.to_excel(writer, sheet_name="04_Model_Columns", index=False)
        canonical_inventory.to_excel(
            writer, sheet_name="05_Canonical_Inventory", index=False
        )
        canonical_checks.to_excel(
            writer, sheet_name="06_Canonical_Integrity", index=False
        )
        file_inventory_df.to_excel(
            writer, sheet_name="07_File_Inventory", index=False
        )

    style_report(OUTPUT_FILE)

    print()
    print("Semantic-model validation completed.")
    print(f"Validation status      : {'PASS' if failures == 0 else 'FAIL'}")
    print(f"Business tables        : {len(actual_tables)}")
    print(f"Relationships          : {len(relationships)}")
    print(f"Inactive               : {inactive}")
    print(f"Bidirectional          : {bidirectional}")
    print(f"Many-to-many           : {many_to_many}")
    print(f"Auto-date tables       : {len(auto_date_refs)}")
    print(f"Measures               : {len(measure_matches)}")
    print(f"Validation checks      : {len(all_checks)}")
    print(f"Validation failures    : {failures}")
    print(f"Review evidence        : {OUTPUT_FILE}")
    print("PBIP and canonical files were read only.")

    return 0 if failures == 0 else 2


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:
        print()
        print("VALIDATION FAILED TO RUN")
        print(str(exc))
        sys.exit(1)
