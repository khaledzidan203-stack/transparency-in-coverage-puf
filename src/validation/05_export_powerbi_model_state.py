from __future__ import annotations

from pathlib import Path
from typing import Any
import hashlib
import os
import re
import sys
from datetime import datetime

import pandas as pd


SCRIPT_VERSION = "1.0.0"


PROJECT_ROOT = Path(__file__).resolve().parents[2]
POWERBI_ROOT = PROJECT_ROOT / "powerbi"


def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, key_path) as key:
                raw_value, _ = winreg.QueryValueEx(key, "Desktop")
            return Path(os.path.expandvars(raw_value)).resolve()
        except Exception:
            pass
    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = Path(os.environ.get("TRANSPARENCY_PUF_REVIEW_DIR", PROJECT_ROOT / ".local-review")).expanduser().resolve()
OUTPUT_FILE = REVIEW_DIR / "05_POWER_BI_MODEL_STATE.xlsx"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "utf-16"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def unquote_tmdl_name(value: str) -> str:
    value = value.strip()
    if len(value) >= 2 and value[0] == "'" and value[-1] == "'":
        return value[1:-1].replace("''", "'")
    return value


def split_column_ref(value: str) -> tuple[str, str]:
    value = value.strip()
    idx = value.rfind("[")
    if idx == -1 or not value.endswith("]"):
        return value, ""
    table = unquote_tmdl_name(value[:idx].strip())
    column = value[idx + 1 : -1].strip()
    return table, column


def first_property(block: str, prop: str) -> str:
    pattern = rf"(?mi)^\s*{re.escape(prop)}\s*:\s*(.*?)\s*$"
    match = re.search(pattern, block)
    return match.group(1).strip() if match else ""


def flag_present(block: str, flag: str) -> bool:
    return bool(
        re.search(
            rf"(?mi)^\s*{re.escape(flag)}(?:\s*:\s*true)?\s*$",
            block,
        )
    )


def extract_blocks(text: str, keyword: str) -> list[tuple[str, str]]:
    lines = text.splitlines()
    blocks: list[tuple[str, str]] = []
    header_pattern = re.compile(
        rf"^(\s*){re.escape(keyword)}\s+(.+?)\s*$",
        re.IGNORECASE,
    )

    i = 0
    while i < len(lines):
        match = header_pattern.match(lines[i])
        if not match:
            i += 1
            continue

        indent = len(match.group(1).replace("\t", "    "))
        header_remainder = match.group(2).strip()
        start = i
        i += 1

        while i < len(lines):
            line = lines[i]
            if line.strip() == "":
                i += 1
                continue

            leading = line[: len(line) - len(line.lstrip(" \t"))]
            current_indent = len(leading.replace("\t", "    "))

            if current_indent <= indent:
                break
            i += 1

        raw = "\n".join(lines[start:i]).rstrip()
        blocks.append((header_remainder, raw))

    return blocks


def find_semantic_model() -> Path:
    candidates = sorted(
        p for p in POWERBI_ROOT.glob("*.SemanticModel") if p.is_dir()
    )
    if not candidates:
        candidates = sorted(
            p for p in POWERBI_ROOT.rglob("*.SemanticModel") if p.is_dir()
        )

    if len(candidates) == 0:
        raise FileNotFoundError(
            f"No *.SemanticModel folder found under: {POWERBI_ROOT}"
        )

    if len(candidates) > 1:
        names = "\n".join(f"- {p}" for p in candidates)
        raise RuntimeError(
            "More than one SemanticModel folder was found.\n" + names
        )

    return candidates[0]


def find_report_folder() -> Path | None:
    candidates = sorted(
        p for p in POWERBI_ROOT.glob("*.Report") if p.is_dir()
    )
    return candidates[0] if len(candidates) == 1 else None


def parse_relationships(tmdl_files: list[Path]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []

    for path in tmdl_files:
        text = read_text(path)
        for relationship_name, block in extract_blocks(text, "relationship"):
            from_ref = first_property(block, "fromColumn")
            to_ref = first_property(block, "toColumn")
            from_table, from_column = split_column_ref(from_ref)
            to_table, to_column = split_column_ref(to_ref)

            explicit_active = first_property(block, "isActive")
            explicit_cross = first_property(block, "crossFilteringBehavior")
            explicit_from_card = first_property(block, "fromCardinality")
            explicit_to_card = first_property(block, "toCardinality")
            security_cross = first_property(block, "securityFilteringBehavior")

            resolved_active = (
                explicit_active.lower()
                if explicit_active
                else "true (default)"
            )
            resolved_from_card = (
                explicit_from_card
                if explicit_from_card
                else "many (default)"
            )
            resolved_to_card = (
                explicit_to_card
                if explicit_to_card
                else "one (default)"
            )
            resolved_cross = (
                explicit_cross
                if explicit_cross
                else "oneDirection (default)"
            )

            rows.append(
                {
                    "RelationshipName": unquote_tmdl_name(relationship_name),
                    "FromTable": from_table,
                    "FromColumn": from_column,
                    "ToTable": to_table,
                    "ToColumn": to_column,
                    "ExplicitIsActive": explicit_active,
                    "ResolvedIsActive": resolved_active,
                    "ExplicitFromCardinality": explicit_from_card,
                    "ResolvedFromCardinality": resolved_from_card,
                    "ExplicitToCardinality": explicit_to_card,
                    "ResolvedToCardinality": resolved_to_card,
                    "ExplicitCrossFiltering": explicit_cross,
                    "ResolvedCrossFiltering": resolved_cross,
                    "SecurityFilteringBehavior": security_cross,
                    "DefinitionFile": str(path.relative_to(PROJECT_ROOT)),
                    "RawDefinition": block,
                }
            )

    return pd.DataFrame(
        rows,
        columns=[
            "RelationshipName",
            "FromTable",
            "FromColumn",
            "ToTable",
            "ToColumn",
            "ExplicitIsActive",
            "ResolvedIsActive",
            "ExplicitFromCardinality",
            "ResolvedFromCardinality",
            "ExplicitToCardinality",
            "ResolvedToCardinality",
            "ExplicitCrossFiltering",
            "ResolvedCrossFiltering",
            "SecurityFilteringBehavior",
            "DefinitionFile",
            "RawDefinition",
        ],
    )


def parse_tables_and_objects(
    tmdl_files: list[Path],
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    table_rows = []
    column_rows = []
    measure_rows = []
    partition_rows = []

    for path in tmdl_files:
        text = read_text(path)
        for table_header, table_block in extract_blocks(text, "table"):
            table_name = unquote_tmdl_name(table_header.split("=", 1)[0].strip())

            columns = extract_blocks(table_block, "column")
            measures = extract_blocks(table_block, "measure")
            partitions = extract_blocks(table_block, "partition")

            table_rows.append(
                {
                    "TableName": table_name,
                    "IsHidden": flag_present(table_block, "isHidden"),
                    "ColumnCount": len(columns),
                    "MeasureCount": len(measures),
                    "PartitionCount": len(partitions),
                    "DefinitionFile": str(path.relative_to(PROJECT_ROOT)),
                }
            )

            for header, block in columns:
                name = header.split("=", 1)[0].strip()
                column_rows.append(
                    {
                        "TableName": table_name,
                        "ColumnName": unquote_tmdl_name(name),
                        "DataType": first_property(block, "dataType"),
                        "SourceColumn": first_property(block, "sourceColumn"),
                        "FormatString": first_property(block, "formatString"),
                        "SummarizeBy": first_property(block, "summarizeBy"),
                        "SortByColumn": first_property(block, "sortByColumn"),
                        "IsHidden": flag_present(block, "isHidden"),
                        "IsKey": flag_present(block, "isKey"),
                        "DataCategory": first_property(block, "dataCategory"),
                        "DefinitionFile": str(path.relative_to(PROJECT_ROOT)),
                    }
                )

            for header, block in measures:
                if "=" in header:
                    name, expression_first = header.split("=", 1)
                else:
                    name, expression_first = header, ""
                measure_rows.append(
                    {
                        "TableName": table_name,
                        "MeasureName": unquote_tmdl_name(name.strip()),
                        "ExpressionFirstLine": expression_first.strip(),
                        "FormatString": first_property(block, "formatString"),
                        "DisplayFolder": first_property(block, "displayFolder"),
                        "IsHidden": flag_present(block, "isHidden"),
                        "DefinitionFile": str(path.relative_to(PROJECT_ROOT)),
                        "RawDefinition": block,
                    }
                )

            for header, block in partitions:
                name = header.split("=", 1)[0].strip()
                partition_rows.append(
                    {
                        "TableName": table_name,
                        "PartitionName": unquote_tmdl_name(name),
                        "Mode": first_property(block, "mode"),
                        "SourceType": (
                            "M"
                            if re.search(r"(?mi)^\s*source\s*=\s*m\s*$", block)
                            else ""
                        ),
                        "DefinitionFile": str(path.relative_to(PROJECT_ROOT)),
                        "RawDefinition": block,
                    }
                )

    return (
        pd.DataFrame(table_rows),
        pd.DataFrame(column_rows),
        pd.DataFrame(measure_rows),
        pd.DataFrame(partition_rows),
    )


def parse_expressions(tmdl_files: list[Path]) -> pd.DataFrame:
    rows = []

    for path in tmdl_files:
        text = read_text(path)
        for header, block in extract_blocks(text, "expression"):
            if "=" in header:
                name, first_line = header.split("=", 1)
            else:
                name, first_line = header, ""

            compact = block.replace(" ", "")
            rows.append(
                {
                    "ExpressionName": unquote_tmdl_name(name.strip()),
                    "ExpressionFirstLine": first_line.strip(),
                    "IsParameterQuery": "IsParameterQuery=true" in compact,
                    "DefinitionFile": str(path.relative_to(PROJECT_ROOT)),
                    "RawDefinition": block,
                }
            )

    return pd.DataFrame(rows)


def file_inventory(
    semantic_model: Path,
    report_folder: Path | None,
) -> pd.DataFrame:
    roots = [semantic_model]
    if report_folder is not None:
        roots.append(report_folder)

    rows = []

    for root in roots:
        for path in sorted(p for p in root.rglob("*") if p.is_file()):
            rows.append(
                {
                    "Area": "SemanticModel" if root == semantic_model else "Report",
                    "RelativePath": str(path.relative_to(PROJECT_ROOT)),
                    "Extension": path.suffix.lower(),
                    "Bytes": path.stat().st_size,
                    "ModifiedLocal": datetime.fromtimestamp(
                        path.stat().st_mtime
                    ).isoformat(timespec="seconds"),
                    "SHA256": sha256(path),
                }
            )

    return pd.DataFrame(rows)


def raw_tmdl_inventory(tmdl_files: list[Path]) -> pd.DataFrame:
    rows = []
    for path in tmdl_files:
        text = read_text(path)
        rows.append(
            {
                "RelativePath": str(path.relative_to(PROJECT_ROOT)),
                "LineCount": len(text.splitlines()),
                "CharacterCount": len(text),
                "SHA256": sha256(path),
                "FullText": text,
            }
        )
    return pd.DataFrame(rows)


def relationship_review(
    relationships: pd.DataFrame,
    tables: pd.DataFrame,
) -> pd.DataFrame:
    table_names = (
        set(tables["TableName"].astype(str))
        if not tables.empty and "TableName" in tables.columns
        else set()
    )
    rows = []

    for _, r in relationships.iterrows():
        issues = []

        if not r["FromTable"] or not r["ToTable"]:
            issues.append("UNPARSED_ENDPOINT")
        if r["FromTable"] and r["FromTable"] not in table_names:
            issues.append("FROM_TABLE_NOT_FOUND")
        if r["ToTable"] and r["ToTable"] not in table_names:
            issues.append("TO_TABLE_NOT_FOUND")

        resolved_from = str(r["ResolvedFromCardinality"]).lower()
        resolved_to = str(r["ResolvedToCardinality"]).lower()
        resolved_cross = str(r["ResolvedCrossFiltering"]).lower()

        if "many" in resolved_from and "many" in resolved_to:
            issues.append("MANY_TO_MANY_REVIEW")
        if "both" in resolved_cross:
            issues.append("BIDIRECTIONAL_REVIEW")

        rows.append(
            {
                "RelationshipName": r["RelationshipName"],
                "From": f"{r['FromTable']}[{r['FromColumn']}]",
                "To": f"{r['ToTable']}[{r['ToColumn']}]",
                "ResolvedCardinality": (
                    f"{r['ResolvedFromCardinality']} -> "
                    f"{r['ResolvedToCardinality']}"
                ),
                "ResolvedCrossFiltering": r["ResolvedCrossFiltering"],
                "ResolvedIsActive": r["ResolvedIsActive"],
                "AutomatedReviewFlag": " | ".join(issues) if issues else "NONE",
            }
        )

    return pd.DataFrame(rows)


def style_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        sample_rows = min(ws.max_row, 200)
        for col_idx in range(1, ws.max_column + 1):
            max_len = 0
            for row_idx in range(1, sample_rows + 1):
                value = ws.cell(row_idx, col_idx).value
                if value is not None:
                    max_len = max(max_len, len(str(value)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(
                max(max_len + 2, 10),
                60,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(path)


def main() -> int:
    print("=" * 78)
    print("Transparency in Coverage PUF — Power BI Model State Export")
    print("=" * 78)
    print(f"Script version : {SCRIPT_VERSION}")
    print(f"Project root   : {PROJECT_ROOT}")
    print(f"Power BI root  : {POWERBI_ROOT}")
    print(f"Output         : {OUTPUT_FILE}")

    if not POWERBI_ROOT.exists():
        raise FileNotFoundError(f"Power BI folder not found: {POWERBI_ROOT}")

    semantic_model = find_semantic_model()
    report_folder = find_report_folder()

    print(f"Semantic model : {semantic_model}")
    if report_folder:
        print(f"Report folder  : {report_folder}")

    tmdl_files = sorted(semantic_model.rglob("*.tmdl"))
    if not tmdl_files:
        raise RuntimeError(
            "No .tmdl files were found. Save the PBIP project in Power BI Desktop first."
        )

    tables, columns, measures, partitions = parse_tables_and_objects(tmdl_files)
    relationships = parse_relationships(tmdl_files)
    expressions = parse_expressions(tmdl_files)
    files = file_inventory(semantic_model, report_folder)
    raw_tmdl = raw_tmdl_inventory(tmdl_files)
    rel_review = relationship_review(relationships, tables)

    summary = pd.DataFrame(
        [
            ("ExportStatus", "PASS"),
            ("ScriptVersion", SCRIPT_VERSION),
            ("ExportedAtLocal", datetime.now().isoformat(timespec="seconds")),
            ("ProjectRoot", str(PROJECT_ROOT)),
            ("PowerBIRoot", str(POWERBI_ROOT)),
            ("SemanticModelFolder", str(semantic_model)),
            ("ReportFolder", str(report_folder) if report_folder else ""),
            ("TMDLFiles", len(tmdl_files)),
            ("Tables", len(tables)),
            ("Columns", len(columns)),
            ("Measures", len(measures)),
            ("Partitions", len(partitions)),
            ("Relationships", len(relationships)),
            ("Expressions", len(expressions)),
            (
                "RelationshipReviewFlags",
                int((rel_review["AutomatedReviewFlag"] != "NONE").sum())
                if not rel_review.empty and "AutomatedReviewFlag" in rel_review.columns
                else 0,
            ),
            ("EvidenceState", "READ_ONLY CURRENT PBIP SOURCE SNAPSHOT"),
        ],
        columns=["Item", "Value"],
    )

    notes = pd.DataFrame(
        [
            {
                "Note": (
                    "This export reports the relationships currently persisted "
                    "in PBIP/TMDL and does not modify the model."
                )
            },
            {
                "Note": (
                    "If Relationships = 0, the saved PBIP source currently has "
                    "no persisted semantic relationships."
                )
            },
            {
                "Note": (
                    "RawDefinition is retained as the authoritative persisted "
                    "relationship evidence."
                )
            },
        ]
    )

    REVIEW_DIR.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="00_Summary", index=False)
        notes.to_excel(writer, sheet_name="01_Notes", index=False)
        relationships.to_excel(writer, sheet_name="02_Relationships", index=False)
        rel_review.to_excel(writer, sheet_name="03_Relationship_Review", index=False)
        tables.to_excel(writer, sheet_name="04_Tables", index=False)
        columns.to_excel(writer, sheet_name="05_Columns", index=False)
        measures.to_excel(writer, sheet_name="06_Measures", index=False)
        partitions.to_excel(writer, sheet_name="07_Partitions", index=False)
        expressions.to_excel(writer, sheet_name="08_Expressions", index=False)
        files.to_excel(writer, sheet_name="09_File_Inventory", index=False)
        raw_tmdl.to_excel(writer, sheet_name="10_Raw_TMDL", index=False)

    style_workbook(OUTPUT_FILE)

    print()
    print("Power BI model-state export completed successfully.")
    print(f"Tables        : {len(tables)}")
    print(f"Columns       : {len(columns)}")
    print(f"Measures      : {len(measures)}")
    print(f"Partitions    : {len(partitions)}")
    print(f"Relationships : {len(relationships)}")
    print(f"Expressions   : {len(expressions)}")
    print(f"Review file   : {OUTPUT_FILE}")
    print("PBIP project was read only; no project file was modified.")

    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:
        print()
        print("EXPORT FAILED")
        print(str(exc))
        sys.exit(1)
