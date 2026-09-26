from __future__ import annotations

from datetime import datetime
from pathlib import Path
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile

import pandas as pd


SCRIPT_VERSION = "1.0.0"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
POWERBI_ROOT = PROJECT_ROOT / "powerbi"
REPORT_DIR = POWERBI_ROOT / "TransparencyInCoverage.Report"
SEMANTIC_DIR = POWERBI_ROOT / "TransparencyInCoverage.SemanticModel"
PAGES_DIR = REPORT_DIR / "definition" / "pages"

MEASURES_TMDL = (
    SEMANTIC_DIR
    / "definition"
    / "tables"
    / "_Measures.tmdl"
)
DIM_STATE_TMDL = (
    SEMANTIC_DIR
    / "definition"
    / "tables"
    / "DimState.tmdl"
)
DIM_ISSUER_TMDL = (
    SEMANTIC_DIR
    / "definition"
    / "tables"
    / "DimIssuer.tmdl"
)

VISUAL_SCHEMA = (
    "https://developer.microsoft.com/json-schemas/fabric/item/report/"
    "definition/visualContainer/2.7.0/schema.json"
)

CANVAS_WIDTH = 1920
CANVAS_HEIGHT = 1080

COLORS = {
    "canvas": "#F5F7FA",
    "surface": "#FFFFFF",
    "surface_alt": "#F9FAFB",
    "border": "#E7ECF2",
    "ink": "#172033",
    "secondary": "#667085",
    "muted": "#98A2B3",
    "navy": "#17324D",
    "slate": "#344054",
    "blue": "#4F6BED",
    "blue_dark": "#425CC7",
    "blue_soft": "#EEF3FF",
    "blue_pale": "#E5ECFF",
    "teal": "#16A6A1",
    "teal_dark": "#0E7774",
    "teal_soft": "#EAF8F6",
    "amber": "#C8872C",
    "amber_dark": "#9C641A",
    "amber_soft": "#FFF4E3",
    "red": "#C64B4B",
    "red_dark": "#9F2F2F",
    "red_soft": "#FCE8E8",
    "white": "#FFFFFF",
}

FONT_REGULAR = (
    "'Segoe UI', wf_segoe-ui_normal, helvetica, arial, sans-serif"
)
FONT_SEMIBOLD = (
    "'Segoe UI Semibold', wf_segoe-ui_semibold, "
    "helvetica, arial, sans-serif"
)

EXPECTED_PAGES = [
    "00 INDEX",
    "01 Executive Overview",
    "02 Network Comparison",
    "03 Denials",
    "04 Denial Reasons",
    "05 Appeals",
    "06 Resubmissions",
    "07 Enrollment",
    "08 Issuer & State Explorer",
    "09 Data Availability",
    "10 Data Quality & Methodology",
]

REQUIRED_MEASURES = {
    "Issuer Resubmission Events per 100 Denied - In Network",
    "Issuer Resubmission Events per 100 Denied - Out of Network",
    "Issuer Comparable Total Resubmission Events",
    "Plan Resubmission Events per 100 Denied - In Network",
    "Plan Resubmission Events per 100 Denied - Out of Network",
    "Open DQ Exception Count",
}

REQUIRED_COLUMNS = {
    "DimState": {
        "StateCode",
    },
    "DimIssuer": {
        "IssuerName",
    },
}

UNSUPPORTED_PAGE_ROOT_PROPERTIES = {
    "horizontalAlignment",
    "verticalAlignment",
}

EXPECTED_BOUND_VISUALS = 10

EXPECTED_VISUAL_NAMES = {
    # Header / navigation / filters.
    "p06_context",
    "p06_page_title",
    "p06_index_surface",
    "p06_index_label",
    "p06_index_nav",
    "p06_state_slicer",
    "p06_issuer_slicer",
    "p06_dq_global",
    "p06_band_dq_global",

    # KPI strip.
    "p06_total_events",
    "p06_band_total_events",
    "p06_issuer_in",
    "p06_band_issuer_in",
    "p06_issuer_out",
    "p06_band_issuer_out",
    "p06_plan_in",
    "p06_band_plan_in",
    "p06_plan_out",
    "p06_band_plan_out",

    # Main state story.
    "p06_band_state_chart",
    "p06_state_chart",
    "p06_state_note",

    # Investigation.
    "p06_band_investigation",
    "p06_rate_legend",
    "p06_issuer_table",

    # Footer.
    "p06_methodology_note",
}


# =====================================================================
# Paths / environment
# =====================================================================

def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )

            with winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                key_path,
            ) as key:
                raw_value, _ = winreg.QueryValueEx(
                    key,
                    "Desktop",
                )

            return Path(
                os.path.expandvars(raw_value)
            ).resolve()

        except Exception:
            pass

    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = (
    get_windows_desktop()
    / "Transparency_PUF_Review"
)

BACKUP_FILE = (
    REVIEW_DIR
    / "35_PRE_RESUBMISSIONS_STORYTELLING_BACKUP.zip"
)

REPORT_FILE = (
    REVIEW_DIR
    / "35_RESUBMISSIONS_STORYTELLING_BUILD_REPORT.xlsx"
)


def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False

    result = subprocess.run(
        [
            "tasklist",
            "/FI",
            "IMAGENAME eq PBIDesktop.exe",
        ],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(
            subprocess,
            "CREATE_NO_WINDOW",
            0,
        ),
    )

    return "PBIDesktop.exe" in result.stdout


def read_json(path: Path):
    return json.loads(
        path.read_text(
            encoding="utf-8-sig",
        )
    )


def write_json(
    path: Path,
    payload,
) -> None:
    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    path.write_text(
        json.dumps(
            payload,
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )


def sha256(path: Path) -> str:
    h = hashlib.sha256()

    with path.open("rb") as f:
        for chunk in iter(
            lambda: f.read(1024 * 1024),
            b"",
        ):
            h.update(chunk)

    return h.hexdigest()


# =====================================================================
# PBIR expression helpers
# =====================================================================

def expr_bool(value: bool) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": "true" if value else "false"
            }
        }
    }


def expr_string(value: str) -> dict:
    escaped = value.replace(
        "'",
        "''",
    )

    return {
        "expr": {
            "Literal": {
                "Value": f"'{escaped}'"
            }
        }
    }


def expr_double(
    value: float | int,
) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": f"{value}D"
            }
        }
    }


def expr_int(
    value: int,
) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": f"{value}L"
            }
        }
    }


def solid_color(
    color: str,
) -> dict:
    return {
        "solid": {
            "color": expr_string(
                color
            )
        }
    }


def measure_expr(
    measure_name: str,
) -> dict:
    return {
        "Measure": {
            "Expression": {
                "SourceRef": {
                    "Entity": "_Measures"
                }
            },
            "Property": measure_name,
        }
    }


def literal_expr(
    value: float | int,
) -> dict:
    return {
        "Literal": {
            "Value": f"{value}D"
        }
    }


# =====================================================================
# TMDL inventory
# =====================================================================

def parse_tmdl_measures(
    path: Path,
) -> set[str]:
    text = path.read_text(
        encoding="utf-8-sig",
    )

    measures: set[str] = set()

    for line in text.splitlines():
        stripped = line.lstrip()

        if not stripped.startswith(
            "measure "
        ):
            continue

        lhs = stripped.split(
            "=",
            1,
        )[0].strip()

        raw = lhs[
            len("measure "):
        ].strip()

        if (
            raw.startswith("'")
            and raw.endswith("'")
        ):
            raw = (
                raw[1:-1]
                .replace(
                    "''",
                    "'",
                )
            )

        measures.add(
            raw
        )

    return measures


def parse_tmdl_columns(
    path: Path,
) -> set[str]:
    text = path.read_text(
        encoding="utf-8-sig",
    )

    columns: set[str] = set()

    for line in text.splitlines():
        stripped = line.lstrip()

        if not stripped.startswith(
            "column "
        ):
            continue

        raw = stripped[
            len("column "):
        ].strip()

        if raw.startswith("'"):
            match = re.match(
                r"'((?:''|[^'])*)'",
                raw,
            )

            if match:
                columns.add(
                    match.group(1)
                    .replace(
                        "''",
                        "'",
                    )
                )
        else:
            columns.add(
                raw.split(
                    "\t",
                    1,
                )[0].strip()
            )

    return columns


# =====================================================================
# Report discovery / backup
# =====================================================================

def find_pages() -> dict[
    str,
    tuple[Path, dict]
]:
    pages = {}

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        display_name = payload.get(
            "displayName"
        )

        if display_name:
            pages[
                display_name
            ] = (
                page_json.parent,
                payload,
            )

    return pages


def create_backup() -> None:
    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if BACKUP_FILE.exists():
        BACKUP_FILE.unlink()

    with zipfile.ZipFile(
        BACKUP_FILE,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zf:
        for path in sorted(
            p
            for p in REPORT_DIR.rglob("*")
            if p.is_file()
        ):
            archive_name = (
                Path(REPORT_DIR.name)
                / path.relative_to(
                    REPORT_DIR
                )
            )

            zf.write(
                path,
                str(
                    archive_name
                ),
            )


def file_hashes(
    root: Path,
) -> dict[str, str]:
    result = {}

    for path in root.rglob("*"):
        if path.is_file():
            result[
                str(
                    path.relative_to(
                        REPORT_DIR
                    )
                )
            ] = sha256(
                path
            )

    return result


def prepare_target_page(
    page_dir: Path,
) -> None:
    page_json = (
        page_dir
        / "page.json"
    )

    payload = read_json(
        page_json
    )

    for prop in (
        UNSUPPORTED_PAGE_ROOT_PROPERTIES
    ):
        payload.pop(
            prop,
            None,
        )

    payload[
        "width"
    ] = CANVAS_WIDTH

    payload[
        "height"
    ] = CANVAS_HEIGHT

    payload[
        "displayOption"
    ] = "FitToPage"

    objects = payload.setdefault(
        "objects",
        {},
    )

    objects[
        "background"
    ] = [
        {
            "properties": {
                "color": solid_color(
                    COLORS["canvas"]
                ),
                "transparency": expr_double(
                    0
                ),
            }
        }
    ]

    write_json(
        page_json,
        payload,
    )


# =====================================================================
# Text / containers
# =====================================================================

def text_run(
    value: str,
    *,
    size: float,
    color: str,
    semibold: bool = False,
) -> dict:
    return {
        "value": value,
        "textStyle": {
            "fontFamily": (
                FONT_SEMIBOLD
                if semibold
                else FONT_REGULAR
            ),
            "fontSize": (
                f"{size}pt"
            ),
            "color": color,
        },
    }


def paragraph(
    *runs: dict,
    align: str = "center",
) -> dict:
    return {
        "textRuns": list(
            runs
        ),
        "horizontalTextAlignment": (
            align
        ),
    }


def container_objects(
    *,
    background: str | None = None,
    border: str | None = None,
    radius: int = 8,
    show_title: bool = False,
    title: str = "",
    title_size: float = 10,
    title_color: str = COLORS["ink"],
) -> dict:
    return {
        "title": [
            {
                "properties": {
                    "show": expr_bool(
                        show_title
                    ),
                    **(
                        {
                            "text": expr_string(
                                title
                            ),
                            "alignment": expr_string(
                                "center"
                            ),
                            "fontColor": solid_color(
                                title_color
                            ),
                            "fontFamily": expr_string(
                                FONT_SEMIBOLD
                            ),
                            "fontSize": expr_double(
                                title_size
                            ),
                            "titleWrap": expr_bool(
                                True
                            ),
                        }
                        if show_title
                        else {}
                    ),
                }
            }
        ],
        "background": [
            {
                "properties": {
                    "show": expr_bool(
                        background is not None
                    ),
                    **(
                        {
                            "color": solid_color(
                                background
                            ),
                            "transparency": expr_double(
                                0
                            ),
                        }
                        if background is not None
                        else {}
                    ),
                }
            }
        ],
        "border": [
            {
                "properties": {
                    "show": expr_bool(
                        border is not None
                    ),
                    **(
                        {
                            "color": solid_color(
                                border
                            ),
                            "width": expr_double(
                                1
                            ),
                            "radius": expr_double(
                                radius
                            ),
                        }
                        if border is not None
                        else {}
                    ),
                }
            }
        ],
        "dropShadow": [
            {
                "properties": {
                    "show": expr_bool(
                        False
                    )
                }
            }
        ],
        "visualHeader": [
            {
                "properties": {
                    "show": expr_bool(
                        False
                    )
                }
            }
        ],
    }


def make_textbox(
    *,
    name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    paragraphs: list[dict],
    background: str | None = None,
    border: str | None = None,
    radius: int = 8,
) -> dict:
    return {
        "$schema": VISUAL_SCHEMA,
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
        },
        "visual": {
            "visualType": "textbox",
            "objects": {
                "general": [
                    {
                        "properties": {
                            "paragraphs": paragraphs
                        }
                    }
                ]
            },
            "visualContainerObjects": (
                container_objects(
                    background=background,
                    border=border,
                    radius=radius,
                )
            ),
            "drillFilterOtherVisuals": True,
        },
    }


def make_title_band(
    *,
    name: str,
    title: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    background: str,
) -> dict:
    return make_textbox(
        name=name,
        x=x,
        y=y,
        width=width,
        height=height,
        z=z,
        paragraphs=[
            paragraph(
                text_run(
                    title,
                    size=9.6,
                    color=COLORS["white"],
                    semibold=True,
                )
            )
        ],
        background=background,
        border=background,
        radius=8,
    )


# =====================================================================
# Query helpers
# =====================================================================

def measure_field(
    measure_name: str,
) -> dict:
    return {
        "Measure": {
            "Expression": {
                "SourceRef": {
                    "Entity": "_Measures"
                }
            },
            "Property": measure_name,
        }
    }


def column_field(
    entity: str,
    column_name: str,
) -> dict:
    return {
        "Column": {
            "Expression": {
                "SourceRef": {
                    "Entity": entity
                }
            },
            "Property": column_name,
        }
    }


def measure_projection(
    measure_name: str,
    *,
    display_name: str | None = None,
) -> dict:
    result = {
        "field": measure_field(
            measure_name
        ),
        "queryRef": (
            f"_Measures.{measure_name}"
        ),
        "nativeQueryRef": measure_name,
    }

    if display_name is not None:
        result[
            "displayName"
        ] = display_name

    return result


def column_projection(
    entity: str,
    column_name: str,
    *,
    active: bool = False,
    display_name: str | None = None,
) -> dict:
    result = {
        "field": column_field(
            entity,
            column_name,
        ),
        "queryRef": (
            f"{entity}.{column_name}"
        ),
        "nativeQueryRef": column_name,
    }

    if active:
        result[
            "active"
        ] = True

    if display_name is not None:
        result[
            "displayName"
        ] = display_name

    return result


# =====================================================================
# Navigation / slicers / cards
# =====================================================================

def make_index_surface(
    *,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
) -> dict:
    return make_textbox(
        name="p06_index_surface",
        x=x,
        y=y,
        width=width,
        height=height,
        z=z,
        paragraphs=[
            paragraph(
                text_run(
                    " ",
                    size=6,
                    color=COLORS["blue"],
                )
            )
        ],
        background=COLORS["blue"],
        border=COLORS["blue_dark"],
        radius=9,
    )


def make_index_label(
    *,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
) -> dict:
    return make_textbox(
        name="p06_index_label",
        x=x,
        y=y,
        width=width,
        height=height,
        z=z,
        paragraphs=[
            paragraph(
                text_run(
                    "INDEX",
                    size=11,
                    color=COLORS["white"],
                    semibold=True,
                )
            )
        ],
        background=None,
        border=None,
        radius=0,
    )


def make_index_nav(
    *,
    target_page_name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
) -> dict:
    return {
        "$schema": VISUAL_SCHEMA,
        "name": "p06_index_nav",
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": 0,
        },
        "visual": {
            "visualType": "actionButton",
            "objects": {
                "icon": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "outline": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "fill": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": {
                "title": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "background": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "border": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "visualHeader": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "visualLink": [
                    {
                        "properties": {
                            "show": expr_bool(
                                True
                            ),
                            "type": expr_string(
                                "PageNavigation"
                            ),
                            "navigationSection": expr_string(
                                target_page_name
                            ),
                            "showDefaultTooltip": expr_bool(
                                False
                            ),
                            "enabledTooltip": expr_string(
                                "Return to INDEX"
                            ),
                        }
                    }
                ],
            },
            "drillFilterOtherVisuals": True,
        },
    }


def make_dropdown_slicer(
    *,
    name: str,
    entity: str,
    column_name: str,
    title: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    tab_order: int,
) -> dict:
    return {
        "$schema": VISUAL_SCHEMA,
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": tab_order,
        },
        "visual": {
            "visualType": "slicer",
            "query": {
                "queryState": {
                    "Values": {
                        "projections": [
                            column_projection(
                                entity,
                                column_name,
                                active=True,
                            )
                        ]
                    }
                }
            },
            "objects": {
                "data": [
                    {
                        "properties": {
                            "mode": expr_string(
                                "Dropdown"
                            )
                        }
                    }
                ],
                "selection": [
                    {
                        "properties": {
                            "singleSelect": expr_bool(
                                True
                            )
                        }
                    }
                ],
                "header": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": (
                container_objects(
                    background=COLORS["surface"],
                    border=COLORS["border"],
                    radius=9,
                    show_title=True,
                    title=title,
                    title_size=9.2,
                )
            ),
            "drillFilterOtherVisuals": True,
        },
    }


def make_value_card(
    *,
    name: str,
    measure_name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    tab_order: int,
    value_color: str,
    background: str = COLORS["surface"],
    value_size: float = 25,
) -> dict:
    return {
        "$schema": VISUAL_SCHEMA,
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": tab_order,
        },
        "visual": {
            "visualType": "card",
            "query": {
                "queryState": {
                    "Values": {
                        "projections": [
                            measure_projection(
                                measure_name
                            )
                        ]
                    }
                },
                "sortDefinition": {
                    "sort": [
                        {
                            "field": measure_field(
                                measure_name
                            ),
                            "direction": "Descending",
                        }
                    ],
                    "isDefaultSort": True,
                },
            },
            "objects": {
                "labels": [
                    {
                        "properties": {
                            "color": solid_color(
                                value_color
                            ),
                            "fontSize": expr_double(
                                value_size
                            ),
                            "bold": expr_bool(
                                True
                            ),
                            "fontFamily": expr_string(
                                FONT_SEMIBOLD
                            ),
                        }
                    }
                ],
                "categoryLabels": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": (
                container_objects(
                    background=background,
                    border=COLORS["border"],
                    radius=9,
                    show_title=False,
                )
            ),
            "drillFilterOtherVisuals": True,
        },
    }


# =====================================================================
# Intensity conditional formatting
# =====================================================================

def intensity_expression(
    measure_name: str,
    *,
    soft: bool,
) -> dict:
    if soft:
        low = "#F4F6FA"
        mid = COLORS["blue_soft"]
        high = COLORS["amber_soft"]
        extreme = COLORS["red_soft"]
    else:
        low = COLORS["secondary"]
        mid = COLORS["blue_dark"]
        high = COLORS["amber_dark"]
        extreme = COLORS["red_dark"]

    def comparison(
        kind: int,
        threshold: float,
    ) -> dict:
        return {
            "Comparison": {
                "ComparisonKind": kind,
                "Left": measure_expr(
                    measure_name
                ),
                "Right": literal_expr(
                    threshold
                ),
            }
        }

    return {
        "Conditional": {
            "Cases": [
                {
                    "Condition": comparison(
                        2,
                        100,
                    ),
                    "Value": {
                        "Literal": {
                            "Value": f"'{extreme}'"
                        }
                    },
                },
                {
                    "Condition": comparison(
                        2,
                        50,
                    ),
                    "Value": {
                        "Literal": {
                            "Value": f"'{high}'"
                        }
                    },
                },
                {
                    "Condition": comparison(
                        2,
                        25,
                    ),
                    "Value": {
                        "Literal": {
                            "Value": f"'{mid}'"
                        }
                    },
                },
                {
                    "Condition": comparison(
                        2,
                        0,
                    ),
                    "Value": {
                        "Literal": {
                            "Value": f"'{low}'"
                        }
                    },
                },
            ]
        }
    }


# =====================================================================
# State chart
# =====================================================================

def make_state_chart(
    *,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    tab_order: int,
) -> dict:
    return {
        "$schema": VISUAL_SCHEMA,
        "name": "p06_state_chart",
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": tab_order,
        },
        "visual": {
            "visualType": "clusteredColumnChart",
            "query": {
                "queryState": {
                    "Category": {
                        "projections": [
                            column_projection(
                                "DimState",
                                "StateCode",
                                active=True,
                            )
                        ]
                    },
                    "Y": {
                        "projections": [
                            measure_projection(
                                "Issuer Resubmission Events per 100 Denied - In Network",
                                display_name="In Network",
                            ),
                            measure_projection(
                                "Issuer Resubmission Events per 100 Denied - Out of Network",
                                display_name="Out of Network",
                            ),
                        ]
                    },
                },
                "sortDefinition": {
                    "sort": [
                        {
                            "field": measure_field(
                                "Issuer Resubmission Events per 100 Denied - In Network"
                            ),
                            "direction": "Descending",
                        }
                    ],
                    "isDefaultSort": True,
                },
            },
            "objects": {
                "legend": [
                    {
                        "properties": {
                            "show": expr_bool(
                                True
                            ),
                            "position": expr_string(
                                "Top"
                            ),
                            "legendMarkerRendering": expr_string(
                                "markerOnly"
                            ),
                            "fontSize": expr_double(
                                8.5
                            ),
                        }
                    }
                ],
                "categoryAxis": [
                    {
                        "properties": {
                            "showAxisTitle": expr_bool(
                                False
                            ),
                            "labelColor": solid_color(
                                COLORS["secondary"]
                            ),
                            "fontFamily": expr_string(
                                FONT_REGULAR
                            ),
                            "fontSize": expr_double(
                                8
                            ),
                            "concatenateLabels": expr_bool(
                                False
                            ),
                            "innerPadding": expr_int(
                                12
                            ),
                        }
                    }
                ],
                "valueAxis": [
                    {
                        "properties": {
                            "show": expr_bool(
                                True
                            ),
                            "showAxisTitle": expr_bool(
                                False
                            ),
                            "labelColor": solid_color(
                                COLORS["muted"]
                            ),
                            "fontFamily": expr_string(
                                FONT_REGULAR
                            ),
                            "fontSize": expr_double(
                                8
                            ),
                            "gridlineShow": expr_bool(
                                False
                            ),
                        }
                    }
                ],
                "dataPoint": [
                    {
                        "properties": {
                            "fill": solid_color(
                                COLORS["teal"]
                            )
                        },
                        "selector": {
                            "metadata": (
                                "_Measures.Issuer Resubmission Events per 100 Denied - In Network"
                            )
                        },
                    },
                    {
                        "properties": {
                            "fill": solid_color(
                                COLORS["blue"]
                            )
                        },
                        "selector": {
                            "metadata": (
                                "_Measures.Issuer Resubmission Events per 100 Denied - Out of Network"
                            )
                        },
                    },
                ],
                "labels": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "zoom": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": (
                container_objects(
                    background=COLORS["surface"],
                    border=COLORS["border"],
                    radius=9,
                    show_title=False,
                )
            ),
            "drillFilterOtherVisuals": True,
        },
    }


# =====================================================================
# Issuer investigation table
# =====================================================================

def make_issuer_table(
    *,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    tab_order: int,
) -> dict:
    in_ref = (
        "_Measures.Issuer Resubmission Events per 100 Denied - In Network"
    )
    out_ref = (
        "_Measures.Issuer Resubmission Events per 100 Denied - Out of Network"
    )

    projections = [
        column_projection(
            "DimIssuer",
            "IssuerName",
            display_name="Issuer",
        ),
        measure_projection(
            "Issuer Resubmission Events per 100 Denied - In Network",
            display_name="IN / 100 Denied",
        ),
        measure_projection(
            "Issuer Resubmission Events per 100 Denied - Out of Network",
            display_name="OON / 100 Denied",
        ),
        measure_projection(
            "Issuer Comparable Total Resubmission Events",
            display_name="Comparable Events",
        ),
    ]

    return {
        "$schema": VISUAL_SCHEMA,
        "name": "p06_issuer_table",
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": tab_order,
        },
        "visual": {
            "visualType": "tableEx",
            "query": {
                "queryState": {
                    "Values": {
                        "projections": projections
                    }
                },
                "sortDefinition": {
                    "sort": [
                        {
                            "field": measure_field(
                                "Issuer Resubmission Events per 100 Denied - In Network"
                            ),
                            "direction": "Descending",
                        }
                    ],
                    "isDefaultSort": True,
                },
            },
            "objects": {
                "grid": [
                    {
                        "properties": {
                            "gridHorizontal": expr_bool(
                                True
                            ),
                            "gridHorizontalColor": solid_color(
                                COLORS["border"]
                            ),
                            "rowPadding": expr_double(
                                2
                            ),
                            "fontSize": expr_double(
                                10
                            ),
                        }
                    }
                ],
                "columnHeaders": [
                    {
                        "properties": {
                            "alignment": expr_string(
                                "Center"
                            ),
                            "fontSize": expr_double(
                                10
                            ),
                            "bold": expr_bool(
                                True
                            ),
                        }
                    }
                ],
                "values": [
                    {
                        "properties": {
                            "backColorSecondary": solid_color(
                                COLORS["surface_alt"]
                            )
                        }
                    },
                    {
                        "properties": {
                            "backColor": {
                                "solid": {
                                    "color": {
                                        "expr": intensity_expression(
                                            "Issuer Resubmission Events per 100 Denied - In Network",
                                            soft=True,
                                        )
                                    }
                                }
                            },
                            "fontColor": {
                                "solid": {
                                    "color": {
                                        "expr": intensity_expression(
                                            "Issuer Resubmission Events per 100 Denied - In Network",
                                            soft=False,
                                        )
                                    }
                                }
                            },
                        },
                        "selector": {
                            "data": [
                                {
                                    "dataViewWildcard": {
                                        "matchingOption": 1
                                    }
                                }
                            ],
                            "metadata": in_ref,
                        },
                    },
                    {
                        "properties": {
                            "backColor": {
                                "solid": {
                                    "color": {
                                        "expr": intensity_expression(
                                            "Issuer Resubmission Events per 100 Denied - Out of Network",
                                            soft=True,
                                        )
                                    }
                                }
                            },
                            "fontColor": {
                                "solid": {
                                    "color": {
                                        "expr": intensity_expression(
                                            "Issuer Resubmission Events per 100 Denied - Out of Network",
                                            soft=False,
                                        )
                                    }
                                }
                            },
                        },
                        "selector": {
                            "data": [
                                {
                                    "dataViewWildcard": {
                                        "matchingOption": 1
                                    }
                                }
                            ],
                            "metadata": out_ref,
                        },
                    },
                ],
                "columnWidth": [
                    {
                        "properties": {
                            "value": expr_double(
                                690
                            )
                        },
                        "selector": {
                            "metadata": "DimIssuer.IssuerName"
                        },
                    },
                    {
                        "properties": {
                            "value": expr_double(
                                310
                            )
                        },
                        "selector": {
                            "metadata": in_ref
                        },
                    },
                    {
                        "properties": {
                            "value": expr_double(
                                310
                            )
                        },
                        "selector": {
                            "metadata": out_ref
                        },
                    },
                    {
                        "properties": {
                            "value": expr_double(
                                320
                            )
                        },
                        "selector": {
                            "metadata": (
                                "_Measures.Issuer Comparable Total Resubmission Events"
                            )
                        },
                    },
                ],
            },
            "visualContainerObjects": (
                container_objects(
                    background=COLORS["surface"],
                    border=COLORS["border"],
                    radius=9,
                    show_title=False,
                )
            ),
            "drillFilterOtherVisuals": True,
        },
    }


# =====================================================================
# Visual IO
# =====================================================================

def write_visual(
    page_dir: Path,
    visual: dict,
) -> Path:
    target = (
        page_dir
        / "visuals"
        / visual["name"]
        / "visual.json"
    )

    write_json(
        target,
        visual,
    )

    return target


# =====================================================================
# Page build
# =====================================================================

def build_resubmissions_page(
    page_dir: Path,
    index_page_name: str,
) -> None:
    visuals_dir = (
        page_dir
        / "visuals"
    )

    if visuals_dir.exists():
        shutil.rmtree(
            visuals_dir
        )

    visuals_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    # ---------------------------------------------------------------
    # Header — approved compact geometry.
    # ---------------------------------------------------------------

    write_visual(
        page_dir,
        make_textbox(
            name="p06_context",
            x=52,
            y=28,
            width=290,
            height=44,
            z=100,
            paragraphs=[
                paragraph(
                    text_run(
                        "2026 PUF  •  EXPERIENCE YEAR 2024",
                        size=9,
                        color=COLORS["blue"],
                        semibold=True,
                    )
                )
            ],
            background=COLORS["blue_soft"],
            border=COLORS["blue_pale"],
            radius=8,
        ),
    )

    write_visual(
        page_dir,
        make_textbox(
            name="p06_page_title",
            x=390,
            y=16,
            width=1140,
            height=76,
            z=110,
            paragraphs=[
                paragraph(
                    text_run(
                        "Resubmissions",
                        size=28,
                        color=COLORS["navy"],
                        semibold=True,
                    )
                ),
                paragraph(
                    text_run(
                        "Events → network → grain → issuers",
                        size=10,
                        color=COLORS["secondary"],
                    )
                ),
            ],
            background=COLORS["blue_soft"],
            border=COLORS["blue_pale"],
            radius=9,
        ),
    )

    index_x = 1648
    index_y = 24
    index_w = 220
    index_h = 56
    label_h = 22
    label_y = (
        index_y
        + (
            index_h
            - label_h
        )
        / 2
    )

    write_visual(
        page_dir,
        make_index_surface(
            x=index_x,
            y=index_y,
            width=index_w,
            height=index_h,
            z=130,
        ),
    )

    write_visual(
        page_dir,
        make_index_label(
            x=index_x,
            y=label_y,
            width=index_w,
            height=label_h,
            z=140,
        ),
    )

    write_visual(
        page_dir,
        make_index_nav(
            target_page_name=index_page_name,
            x=index_x,
            y=index_y,
            width=index_w,
            height=index_h,
            z=5000,
        ),
    )

    write_visual(
        page_dir,
        make_dropdown_slicer(
            name="p06_state_slicer",
            entity="DimState",
            column_name="StateCode",
            title="State",
            x=548,
            y=104,
            width=300,
            height=62,
            z=220,
            tab_order=1,
        ),
    )

    write_visual(
        page_dir,
        make_dropdown_slicer(
            name="p06_issuer_slicer",
            entity="DimIssuer",
            column_name="IssuerName",
            title="Issuer",
            x=864,
            y=104,
            width=508,
            height=62,
            z=220,
            tab_order=2,
        ),
    )

    # DQ title and body are separate from the start.
    write_visual(
        page_dir,
        make_title_band(
            name="p06_band_dq_global",
            title="GLOBAL OPEN DQ EXCEPTIONS",
            x=1388,
            y=104,
            width=480,
            height=24,
            z=1600,
            background=COLORS["amber"],
        ),
    )

    write_visual(
        page_dir,
        make_value_card(
            name="p06_dq_global",
            measure_name="Open DQ Exception Count",
            x=1388,
            y=130,
            width=480,
            height=36,
            z=230,
            tab_order=3,
            value_color=COLORS["amber_dark"],
            background=COLORS["surface"],
            value_size=14,
        ),
    )

    # ---------------------------------------------------------------
    # KPI strip — five concise cards.
    # ---------------------------------------------------------------

    margin = 52
    gap = 16
    total_width = (
        CANVAS_WIDTH
        - 2 * margin
    )
    card_width = (
        total_width
        - 4 * gap
    ) / 5

    kpis = [
        (
            "p06_total_events",
            "p06_band_total_events",
            "Issuer Comparable Total Resubmission Events",
            "COMPARABLE RESUBMISSION EVENTS",
            COLORS["amber_dark"],
            COLORS["amber_soft"],
            COLORS["amber"],
        ),
        (
            "p06_issuer_in",
            "p06_band_issuer_in",
            "Issuer Resubmission Events per 100 Denied - In Network",
            "ISSUER IN / 100 DENIED",
            COLORS["teal_dark"],
            COLORS["teal_soft"],
            COLORS["teal_dark"],
        ),
        (
            "p06_issuer_out",
            "p06_band_issuer_out",
            "Issuer Resubmission Events per 100 Denied - Out of Network",
            "ISSUER OON / 100 DENIED",
            COLORS["blue"],
            COLORS["blue_soft"],
            COLORS["blue"],
        ),
        (
            "p06_plan_in",
            "p06_band_plan_in",
            "Plan Resubmission Events per 100 Denied - In Network",
            "PLAN IN / 100 DENIED",
            COLORS["teal_dark"],
            COLORS["teal_soft"],
            COLORS["teal"],
        ),
        (
            "p06_plan_out",
            "p06_band_plan_out",
            "Plan Resubmission Events per 100 Denied - Out of Network",
            "PLAN OON / 100 DENIED",
            COLORS["blue"],
            COLORS["blue_soft"],
            COLORS["blue_dark"],
        ),
    ]

    for idx, (
        card_name,
        band_name,
        measure_name,
        band_title,
        value_color,
        background,
        band_color,
    ) in enumerate(
        kpis
    ):
        x = (
            margin
            + idx
            * (
                card_width
                + gap
            )
        )

        write_visual(
            page_dir,
            make_title_band(
                name=band_name,
                title=band_title,
                x=x,
                y=182,
                width=card_width,
                height=28,
                z=1700 + idx,
                background=band_color,
            ),
        )

        write_visual(
            page_dir,
            make_value_card(
                name=card_name,
                measure_name=measure_name,
                x=x,
                y=212,
                width=card_width,
                height=94,
                z=400 + idx,
                tab_order=10 + idx,
                value_color=value_color,
                background=background,
                value_size=24,
            ),
        )

    # ---------------------------------------------------------------
    # Main chart — one clean state story.
    # ---------------------------------------------------------------

    write_visual(
        page_dir,
        make_title_band(
            name="p06_band_state_chart",
            title="RESUBMISSION EVENTS PER 100 DENIED BY STATE",
            x=52,
            y=330,
            width=1816,
            height=34,
            z=1800,
            background=COLORS["navy"],
        ),
    )

    write_visual(
        page_dir,
        make_state_chart(
            x=52,
            y=366,
            width=1816,
            height=360,
            z=700,
            tab_order=30,
        ),
    )

    write_visual(
        page_dir,
        make_textbox(
            name="p06_state_note",
            x=52,
            y=728,
            width=1816,
            height=24,
            z=1800,
            paragraphs=[
                paragraph(
                    text_run(
                        "Events per 100 denied  •  Values above 100 can occur",
                        size=8.5,
                        color=COLORS["secondary"],
                        semibold=True,
                    )
                )
            ],
            background=COLORS["surface"],
            border=COLORS["border"],
            radius=7,
        ),
    )

    # ---------------------------------------------------------------
    # Issuer investigation.
    # ---------------------------------------------------------------

    write_visual(
        page_dir,
        make_title_band(
            name="p06_band_investigation",
            title="ISSUER RESUBMISSION INVESTIGATION",
            x=52,
            y=772,
            width=1816,
            height=34,
            z=1800,
            background=COLORS["navy"],
        ),
    )

    write_visual(
        page_dir,
        make_textbox(
            name="p06_rate_legend",
            x=52,
            y=808,
            width=1816,
            height=24,
            z=1800,
            paragraphs=[
                paragraph(
                    text_run(
                        "■ <25",
                        size=8.2,
                        color=COLORS["secondary"],
                        semibold=True,
                    ),
                    text_run(
                        "     ■ 25–50",
                        size=8.2,
                        color=COLORS["blue_dark"],
                        semibold=True,
                    ),
                    text_run(
                        "     ■ 50–100",
                        size=8.2,
                        color=COLORS["amber_dark"],
                        semibold=True,
                    ),
                    text_run(
                        "     ■ ≥100",
                        size=8.2,
                        color=COLORS["red_dark"],
                        semibold=True,
                    ),
                    text_run(
                        "     Magnitude only",
                        size=8.2,
                        color=COLORS["secondary"],
                    ),
                )
            ],
            background=COLORS["surface"],
            border=COLORS["border"],
            radius=7,
        ),
    )

    write_visual(
        page_dir,
        make_issuer_table(
            x=52,
            y=834,
            width=1816,
            height=194,
            z=900,
            tab_order=50,
        ),
    )

    # ---------------------------------------------------------------
    # Minimal methodology footer.
    # ---------------------------------------------------------------

    write_visual(
        page_dir,
        make_textbox(
            name="p06_methodology_note",
            x=52,
            y=1038,
            width=1816,
            height=22,
            z=200,
            paragraphs=[
                paragraph(
                    text_run(
                        (
                            "Events per 100 denied • Ratio-of-totals • "
                            "Separate issuer/plan grains • >100 is not automatically a DQ failure"
                        ),
                        size=8.2,
                        color=COLORS["muted"],
                        semibold=True,
                    )
                )
            ],
            background=None,
            border=None,
            radius=0,
        ),
    )


# =====================================================================
# Validation
# =====================================================================

def walk_json(
    obj,
):
    if isinstance(
        obj,
        dict,
    ):
        yield obj

        for value in obj.values():
            yield from walk_json(
                value
            )

    elif isinstance(
        obj,
        list,
    ):
        for value in obj:
            yield from walk_json(
                value
            )


def validate_json_tree() -> list[dict]:
    failures = []

    for path in sorted(
        (
            REPORT_DIR
            / "definition"
        ).rglob(
            "*.json"
        )
    ):
        try:
            read_json(
                path
            )

        except Exception as exc:
            failures.append(
                {
                    "RelativePath": str(
                        path.relative_to(
                            PROJECT_ROOT
                        )
                    ),
                    "Error": str(
                        exc
                    ),
                }
            )

    return failures


def validate_no_unsupported_page_roots(
    pages: dict,
) -> list[dict]:
    offenders = []

    for display_name, (
        page_dir,
        _,
    ) in pages.items():
        payload = read_json(
            page_dir
            / "page.json"
        )

        for prop in sorted(
            UNSUPPORTED_PAGE_ROOT_PROPERTIES
        ):
            if prop in payload:
                offenders.append(
                    {
                        "Page": display_name,
                        "UnsupportedProperty": prop,
                    }
                )

    return offenders


def validate_no_full_page_visuals(
    pages: dict,
) -> list[dict]:
    offenders = []

    for display_name, (
        page_dir,
        _,
    ) in pages.items():
        visuals_dir = (
            page_dir
            / "visuals"
        )

        if not visuals_dir.exists():
            continue

        for path in visuals_dir.rglob(
            "visual.json"
        ):
            payload = read_json(
                path
            )

            pos = payload.get(
                "position",
                {},
            )

            width = float(
                pos.get(
                    "width",
                    0,
                )
                or 0
            )

            height = float(
                pos.get(
                    "height",
                    0,
                )
                or 0
            )

            if (
                width
                >= CANVAS_WIDTH * 0.94
                and height
                >= CANVAS_HEIGHT * 0.94
            ):
                offenders.append(
                    {
                        "Page": display_name,
                        "VisualName": payload.get(
                            "name"
                        ),
                        "Width": width,
                        "Height": height,
                    }
                )

    return offenders


def count_bound_visuals(
    page_dir: Path,
) -> int:
    count = 0

    for visual_json in (
        page_dir
        / "visuals"
    ).rglob(
        "visual.json"
    ):
        payload = read_json(
            visual_json
        )

        if (
            payload.get(
                "visual",
                {},
            ).get(
                "query"
            )
        ):
            count += 1

    return count


def validate_state_chart(
    page_dir: Path,
) -> dict:
    payload = read_json(
        page_dir
        / "visuals"
        / "p06_state_chart"
        / "visual.json"
    )

    visual = payload.get(
        "visual",
        {},
    )

    visual_type = visual.get(
        "visualType"
    )

    projections = (
        visual.get(
            "query",
            {},
        )
        .get(
            "queryState",
            {},
        )
        .get(
            "Y",
            {},
        )
        .get(
            "projections",
            [],
        )
    )

    refs = {
        item.get(
            "queryRef"
        )
        for item in projections
    }

    expected = {
        (
            "_Measures.Issuer Resubmission Events per 100 Denied - In Network"
        ),
        (
            "_Measures.Issuer Resubmission Events per 100 Denied - Out of Network"
        ),
    }

    passed = (
        visual_type
        == "clusteredColumnChart"
        and expected.issubset(
            refs
        )
    )

    return {
        "CheckID": "RS-001",
        "TestName": (
            "State intensity comparison"
        ),
        "Expected": (
            "clusteredColumnChart + IN/OON issuer intensity"
        ),
        "Actual": (
            f"type={visual_type}; measures={len(refs & expected)}"
        ),
        "Status": (
            "PASS"
            if passed
            else "FAIL"
        ),
    }


def validate_table_story(
    page_dir: Path,
) -> dict:
    payload = read_json(
        page_dir
        / "visuals"
        / "p06_issuer_table"
        / "visual.json"
    )

    objects = (
        payload.get(
            "visual",
            {},
        )
        .get(
            "objects",
            {},
        )
    )

    conditionals = sum(
        1
        for node in walk_json(
            objects.get(
                "values",
                [],
            )
        )
        if "Conditional" in node
    )

    return {
        "CheckID": "RS-002",
        "TestName": (
            "Issuer intensity conditional formatting"
        ),
        "Expected": ">=2",
        "Actual": conditionals,
        "Status": (
            "PASS"
            if conditionals >= 2
            else "FAIL"
        ),
    }


def validate_index_geometry(
    page_dir: Path,
) -> dict:
    surface = read_json(
        page_dir
        / "visuals"
        / "p06_index_surface"
        / "visual.json"
    )

    label = read_json(
        page_dir
        / "visuals"
        / "p06_index_label"
        / "visual.json"
    )

    s = surface[
        "position"
    ]
    l = label[
        "position"
    ]

    expected_y = (
        s["y"]
        + (
            s["height"]
            - l["height"]
        )
        / 2
    )

    passed = (
        abs(
            float(
                l["y"]
            )
            - float(
                expected_y
            )
        )
        < 0.001
        and float(
            l["x"]
        )
        == float(
            s["x"]
        )
        and float(
            l["width"]
        )
        == float(
            s["width"]
        )
    )

    return {
        "CheckID": "RS-003",
        "TestName": (
            "INDEX geometric centering"
        ),
        "Expected": "centered",
        "Actual": (
            "centered"
            if passed
            else "not-centered"
        ),
        "Status": (
            "PASS"
            if passed
            else "FAIL"
        ),
    }


def style_report(
    path: Path,
) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import (
        Alignment,
        Font,
        PatternFill,
    )
    from openpyxl.utils import (
        get_column_letter,
    )

    wb = load_workbook(
        path
    )

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if (
            ws.max_row
            and ws.max_column
        ):
            ws.auto_filter.ref = (
                ws.dimensions
            )

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            max_len = 0

            for row_idx in range(
                1,
                min(
                    ws.max_row,
                    220,
                )
                + 1,
            ):
                value = ws.cell(
                    row=row_idx,
                    column=col_idx,
                ).value

                if value is not None:
                    max_len = max(
                        max_len,
                        len(
                            str(
                                value
                            )
                        ),
                    )

            ws.column_dimensions[
                get_column_letter(
                    col_idx
                )
            ].width = min(
                max(
                    max_len + 2,
                    10,
                ),
                75,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    wb.save(
        path
    )


# =====================================================================
# Main
# =====================================================================

def main() -> int:
    print("=" * 78)
    print(
        "Transparency in Coverage PUF — "
        "06 Resubmissions Storytelling Build"
    )
    print("=" * 78)

    print(
        f"Script version : {SCRIPT_VERSION}"
    )
    print(
        f"Report folder  : {REPORT_DIR}"
    )
    print(
        f"Backup         : {BACKUP_FILE}"
    )
    print(
        f"Review file    : {REPORT_FILE}"
    )

    if power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is running. "
            "Save the PBIP and close Power BI Desktop completely "
            "before changing PBIR files."
        )

    if not REPORT_DIR.exists():
        raise FileNotFoundError(
            f"Report folder not found: {REPORT_DIR}"
        )

    for required_file in (
        MEASURES_TMDL,
        DIM_STATE_TMDL,
        DIM_ISSUER_TMDL,
    ):
        if not required_file.exists():
            raise FileNotFoundError(
                f"Required semantic-model file not found: "
                f"{required_file}"
            )

    pages = find_pages()

    if set(
        pages
    ) != set(
        EXPECTED_PAGES
    ):
        raise RuntimeError(
            "Current report page inventory does not match "
            "the approved 11-page shell. No changes made."
        )

    installed_measures = (
        parse_tmdl_measures(
            MEASURES_TMDL
        )
    )

    missing_measures = sorted(
        REQUIRED_MEASURES
        - installed_measures
    )

    if missing_measures:
        raise RuntimeError(
            "Required governed measures are missing: "
            + ", ".join(
                missing_measures
            )
        )

    installed_columns = {
        "DimState": (
            parse_tmdl_columns(
                DIM_STATE_TMDL
            )
        ),
        "DimIssuer": (
            parse_tmdl_columns(
                DIM_ISSUER_TMDL
            )
        ),
    }

    missing_columns = []

    for table_name, required in (
        REQUIRED_COLUMNS.items()
    ):
        for column_name in sorted(
            required
            - installed_columns[
                table_name
            ]
        ):
            missing_columns.append(
                f"{table_name}.{column_name}"
            )

    if missing_columns:
        raise RuntimeError(
            "Required semantic columns are missing: "
            + ", ".join(
                missing_columns
            )
        )

    unsupported_before = (
        validate_no_unsupported_page_roots(
            pages
        )
    )

    if unsupported_before:
        raise RuntimeError(
            "Unsupported page-root properties are present before build: "
            + "; ".join(
                f"{row['Page']}:{row['UnsupportedProperty']}"
                for row in unsupported_before
            )
            + ". No changes made."
        )

    _, index_page = (
        pages[
            "00 INDEX"
        ]
    )

    resubmissions_dir, _ = (
        pages[
            "06 Resubmissions"
        ]
    )

    index_page_name = (
        index_page[
            "name"
        ]
    )

    before_hashes = (
        file_hashes(
            REPORT_DIR
        )
    )

    target_prefix = str(
        resubmissions_dir.relative_to(
            REPORT_DIR
        )
    ).replace(
        "\\",
        "/",
    )

    create_backup()

    prepare_target_page(
        resubmissions_dir
    )

    build_resubmissions_page(
        resubmissions_dir,
        index_page_name,
    )

    pages = find_pages()

    after_hashes = (
        file_hashes(
            REPORT_DIR
        )
    )

    non_target_changes = []

    for rel_path in sorted(
        set(
            before_hashes
        )
        | set(
            after_hashes
        )
    ):
        normalized = rel_path.replace(
            "\\",
            "/",
        )

        if normalized.startswith(
            target_prefix
        ):
            continue

        if (
            before_hashes.get(
                rel_path
            )
            != after_hashes.get(
                rel_path
            )
        ):
            non_target_changes.append(
                {
                    "RelativePath": rel_path,
                    "BeforeSHA256": before_hashes.get(
                        rel_path
                    ),
                    "AfterSHA256": after_hashes.get(
                        rel_path
                    ),
                }
            )

    json_failures = (
        validate_json_tree()
    )

    unsupported_after = (
        validate_no_unsupported_page_roots(
            pages
        )
    )

    full_page_offenders = (
        validate_no_full_page_visuals(
            pages
        )
    )

    visual_files = sorted(
        (
            resubmissions_dir
            / "visuals"
        ).rglob(
            "visual.json"
        )
    )

    actual_names = {
        read_json(
            path
        ).get(
            "name"
        )
        for path in visual_files
    }

    bound_after = (
        count_bound_visuals(
            resubmissions_dir
        )
    )

    validation_rows = [
        {
            "CheckID": "RS-000",
            "TestName": (
                "Resubmissions visual inventory"
            ),
            "Expected": len(
                EXPECTED_VISUAL_NAMES
            ),
            "Actual": len(
                actual_names
            ),
            "Status": (
                "PASS"
                if actual_names
                == EXPECTED_VISUAL_NAMES
                else "FAIL"
            ),
        },
        {
            "CheckID": "RS-BOUND",
            "TestName": (
                "Bound/query visuals"
            ),
            "Expected": EXPECTED_BOUND_VISUALS,
            "Actual": bound_after,
            "Status": (
                "PASS"
                if bound_after
                == EXPECTED_BOUND_VISUALS
                else "FAIL"
            ),
        },
        validate_state_chart(
            resubmissions_dir
        ),
        validate_table_story(
            resubmissions_dir
        ),
        validate_index_geometry(
            resubmissions_dir
        ),
        {
            "CheckID": "RS-LOCK",
            "TestName": (
                "All non-Resubmissions report files preserved"
            ),
            "Expected": 0,
            "Actual": len(
                non_target_changes
            ),
            "Status": (
                "PASS"
                if not non_target_changes
                else "FAIL"
            ),
        },
        {
            "CheckID": "RS-SCHEMA",
            "TestName": (
                "Unsupported page-root properties"
            ),
            "Expected": 0,
            "Actual": len(
                unsupported_after
            ),
            "Status": (
                "PASS"
                if not unsupported_after
                else "FAIL"
            ),
        },
        {
            "CheckID": "RS-BG",
            "TestName": (
                "Full-page selectable visuals"
            ),
            "Expected": 0,
            "Actual": len(
                full_page_offenders
            ),
            "Status": (
                "PASS"
                if not full_page_offenders
                else "FAIL"
            ),
        },
        {
            "CheckID": "RS-JSON",
            "TestName": (
                "PBIR JSON parse failures"
            ),
            "Expected": 0,
            "Actual": len(
                json_failures
            ),
            "Status": (
                "PASS"
                if not json_failures
                else "FAIL"
            ),
        },
    ]

    failures = sum(
        row[
            "Status"
        ] == "FAIL"
        for row in validation_rows
    )

    summary_df = pd.DataFrame(
        [
            (
                "BuildStatus",
                (
                    "PASS"
                    if failures == 0
                    else "FAIL"
                ),
            ),
            (
                "ScriptVersion",
                SCRIPT_VERSION,
            ),
            (
                "BuiltAtLocal",
                datetime.now()
                .isoformat(
                    timespec="seconds"
                ),
            ),
            (
                "PageBuilt",
                "06 Resubmissions",
            ),
            (
                "StoryFlow",
                (
                    "Events → network → grain → issuers"
                ),
            ),
            (
                "VisualsAfter",
                len(
                    visual_files
                ),
            ),
            (
                "BoundVisualsAfter",
                bound_after,
            ),
            (
                "MainChart",
                (
                    "Issuer IN vs OON resubmission events per 100 denied by state"
                ),
            ),
            (
                "InvestigationTable",
                (
                    "Issuer + IN intensity + OON intensity + comparable events"
                ),
            ),
            (
                "NonTargetFilesChanged",
                len(
                    non_target_changes
                ),
            ),
            (
                "FullPageSelectableVisuals",
                len(
                    full_page_offenders
                ),
            ),
            (
                "UnsupportedPageProperties",
                len(
                    unsupported_after
                ),
            ),
            (
                "ValidationFailures",
                failures,
            ),
            (
                "BackupFile",
                str(
                    BACKUP_FILE
                ),
            ),
            (
                "GovernanceRule",
                (
                    "Resubmission KPI is events per 100 denied claims, not a "
                    "bounded probability/rate; values above 100 remain visible."
                ),
            ),
        ],
        columns=[
            "Item",
            "Value",
        ],
    )

    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        REPORT_FILE,
        engine="openpyxl",
    ) as writer:
        summary_df.to_excel(
            writer,
            sheet_name="00_Summary",
            index=False,
        )

        pd.DataFrame(
            validation_rows
        ).to_excel(
            writer,
            sheet_name="01_Validation",
            index=False,
        )

        if non_target_changes:
            pd.DataFrame(
                non_target_changes
            ).to_excel(
                writer,
                sheet_name="02_NonTarget_Changes",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No files outside 06 Resubmissions changed. "
                            "Approved pages 01-05 were preserved."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="02_NonTarget_Changes",
                index=False,
            )

        if full_page_offenders:
            pd.DataFrame(
                full_page_offenders
            ).to_excel(
                writer,
                sheet_name="03_FullPage_Offenders",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No full-page selectable visuals."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="03_FullPage_Offenders",
                index=False,
            )

        if unsupported_after:
            pd.DataFrame(
                unsupported_after
            ).to_excel(
                writer,
                sheet_name="04_Schema_Offenders",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No unsupported page-root properties."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="04_Schema_Offenders",
                index=False,
            )

        if json_failures:
            pd.DataFrame(
                json_failures
            ).to_excel(
                writer,
                sheet_name="05_JSON_Failures",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No JSON parse failures."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="05_JSON_Failures",
                index=False,
            )

    style_report(
        REPORT_FILE
    )

    print()
    print(
        "06 Resubmissions storytelling build completed."
    )
    print(
        f"Build status                 : "
        f"{'PASS' if failures == 0 else 'FAIL'}"
    )
    print(
        f"Visuals after                : "
        f"{len(visual_files)}"
    )
    print(
        f"Bound/query visuals          : "
        f"{bound_after}"
    )
    print(
        "Business filters             : State + Issuer"
    )
    print(
        "Main comparison              : Issuer IN vs OON intensity by state"
    )
    print(
        "Grain KPIs                   : Issuer + Plan"
    )
    print(
        "Issuer table                 : IN/OON intensity + comparable events"
    )
    print(
        f"Non-target files changed     : "
        f"{len(non_target_changes)}"
    )
    print(
        f"Full-page selectable visuals : "
        f"{len(full_page_offenders)}"
    )
    print(
        f"Unsupported page properties  : "
        f"{len(unsupported_after)}"
    )
    print(
        f"Validation failures          : "
        f"{failures}"
    )
    print(
        f"Backup                       : "
        f"{BACKUP_FILE}"
    )
    print(
        f"Review evidence              : "
        f"{REPORT_FILE}"
    )

    if failures:
        print()
        print(
            "Offline validation failed. "
            "Do not open/save the PBIP until the evidence report is reviewed."
        )
        return 2

    print()
    print(
        "Next: open TransparencyInCoverage.pbip and inspect 06 Resubmissions. "
        "The page should be deliberately light: five KPIs, one state comparison, "
        "one issuer table, and only the minimum interpretation text."
    )

    return 0


if __name__ == "__main__":
    try:
        sys.exit(
            main()
        )

    except Exception as exc:
        print()
        print(
            "RESUBMISSIONS STORYTELLING BUILD FAILED"
        )
        print(
            str(
                exc
            )
        )
        sys.exit(
            1
        )
