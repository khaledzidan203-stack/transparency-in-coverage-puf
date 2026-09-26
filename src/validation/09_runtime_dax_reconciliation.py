from __future__ import annotations

from pathlib import Path
from datetime import datetime
import math
import os
import shutil
import subprocess
import sys
import tempfile

import pandas as pd


SCRIPT_VERSION = "1.0.1"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
CANONICAL_DIR = PROJECT_ROOT / "data" / "canonical"
PBIP_FILE = PROJECT_ROOT / "powerbi" / "TransparencyInCoverage.pbip"
PBIP_SERVER_NAME = PBIP_FILE.name


def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, key_path) as key:
                raw_value, _ = winreg.QueryValueEx(key, "Desktop")
            return Path(os.path.expandvars(raw_value)).resolve()
        except Exception:
            pass

    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = Path(os.environ.get("TRANSPARENCY_PUF_REVIEW_DIR", PROJECT_ROOT / ".local-review")).expanduser().resolve()
OUTPUT_FILE = REVIEW_DIR / "09_DAX_RUNTIME_RECONCILIATION.xlsx"


CORE_TOLERANCES = {
    "Issuer Count": 0.0,
    "Plan Count": 0.0,
    "Issuer Claims Received - In Network": 0.5,
    "Issuer Claims Received - Out of Network": 0.5,
    "Issuer Comparable Claims Received - Total": 0.5,
    "Issuer Claims Denied - In Network": 0.5,
    "Issuer Claims Denied - Out of Network": 0.5,
    "Issuer Comparable Claims Denied - Total": 0.5,
    "Issuer In-Network Denial Rate": 1e-10,
    "Issuer Out-of-Network Denial Rate": 1e-10,
    "Issuer Comparable Overall Denial Rate": 1e-10,
    "Issuer Network Denial Rate Gap": 1e-10,
    "Issuer Resubmission Events per 100 Denied - In Network": 1e-8,
    "Issuer Resubmission Events per 100 Denied - Out of Network": 1e-8,
    "Internal Appeals Filed": 0.5,
    "Internal Appeals Overturned": 0.5,
    "Internal Appeal Overturn Rate": 1e-10,
    "External Appeals Filed": 0.5,
    "External Appeals Overturned": 0.5,
    "External Appeal Overturn Rate": 1e-10,
    "Reported Average Monthly Enrollment": 1e-8,
    "Reported Average Monthly Disenrollment": 1e-8,
    "Comparable Disenrollment-to-Enrollment Ratio": 1e-10,
    "Open DQ Exception Count": 0.0,
    "Known Source Exception Row Count": 0.0,
    "Known Source Exception Entity Count": 0.0,
    "Plan In-Network Denial Rate": 1e-10,
    "Plan Out-of-Network Denial Rate": 1e-10,
    "Plan Comparable Overall Denial Rate": 1e-10,
    "Plan Resubmission Events per 100 Denied - In Network": 1e-8,
    "Plan Resubmission Events per 100 Denied - Out of Network": 1e-8,
    "Internal Appeal Overturn Rate - Excluding Known Source Exception": 1e-10,
}


def find_dscmd() -> Path:
    candidates = []

    in_path = shutil.which("dscmd.exe") or shutil.which("dscmd")
    if in_path:
        candidates.append(Path(in_path))

    localappdata = os.environ.get("LOCALAPPDATA")
    if localappdata:
        candidates.append(
            Path(localappdata) / "Programs" / "DAX Studio" / "dscmd.exe"
        )

    candidates.append(Path(r"C:\Program Files\DAX Studio\dscmd.exe"))

    for candidate in candidates:
        if candidate.exists():
            return candidate.resolve()

    raise FileNotFoundError(
        "dscmd.exe was not found. No installation is requested; "
        "the script only uses an existing DAX Studio installation."
    )


def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False

    result = subprocess.run(
        ["tasklist", "/FI", "IMAGENAME eq PBIDesktop.exe"],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    return "PBIDesktop.exe" in result.stdout


def read_csv(name: str, dtype=None) -> pd.DataFrame:
    path = CANONICAL_DIR / name
    if not path.exists():
        raise FileNotFoundError(f"Missing canonical file: {path}")
    return pd.read_csv(path, dtype=dtype, low_memory=False)


def numeric(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series, errors="coerce")


def sum_numeric(series: pd.Series) -> float:
    return float(numeric(series).sum(skipna=True))


def ratio_of_totals(
    df: pd.DataFrame,
    numerator_cols: list[str],
    denominator_cols: list[str],
) -> float:
    required = numerator_cols + denominator_cols
    work = df.copy()

    for col in required:
        work[col] = numeric(work[col])

    comparable = work[required].notna().all(axis=1)
    work = work.loc[comparable]

    numerator = sum(float(work[c].sum()) for c in numerator_cols)
    denominator = sum(float(work[c].sum()) for c in denominator_cols)

    if denominator == 0:
        return math.nan

    return numerator / denominator


def events_per_100(
    df: pd.DataFrame,
    events_col: str,
    denominator_col: str,
) -> float:
    ratio = ratio_of_totals(df, [events_col], [denominator_col])
    return ratio * 100 if not math.isnan(ratio) else math.nan


def normalize_dscmd_column(name: str) -> str:
    text = str(name).strip()
    if text.startswith("[") and text.endswith("]"):
        text = text[1:-1]
    return text


def run_dax_csv(
    dscmd: Path,
    query: str,
    output_csv: Path,
    label: str,
) -> tuple[pd.DataFrame, dict]:
    query_file = output_csv.with_suffix(".dax")
    query_file.write_text(query, encoding="utf-8")

    command = [
        str(dscmd),
        "csv",
        str(output_csv),
        "--server",
        PBIP_SERVER_NAME,
        "--file",
        str(query_file),
    ]

    result = subprocess.run(
        command,
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )

    evidence = {
        "Label": label,
        "ExitCode": result.returncode,
        "StdOut": result.stdout.strip(),
        "StdErr": result.stderr.strip(),
        "OutputFile": str(output_csv),
    }

    if result.returncode != 0:
        raise RuntimeError(
            f"DSCMD failed for {label}.\n"
            f"Exit code: {result.returncode}\n"
            f"STDOUT:\n{result.stdout}\n"
            f"STDERR:\n{result.stderr}"
        )

    if not output_csv.exists():
        raise RuntimeError(
            f"DSCMD reported success but did not create: {output_csv}"
        )

    df = pd.read_csv(output_csv)
    df.columns = [normalize_dscmd_column(c) for c in df.columns]
    return df, evidence


def core_dax_query() -> str:
    measure_names = [
        "Issuer Count",
        "Plan Count",
        "Issuer Claims Received - In Network",
        "Issuer Claims Received - Out of Network",
        "Issuer Comparable Claims Received - Total",
        "Issuer Claims Denied - In Network",
        "Issuer Claims Denied - Out of Network",
        "Issuer Comparable Claims Denied - Total",
        "Issuer In-Network Denial Rate",
        "Issuer Out-of-Network Denial Rate",
        "Issuer Comparable Overall Denial Rate",
        "Issuer Network Denial Rate Gap",
        "Issuer Resubmission Events per 100 Denied - In Network",
        "Issuer Resubmission Events per 100 Denied - Out of Network",
        "Internal Appeals Filed",
        "Internal Appeals Overturned",
        "Internal Appeal Overturn Rate",
        "External Appeals Filed",
        "External Appeals Overturned",
        "External Appeal Overturn Rate",
        "Reported Average Monthly Enrollment",
        "Reported Average Monthly Disenrollment",
        "Comparable Disenrollment-to-Enrollment Ratio",
        "Open DQ Exception Count",
        "Known Source Exception Row Count",
        "Known Source Exception Entity Count",
        "Plan In-Network Denial Rate",
        "Plan Out-of-Network Denial Rate",
        "Plan Comparable Overall Denial Rate",
        "Plan Resubmission Events per 100 Denied - In Network",
        "Plan Resubmission Events per 100 Denied - Out of Network",
        "Internal Appeal Overturn Rate - Excluding Known Source Exception",
    ]

    pairs = []
    for name in measure_names:
        escaped = name.replace('"', '""')
        pairs.append(f'    "{escaped}", [{name}]')

    return "EVALUATE\nROW(\n" + ",\n".join(pairs) + "\n)\n"


def reason_dax_query() -> str:
    rows = []
    for metric_key in range(19, 29):
        rows.append(
            "ROW("
            f'"MetricKey", {metric_key}, '
            '"DAXReasonComposition", '
            "CALCULATE("
            "[Comparable Denial Reason Composition %], "
            f"TREATAS({{{metric_key}}}, DimMetric[MetricKey])"
            ")"
            ")"
        )

    return "EVALUATE\nUNION(\n    " + ",\n    ".join(rows) + "\n)\n"


def availability_dax_query(metric_keys: list[int]) -> str:
    rows = []

    for metric_key in metric_keys:
        rows.append(
            "ROW("
            f'"MetricKey", {metric_key}, '
            '"AvailableEntityCount", '
            "CALCULATE([Available Entity Count], "
            f"TREATAS({{{metric_key}}}, DimMetric[MetricKey])), "
            '"ApplicableEntityCount", '
            "CALCULATE([Applicable Entity Count], "
            f"TREATAS({{{metric_key}}}, DimMetric[MetricKey])), "
            '"DataAvailabilityRate", '
            "CALCULATE([Data Availability Rate], "
            f"TREATAS({{{metric_key}}}, DimMetric[MetricKey])), "
            '"SuppressionRate", '
            "CALCULATE([Suppression Rate], "
            f"TREATAS({{{metric_key}}}, DimMetric[MetricKey])), "
            '"UnavailableRate", '
            "CALCULATE([Unavailable Rate], "
            f"TREATAS({{{metric_key}}}, DimMetric[MetricKey])), "
            '"StructuralNonApplicabilityCount", '
            "CALCULATE([Structural Non-Applicability Count], "
            f"TREATAS({{{metric_key}}}, DimMetric[MetricKey]))"
            ")"
        )

    return "EVALUATE\nUNION(\n    " + ",\n    ".join(rows) + "\n)\n"


def escape_dax_string(value: str) -> str:
    return value.replace('"', '""')


def dq_context_query(plan_id: str, issuer_id: str) -> str:
    plan = escape_dax_string(plan_id)
    issuer = escape_dax_string(issuer_id)

    return f'''EVALUATE
ROW(
    "TotalContextDQ", [Current Context DQ Exception Count],
    "PlanID", "{plan}",
    "PlanContextDQ",
        CALCULATE(
            [Current Context DQ Exception Count],
            TREATAS({{"{plan}"}}, DimPlan[PlanID])
        ),
    "IssuerID", "{issuer}",
    "IssuerContextDQ",
        CALCULATE(
            [Current Context DQ Exception Count],
            TREATAS({{"{issuer}"}}, DimIssuer[IssuerID])
        )
)
'''


def metadata_query(table_name: str) -> str:
    return f"EVALUATE {table_name}\n"


def build_python_baselines(
    dim_issuer: pd.DataFrame,
    dim_plan: pd.DataFrame,
    fact_issuer: pd.DataFrame,
    fact_plan: pd.DataFrame,
    dq: pd.DataFrame,
) -> dict[str, float]:
    baseline: dict[str, float] = {}

    baseline["Issuer Count"] = float(dim_issuer["IssuerID"].nunique())
    baseline["Plan Count"] = float(dim_plan["PlanID"].nunique())

    baseline["Issuer Claims Received - In Network"] = sum_numeric(
        fact_issuer["ClaimsReceivedInNetwork"]
    )
    baseline["Issuer Claims Received - Out of Network"] = sum_numeric(
        fact_issuer["ClaimsReceivedOutOfNetwork"]
    )
    baseline["Issuer Claims Denied - In Network"] = sum_numeric(
        fact_issuer["ClaimsDeniedInNetwork"]
    )
    baseline["Issuer Claims Denied - Out of Network"] = sum_numeric(
        fact_issuer["ClaimsDeniedOutOfNetwork"]
    )

    claims_cols = ["ClaimsReceivedInNetwork", "ClaimsReceivedOutOfNetwork"]
    denied_cols = ["ClaimsDeniedInNetwork", "ClaimsDeniedOutOfNetwork"]

    issuer_work = fact_issuer.copy()
    for col in claims_cols + denied_cols:
        issuer_work[col] = numeric(issuer_work[col])

    comparable = issuer_work[claims_cols + denied_cols].notna().all(axis=1)
    issuer_comp = issuer_work.loc[comparable]

    baseline["Issuer Comparable Claims Received - Total"] = float(
        issuer_comp["ClaimsReceivedInNetwork"].sum()
        + issuer_comp["ClaimsReceivedOutOfNetwork"].sum()
    )
    baseline["Issuer Comparable Claims Denied - Total"] = float(
        issuer_comp["ClaimsDeniedInNetwork"].sum()
        + issuer_comp["ClaimsDeniedOutOfNetwork"].sum()
    )

    baseline["Issuer In-Network Denial Rate"] = ratio_of_totals(
        fact_issuer,
        ["ClaimsDeniedInNetwork"],
        ["ClaimsReceivedInNetwork"],
    )
    baseline["Issuer Out-of-Network Denial Rate"] = ratio_of_totals(
        fact_issuer,
        ["ClaimsDeniedOutOfNetwork"],
        ["ClaimsReceivedOutOfNetwork"],
    )
    baseline["Issuer Comparable Overall Denial Rate"] = ratio_of_totals(
        fact_issuer,
        denied_cols,
        claims_cols,
    )
    baseline["Issuer Network Denial Rate Gap"] = (
        baseline["Issuer Out-of-Network Denial Rate"]
        - baseline["Issuer In-Network Denial Rate"]
    )

    baseline["Issuer Resubmission Events per 100 Denied - In Network"] = (
        events_per_100(
            fact_issuer,
            "ClaimsResubmittedInNetwork",
            "ClaimsDeniedInNetwork",
        )
    )
    baseline["Issuer Resubmission Events per 100 Denied - Out of Network"] = (
        events_per_100(
            fact_issuer,
            "ClaimsResubmittedOutOfNetwork",
            "ClaimsDeniedOutOfNetwork",
        )
    )

    baseline["Internal Appeals Filed"] = sum_numeric(
        fact_issuer["InternalAppealsFiled"]
    )
    baseline["Internal Appeals Overturned"] = sum_numeric(
        fact_issuer["InternalAppealsOverturned"]
    )
    baseline["Internal Appeal Overturn Rate"] = ratio_of_totals(
        fact_issuer,
        ["InternalAppealsOverturned"],
        ["InternalAppealsFiled"],
    )
    baseline["External Appeals Filed"] = sum_numeric(
        fact_issuer["ExternalAppealsFiled"]
    )
    baseline["External Appeals Overturned"] = sum_numeric(
        fact_issuer["ExternalAppealsOverturned"]
    )
    baseline["External Appeal Overturn Rate"] = ratio_of_totals(
        fact_issuer,
        ["ExternalAppealsOverturned"],
        ["ExternalAppealsFiled"],
    )

    baseline["Reported Average Monthly Enrollment"] = sum_numeric(
        fact_plan["AverageMonthlyEnrollment"]
    )
    baseline["Reported Average Monthly Disenrollment"] = sum_numeric(
        fact_plan["AverageMonthlyDisenrollment"]
    )
    baseline["Comparable Disenrollment-to-Enrollment Ratio"] = ratio_of_totals(
        fact_plan,
        ["AverageMonthlyDisenrollment"],
        ["AverageMonthlyEnrollment"],
    )

    baseline["Open DQ Exception Count"] = float(
        (dq["ResolutionState"].astype(str) == "OPEN_REVIEW").sum()
    )

    known_mask = (
        dq["IsKnownSourceException"]
        .astype(str)
        .str.strip()
        .str.lower()
        .isin(["true", "1", "yes"])
    )

    baseline["Known Source Exception Row Count"] = float(known_mask.sum())
    baseline["Known Source Exception Entity Count"] = float(
        dq.loc[known_mask, "EntityKey"].astype(str).nunique()
    )

    baseline["Plan In-Network Denial Rate"] = ratio_of_totals(
        fact_plan,
        ["ClaimsDeniedInNetwork"],
        ["ClaimsReceivedInNetwork"],
    )
    baseline["Plan Out-of-Network Denial Rate"] = ratio_of_totals(
        fact_plan,
        ["ClaimsDeniedOutOfNetwork"],
        ["ClaimsReceivedOutOfNetwork"],
    )
    baseline["Plan Comparable Overall Denial Rate"] = ratio_of_totals(
        fact_plan,
        ["ClaimsDeniedInNetwork", "ClaimsDeniedOutOfNetwork"],
        ["ClaimsReceivedInNetwork", "ClaimsReceivedOutOfNetwork"],
    )
    baseline["Plan Resubmission Events per 100 Denied - In Network"] = (
        events_per_100(
            fact_plan,
            "ClaimsResubmittedInNetwork",
            "ClaimsDeniedInNetwork",
        )
    )
    baseline["Plan Resubmission Events per 100 Denied - Out of Network"] = (
        events_per_100(
            fact_plan,
            "ClaimsResubmittedOutOfNetwork",
            "ClaimsDeniedOutOfNetwork",
        )
    )

    exception_issuer_ids = set(
        dq.loc[
            known_mask & (dq["Scope"].astype(str) == "ISSUER"),
            "EntityKey",
        ].astype(str).str.strip()
    )

    issuer_map = dim_issuer.copy()
    issuer_map["IssuerID"] = issuer_map["IssuerID"].astype(str).str.strip()

    exception_keys = set(
        pd.to_numeric(
            issuer_map.loc[
                issuer_map["IssuerID"].isin(exception_issuer_ids),
                "IssuerKey",
            ],
            errors="coerce",
        ).dropna().astype(int)
    )

    sensitivity = fact_issuer.copy()
    sensitivity["IssuerKey"] = pd.to_numeric(
        sensitivity["IssuerKey"], errors="coerce"
    )
    sensitivity = sensitivity[
        ~sensitivity["IssuerKey"].isin(exception_keys)
    ]

    baseline[
        "Internal Appeal Overturn Rate - Excluding Known Source Exception"
    ] = ratio_of_totals(
        sensitivity,
        ["InternalAppealsOverturned"],
        ["InternalAppealsFiled"],
    )

    return baseline


def build_reason_baselines(fact_plan: pd.DataFrame) -> pd.DataFrame:
    reason_map = {
        19: "DeniedReferralOrPriorAuthorization",
        20: "DeniedDueToOutOfNetworkProvider",
        21: "DeniedServicesExcluded",
        22: "DeniedNotMedicallyNecessaryExclBH",
        23: "DeniedNotMedicallyNecessaryBHOnly",
        24: "DeniedBenefitLimitReached",
        25: "DeniedMemberNotCovered",
        26: "DeniedInvestigationalExperimentalCosmetic",
        27: "DeniedAdministrativeReason",
        28: "DeniedOther",
    }

    work = fact_plan.copy()
    reason_cols = list(reason_map.values())

    for col in reason_cols:
        work[col] = numeric(work[col])

    fully_comparable = work[reason_cols].notna().all(axis=1)
    work = work.loc[fully_comparable]

    total = float(sum(work[col].sum() for col in reason_cols))

    rows = []
    for metric_key, col in reason_map.items():
        value = float(work[col].sum())
        rows.append(
            {
                "MetricKey": metric_key,
                "CanonicalColumn": col,
                "FullyComparablePlanCount": len(work),
                "PythonReasonCount": value,
                "PythonReasonComposition": (
                    value / total if total else math.nan
                ),
            }
        )

    return pd.DataFrame(rows)


def build_availability_baselines(
    dim_metric: pd.DataFrame,
    dim_status: pd.DataFrame,
    issuer_avail: pd.DataFrame,
    plan_avail: pd.DataFrame,
    metric_keys: list[int],
) -> pd.DataFrame:
    status_lookup = dim_status[["StatusKey", "StatusCode"]].copy()
    status_lookup["StatusKey"] = pd.to_numeric(
        status_lookup["StatusKey"], errors="coerce"
    )
    status_lookup["StatusCode"] = (
        status_lookup["StatusCode"].astype(str).str.strip()
    )

    metric_lookup = dim_metric[
        ["MetricKey", "MetricCode", "MetricDisplayName", "EntityLevel"]
    ].copy()
    metric_lookup["MetricKey"] = pd.to_numeric(
        metric_lookup["MetricKey"], errors="coerce"
    )

    structural_codes = {
        "NOT_REQUIRED_PLAN_TYPE",
        "NOT_APPLICABLE_NEW_ENTITY",
    }

    rows = []

    for metric_key in metric_keys:
        metric_row = metric_lookup[
            metric_lookup["MetricKey"] == metric_key
        ]

        if metric_row.empty:
            rows.append(
                {
                    "MetricKey": metric_key,
                    "MetricDisplayName": "MISSING",
                    "EntityLevel": "",
                    "PythonAvailable": math.nan,
                    "PythonApplicable": math.nan,
                    "PythonAvailabilityRate": math.nan,
                    "PythonSuppressionRate": math.nan,
                    "PythonUnavailableRate": math.nan,
                    "PythonStructuralNonApplicability": math.nan,
                }
            )
            continue

        entity_level = str(metric_row.iloc[0]["EntityLevel"]).strip()
        display_name = str(
            metric_row.iloc[0]["MetricDisplayName"]
        ).strip()

        if entity_level.upper() == "ISSUER":
            source = issuer_avail.copy()
        elif entity_level.upper() == "PLAN":
            source = plan_avail.copy()
        else:
            source = pd.DataFrame()

        source["MetricKey"] = pd.to_numeric(
            source["MetricKey"], errors="coerce"
        )
        source["AvailabilityStatusKey"] = pd.to_numeric(
            source["AvailabilityStatusKey"], errors="coerce"
        )

        sample = source[source["MetricKey"] == metric_key].merge(
            status_lookup,
            left_on="AvailabilityStatusKey",
            right_on="StatusKey",
            how="left",
        )

        available = int((sample["StatusCode"] == "AVAILABLE").sum())
        structural = int(
            sample["StatusCode"].isin(structural_codes).sum()
        )
        applicable = int(len(sample) - structural)
        suppressed = int(
            (sample["StatusCode"] == "SUPPRESSED_SMALL_CELL").sum()
        )
        unavailable = int(
            (sample["StatusCode"] == "NOT_AVAILABLE").sum()
        )

        rows.append(
            {
                "MetricKey": metric_key,
                "MetricDisplayName": display_name,
                "EntityLevel": entity_level,
                "PythonAvailable": available,
                "PythonApplicable": applicable,
                "PythonAvailabilityRate": (
                    available / applicable if applicable else math.nan
                ),
                "PythonSuppressionRate": (
                    suppressed / applicable if applicable else math.nan
                ),
                "PythonUnavailableRate": (
                    unavailable / applicable if applicable else math.nan
                ),
                "PythonStructuralNonApplicability": structural,
            }
        )

    return pd.DataFrame(rows)


def compare_values(
    name: str,
    expected: float,
    actual: float,
    tolerance: float,
) -> dict:
    if pd.isna(expected) and pd.isna(actual):
        diff = 0.0
        passed = True
    elif pd.isna(expected) or pd.isna(actual):
        diff = math.nan
        passed = False
    else:
        diff = float(actual) - float(expected)
        passed = abs(diff) <= tolerance

    return {
        "CheckName": name,
        "PythonExpected": expected,
        "DAXActual": actual,
        "Difference": diff,
        "Tolerance": tolerance,
        "Status": "PASS" if passed else "FAIL",
    }


def expected_dq_context(
    dq: pd.DataFrame,
    dim_plan: pd.DataFrame,
    dim_issuer: pd.DataFrame,
    plan_id: str,
    issuer_id: str,
) -> tuple[int, int, int]:
    dq2 = dq.copy()
    dq2["EntityKey"] = dq2["EntityKey"].astype(str).str.strip()
    dq2["Scope"] = dq2["Scope"].astype(str).str.strip()

    plans = dim_plan.copy()
    plans["PlanID"] = plans["PlanID"].astype(str).str.strip()
    plans["IssuerKey"] = pd.to_numeric(
        plans["IssuerKey"], errors="coerce"
    )

    issuers = dim_issuer.copy()
    issuers["IssuerID"] = issuers["IssuerID"].astype(str).str.strip()
    issuers["IssuerKey"] = pd.to_numeric(
        issuers["IssuerKey"], errors="coerce"
    )

    total = len(dq2)

    selected_plan = plans[plans["PlanID"] == plan_id]
    if selected_plan.empty:
        raise RuntimeError(f"Plan ID not found in DimPlan: {plan_id}")

    plan_issuer_key = int(selected_plan.iloc[0]["IssuerKey"])
    plan_issuer = issuers[issuers["IssuerKey"] == plan_issuer_key]
    if plan_issuer.empty:
        raise RuntimeError(
            f"Issuer key {plan_issuer_key} for plan {plan_id} not found."
        )
    plan_issuer_id = str(plan_issuer.iloc[0]["IssuerID"])

    plan_expected = int(
        (
            ((dq2["Scope"] == "PLAN") & (dq2["EntityKey"] == plan_id))
            | (
                (dq2["Scope"] == "ISSUER")
                & (dq2["EntityKey"] == plan_issuer_id)
            )
        ).sum()
    )

    issuer_row = issuers[issuers["IssuerID"] == issuer_id]
    if issuer_row.empty:
        raise RuntimeError(f"Issuer ID not found: {issuer_id}")

    issuer_key = int(issuer_row.iloc[0]["IssuerKey"])
    issuer_plan_ids = set(
        plans.loc[
            plans["IssuerKey"] == issuer_key,
            "PlanID",
        ].astype(str)
    )

    issuer_expected = int(
        (
            ((dq2["Scope"] == "ISSUER") & (dq2["EntityKey"] == issuer_id))
            | (
                (dq2["Scope"] == "PLAN")
                & dq2["EntityKey"].isin(issuer_plan_ids)
            )
        ).sum()
    )

    return total, plan_expected, issuer_expected


def style_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(1, ws.max_column + 1):
            max_len = 0
            for row_idx in range(
                1,
                min(ws.max_row, 250) + 1,
            ):
                value = ws.cell(row_idx, col_idx).value
                if value is not None:
                    max_len = max(max_len, len(str(value)))

            ws.column_dimensions[
                get_column_letter(col_idx)
            ].width = min(max(max_len + 2, 10), 60)

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    wb.save(path)


def main() -> int:
    print("=" * 78)
    print("Transparency in Coverage PUF — Runtime DAX Reconciliation")
    print("=" * 78)
    print(f"Script version : {SCRIPT_VERSION}")
    print(f"PBIP           : {PBIP_FILE}")
    print(f"Canonical dir  : {CANONICAL_DIR}")
    print(f"Review file    : {OUTPUT_FILE}")

    if not PBIP_FILE.exists():
        raise FileNotFoundError(f"PBIP file not found: {PBIP_FILE}")

    if not power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is not running. Open TransparencyInCoverage.pbip "
            "and keep it open while this validation runs."
        )

    dscmd = find_dscmd()
    print(f"DSCMD          : {dscmd}")

    dim_issuer = read_csv(
        "dim_issuer.csv",
        dtype={"IssuerID": "string"},
    )
    dim_plan = read_csv(
        "dim_plan.csv",
        dtype={"PlanID": "string"},
    )
    dim_metric = read_csv("dim_metric.csv")
    dim_status = read_csv("dim_availability_status.csv")
    fact_issuer = read_csv("fact_issuer_transparency.csv")
    fact_plan = read_csv("fact_plan_transparency.csv")
    issuer_avail = read_csv(
        "fact_issuer_metric_availability.csv"
    )
    plan_avail = read_csv(
        "fact_plan_metric_availability.csv"
    )
    dq = read_csv(
        "dq_exception_register.csv",
        dtype={"EntityKey": "string"},
    )

    baselines = build_python_baselines(
        dim_issuer=dim_issuer,
        dim_plan=dim_plan,
        fact_issuer=fact_issuer,
        fact_plan=fact_plan,
        dq=dq,
    )

    reason_baseline = build_reason_baselines(fact_plan)

    availability_metric_keys = [2, 29, 30]
    availability_baseline = build_availability_baselines(
        dim_metric,
        dim_status,
        issuer_avail,
        plan_avail,
        availability_metric_keys,
    )

    dq2 = dq.copy()
    dq2["Scope"] = dq2["Scope"].astype(str).str.strip()
    dq2["EntityKey"] = dq2["EntityKey"].astype(str).str.strip()

    plan_dq_rows = dq2[dq2["Scope"] == "PLAN"]
    issuer_dq_rows = dq2[dq2["Scope"] == "ISSUER"]

    if plan_dq_rows.empty or issuer_dq_rows.empty:
        raise RuntimeError(
            "DQ context validation needs at least one PLAN and one ISSUER DQ row."
        )

    test_plan_id = str(plan_dq_rows.iloc[0]["EntityKey"])
    test_issuer_id = str(issuer_dq_rows.iloc[0]["EntityKey"])

    (
        expected_total_dq,
        expected_plan_context_dq,
        expected_issuer_context_dq,
    ) = expected_dq_context(
        dq,
        dim_plan,
        dim_issuer,
        test_plan_id,
        test_issuer_id,
    )

    dscmd_evidence = []

    with tempfile.TemporaryDirectory(
        prefix="tic_puf_dax_runtime_"
    ) as temp_dir:
        temp = Path(temp_dir)

        core_df, evidence = run_dax_csv(
            dscmd,
            core_dax_query(),
            temp / "core.csv",
            "Core measures",
        )
        dscmd_evidence.append(evidence)

        reason_df, evidence = run_dax_csv(
            dscmd,
            reason_dax_query(),
            temp / "reasons.csv",
            "Denial reason composition",
        )
        dscmd_evidence.append(evidence)

        availability_df, evidence = run_dax_csv(
            dscmd,
            availability_dax_query(
                availability_metric_keys
            ),
            temp / "availability.csv",
            "Availability metrics",
        )
        dscmd_evidence.append(evidence)

        dq_context_df, evidence = run_dax_csv(
            dscmd,
            dq_context_query(
                test_plan_id,
                test_issuer_id,
            ),
            temp / "dq_context.csv",
            "DQ context behavior",
        )
        dscmd_evidence.append(evidence)

        metric_df, evidence = run_dax_csv(
            dscmd,
            metadata_query("DimMetric"),
            temp / "dim_metric.csv",
            "Live DimMetric",
        )
        dscmd_evidence.append(evidence)

        status_df, evidence = run_dax_csv(
            dscmd,
            metadata_query("DimAvailabilityStatus"),
            temp / "dim_status.csv",
            "Live DimAvailabilityStatus",
        )
        dscmd_evidence.append(evidence)

    if len(core_df) != 1:
        raise RuntimeError(
            f"Expected one row from core DAX query; got {len(core_df)}."
        )

    core_actual = {
        key: pd.to_numeric(
            pd.Series([core_df.iloc[0][key]]),
            errors="coerce",
        ).iloc[0]
        for key in baselines
        if key in core_df.columns
    }

    missing_core = sorted(set(baselines) - set(core_actual))
    if missing_core:
        raise RuntimeError(
            "Core DAX output is missing expected columns: "
            + ", ".join(missing_core)
        )

    core_checks = []
    for measure_name, expected in baselines.items():
        tolerance = CORE_TOLERANCES[measure_name]
        actual = float(core_actual[measure_name])
        core_checks.append(
            compare_values(
                measure_name,
                float(expected),
                actual,
                tolerance,
            )
        )

    core_checks_df = pd.DataFrame(core_checks)

    reason_df["MetricKey"] = pd.to_numeric(
        reason_df["MetricKey"],
        errors="coerce",
    )
    reason_df["DAXReasonComposition"] = pd.to_numeric(
        reason_df["DAXReasonComposition"],
        errors="coerce",
    )

    reason_recon = reason_baseline.merge(
        reason_df[["MetricKey", "DAXReasonComposition"]],
        on="MetricKey",
        how="left",
    )
    reason_recon["Difference"] = (
        reason_recon["DAXReasonComposition"]
        - reason_recon["PythonReasonComposition"]
    )
    reason_recon["Tolerance"] = 1e-10
    reason_recon["Status"] = reason_recon["Difference"].abs().le(
        reason_recon["Tolerance"]
    ).map({True: "PASS", False: "FAIL"})

    availability_df["MetricKey"] = pd.to_numeric(
        availability_df["MetricKey"],
        errors="coerce",
    )

    for col in [
        "AvailableEntityCount",
        "ApplicableEntityCount",
        "DataAvailabilityRate",
        "SuppressionRate",
        "UnavailableRate",
        "StructuralNonApplicabilityCount",
    ]:
        availability_df[col] = pd.to_numeric(
            availability_df[col],
            errors="coerce",
        )

    availability_recon = availability_baseline.merge(
        availability_df,
        on="MetricKey",
        how="left",
    )

    availability_specs = [
        (
            "Available",
            "PythonAvailable",
            "AvailableEntityCount",
            0.0,
        ),
        (
            "Applicable",
            "PythonApplicable",
            "ApplicableEntityCount",
            0.0,
        ),
        (
            "AvailabilityRate",
            "PythonAvailabilityRate",
            "DataAvailabilityRate",
            1e-10,
        ),
        (
            "SuppressionRate",
            "PythonSuppressionRate",
            "SuppressionRate",
            1e-10,
        ),
        (
            "UnavailableRate",
            "PythonUnavailableRate",
            "UnavailableRate",
            1e-10,
        ),
        (
            "StructuralNonApplicability",
            "PythonStructuralNonApplicability",
            "StructuralNonApplicabilityCount",
            0.0,
        ),
    ]

    availability_long_rows = []
    for _, row in availability_recon.iterrows():
        for label, expected_col, actual_col, tolerance in availability_specs:
            expected = row[expected_col]
            actual = row[actual_col]

            if pd.isna(expected) and pd.isna(actual):
                difference = 0.0
                passed = True
            elif pd.isna(expected) or pd.isna(actual):
                difference = math.nan
                passed = False
            else:
                difference = float(actual) - float(expected)
                passed = abs(difference) <= tolerance

            availability_long_rows.append(
                {
                    "MetricKey": row["MetricKey"],
                    "MetricDisplayName": row["MetricDisplayName"],
                    "EntityLevel": row["EntityLevel"],
                    "Check": label,
                    "PythonExpected": expected,
                    "DAXActual": actual,
                    "Difference": difference,
                    "Tolerance": tolerance,
                    "Status": "PASS" if passed else "FAIL",
                }
            )

    availability_checks_df = pd.DataFrame(
        availability_long_rows
    )

    if len(dq_context_df) != 1:
        raise RuntimeError(
            "Expected one row from DQ context query."
        )

    dq_row = dq_context_df.iloc[0]

    dq_context_checks = pd.DataFrame(
        [
            compare_values(
                "DQ Total Context",
                float(expected_total_dq),
                float(
                    pd.to_numeric(
                        pd.Series([dq_row["TotalContextDQ"]]),
                        errors="coerce",
                    ).iloc[0]
                ),
                0.0,
            ),
            compare_values(
                f"DQ Plan Context ({test_plan_id})",
                float(expected_plan_context_dq),
                float(
                    pd.to_numeric(
                        pd.Series([dq_row["PlanContextDQ"]]),
                        errors="coerce",
                    ).iloc[0]
                ),
                0.0,
            ),
            compare_values(
                f"DQ Issuer Context ({test_issuer_id})",
                float(expected_issuer_context_dq),
                float(
                    pd.to_numeric(
                        pd.Series([dq_row["IssuerContextDQ"]]),
                        errors="coerce",
                    ).iloc[0]
                ),
                0.0,
            ),
        ]
    )

    all_statuses = pd.concat(
        [
            core_checks_df[["CheckName", "Status"]],
            reason_recon.rename(
                columns={"MetricKey": "CheckName"}
            )[["CheckName", "Status"]],
            availability_checks_df[["Check", "Status"]].rename(
                columns={"Check": "CheckName"}
            ),
            dq_context_checks[["CheckName", "Status"]],
        ],
        ignore_index=True,
    )

    failures = int((all_statuses["Status"] == "FAIL").sum())

    summary_df = pd.DataFrame(
        [
            ("ValidationStatus", "PASS" if failures == 0 else "FAIL"),
            ("ScriptVersion", SCRIPT_VERSION),
            ("ValidatedAtLocal", datetime.now().isoformat(timespec="seconds")),
            ("PBIP", str(PBIP_FILE)),
            ("DSCMD", str(dscmd)),
            ("CoreMeasureChecks", len(core_checks_df)),
            ("ReasonChecks", len(reason_recon)),
            ("AvailabilityChecks", len(availability_checks_df)),
            ("DQContextChecks", len(dq_context_checks)),
            ("TotalFailures", failures),
            (
                "EvidenceState",
                "LIVE POWER BI DESKTOP DAX EXECUTION + PYTHON CANONICAL RECONCILIATION",
            ),
        ],
        columns=["Item", "Value"],
    )

    python_baselines_df = pd.DataFrame(
        [
            {
                "MeasureName": name,
                "PythonBaseline": value,
            }
            for name, value in baselines.items()
        ]
    )

    dscmd_evidence_df = pd.DataFrame(dscmd_evidence)

    REVIEW_DIR.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(
        OUTPUT_FILE,
        engine="openpyxl",
    ) as writer:
        summary_df.to_excel(
            writer,
            sheet_name="00_Summary",
            index=False,
        )
        core_checks_df.to_excel(
            writer,
            sheet_name="01_Core_Reconciliation",
            index=False,
        )
        python_baselines_df.to_excel(
            writer,
            sheet_name="02_Python_Baselines",
            index=False,
        )
        core_df.to_excel(
            writer,
            sheet_name="03_Live_DAX_Core",
            index=False,
        )
        reason_recon.to_excel(
            writer,
            sheet_name="04_Reason_Reconciliation",
            index=False,
        )
        availability_checks_df.to_excel(
            writer,
            sheet_name="05_Availability_Recon",
            index=False,
        )
        dq_context_checks.to_excel(
            writer,
            sheet_name="06_DQ_Context",
            index=False,
        )
        metric_df.to_excel(
            writer,
            sheet_name="07_Live_DimMetric",
            index=False,
        )
        status_df.to_excel(
            writer,
            sheet_name="08_Live_Status_Dim",
            index=False,
        )
        dscmd_evidence_df.to_excel(
            writer,
            sheet_name="09_DSCMD_Evidence",
            index=False,
        )

    style_workbook(OUTPUT_FILE)

    print()
    print("Runtime DAX reconciliation completed.")
    print(f"Validation status     : {'PASS' if failures == 0 else 'FAIL'}")
    print(f"Core measure checks   : {len(core_checks_df)}")
    print(f"Reason checks         : {len(reason_recon)}")
    print(f"Availability checks   : {len(availability_checks_df)}")
    print(f"DQ context checks     : {len(dq_context_checks)}")
    print(f"Total failures        : {failures}")
    print(f"Review evidence       : {OUTPUT_FILE}")
    print("Power BI and project source files were read only.")

    if failures:
        print()
        print(
            "One or more runtime semantic checks failed. "
            "Do not build visuals yet; review the workbook first."
        )
        return 2

    print()
    print(
        "Runtime DAX gate passed. "
        "The semantic model is ready for report-page engineering."
    )
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:
        print()
        print("RUNTIME DAX VALIDATION FAILED TO RUN")
        print(str(exc))

        try:
            REVIEW_DIR.mkdir(parents=True, exist_ok=True)
            failure_df = pd.DataFrame(
                [
                    ("ValidationStatus", "FAILED_TO_RUN"),
                    ("ScriptVersion", SCRIPT_VERSION),
                    ("FailedAtLocal", datetime.now().isoformat(timespec="seconds")),
                    ("PBIP", str(PBIP_FILE)),
                    ("ErrorType", type(exc).__name__),
                    ("ErrorMessage", str(exc)),
                    (
                        "EvidenceState",
                        "RUNTIME VALIDATION ABORTED; NO PASS CLAIM",
                    ),
                ],
                columns=["Item", "Value"],
            )
            with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
                failure_df.to_excel(
                    writer,
                    sheet_name="00_Failure",
                    index=False,
                )
            style_workbook(OUTPUT_FILE)
            print(f"Failure evidence       : {OUTPUT_FILE}")
        except Exception as report_exc:
            print(f"Could not write failure workbook: {report_exc}")

        sys.exit(1)
