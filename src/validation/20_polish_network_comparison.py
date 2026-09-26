from __future__ import annotations

from datetime import datetime
from pathlib import Path
import json
import os
import subprocess
import sys
import zipfile

import pandas as pd


SCRIPT_VERSION = "1.0.0"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
REPORT_DIR = (
    PROJECT_ROOT
    / "powerbi"
    / "TransparencyInCoverage.Report"
)
SEMANTIC_DIR = (
    PROJECT_ROOT
    / "powerbi"
    / "TransparencyInCoverage.SemanticModel"
)
PAGES_DIR = REPORT_DIR / "definition" / "pages"
MEASURES_TMDL = (
    SEMANTIC_DIR
    / "definition"
    / "tables"
    / "_Measures.tmdl"
)

EXPECTED_CURRENT_VISUALS = 20
EXPECTED_CURRENT_BOUND_VISUALS = 13
EXPECTED_AFTER_VISUALS = 21
EXPECTED_AFTER_BOUND_VISUALS = 14

CANVAS_WIDTH = 1920
CANVAS_HEIGHT = 1080

COLORS = {
    "surface": "#FFFFFF",
    "surface_alt": "#F9FAFB",
    "border": "#E7ECF2",
    "ink": "#172033",
    "secondary": "#667085",
    "muted": "#98A2B3",
    "amber": "#C8872C",
}

FONT_REGULAR = (
    "'Segoe UI', wf_segoe-ui_normal, helvetica, arial, sans-serif"
)
FONT_SEMIBOLD = (
    "'Segoe UI Semibold', wf_segoe-ui_semibold, "
    "helvetica, arial, sans-serif"
)

REQUIRED_MEASURE = "Current Context DQ Exception Count"

UNSUPPORTED_PAGE_ROOT_PROPERTIES = {
    "horizontalAlignment",
    "verticalAlignment",
}


# =====================================================================
# Paths / environment
# =====================================================================

def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )

            with winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                key_path,
            ) as key:
                raw_value, _ = winreg.QueryValueEx(
                    key,
                    "Desktop",
                )

            return Path(
                os.path.expandvars(raw_value)
            ).resolve()

        except Exception:
            pass

    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = (
    get_windows_desktop()
    / "Transparency_PUF_Review"
)

BACKUP_FILE = (
    REVIEW_DIR
    / "20_PRE_NETWORK_COMPARISON_POLISH_BACKUP.zip"
)

REPORT_FILE = (
    REVIEW_DIR
    / "20_NETWORK_COMPARISON_POLISH_REPORT.xlsx"
)


def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False

    result = subprocess.run(
        [
            "tasklist",
            "/FI",
            "IMAGENAME eq PBIDesktop.exe",
        ],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(
            subprocess,
            "CREATE_NO_WINDOW",
            0,
        ),
    )

    return "PBIDesktop.exe" in result.stdout


def read_json(path: Path):
    return json.loads(
        path.read_text(
            encoding="utf-8-sig",
        )
    )


def write_json(
    path: Path,
    payload,
) -> None:
    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    path.write_text(
        json.dumps(
            payload,
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )


# =====================================================================
# PBIR literal helpers
# =====================================================================

def expr_bool(value: bool) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": "true" if value else "false"
            }
        }
    }


def expr_string(value: str) -> dict:
    escaped = value.replace(
        "'",
        "''",
    )

    return {
        "expr": {
            "Literal": {
                "Value": f"'{escaped}'"
            }
        }
    }


def expr_double(
    value: float | int,
) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": f"{value}D"
            }
        }
    }


def solid_color(
    color: str,
) -> dict:
    return {
        "solid": {
            "color": expr_string(
                color
            )
        }
    }


# =====================================================================
# TMDL / page discovery
# =====================================================================

def parse_tmdl_measures(
    path: Path,
) -> set[str]:
    text = path.read_text(
        encoding="utf-8-sig",
    )

    measures: set[str] = set()

    for line in text.splitlines():
        stripped = line.lstrip()

        if not stripped.startswith(
            "measure "
        ):
            continue

        lhs = stripped.split(
            "=",
            1,
        )[0].strip()

        raw = lhs[
            len("measure "):
        ].strip()

        if (
            raw.startswith("'")
            and raw.endswith("'")
        ):
            raw = (
                raw[1:-1]
                .replace(
                    "''",
                    "'",
                )
            )

        measures.add(
            raw
        )

    return measures


def find_network_page() -> tuple[Path, dict]:
    matches = []

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        if payload.get(
            "displayName"
        ) == "02 Network Comparison":
            matches.append(
                (
                    page_json.parent,
                    payload,
                )
            )

    if len(matches) != 1:
        raise RuntimeError(
            "Expected exactly one '02 Network Comparison' page; "
            f"found {len(matches)}."
        )

    return matches[0]


def create_backup() -> None:
    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if BACKUP_FILE.exists():
        BACKUP_FILE.unlink()

    with zipfile.ZipFile(
        BACKUP_FILE,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zf:
        for path in sorted(
            p
            for p in REPORT_DIR.rglob("*")
            if p.is_file()
        ):
            archive_name = (
                Path(
                    REPORT_DIR.name
                )
                / path.relative_to(
                    REPORT_DIR
                )
            )

            zf.write(
                path,
                str(
                    archive_name
                ),
            )


# =====================================================================
# Visual helpers
# =====================================================================

def measure_field(
    measure_name: str,
) -> dict:
    return {
        "Measure": {
            "Expression": {
                "SourceRef": {
                    "Entity": "_Measures"
                }
            },
            "Property": measure_name,
        }
    }


def measure_projection(
    measure_name: str,
) -> dict:
    return {
        "field": measure_field(
            measure_name
        ),
        "queryRef": (
            f"_Measures.{measure_name}"
        ),
        "nativeQueryRef": (
            measure_name
        ),
    }


def container_objects(
    *,
    background: str,
    border: str,
    title: str,
    title_size: float,
) -> dict:
    return {
        "title": [
            {
                "properties": {
                    "show": expr_bool(
                        True
                    ),
                    "text": expr_string(
                        title
                    ),
                    "alignment": expr_string(
                        "center"
                    ),
                    "fontColor": solid_color(
                        COLORS["ink"]
                    ),
                    "fontFamily": expr_string(
                        FONT_SEMIBOLD
                    ),
                    "fontSize": expr_double(
                        title_size
                    ),
                    "titleWrap": expr_bool(
                        True
                    ),
                }
            }
        ],
        "background": [
            {
                "properties": {
                    "show": expr_bool(
                        True
                    ),
                    "color": solid_color(
                        background
                    ),
                    "transparency": expr_double(
                        0
                    ),
                }
            }
        ],
        "border": [
            {
                "properties": {
                    "show": expr_bool(
                        True
                    ),
                    "color": solid_color(
                        border
                    ),
                    "width": expr_double(
                        1
                    ),
                    "radius": expr_double(
                        9
                    ),
                }
            }
        ],
        "dropShadow": [
            {
                "properties": {
                    "show": expr_bool(
                        False
                    )
                }
            }
        ],
        "visualHeader": [
            {
                "properties": {
                    "show": expr_bool(
                        False
                    )
                }
            }
        ],
    }


def make_dq_card() -> dict:
    return {
        "$schema": (
            "https://developer.microsoft.com/json-schemas/fabric/item/report/"
            "definition/visualContainer/2.7.0/schema.json"
        ),
        "name": "p02_dq_context",
        "position": {
            "x": 1388,
            "y": 102,
            "z": 250,
            "width": 480,
            "height": 62,
            "tabOrder": 3,
        },
        "visual": {
            "visualType": "card",
            "query": {
                "queryState": {
                    "Values": {
                        "projections": [
                            measure_projection(
                                REQUIRED_MEASURE
                            )
                        ]
                    }
                },
                "sortDefinition": {
                    "sort": [
                        {
                            "field": measure_field(
                                REQUIRED_MEASURE
                            ),
                            "direction": "Descending",
                        }
                    ],
                    "isDefaultSort": True,
                },
            },
            "objects": {
                "labels": [
                    {
                        "properties": {
                            "color": solid_color(
                                COLORS["amber"]
                            ),
                            "fontSize": expr_double(
                                17
                            ),
                            "bold": expr_bool(
                                True
                            ),
                            "fontFamily": expr_string(
                                FONT_SEMIBOLD
                            ),
                        }
                    }
                ],
                "categoryLabels": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": (
                container_objects(
                    background=COLORS["surface"],
                    border=COLORS["border"],
                    title="DQ Flags in Context",
                    title_size=9.5,
                )
            ),
            "drillFilterOtherVisuals": True,
        },
    }


# =====================================================================
# Patch actions
# =====================================================================

TABLE_ALIAS_MAP = {
    "DimIssuer.IssuerName": "Issuer",
    "_Measures.Issuer In-Network Denial Rate": "IN Rate",
    "_Measures.Issuer Out-of-Network Denial Rate": "OON Rate",
    "_Measures.Issuer Network Denial Rate Gap": "Gap",
    "_Measures.Issuer Claims Received - In Network": "IN Claims",
    "_Measures.Issuer Claims Received - Out of Network": "OON Claims",
}


def patch_table_headers(
    table_path: Path,
) -> dict:
    payload = read_json(
        table_path
    )

    if (
        payload.get(
            "visual",
            {},
        ).get(
            "visualType"
        )
        != "tableEx"
    ):
        raise RuntimeError(
            "p02_issuer_table is not a tableEx visual."
        )

    projections = (
        payload[
            "visual"
        ][
            "query"
        ][
            "queryState"
        ][
            "Values"
        ][
            "projections"
        ]
    )

    patched = []

    for projection in projections:
        query_ref = projection.get(
            "queryRef"
        )

        if query_ref in TABLE_ALIAS_MAP:
            old = projection.get(
                "nativeQueryRef"
            )

            new = (
                TABLE_ALIAS_MAP[
                    query_ref
                ]
            )

            projection[
                "nativeQueryRef"
            ] = new

            patched.append(
                (
                    query_ref,
                    old,
                    new,
                )
            )

    if len(patched) != len(
        TABLE_ALIAS_MAP
    ):
        found = {
            item[0]
            for item in patched
        }

        missing = sorted(
            set(
                TABLE_ALIAS_MAP
            )
            - found
        )

        raise RuntimeError(
            "Could not alias every intended table column. Missing: "
            + ", ".join(
                missing
            )
        )

    write_json(
        table_path,
        payload,
    )

    return {
        "VisualName": "p02_issuer_table",
        "Change": (
            "Shortened table display headers via nativeQueryRef aliases "
            "without changing semantic field/queryRef bindings."
        ),
        "ColumnsPatched": len(
            patched
        ),
        "Status": "PATCHED",
    }


def add_dq_context_card(
    page_dir: Path,
) -> dict:
    target = (
        page_dir
        / "visuals"
        / "p02_dq_context"
        / "visual.json"
    )

    write_json(
        target,
        make_dq_card(),
    )

    return {
        "VisualName": "p02_dq_context",
        "Change": (
            "Added dynamic Current Context DQ Exception Count card "
            "next to State/Issuer slicers."
        ),
        "ColumnsPatched": "",
        "Status": "CREATED",
    }


def patch_methodology_note(
    note_path: Path,
) -> dict:
    payload = read_json(
        note_path
    )

    if (
        payload.get(
            "visual",
            {},
        ).get(
            "visualType"
        )
        != "textbox"
    ):
        raise RuntimeError(
            "p02_methodology_note is not a textbox."
        )

    paragraphs = (
        payload[
            "visual"
        ][
            "objects"
        ][
            "general"
        ][0][
            "properties"
        ][
            "paragraphs"
        ]
    )

    paragraphs[:] = [
        {
            "textRuns": [
                {
                    "value": (
                        "Network Gap = OON denial rate − IN denial rate  •  "
                        "Availability-aware ratio-of-totals  •  "
                        "Source-reported anomalies are preserved as published; "
                        "extreme values are not clipped or imputed"
                    ),
                    "textStyle": {
                        "fontFamily": FONT_SEMIBOLD,
                        "fontSize": "9.2pt",
                        "color": COLORS["muted"],
                    },
                }
            ],
            "horizontalTextAlignment": "center",
        }
    ]

    write_json(
        note_path,
        payload,
    )

    return {
        "VisualName": "p02_methodology_note",
        "Change": (
            "Clarified gap definition and explicit source-anomaly/no-clipping methodology."
        ),
        "ColumnsPatched": "",
        "Status": "PATCHED",
    }


# =====================================================================
# Validation
# =====================================================================

def walk_json(obj):
    if isinstance(
        obj,
        dict,
    ):
        yield obj

        for value in obj.values():
            yield from walk_json(
                value
            )

    elif isinstance(
        obj,
        list,
    ):
        for value in obj:
            yield from walk_json(
                value
            )


def count_bound_visuals(
    page_dir: Path,
) -> int:
    count = 0

    for visual_json in (
        page_dir
        / "visuals"
    ).rglob(
        "visual.json"
    ):
        payload = read_json(
            visual_json
        )

        if (
            payload.get(
                "visual",
                {},
            ).get(
                "query"
            )
        ):
            count += 1

    return count


def validate_table_aliases(
    table_path: Path,
) -> dict:
    payload = read_json(
        table_path
    )

    projections = (
        payload[
            "visual"
        ][
            "query"
        ][
            "queryState"
        ][
            "Values"
        ][
            "projections"
        ]
    )

    actual = {
        projection.get(
            "queryRef"
        ): projection.get(
            "nativeQueryRef"
        )
        for projection in projections
    }

    passed = all(
        actual.get(
            query_ref
        )
        == alias
        for query_ref, alias
        in TABLE_ALIAS_MAP.items()
    )

    return {
        "CheckID": "POL-001",
        "TestName": "Short issuer-table display headers",
        "Expected": len(
            TABLE_ALIAS_MAP
        ),
        "Actual": sum(
            actual.get(
                query_ref
            )
            == alias
            for query_ref, alias
            in TABLE_ALIAS_MAP.items()
        ),
        "Status": (
            "PASS"
            if passed
            else "FAIL"
        ),
    }


def validate_dq_card(
    dq_path: Path,
) -> dict:
    payload = read_json(
        dq_path
    )

    bindings = []

    for node in walk_json(
        payload.get(
            "visual",
            {},
        ).get(
            "query",
            {},
        )
    ):
        if "Measure" in node:
            prop = (
                node[
                    "Measure"
                ].get(
                    "Property"
                )
            )

            if prop:
                bindings.append(
                    prop
                )

    passed = (
        REQUIRED_MEASURE
        in bindings
    )

    return {
        "CheckID": "POL-002",
        "TestName": "Dynamic DQ context card binding",
        "Expected": REQUIRED_MEASURE,
        "Actual": (
            ", ".join(
                bindings
            )
        ),
        "Status": (
            "PASS"
            if passed
            else "FAIL"
        ),
    }


def validate_no_unsupported_page_roots() -> list[dict]:
    offenders = []

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        for prop in sorted(
            UNSUPPORTED_PAGE_ROOT_PROPERTIES
        ):
            if prop in payload:
                offenders.append(
                    {
                        "Page": payload.get(
                            "displayName"
                        ),
                        "Property": prop,
                    }
                )

    return offenders


def validate_no_full_page_visuals() -> list[dict]:
    offenders = []

    for visual_json in PAGES_DIR.rglob(
        "visual.json"
    ):
        payload = read_json(
            visual_json
        )

        pos = payload.get(
            "position",
            {},
        )

        width = float(
            pos.get(
                "width",
                0,
            )
            or 0
        )

        height = float(
            pos.get(
                "height",
                0,
            )
            or 0
        )

        if (
            width
            >= CANVAS_WIDTH * 0.94
            and height
            >= CANVAS_HEIGHT * 0.94
        ):
            offenders.append(
                {
                    "VisualName": payload.get(
                        "name"
                    ),
                    "Width": width,
                    "Height": height,
                    "RelativePath": str(
                        visual_json.relative_to(
                            PROJECT_ROOT
                        )
                    ),
                }
            )

    return offenders


def validate_json_tree() -> list[dict]:
    failures = []

    for path in (
        REPORT_DIR
        / "definition"
    ).rglob(
        "*.json"
    ):
        try:
            read_json(
                path
            )

        except Exception as exc:
            failures.append(
                {
                    "RelativePath": str(
                        path.relative_to(
                            PROJECT_ROOT
                        )
                    ),
                    "Error": str(
                        exc
                    ),
                }
            )

    return failures


# =====================================================================
# Evidence formatting
# =====================================================================

def style_report(
    path: Path,
) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import (
        Alignment,
        Font,
        PatternFill,
    )
    from openpyxl.utils import (
        get_column_letter,
    )

    wb = load_workbook(
        path
    )

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if (
            ws.max_row
            and ws.max_column
        ):
            ws.auto_filter.ref = (
                ws.dimensions
            )

        for cell in ws[1]:
            cell.fill = (
                header_fill
            )
            cell.font = (
                header_font
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            max_len = 0

            for row_idx in range(
                1,
                min(
                    ws.max_row,
                    200,
                )
                + 1,
            ):
                value = ws.cell(
                    row=row_idx,
                    column=col_idx,
                ).value

                if value is not None:
                    max_len = max(
                        max_len,
                        len(
                            str(
                                value
                            )
                        ),
                    )

            ws.column_dimensions[
                get_column_letter(
                    col_idx
                )
            ].width = min(
                max(
                    max_len + 2,
                    10,
                ),
                70,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    wb.save(
        path
    )


# =====================================================================
# Main
# =====================================================================

def main() -> int:
    print("=" * 78)
    print(
        "Transparency in Coverage PUF — "
        "02 Network Comparison Polish"
    )
    print("=" * 78)

    print(
        f"Script version : "
        f"{SCRIPT_VERSION}"
    )
    print(
        f"Report folder  : "
        f"{REPORT_DIR}"
    )
    print(
        f"Backup         : "
        f"{BACKUP_FILE}"
    )
    print(
        f"Review file    : "
        f"{REPORT_FILE}"
    )

    if power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is running. "
            "Save and close it completely before applying this PBIR polish."
        )

    if not REPORT_DIR.exists():
        raise FileNotFoundError(
            f"Report folder not found: "
            f"{REPORT_DIR}"
        )

    if not MEASURES_TMDL.exists():
        raise FileNotFoundError(
            f"Measure table not found: "
            f"{MEASURES_TMDL}"
        )

    measures = parse_tmdl_measures(
        MEASURES_TMDL
    )

    if REQUIRED_MEASURE not in measures:
        raise RuntimeError(
            f"Required measure not found: "
            f"{REQUIRED_MEASURE}"
        )

    page_dir, _ = find_network_page()

    current_visuals = len(
        list(
            (
                page_dir
                / "visuals"
            ).rglob(
                "visual.json"
            )
        )
    )

    current_bound = (
        count_bound_visuals(
            page_dir
        )
    )

    if (
        current_visuals
        != EXPECTED_CURRENT_VISUALS
    ):
        raise RuntimeError(
            f"Expected current Network Comparison to have "
            f"{EXPECTED_CURRENT_VISUALS} visuals; found {current_visuals}. "
            "No changes made."
        )

    if (
        current_bound
        != EXPECTED_CURRENT_BOUND_VISUALS
    ):
        raise RuntimeError(
            f"Expected current Network Comparison to have "
            f"{EXPECTED_CURRENT_BOUND_VISUALS} bound/query visuals; "
            f"found {current_bound}. No changes made."
        )

    table_path = (
        page_dir
        / "visuals"
        / "p02_issuer_table"
        / "visual.json"
    )

    note_path = (
        page_dir
        / "visuals"
        / "p02_methodology_note"
        / "visual.json"
    )

    if not table_path.exists():
        raise FileNotFoundError(
            "p02_issuer_table not found."
        )

    if not note_path.exists():
        raise FileNotFoundError(
            "p02_methodology_note not found."
        )

    create_backup()

    patch_rows = [
        patch_table_headers(
            table_path
        ),
        add_dq_context_card(
            page_dir
        ),
        patch_methodology_note(
            note_path
        ),
    ]

    dq_path = (
        page_dir
        / "visuals"
        / "p02_dq_context"
        / "visual.json"
    )

    visual_count_after = len(
        list(
            (
                page_dir
                / "visuals"
            ).rglob(
                "visual.json"
            )
        )
    )

    bound_after = (
        count_bound_visuals(
            page_dir
        )
    )

    unsupported = (
        validate_no_unsupported_page_roots()
    )

    full_page = (
        validate_no_full_page_visuals()
    )

    json_failures = (
        validate_json_tree()
    )

    validation_rows = [
        validate_table_aliases(
            table_path
        ),
        validate_dq_card(
            dq_path
        ),
        {
            "CheckID": "POL-003",
            "TestName": "Visual count after polish",
            "Expected": EXPECTED_AFTER_VISUALS,
            "Actual": visual_count_after,
            "Status": (
                "PASS"
                if visual_count_after
                == EXPECTED_AFTER_VISUALS
                else "FAIL"
            ),
        },
        {
            "CheckID": "POL-004",
            "TestName": "Bound/query visuals after polish",
            "Expected": EXPECTED_AFTER_BOUND_VISUALS,
            "Actual": bound_after,
            "Status": (
                "PASS"
                if bound_after
                == EXPECTED_AFTER_BOUND_VISUALS
                else "FAIL"
            ),
        },
        {
            "CheckID": "POL-005",
            "TestName": "Unsupported page-root properties",
            "Expected": 0,
            "Actual": len(
                unsupported
            ),
            "Status": (
                "PASS"
                if not unsupported
                else "FAIL"
            ),
        },
        {
            "CheckID": "POL-006",
            "TestName": "Full-page selectable visuals",
            "Expected": 0,
            "Actual": len(
                full_page
            ),
            "Status": (
                "PASS"
                if not full_page
                else "FAIL"
            ),
        },
        {
            "CheckID": "POL-007",
            "TestName": "PBIR JSON parse failures",
            "Expected": 0,
            "Actual": len(
                json_failures
            ),
            "Status": (
                "PASS"
                if not json_failures
                else "FAIL"
            ),
        },
    ]

    failures = sum(
        row[
            "Status"
        ]
        == "FAIL"
        for row
        in validation_rows
    )

    summary_df = pd.DataFrame(
        [
            (
                "PolishStatus",
                "PASS"
                if failures == 0
                else "FAIL",
            ),
            (
                "ScriptVersion",
                SCRIPT_VERSION,
            ),
            (
                "PolishedAtLocal",
                datetime.now()
                .isoformat(
                    timespec="seconds"
                ),
            ),
            (
                "Page",
                "02 Network Comparison",
            ),
            (
                "TableHeaders",
                (
                    "Issuer | IN Rate | OON Rate | Gap | "
                    "IN Claims | OON Claims"
                ),
            ),
            (
                "DQContextMeasure",
                REQUIRED_MEASURE,
            ),
            (
                "VisualsAfter",
                visual_count_after,
            ),
            (
                "BoundVisualsAfter",
                bound_after,
            ),
            (
                "ValidationFailures",
                failures,
            ),
            (
                "BackupFile",
                str(
                    BACKUP_FILE
                ),
            ),
            (
                "MethodologyMessage",
                (
                    "Extreme source-reported values are preserved; "
                    "no clipping or imputation."
                ),
            ),
        ],
        columns=[
            "Item",
            "Value",
        ],
    )

    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        REPORT_FILE,
        engine="openpyxl",
    ) as writer:
        summary_df.to_excel(
            writer,
            sheet_name="00_Summary",
            index=False,
        )

        pd.DataFrame(
            patch_rows
        ).to_excel(
            writer,
            sheet_name="01_Patches",
            index=False,
        )

        pd.DataFrame(
            validation_rows
        ).to_excel(
            writer,
            sheet_name="02_Validation",
            index=False,
        )

        if unsupported:
            pd.DataFrame(
                unsupported
            ).to_excel(
                writer,
                sheet_name="03_Schema_Offenders",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No unsupported page-root properties."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="03_Schema_Offenders",
                index=False,
            )

        if full_page:
            pd.DataFrame(
                full_page
            ).to_excel(
                writer,
                sheet_name="04_FullPage_Offenders",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No full-page selectable visuals."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="04_FullPage_Offenders",
                index=False,
            )

        if json_failures:
            pd.DataFrame(
                json_failures
            ).to_excel(
                writer,
                sheet_name="05_JSON_Failures",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No JSON parse failures."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="05_JSON_Failures",
                index=False,
            )

    style_report(
        REPORT_FILE
    )

    print()
    print(
        "02 Network Comparison polish completed."
    )
    print(
        f"Polish status                 : "
        f"{'PASS' if failures == 0 else 'FAIL'}"
    )
    print(
        f"Table headers shortened       : "
        f"{len(TABLE_ALIAS_MAP)}"
    )
    print(
        "DQ context card               : added"
    )
    print(
        f"Visuals after                 : "
        f"{visual_count_after}"
    )
    print(
        f"Bound/query visuals after     : "
        f"{bound_after}"
    )
    print(
        f"Full-page selectable visuals  : "
        f"{len(full_page)}"
    )
    print(
        f"Unsupported page properties   : "
        f"{len(unsupported)}"
    )
    print(
        f"Validation failures           : "
        f"{failures}"
    )
    print(
        f"Backup                        : "
        f"{BACKUP_FILE}"
    )
    print(
        f"Review evidence               : "
        f"{REPORT_FILE}"
    )

    if failures:
        print()
        print(
            "Offline validation failed. "
            "Do not open/save the PBIP until the evidence report is reviewed."
        )
        return 2

    print()
    print(
        "Next: open TransparencyInCoverage.pbip and inspect "
        "02 Network Comparison. Confirm the short table headers, "
        "the DQ Flags in Context card, and the clarified methodology note."
    )

    return 0


if __name__ == "__main__":
    try:
        sys.exit(
            main()
        )

    except Exception as exc:
        print()
        print(
            "NETWORK COMPARISON POLISH FAILED"
        )
        print(
            str(
                exc
            )
        )
        sys.exit(
            1
        )
