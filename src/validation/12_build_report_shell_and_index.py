from __future__ import annotations

from copy import deepcopy
from datetime import datetime
from pathlib import Path
import hashlib
import json
import os
import shutil
import subprocess
import sys
import zipfile

import pandas as pd


SCRIPT_VERSION = "1.0.0"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
POWERBI_ROOT = PROJECT_ROOT / "powerbi"
REPORT_DIR = POWERBI_ROOT / "TransparencyInCoverage.Report"
DEFINITION_DIR = REPORT_DIR / "definition"
PAGES_DIR = DEFINITION_DIR / "pages"
PAGES_FILE = PAGES_DIR / "pages.json"

VISUAL_SCHEMA = (
    "https://developer.microsoft.com/json-schemas/fabric/item/report/"
    "definition/visualContainer/2.7.0/schema.json"
)

CANVAS_WIDTH = 1920
CANVAS_HEIGHT = 1080

# ---------------------------------------------------------------------
# Design tokens — approved visual direction
# ---------------------------------------------------------------------

COLORS = {
    "canvas": "#F5F7FA",
    "surface": "#FFFFFF",
    "border": "#E7ECF2",
    "ink": "#172033",
    "secondary": "#667085",
    "muted": "#98A2B3",
    "navy": "#17324D",
    "blue": "#4F6BED",
    "blue_soft": "#EEF3FF",
    "blue_border": "#D8E2FF",
    "teal": "#16A6A1",
    "teal_soft": "#EAF8F6",
    "amber": "#C8872C",
    "red": "#C64B4B",
    "neutral_soft": "#F9FAFB",
}

FONT_REGULAR = "'Segoe UI', wf_segoe-ui_normal, helvetica, arial, sans-serif"
FONT_SEMIBOLD = (
    "'Segoe UI Semibold', wf_segoe-ui_semibold, helvetica, arial, sans-serif"
)

PAGE_SPECS = [
    (
        "01 Executive Overview",
        "Market size, claims, denials, appeals and governance at a glance.",
    ),
    (
        "02 Network Comparison",
        "Compare in-network and out-of-network denial behavior on governed populations.",
    ),
    (
        "03 Denials",
        "Understand claim denial volume, rates and comparable populations.",
    ),
    (
        "04 Denial Reasons",
        "Rank reported denial reasons across fully comparable plans.",
    ),
    (
        "05 Appeals",
        "Track internal and external appeals, overturns and sensitivity.",
    ),
    (
        "06 Resubmissions",
        "Measure resubmission events per 100 denied across network settings.",
    ),
    (
        "07 Enrollment",
        "Review reported enrollment, disenrollment and comparability.",
    ),
    (
        "08 Issuer & State Explorer",
        "Explore state, issuer and plan-level patterns.",
    ),
    (
        "09 Data Availability",
        "Separate available, suppressed, unavailable and structurally non-applicable data.",
    ),
    (
        "10 Data Quality & Methodology",
        "Review data-quality exceptions, source anomalies and governed analytical rules.",
    ),
]

INDEX_TILES = [
    {
        "page": "01 Executive Overview",
        "eyebrow": "01  OVERVIEW",
        "title": "Executive Overview",
        "description": (
            "Market size, claims volume, denial disparity, appeals and governance context."
        ),
        "x": 64,
        "y": 265,
        "w": 1792,
        "h": 132,
        "fill": COLORS["blue_soft"],
        "border": COLORS["blue_border"],
    },
    {
        "page": "02 Network Comparison",
        "eyebrow": "02  CLAIMS & DENIALS",
        "title": "Network Comparison",
        "description": "In-network vs out-of-network denial behavior.",
        "x": 64,
        "y": 495,
        "w": 581,
        "h": 150,
    },
    {
        "page": "03 Denials",
        "eyebrow": "03  CLAIMS & DENIALS",
        "title": "Denials",
        "description": "Volume, rates and governed comparable populations.",
        "x": 669,
        "y": 495,
        "w": 581,
        "h": 150,
    },
    {
        "page": "04 Denial Reasons",
        "eyebrow": "04  CLAIMS & DENIALS",
        "title": "Denial Reasons",
        "description": "Ranked reasons across fully comparable plans.",
        "x": 1274,
        "y": 495,
        "w": 582,
        "h": 150,
    },
    {
        "page": "05 Appeals",
        "eyebrow": "05  POST-DENIAL",
        "title": "Appeals",
        "description": "Filed, overturned and governed sensitivity view.",
        "x": 64,
        "y": 750,
        "w": 430,
        "h": 145,
    },
    {
        "page": "06 Resubmissions",
        "eyebrow": "06  POST-DENIAL",
        "title": "Resubmissions",
        "description": "Events per 100 denied across network settings.",
        "x": 518,
        "y": 750,
        "w": 430,
        "h": 145,
    },
    {
        "page": "07 Enrollment",
        "eyebrow": "07  MARKET & POPULATION",
        "title": "Enrollment",
        "description": "Reported enrollment and disenrollment context.",
        "x": 972,
        "y": 750,
        "w": 430,
        "h": 145,
    },
    {
        "page": "08 Issuer & State Explorer",
        "eyebrow": "08  MARKET & POPULATION",
        "title": "Issuer & State Explorer",
        "description": "Explore issuers and plans by state.",
        "x": 1426,
        "y": 750,
        "w": 430,
        "h": 145,
    },
    {
        "page": "09 Data Availability",
        "eyebrow": "09  GOVERNANCE",
        "title": "Data Availability",
        "description": "Availability, suppression and structural applicability.",
        "x": 64,
        "y": 950,
        "w": 884,
        "h": 100,
        "fill": COLORS["neutral_soft"],
    },
    {
        "page": "10 Data Quality & Methodology",
        "eyebrow": "10  GOVERNANCE",
        "title": "Data Quality & Methodology",
        "description": "Exceptions, known-source issues and analytical rules.",
        "x": 972,
        "y": 950,
        "w": 884,
        "h": 100,
        "fill": COLORS["neutral_soft"],
    },
]


# ---------------------------------------------------------------------
# Paths / environment
# ---------------------------------------------------------------------

def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, key_path) as key:
                raw_value, _ = winreg.QueryValueEx(key, "Desktop")
            return Path(os.path.expandvars(raw_value)).resolve()
        except Exception:
            pass

    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = get_windows_desktop() / "Transparency_PUF_Review"
BACKUP_FILE = REVIEW_DIR / "12_PRE_REPORT_SHELL_BACKUP.zip"
REPORT_FILE = REVIEW_DIR / "12_REPORT_SHELL_BUILD_REPORT.xlsx"


def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False

    result = subprocess.run(
        ["tasklist", "/FI", "IMAGENAME eq PBIDesktop.exe"],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    return "PBIDesktop.exe" in result.stdout


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def write_json(path: Path, payload) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def deterministic_id(label: str, length: int = 20) -> str:
    return hashlib.sha1(label.encode("utf-8")).hexdigest()[:length]


# ---------------------------------------------------------------------
# PBIR expression helpers
# ---------------------------------------------------------------------

def expr_bool(value: bool) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": "true" if value else "false"
            }
        }
    }


def expr_string(value: str) -> dict:
    escaped = value.replace("'", "''")
    return {
        "expr": {
            "Literal": {
                "Value": f"'{escaped}'"
            }
        }
    }


def expr_double(value: float | int) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": f"{value}D"
            }
        }
    }


def solid_color(value: str) -> dict:
    return {
        "solid": {
            "color": expr_string(value)
        }
    }


# ---------------------------------------------------------------------
# Visual builders
# ---------------------------------------------------------------------

def make_textbox(
    *,
    name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    paragraphs: list[dict],
    background: str | None = None,
    border: str | None = None,
    radius: int = 0,
    tab_order: int | None = None,
) -> dict:
    vco = {
        "title": [
            {
                "properties": {
                    "show": expr_bool(False)
                }
            }
        ],
        "background": [
            {
                "properties": {
                    "show": expr_bool(background is not None),
                    **(
                        {
                            "color": solid_color(background),
                            "transparency": expr_double(0),
                        }
                        if background is not None
                        else {}
                    ),
                }
            }
        ],
        "border": [
            {
                "properties": {
                    "show": expr_bool(border is not None),
                    **(
                        {
                            "color": solid_color(border),
                            "width": expr_double(1),
                            "radius": expr_double(radius),
                        }
                        if border is not None
                        else {}
                    ),
                }
            }
        ],
        "dropShadow": [
            {
                "properties": {
                    "show": expr_bool(False)
                }
            }
        ],
        "visualHeader": [
            {
                "properties": {
                    "show": expr_bool(False)
                }
            }
        ],
    }

    position = {
        "x": x,
        "y": y,
        "z": z,
        "width": width,
        "height": height,
    }

    if tab_order is not None:
        position["tabOrder"] = tab_order

    return {
        "$schema": VISUAL_SCHEMA,
        "name": name,
        "position": position,
        "visual": {
            "visualType": "textbox",
            "objects": {
                "general": [
                    {
                        "properties": {
                            "paragraphs": paragraphs
                        }
                    }
                ]
            },
            "visualContainerObjects": vco,
            "drillFilterOtherVisuals": True,
        },
    }


def text_run(
    value: str,
    *,
    size: float,
    color: str,
    semibold: bool = False,
) -> dict:
    return {
        "value": value,
        "textStyle": {
            "fontFamily": FONT_SEMIBOLD if semibold else FONT_REGULAR,
            "fontSize": f"{size}pt",
            "color": color,
        },
    }


def paragraph(*runs: dict) -> dict:
    return {
        "textRuns": list(runs)
    }


def make_transparent_nav_button(
    *,
    name: str,
    target_page_name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    tab_order: int,
    tooltip: str,
) -> dict:
    return {
        "$schema": VISUAL_SCHEMA,
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": tab_order,
        },
        "visual": {
            "visualType": "actionButton",
            "objects": {
                "icon": [
                    {
                        "properties": {
                            "shapeType": expr_string("blank")
                        },
                        "selector": {
                            "id": "default"
                        },
                    },
                    {
                        "properties": {
                            "show": expr_bool(False)
                        }
                    },
                ],
                "outline": [
                    {
                        "properties": {
                            "show": expr_bool(False)
                        }
                    }
                ],
                "fill": [
                    {
                        "properties": {
                            "show": expr_bool(False)
                        }
                    }
                ],
            },
            "visualContainerObjects": {
                "title": [
                    {
                        "properties": {
                            "show": expr_bool(False)
                        }
                    }
                ],
                "background": [
                    {
                        "properties": {
                            "show": expr_bool(False)
                        }
                    }
                ],
                "border": [
                    {
                        "properties": {
                            "show": expr_bool(False)
                        }
                    }
                ],
                "dropShadow": [
                    {
                        "properties": {
                            "show": expr_bool(False)
                        }
                    }
                ],
                "visualHeader": [
                    {
                        "properties": {
                            "show": expr_bool(False)
                        }
                    }
                ],
                "visualLink": [
                    {
                        "properties": {
                            "show": expr_bool(True),
                            "type": expr_string("PageNavigation"),
                            "navigationSection": expr_string(target_page_name),
                            "showDefaultTooltip": expr_bool(False),
                            "enabledTooltip": expr_string(tooltip),
                        }
                    }
                ],
            },
            "drillFilterOtherVisuals": True,
        },
    }


def write_visual(page_dir: Path, visual: dict) -> Path:
    visual_name = visual["name"]
    target = page_dir / "visuals" / visual_name / "visual.json"
    write_json(target, visual)
    return target


# ---------------------------------------------------------------------
# Page builders
# ---------------------------------------------------------------------

def create_page_payload(
    template: dict,
    *,
    page_name: str,
    display_name: str,
) -> dict:
    payload = deepcopy(template)
    payload["name"] = page_name
    payload["displayName"] = display_name
    payload["width"] = CANVAS_WIDTH
    payload["height"] = CANVAS_HEIGHT

    # Do not inherit interactions or filters from the empty template if they
    # unexpectedly appear later.
    payload.pop("visualInteractions", None)

    if isinstance(payload.get("filterConfig"), dict):
        payload["filterConfig"] = {
            **payload["filterConfig"],
            "filters": [],
        }

    return payload


def add_canvas_background(page_dir: Path, prefix: str) -> list[Path]:
    visual = make_textbox(
        name=f"{prefix}_canvas",
        x=0,
        y=0,
        width=CANVAS_WIDTH,
        height=CANVAS_HEIGHT,
        z=0,
        paragraphs=[
            paragraph(
                text_run(
                    " ",
                    size=6,
                    color=COLORS["canvas"],
                )
            )
        ],
        background=COLORS["canvas"],
    )
    return [write_visual(page_dir, visual)]


def add_index_content(
    page_dir: Path,
    page_name_by_display: dict[str, str],
) -> list[Path]:
    files: list[Path] = []
    files.extend(add_canvas_background(page_dir, "idx"))

    files.append(
        write_visual(
            page_dir,
            make_textbox(
                name="idx_overline",
                x=64,
                y=50,
                width=900,
                height=30,
                z=100,
                paragraphs=[
                    paragraph(
                        text_run(
                            "HEALTHCARE TRANSPARENCY ANALYTICS",
                            size=10,
                            color=COLORS["blue"],
                            semibold=True,
                        )
                    )
                ],
            ),
        )
    )

    files.append(
        write_visual(
            page_dir,
            make_textbox(
                name="idx_title",
                x=64,
                y=80,
                width=1180,
                height=118,
                z=110,
                paragraphs=[
                    paragraph(
                        text_run(
                            "Transparency in Coverage",
                            size=32,
                            color=COLORS["ink"],
                            semibold=True,
                        )
                    ),
                    paragraph(
                        text_run(
                            (
                                "Claims, denials, appeals, availability and "
                                "data-quality analytics"
                            ),
                            size=13,
                            color=COLORS["secondary"],
                        )
                    ),
                ],
            ),
        )
    )

    files.append(
        write_visual(
            page_dir,
            make_textbox(
                name="idx_context_badge",
                x=1470,
                y=72,
                width=386,
                height=58,
                z=120,
                paragraphs=[
                    paragraph(
                        text_run(
                            "2026 PUF   •   Experience Year 2024",
                            size=10.5,
                            color=COLORS["navy"],
                            semibold=True,
                        )
                    )
                ],
                background=COLORS["surface"],
                border=COLORS["border"],
                radius=8,
            ),
        )
    )

    section_labels = [
        ("idx_label_overview", "OVERVIEW", 64, 225, 400),
        ("idx_label_claims", "CLAIMS & DENIALS", 64, 455, 500),
        ("idx_label_post", "POST-DENIAL JOURNEY", 64, 710, 500),
        ("idx_label_market", "MARKET & POPULATION", 972, 710, 500),
        ("idx_label_gov", "GOVERNANCE", 64, 910, 500),
    ]

    for name, label, x, y, width in section_labels:
        files.append(
            write_visual(
                page_dir,
                make_textbox(
                    name=name,
                    x=x,
                    y=y,
                    width=width,
                    height=28,
                    z=150,
                    paragraphs=[
                        paragraph(
                            text_run(
                                label,
                                size=9,
                                color=COLORS["muted"],
                                semibold=True,
                            )
                        )
                    ],
                ),
            )
        )

    tab_order = 0

    for idx, tile in enumerate(INDEX_TILES, start=1):
        target_page_name = page_name_by_display[tile["page"]]
        fill = tile.get("fill", COLORS["surface"])
        border = tile.get("border", COLORS["border"])

        text_name = f"idx_tile_{idx:02d}_surface"
        nav_name = f"idx_tile_{idx:02d}_nav"

        files.append(
            write_visual(
                page_dir,
                make_textbox(
                    name=text_name,
                    x=tile["x"],
                    y=tile["y"],
                    width=tile["w"],
                    height=tile["h"],
                    z=1000 + idx,
                    paragraphs=[
                        paragraph(
                            text_run(
                                tile["eyebrow"],
                                size=8.5,
                                color=COLORS["blue"],
                                semibold=True,
                            )
                        ),
                        paragraph(
                            text_run(
                                tile["title"],
                                size=17 if tile["h"] >= 130 else 15,
                                color=COLORS["ink"],
                                semibold=True,
                            )
                        ),
                        paragraph(
                            text_run(
                                tile["description"],
                                size=9.5,
                                color=COLORS["secondary"],
                            )
                        ),
                    ],
                    background=fill,
                    border=border,
                    radius=10,
                ),
            )
        )

        files.append(
            write_visual(
                page_dir,
                make_transparent_nav_button(
                    name=nav_name,
                    target_page_name=target_page_name,
                    x=tile["x"],
                    y=tile["y"],
                    width=tile["w"],
                    height=tile["h"],
                    z=3000 + idx,
                    tab_order=tab_order,
                    tooltip=f"Open {tile['page']}",
                ),
            )
        )

        tab_order += 1

    files.append(
        write_visual(
            page_dir,
            make_textbox(
                name="idx_footer",
                x=64,
                y=1052,
                width=1792,
                height=20,
                z=200,
                paragraphs=[
                    paragraph(
                        text_run(
                            (
                                "Governed semantic model • Availability-aware KPIs "
                                "• Data quality preserved, not imputed"
                            ),
                            size=8.5,
                            color=COLORS["muted"],
                        )
                    )
                ],
            ),
        )
    )

    return files


def add_analytical_shell(
    page_dir: Path,
    *,
    page_number: int,
    display_name: str,
    subtitle: str,
    index_page_name: str,
) -> list[Path]:
    prefix = f"p{page_number:02d}"
    files: list[Path] = []
    files.extend(add_canvas_background(page_dir, prefix))

    # Accent line
    files.append(
        write_visual(
            page_dir,
            make_textbox(
                name=f"{prefix}_accent",
                x=0,
                y=0,
                width=CANVAS_WIDTH,
                height=7,
                z=20,
                paragraphs=[
                    paragraph(
                        text_run(
                            " ",
                            size=6,
                            color=COLORS["blue"],
                        )
                    )
                ],
                background=COLORS["blue"],
            ),
        )
    )

    files.append(
        write_visual(
            page_dir,
            make_textbox(
                name=f"{prefix}_header",
                x=64,
                y=45,
                width=1420,
                height=112,
                z=100,
                paragraphs=[
                    paragraph(
                        text_run(
                            display_name,
                            size=27,
                            color=COLORS["ink"],
                            semibold=True,
                        )
                    ),
                    paragraph(
                        text_run(
                            subtitle,
                            size=11,
                            color=COLORS["secondary"],
                        )
                    ),
                ],
            ),
        )
    )

    files.append(
        write_visual(
            page_dir,
            make_textbox(
                name=f"{prefix}_home_surface",
                x=1660,
                y=50,
                width=196,
                height=54,
                z=120,
                paragraphs=[
                    paragraph(
                        text_run(
                            "⌂   INDEX",
                            size=10.5,
                            color=COLORS["navy"],
                            semibold=True,
                        )
                    )
                ],
                background=COLORS["surface"],
                border=COLORS["border"],
                radius=8,
            ),
        )
    )

    files.append(
        write_visual(
            page_dir,
            make_transparent_nav_button(
                name=f"{prefix}_home_nav",
                target_page_name=index_page_name,
                x=1660,
                y=50,
                width=196,
                height=54,
                z=3000,
                tab_order=0,
                tooltip="Return to 00 INDEX",
            ),
        )
    )

    files.append(
        write_visual(
            page_dir,
            make_textbox(
                name=f"{prefix}_header_rule",
                x=64,
                y=172,
                width=1792,
                height=2,
                z=90,
                paragraphs=[
                    paragraph(
                        text_run(
                            " ",
                            size=6,
                            color=COLORS["border"],
                        )
                    )
                ],
                background=COLORS["border"],
            ),
        )
    )

    files.append(
        write_visual(
            page_dir,
            make_textbox(
                name=f"{prefix}_footer",
                x=64,
                y=1045,
                width=1792,
                height=24,
                z=90,
                paragraphs=[
                    paragraph(
                        text_run(
                            "Transparency in Coverage PUF  •  Experience Year 2024",
                            size=8.5,
                            color=COLORS["muted"],
                        )
                    )
                ],
            ),
        )
    )

    return files


# ---------------------------------------------------------------------
# Backup / validation / evidence
# ---------------------------------------------------------------------

def create_backup() -> None:
    REVIEW_DIR.mkdir(parents=True, exist_ok=True)

    if BACKUP_FILE.exists():
        BACKUP_FILE.unlink()

    with zipfile.ZipFile(
        BACKUP_FILE,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zf:
        for path in sorted(
            p for p in REPORT_DIR.rglob("*") if p.is_file()
        ):
            archive_name = (
                Path(REPORT_DIR.name)
                / path.relative_to(REPORT_DIR)
            )
            zf.write(path, str(archive_name))


def list_report_hashes() -> dict[str, str]:
    return {
        str(path.relative_to(REPORT_DIR)): sha256(path)
        for path in REPORT_DIR.rglob("*")
        if path.is_file()
    }


def count_existing_visuals() -> int:
    return len(list(PAGES_DIR.rglob("visual.json")))


def get_page_json_files() -> list[Path]:
    return sorted(PAGES_DIR.glob("*/page.json"))


def validate_pages_metadata(
    pages_payload: dict,
    expected_order: list[str],
    index_name: str,
) -> list[dict]:
    rows = []

    actual_order = pages_payload.get("pageOrder")
    rows.append(
        {
            "CheckID": "RPT-001",
            "TestName": "pages.json pageOrder",
            "Expected": len(expected_order),
            "Actual": len(actual_order) if isinstance(actual_order, list) else "MISSING",
            "Status": (
                "PASS"
                if actual_order == expected_order
                else "FAIL"
            ),
        }
    )

    rows.append(
        {
            "CheckID": "RPT-002",
            "TestName": "INDEX active page",
            "Expected": index_name,
            "Actual": pages_payload.get("activePageName"),
            "Status": (
                "PASS"
                if pages_payload.get("activePageName") == index_name
                else "FAIL"
            ),
        }
    )

    return rows


def validate_visual_files(
    display_by_name: dict[str, str],
    index_page_name: str,
) -> tuple[list[dict], list[dict]]:
    validation = []
    inventory = []

    page_dirs = sorted(PAGES_DIR.iterdir())

    for page_dir in page_dirs:
        if not page_dir.is_dir():
            continue

        page_json_file = page_dir / "page.json"
        if not page_json_file.exists():
            continue

        page_payload = read_json(page_json_file)
        page_name = page_payload.get("name", page_dir.name)
        page_display = page_payload.get("displayName", page_name)

        visual_files = sorted(
            (page_dir / "visuals").rglob("visual.json")
        ) if (page_dir / "visuals").exists() else []

        home_targets = []
        nav_targets = []
        names = []

        for visual_file in visual_files:
            payload = read_json(visual_file)
            names.append(payload.get("name"))

            visual = payload.get("visual", {})
            vtype = visual.get("visualType", "UNKNOWN")

            target = None

            for item in (
                visual.get("visualContainerObjects", {})
                .get("visualLink", [])
            ):
                props = item.get("properties", {})
                nav = (
                    props.get("navigationSection", {})
                    .get("expr", {})
                    .get("Literal", {})
                    .get("Value")
                )

                if isinstance(nav, str):
                    target = nav.strip("'")
                    nav_targets.append(target)

            if target == index_page_name:
                home_targets.append(target)

            inventory.append(
                {
                    "PageName": page_name,
                    "PageDisplayName": page_display,
                    "VisualName": payload.get("name"),
                    "VisualType": vtype,
                    "X": payload.get("position", {}).get("x"),
                    "Y": payload.get("position", {}).get("y"),
                    "Width": payload.get("position", {}).get("width"),
                    "Height": payload.get("position", {}).get("height"),
                    "NavigationTarget": target,
                    "RelativePath": str(
                        visual_file.relative_to(PROJECT_ROOT)
                    ),
                }
            )

        validation.append(
            {
                "CheckID": f"PAGE-{page_name}",
                "TestName": f"Unique visual names: {page_display}",
                "Expected": len(names),
                "Actual": len(set(names)),
                "Status": "PASS" if len(names) == len(set(names)) else "FAIL",
            }
        )

        if page_name != index_page_name:
            validation.append(
                {
                    "CheckID": f"HOME-{page_name}",
                    "TestName": f"Home navigation: {page_display}",
                    "Expected": 1,
                    "Actual": len(home_targets),
                    "Status": "PASS" if len(home_targets) == 1 else "FAIL",
                }
            )
        else:
            expected_targets = set(display_by_name.keys()) - {index_page_name}
            actual_targets = set(nav_targets)

            validation.append(
                {
                    "CheckID": "INDEX-NAV",
                    "TestName": "INDEX navigation targets",
                    "Expected": len(expected_targets),
                    "Actual": len(actual_targets),
                    "Status": (
                        "PASS"
                        if actual_targets == expected_targets
                        else "FAIL"
                    ),
                }
            )

    return validation, inventory


def style_report(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if ws.max_row and ws.max_column:
            ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(1, ws.max_column + 1):
            max_len = 0

            for row_idx in range(
                1,
                min(ws.max_row, 250) + 1,
            ):
                value = ws.cell(
                    row=row_idx,
                    column=col_idx,
                ).value

                if value is not None:
                    max_len = max(
                        max_len,
                        len(str(value)),
                    )

            ws.column_dimensions[
                get_column_letter(col_idx)
            ].width = min(
                max(max_len + 2, 10),
                60,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    wb.save(path)


# ---------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------

def main() -> int:
    print("=" * 78)
    print("Transparency in Coverage PUF — Report Shell + INDEX Build")
    print("=" * 78)
    print(f"Script version : {SCRIPT_VERSION}")
    print(f"Report folder  : {REPORT_DIR}")
    print(f"Canvas         : {CANVAS_WIDTH} x {CANVAS_HEIGHT}")
    print(f"Backup         : {BACKUP_FILE}")
    print(f"Review file    : {REPORT_FILE}")

    if power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is running. Save the PBIP and close Power BI "
            "Desktop completely before changing PBIR report files."
        )

    if not REPORT_DIR.exists():
        raise FileNotFoundError(
            f"Report folder not found: {REPORT_DIR}"
        )

    if not PAGES_FILE.exists():
        raise FileNotFoundError(
            f"pages.json not found: {PAGES_FILE}"
        )

    page_files = get_page_json_files()
    existing_visuals = count_existing_visuals()

    # Governance gate: this script intentionally targets the audited empty
    # report state only. It refuses to overwrite later work.
    if len(page_files) != 1:
        raise RuntimeError(
            f"Expected exactly 1 current page before shell build; found "
            f"{len(page_files)}. No changes made."
        )

    if existing_visuals != 0:
        raise RuntimeError(
            f"Expected 0 current visuals before shell build; found "
            f"{existing_visuals}. No changes made."
        )

    template_page_file = page_files[0]
    template_page_dir = template_page_file.parent
    template_page = read_json(template_page_file)

    index_page_name = template_page.get(
        "name",
        template_page_dir.name,
    )

    if not index_page_name:
        raise RuntimeError(
            "Could not determine the existing page name."
        )

    pages_payload = read_json(PAGES_FILE)

    if "pageOrder" not in pages_payload:
        raise RuntimeError(
            "Unsupported pages.json shape: pageOrder is missing. "
            "No changes made."
        )

    create_backup()
    before_hashes = list_report_hashes()

    # -----------------------------------------------------------------
    # Page IDs / display-name map
    # -----------------------------------------------------------------

    page_name_by_display = {
        "00 INDEX": index_page_name
    }

    for display_name, _subtitle in PAGE_SPECS:
        page_name_by_display[display_name] = deterministic_id(
            "TransparencyInCoverage:" + display_name
        )

    # -----------------------------------------------------------------
    # Convert current empty Page 1 into 00 INDEX.
    # -----------------------------------------------------------------

    index_payload = create_page_payload(
        template_page,
        page_name=index_page_name,
        display_name="00 INDEX",
    )
    write_json(template_page_file, index_payload)

    index_visuals_dir = template_page_dir / "visuals"
    if index_visuals_dir.exists():
        shutil.rmtree(index_visuals_dir)

    add_index_content(
        template_page_dir,
        page_name_by_display,
    )

    # -----------------------------------------------------------------
    # Create analytical page shells.
    # -----------------------------------------------------------------

    for page_number, (display_name, subtitle) in enumerate(
        PAGE_SPECS,
        start=1,
    ):
        page_name = page_name_by_display[display_name]
        page_dir = PAGES_DIR / page_name

        if page_dir.exists():
            raise RuntimeError(
                f"Target page directory already exists: {page_dir}. "
                "Restore backup before retrying."
            )

        page_dir.mkdir(parents=True, exist_ok=False)

        page_payload = create_page_payload(
            template_page,
            page_name=page_name,
            display_name=display_name,
        )

        write_json(
            page_dir / "page.json",
            page_payload,
        )

        add_analytical_shell(
            page_dir,
            page_number=page_number,
            display_name=display_name,
            subtitle=subtitle,
            index_page_name=index_page_name,
        )

    # -----------------------------------------------------------------
    # Rewrite page order / active page.
    # -----------------------------------------------------------------

    desired_order = [
        page_name_by_display["00 INDEX"],
        *[
            page_name_by_display[display_name]
            for display_name, _ in PAGE_SPECS
        ],
    ]

    pages_payload["pageOrder"] = desired_order
    pages_payload["activePageName"] = index_page_name

    write_json(PAGES_FILE, pages_payload)

    # -----------------------------------------------------------------
    # Offline validation.
    # -----------------------------------------------------------------

    current_page_files = get_page_json_files()

    validation_rows = []

    validation_rows.append(
        {
            "CheckID": "RPT-000",
            "TestName": "Page count",
            "Expected": 11,
            "Actual": len(current_page_files),
            "Status": "PASS" if len(current_page_files) == 11 else "FAIL",
        }
    )

    validation_rows.extend(
        validate_pages_metadata(
            read_json(PAGES_FILE),
            desired_order,
            index_page_name,
        )
    )

    display_by_name = {
        page_name: display_name
        for display_name, page_name in page_name_by_display.items()
    }

    visual_validation, visual_inventory = validate_visual_files(
        display_by_name,
        index_page_name,
    )

    validation_rows.extend(visual_validation)

    # All JSON in report definition must parse.
    parse_failures = []

    for json_file in sorted(
        DEFINITION_DIR.rglob("*.json")
    ):
        try:
            read_json(json_file)
        except Exception as exc:
            parse_failures.append(
                {
                    "File": str(
                        json_file.relative_to(PROJECT_ROOT)
                    ),
                    "Error": str(exc),
                }
            )

    validation_rows.append(
        {
            "CheckID": "RPT-JSON",
            "TestName": "PBIR JSON parse failures",
            "Expected": 0,
            "Actual": len(parse_failures),
            "Status": "PASS" if not parse_failures else "FAIL",
        }
    )

    failures = sum(
        row["Status"] == "FAIL"
        for row in validation_rows
    )

    # -----------------------------------------------------------------
    # Evidence
    # -----------------------------------------------------------------

    after_hashes = list_report_hashes()

    file_changes = []

    for rel_path in sorted(
        set(before_hashes) | set(after_hashes)
    ):
        before = before_hashes.get(rel_path)
        after = after_hashes.get(rel_path)

        if before == after:
            continue

        if before and after:
            state = "MODIFIED"
        elif after:
            state = "CREATED"
        else:
            state = "DELETED"

        file_changes.append(
            {
                "RelativePath": rel_path,
                "ChangeState": state,
                "BeforeSHA256": before,
                "AfterSHA256": after,
            }
        )

    page_inventory = []

    for page_json_file in get_page_json_files():
        payload = read_json(page_json_file)
        page_name = payload.get("name", page_json_file.parent.name)
        visual_count = len(
            list(
                (page_json_file.parent / "visuals")
                .rglob("visual.json")
            )
        ) if (page_json_file.parent / "visuals").exists() else 0

        page_inventory.append(
            {
                "PageName": page_name,
                "DisplayName": payload.get("displayName"),
                "Width": payload.get("width"),
                "Height": payload.get("height"),
                "VisualCount": visual_count,
                "IsIndex": page_name == index_page_name,
            }
        )

    page_inventory_df = pd.DataFrame(page_inventory)

    page_order_position = {
        name: idx
        for idx, name in enumerate(
            desired_order,
            start=1,
        )
    }

    page_inventory_df["PageOrder"] = (
        page_inventory_df["PageName"]
        .map(page_order_position)
    )

    page_inventory_df = page_inventory_df.sort_values(
        "PageOrder"
    )

    summary_df = pd.DataFrame(
        [
            (
                "BuildStatus",
                "PASS" if failures == 0 else "FAIL",
            ),
            ("ScriptVersion", SCRIPT_VERSION),
            (
                "BuiltAtLocal",
                datetime.now().isoformat(timespec="seconds"),
            ),
            ("ReportFolder", str(REPORT_DIR)),
            ("Canvas", f"{CANVAS_WIDTH} x {CANVAS_HEIGHT}"),
            ("PagesBefore", 1),
            ("PagesAfter", len(current_page_files)),
            ("IndexPage", "00 INDEX"),
            (
                "IndexNavigationTiles",
                len(INDEX_TILES),
            ),
            (
                "AnalyticalPagesWithHome",
                len(PAGE_SPECS),
            ),
            (
                "VisualsAfter",
                len(list(PAGES_DIR.rglob("visual.json"))),
            ),
            ("ValidationChecks", len(validation_rows)),
            ("ValidationFailures", failures),
            ("BackupFile", str(BACKUP_FILE)),
            (
                "EvidenceState",
                "PBIR SOURCE MODIFIED + OFFLINE STRUCTURAL VALIDATION",
            ),
        ],
        columns=["Item", "Value"],
    )

    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        REPORT_FILE,
        engine="openpyxl",
    ) as writer:
        summary_df.to_excel(
            writer,
            sheet_name="00_Summary",
            index=False,
        )

        pd.DataFrame(validation_rows).to_excel(
            writer,
            sheet_name="01_Validation",
            index=False,
        )

        page_inventory_df.to_excel(
            writer,
            sheet_name="02_Pages",
            index=False,
        )

        pd.DataFrame(visual_inventory).to_excel(
            writer,
            sheet_name="03_Visuals",
            index=False,
        )

        pd.DataFrame(file_changes).to_excel(
            writer,
            sheet_name="04_File_Changes",
            index=False,
        )

        if parse_failures:
            pd.DataFrame(parse_failures).to_excel(
                writer,
                sheet_name="05_Parse_Failures",
                index=False,
            )
        else:
            pd.DataFrame(
                [{"Info": "No parse failures"}]
            ).to_excel(
                writer,
                sheet_name="05_Parse_Failures",
                index=False,
            )

    style_report(REPORT_FILE)

    print()
    print("Report shell + INDEX build completed.")
    print(
        f"Build status          : "
        f"{'PASS' if failures == 0 else 'FAIL'}"
    )
    print(f"Pages after           : {len(current_page_files)}")
    print(
        f"INDEX navigation tiles: {len(INDEX_TILES)}"
    )
    print(
        f"Pages with Home       : {len(PAGE_SPECS)}"
    )
    print(
        f"Visuals after         : "
        f"{len(list(PAGES_DIR.rglob('visual.json')))}"
    )
    print(f"Validation failures   : {failures}")
    print(f"Backup                : {BACKUP_FILE}")
    print(f"Review evidence       : {REPORT_FILE}")

    if failures:
        print()
        print(
            "Offline validation failed. Do not open/save the PBIP "
            "until the evidence workbook is reviewed."
        )
        return 2

    print()
    print(
        "Next: open TransparencyInCoverage.pbip in Power BI Desktop. "
        "If it opens without a PBIR error, inspect 00 INDEX, test one "
        "navigation tile and one Home button, then Save (Ctrl+S)."
    )

    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())

    except Exception as exc:
        print()
        print("REPORT SHELL BUILD FAILED")
        print(str(exc))
        sys.exit(1)
