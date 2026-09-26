from __future__ import annotations

from datetime import datetime
from pathlib import Path
import hashlib
import json
import os
import subprocess
import sys
import zipfile

import pandas as pd


SCRIPT_VERSION = "1.0.0"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
REPORT_DIR = (
    PROJECT_ROOT
    / "powerbi"
    / "TransparencyInCoverage.Report"
)
PAGES_DIR = REPORT_DIR / "definition" / "pages"


def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, key_path) as key:
                raw_value, _ = winreg.QueryValueEx(key, "Desktop")
            return Path(os.path.expandvars(raw_value)).resolve()
        except Exception:
            pass

    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = get_windows_desktop() / "Transparency_PUF_Review"
BACKUP_FILE = REVIEW_DIR / "13_PRE_INDEX_TEXT_ALIGNMENT_BACKUP.zip"
REPORT_FILE = REVIEW_DIR / "13_INDEX_TEXT_ALIGNMENT_REPORT.xlsx"


def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False

    result = subprocess.run(
        ["tasklist", "/FI", "IMAGENAME eq PBIDesktop.exe"],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    return "PBIDesktop.exe" in result.stdout


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def write_json(path: Path, payload) -> None:
    path.write_text(
        json.dumps(payload, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def expr_double(value: float | int) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": f"{value}D"
            }
        }
    }


def create_backup() -> None:
    REVIEW_DIR.mkdir(parents=True, exist_ok=True)

    if BACKUP_FILE.exists():
        BACKUP_FILE.unlink()

    with zipfile.ZipFile(
        BACKUP_FILE,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zf:
        for path in sorted(
            p for p in REPORT_DIR.rglob("*") if p.is_file()
        ):
            archive_name = (
                Path(REPORT_DIR.name)
                / path.relative_to(REPORT_DIR)
            )
            zf.write(path, str(archive_name))


def find_index_page_dir() -> Path:
    matches = []

    for page_json in PAGES_DIR.glob("*/page.json"):
        payload = read_json(page_json)

        if payload.get("displayName") == "00 INDEX":
            matches.append(page_json.parent)

    if len(matches) != 1:
        raise RuntimeError(
            f"Expected exactly one page with displayName '00 INDEX'; "
            f"found {len(matches)}."
        )

    return matches[0]


def vertical_padding_for_height(height: float) -> tuple[int, int]:
    # The text block contains eyebrow + title + description.
    # These values visually center that block while retaining breathing room.
    if height >= 145:
        return 22, 16

    if height >= 125:
        return 17, 12

    if height >= 95:
        return 10, 8

    return 6, 6


def patch_tile(path: Path) -> dict:
    payload = read_json(path)

    name = payload.get("name", "")
    if not (
        name.startswith("idx_tile_")
        and name.endswith("_surface")
    ):
        raise ValueError(
            f"Not an INDEX tile surface visual: {path}"
        )

    visual = payload.get("visual", {})
    if visual.get("visualType") != "textbox":
        raise RuntimeError(
            f"Expected textbox for {name}; "
            f"found {visual.get('visualType')}"
        )

    general = (
        visual.get("objects", {})
        .get("general", [])
    )

    if not general:
        raise RuntimeError(
            f"Textbox {name} has no general object."
        )

    paragraphs = (
        general[0]
        .get("properties", {})
        .get("paragraphs", [])
    )

    if not paragraphs:
        raise RuntimeError(
            f"Textbox {name} has no paragraphs."
        )

    # Center every line horizontally.
    for p in paragraphs:
        p["horizontalTextAlignment"] = "center"

    height = float(
        payload.get("position", {}).get("height", 0)
    )
    top_padding, bottom_padding = vertical_padding_for_height(
        height
    )

    vco = visual.setdefault(
        "visualContainerObjects",
        {},
    )

    vco["padding"] = [
        {
            "properties": {
                "top": expr_double(top_padding),
                "bottom": expr_double(bottom_padding),
                "left": expr_double(18),
                "right": expr_double(18),
            }
        }
    ]

    before_hash = sha256(path)
    write_json(path, payload)
    after_hash = sha256(path)

    return {
        "VisualName": name,
        "Height": height,
        "HorizontalAlignment": "center",
        "TopPadding": top_padding,
        "BottomPadding": bottom_padding,
        "LeftPadding": 18,
        "RightPadding": 18,
        "BeforeSHA256": before_hash,
        "AfterSHA256": after_hash,
        "Status": "PATCHED",
    }


def validate_tile(path: Path) -> dict:
    payload = read_json(path)
    name = payload.get("name", "")

    paragraphs = (
        payload.get("visual", {})
        .get("objects", {})
        .get("general", [{}])[0]
        .get("properties", {})
        .get("paragraphs", [])
    )

    aligned = (
        bool(paragraphs)
        and all(
            p.get("horizontalTextAlignment") == "center"
            for p in paragraphs
        )
    )

    padding = (
        payload.get("visual", {})
        .get("visualContainerObjects", {})
        .get("padding", [])
    )

    has_padding = bool(padding)

    return {
        "VisualName": name,
        "Paragraphs": len(paragraphs),
        "AllCentered": aligned,
        "PaddingPresent": has_padding,
        "Status": (
            "PASS"
            if aligned and has_padding
            else "FAIL"
        ),
    }


def style_report(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if ws.max_row and ws.max_column:
            ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(1, ws.max_column + 1):
            max_len = 0

            for row_idx in range(
                1,
                min(ws.max_row, 200) + 1,
            ):
                value = ws.cell(
                    row=row_idx,
                    column=col_idx,
                ).value

                if value is not None:
                    max_len = max(
                        max_len,
                        len(str(value)),
                    )

            ws.column_dimensions[
                get_column_letter(col_idx)
            ].width = min(
                max(max_len + 2, 10),
                60,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    wb.save(path)


def main() -> int:
    print("=" * 78)
    print("Transparency in Coverage PUF — INDEX Tile Text Alignment")
    print("=" * 78)
    print(f"Script version : {SCRIPT_VERSION}")
    print(f"Report folder  : {REPORT_DIR}")
    print(f"Backup         : {BACKUP_FILE}")
    print(f"Review file    : {REPORT_FILE}")

    if power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is running. Save and close Power BI "
            "Desktop completely before applying this PBIR patch."
        )

    if not REPORT_DIR.exists():
        raise FileNotFoundError(
            f"Report folder not found: {REPORT_DIR}"
        )

    index_dir = find_index_page_dir()
    visuals_dir = index_dir / "visuals"

    tile_files = []

    for visual_json in sorted(
        visuals_dir.rglob("visual.json")
    ):
        payload = read_json(visual_json)
        name = payload.get("name", "")

        if (
            name.startswith("idx_tile_")
            and name.endswith("_surface")
        ):
            tile_files.append(visual_json)

    if len(tile_files) != 10:
        raise RuntimeError(
            f"Expected 10 INDEX tile surface visuals; "
            f"found {len(tile_files)}. No changes made."
        )

    create_backup()

    patch_rows = [
        patch_tile(path)
        for path in tile_files
    ]

    validation_rows = [
        validate_tile(path)
        for path in tile_files
    ]

    failures = sum(
        row["Status"] == "FAIL"
        for row in validation_rows
    )

    summary_df = pd.DataFrame(
        [
            (
                "PatchStatus",
                "PASS" if failures == 0 else "FAIL",
            ),
            ("ScriptVersion", SCRIPT_VERSION),
            (
                "PatchedAtLocal",
                datetime.now().isoformat(timespec="seconds"),
            ),
            ("IndexTilesPatched", len(tile_files)),
            ("HorizontalAlignment", "center"),
            (
                "VerticalTreatment",
                "height-aware top/bottom padding",
            ),
            ("ValidationFailures", failures),
            ("BackupFile", str(BACKUP_FILE)),
            (
                "EvidenceState",
                "PBIR INDEX VISUAL ALIGNMENT PATCH",
            ),
        ],
        columns=["Item", "Value"],
    )

    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        REPORT_FILE,
        engine="openpyxl",
    ) as writer:
        summary_df.to_excel(
            writer,
            sheet_name="00_Summary",
            index=False,
        )

        pd.DataFrame(patch_rows).to_excel(
            writer,
            sheet_name="01_Patched_Tiles",
            index=False,
        )

        pd.DataFrame(validation_rows).to_excel(
            writer,
            sheet_name="02_Validation",
            index=False,
        )

    style_report(REPORT_FILE)

    print()
    print("INDEX tile alignment patch completed.")
    print(
        f"Patch status        : "
        f"{'PASS' if failures == 0 else 'FAIL'}"
    )
    print(f"Tiles patched       : {len(tile_files)}")
    print("Horizontal alignment: center")
    print("Vertical treatment  : centered padding by tile height")
    print(f"Validation failures : {failures}")
    print(f"Backup              : {BACKUP_FILE}")
    print(f"Review evidence     : {REPORT_FILE}")

    if failures:
        print()
        print(
            "Validation failed. Do not open/save the PBIP "
            "until reviewed."
        )
        return 2

    print()
    print(
        "Next: open the PBIP, inspect 00 INDEX and Save (Ctrl+S) "
        "if the centered tile text looks correct."
    )

    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())

    except Exception as exc:
        print()
        print("INDEX ALIGNMENT PATCH FAILED")
        print(str(exc))
        sys.exit(1)
