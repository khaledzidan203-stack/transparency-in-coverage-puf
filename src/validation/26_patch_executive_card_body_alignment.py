from __future__ import annotations

from datetime import datetime
from pathlib import Path
import copy
import json
import os
import subprocess
import sys
import zipfile

import pandas as pd


SCRIPT_VERSION = "1.0.0"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
REPORT_DIR = (
    PROJECT_ROOT
    / "powerbi"
    / "TransparencyInCoverage.Report"
)
PAGES_DIR = REPORT_DIR / "definition" / "pages"

EXPECTED_VISUALS = 48
EXPECTED_BOUND_VISUALS = 20

# The user's current saved layout is now the baseline.
# This patch only separates card bodies from their title-band textboxes.
BODY_GAP = 2.0
POSITION_TOLERANCE = 0.01

CANVAS_WIDTH = 1920
CANVAS_HEIGHT = 1080

UNSUPPORTED_PAGE_ROOT_PROPERTIES = {
    "horizontalAlignment",
    "verticalAlignment",
}

# Every body/title pair on Executive Overview.
# If the user already fixed a pair manually, the script detects that
# the card body is below the title band and leaves it untouched.
CARD_BAND_PAIRS = [
    ("p01_dq_global", "p01_band_dq_global"),

    ("p01_kpi_issuers", "p01_band_kpi_issuers"),
    ("p01_kpi_plans", "p01_band_kpi_plans"),
    ("p01_kpi_claims", "p01_band_kpi_claims"),
    ("p01_kpi_denial_rate", "p01_band_kpi_denial_rate"),
    ("p01_kpi_appeal_rate", "p01_band_kpi_appeal_rate"),

    ("p01_network_in", "p01_band_network_in"),
    ("p01_network_out", "p01_band_network_out"),
    ("p01_network_gap", "p01_band_network_gap"),

    ("p01_claims_received", "p01_band_claims_received"),
    ("p01_claims_denied", "p01_band_claims_denied"),
    ("p01_claims_rate", "p01_band_claims_rate"),

    ("p01_appeal_internal_rate", "p01_band_appeal_internal_rate"),
    ("p01_appeal_external_rate", "p01_band_appeal_external_rate"),
    ("p01_appeal_internal_filed", "p01_band_appeal_internal_filed"),
    ("p01_appeal_external_filed", "p01_band_appeal_external_filed"),

    ("p01_known_source", "p01_band_known_source"),
]


# =====================================================================
# Paths / environment
# =====================================================================

def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )

            with winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                key_path,
            ) as key:
                raw_value, _ = winreg.QueryValueEx(
                    key,
                    "Desktop",
                )

            return Path(
                os.path.expandvars(raw_value)
            ).resolve()

        except Exception:
            pass

    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = (
    get_windows_desktop()
    / "Transparency_PUF_Review"
)

BACKUP_FILE = (
    REVIEW_DIR
    / "26_PRE_EXECUTIVE_CARD_BODY_ALIGNMENT_BACKUP.zip"
)

REPORT_FILE = (
    REVIEW_DIR
    / "26_EXECUTIVE_CARD_BODY_ALIGNMENT_REPORT.xlsx"
)


def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False

    result = subprocess.run(
        [
            "tasklist",
            "/FI",
            "IMAGENAME eq PBIDesktop.exe",
        ],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(
            subprocess,
            "CREATE_NO_WINDOW",
            0,
        ),
    )

    return "PBIDesktop.exe" in result.stdout


def read_json(path: Path):
    return json.loads(
        path.read_text(
            encoding="utf-8-sig",
        )
    )


def write_json(
    path: Path,
    payload,
) -> None:
    path.write_text(
        json.dumps(
            payload,
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )


def find_executive_page() -> tuple[Path, dict]:
    matches = []

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        if payload.get(
            "displayName"
        ) == "01 Executive Overview":
            matches.append(
                (
                    page_json.parent,
                    payload,
                )
            )

    if len(matches) != 1:
        raise RuntimeError(
            "Expected exactly one '01 Executive Overview' page; "
            f"found {len(matches)}."
        )

    return matches[0]


def create_backup() -> None:
    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if BACKUP_FILE.exists():
        BACKUP_FILE.unlink()

    with zipfile.ZipFile(
        BACKUP_FILE,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zf:
        for path in sorted(
            p
            for p in REPORT_DIR.rglob("*")
            if p.is_file()
        ):
            archive_name = (
                Path(REPORT_DIR.name)
                / path.relative_to(
                    REPORT_DIR
                )
            )

            zf.write(
                path,
                str(
                    archive_name
                ),
            )


# =====================================================================
# Inventory / geometry helpers
# =====================================================================

def visual_path(
    page_dir: Path,
    visual_name: str,
) -> Path:
    return (
        page_dir
        / "visuals"
        / visual_name
        / "visual.json"
    )


def get_position(
    page_dir: Path,
    visual_name: str,
) -> dict:
    path = visual_path(
        page_dir,
        visual_name,
    )

    if not path.exists():
        raise FileNotFoundError(
            f"Visual not found: {visual_name}"
        )

    payload = read_json(
        path
    )

    return copy.deepcopy(
        payload.get(
            "position",
            {},
        )
    )


def capture_positions(
    page_dir: Path,
) -> dict[str, dict]:
    result = {}

    for path in (
        page_dir
        / "visuals"
    ).rglob(
        "visual.json"
    ):
        payload = read_json(
            path
        )

        name = payload.get(
            "name"
        )

        if name:
            result[
                name
            ] = copy.deepcopy(
                payload.get(
                    "position",
                    {},
                )
            )

    return result


def count_bound_visuals(
    page_dir: Path,
) -> int:
    count = 0

    for path in (
        page_dir
        / "visuals"
    ).rglob(
        "visual.json"
    ):
        payload = read_json(
            path
        )

        if (
            payload.get(
                "visual",
                {},
            ).get(
                "query"
            )
        ):
            count += 1

    return count


# =====================================================================
# Patch logic
# =====================================================================

def align_card_body_below_band(
    page_dir: Path,
    card_name: str,
    band_name: str,
) -> dict:
    card_path = visual_path(
        page_dir,
        card_name,
    )
    band_path = visual_path(
        page_dir,
        band_name,
    )

    card_payload = read_json(
        card_path
    )
    band_payload = read_json(
        band_path
    )

    if (
        card_payload.get(
            "visual",
            {},
        ).get(
            "visualType"
        )
        != "card"
    ):
        raise RuntimeError(
            f"{card_name} is not a card visual."
        )

    if (
        band_payload.get(
            "visual",
            {},
        ).get(
            "visualType"
        )
        != "textbox"
    ):
        raise RuntimeError(
            f"{band_name} is not a textbox title band."
        )

    card_pos = card_payload[
        "position"
    ]
    band_pos = band_payload[
        "position"
    ]

    old_y = float(
        card_pos[
            "y"
        ]
    )
    old_height = float(
        card_pos[
            "height"
        ]
    )
    old_bottom = (
        old_y
        + old_height
    )

    band_bottom = (
        float(
            band_pos[
                "y"
            ]
        )
        + float(
            band_pos[
                "height"
            ]
        )
    )

    # User already fixed it manually -> preserve exactly.
    if (
        old_y
        >= band_bottom
        - POSITION_TOLERANCE
    ):
        return {
            "Card": card_name,
            "Band": band_name,
            "Action": "PRESERVED_USER_LAYOUT",
            "OldY": old_y,
            "OldHeight": old_height,
            "NewY": old_y,
            "NewHeight": old_height,
            "BandBottom": band_bottom,
            "OldBottom": old_bottom,
            "NewBottom": old_bottom,
            "Status": "UNCHANGED",
        }

    desired_y = (
        band_bottom
        + BODY_GAP
    )

    desired_height = (
        old_bottom
        - desired_y
    )

    # Do not create a body too small to render the callout value.
    if desired_height < 30:
        desired_y = band_bottom
        desired_height = (
            old_bottom
            - desired_y
        )

    if desired_height < 24:
        raise RuntimeError(
            f"Cannot safely separate {card_name} from {band_name}: "
            f"resulting body height would be {desired_height:.2f}px."
        )

    card_pos[
        "y"
    ] = desired_y

    card_pos[
        "height"
    ] = desired_height

    write_json(
        card_path,
        card_payload,
    )

    new_bottom = (
        desired_y
        + desired_height
    )

    return {
        "Card": card_name,
        "Band": band_name,
        "Action": "ALIGNED_BELOW_TITLE_BAND",
        "OldY": old_y,
        "OldHeight": old_height,
        "NewY": desired_y,
        "NewHeight": desired_height,
        "BandBottom": band_bottom,
        "OldBottom": old_bottom,
        "NewBottom": new_bottom,
        "Status": "PATCHED",
    }


# =====================================================================
# Validation
# =====================================================================

def validate_pair_alignment(
    page_dir: Path,
) -> tuple[dict, list[dict]]:
    failures = []
    rows = []

    for card_name, band_name in (
        CARD_BAND_PAIRS
    ):
        card_pos = get_position(
            page_dir,
            card_name,
        )
        band_pos = get_position(
            page_dir,
            band_name,
        )

        card_y = float(
            card_pos[
                "y"
            ]
        )

        card_bottom = (
            card_y
            + float(
                card_pos[
                    "height"
                ]
            )
        )

        band_bottom = (
            float(
                band_pos[
                    "y"
                ]
            )
            + float(
                band_pos[
                    "height"
                ]
            )
        )

        aligned = (
            card_y
            >= band_bottom
            - POSITION_TOLERANCE
        )

        row = {
            "Card": card_name,
            "Band": band_name,
            "CardY": card_y,
            "BandBottom": band_bottom,
            "CardBottom": card_bottom,
            "Aligned": aligned,
        }

        rows.append(
            row
        )

        if not aligned:
            failures.append(
                row
            )

    return (
        {
            "CheckID": "ALIGN-001",
            "TestName": (
                "Every card body begins below its title band"
            ),
            "Expected": 0,
            "Actual": len(
                failures
            ),
            "Status": (
                "PASS"
                if not failures
                else "FAIL"
            ),
        },
        rows,
    )


def validate_non_target_geometry(
    before: dict[str, dict],
    after: dict[str, dict],
) -> tuple[dict, list[dict]]:
    allowed_cards = {
        card
        for card, _
        in CARD_BAND_PAIRS
    }

    changes = []

    for name in sorted(
        set(
            before
        )
        | set(
            after
        )
    ):
        if name in allowed_cards:
            continue

        if (
            before.get(
                name
            )
            != after.get(
                name
            )
        ):
            changes.append(
                {
                    "VisualName": name,
                    "Before": json.dumps(
                        before.get(
                            name
                        ),
                        sort_keys=True,
                    ),
                    "After": json.dumps(
                        after.get(
                            name
                        ),
                        sort_keys=True,
                    ),
                }
            )

    return (
        {
            "CheckID": "ALIGN-002",
            "TestName": (
                "Non-card geometry preserved"
            ),
            "Expected": 0,
            "Actual": len(
                changes
            ),
            "Status": (
                "PASS"
                if not changes
                else "FAIL"
            ),
        },
        changes,
    )


def validate_bottom_preservation(
    patch_rows: list[dict],
) -> dict:
    changed_rows = [
        row
        for row in patch_rows
        if row[
            "Action"
        ]
        == "ALIGNED_BELOW_TITLE_BAND"
    ]

    bad = [
        row
        for row in changed_rows
        if abs(
            float(
                row[
                    "OldBottom"
                ]
            )
            - float(
                row[
                    "NewBottom"
                ]
            )
        )
        > POSITION_TOLERANCE
    ]

    return {
        "CheckID": "ALIGN-003",
        "TestName": (
            "Patched card bottoms preserved"
        ),
        "Expected": 0,
        "Actual": len(
            bad
        ),
        "Status": (
            "PASS"
            if not bad
            else "FAIL"
        ),
    }


def validate_no_unsupported_page_roots() -> list[dict]:
    offenders = []

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        for prop in sorted(
            UNSUPPORTED_PAGE_ROOT_PROPERTIES
        ):
            if prop in payload:
                offenders.append(
                    {
                        "Page": payload.get(
                            "displayName"
                        ),
                        "Property": prop,
                    }
                )

    return offenders


def validate_no_full_page_visuals() -> list[dict]:
    offenders = []

    for visual_json in PAGES_DIR.rglob(
        "visual.json"
    ):
        payload = read_json(
            visual_json
        )

        pos = payload.get(
            "position",
            {},
        )

        width = float(
            pos.get(
                "width",
                0,
            )
            or 0
        )

        height = float(
            pos.get(
                "height",
                0,
            )
            or 0
        )

        if (
            width
            >= CANVAS_WIDTH * 0.94
            and height
            >= CANVAS_HEIGHT * 0.94
        ):
            offenders.append(
                {
                    "VisualName": payload.get(
                        "name"
                    ),
                    "Width": width,
                    "Height": height,
                    "RelativePath": str(
                        visual_json.relative_to(
                            PROJECT_ROOT
                        )
                    ),
                }
            )

    return offenders


def validate_json_tree() -> list[dict]:
    failures = []

    for path in (
        REPORT_DIR
        / "definition"
    ).rglob(
        "*.json"
    ):
        try:
            read_json(
                path
            )

        except Exception as exc:
            failures.append(
                {
                    "RelativePath": str(
                        path.relative_to(
                            PROJECT_ROOT
                        )
                    ),
                    "Error": str(
                        exc
                    ),
                }
            )

    return failures


# =====================================================================
# Evidence workbook
# =====================================================================

def style_report(
    path: Path,
) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import (
        Alignment,
        Font,
        PatternFill,
    )
    from openpyxl.utils import (
        get_column_letter,
    )

    wb = load_workbook(
        path
    )

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if (
            ws.max_row
            and ws.max_column
        ):
            ws.auto_filter.ref = (
                ws.dimensions
            )

        for cell in ws[1]:
            cell.fill = (
                header_fill
            )
            cell.font = (
                header_font
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            max_len = 0

            for row_idx in range(
                1,
                min(
                    ws.max_row,
                    250,
                )
                + 1,
            ):
                value = ws.cell(
                    row=row_idx,
                    column=col_idx,
                ).value

                if value is not None:
                    max_len = max(
                        max_len,
                        len(
                            str(
                                value
                            )
                        ),
                    )

            ws.column_dimensions[
                get_column_letter(
                    col_idx
                )
            ].width = min(
                max(
                    max_len + 2,
                    10,
                ),
                75,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    wb.save(
        path
    )


# =====================================================================
# Main
# =====================================================================

def main() -> int:
    print("=" * 78)
    print(
        "Transparency in Coverage PUF — "
        "01 Executive Card-Body Alignment Patch"
    )
    print("=" * 78)

    print(
        f"Script version : {SCRIPT_VERSION}"
    )
    print(
        f"Report folder  : {REPORT_DIR}"
    )
    print(
        f"Backup         : {BACKUP_FILE}"
    )
    print(
        f"Review file    : {REPORT_FILE}"
    )

    if power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is running. "
            "Save the PBIP and close Desktop completely before applying the patch."
        )

    if not REPORT_DIR.exists():
        raise FileNotFoundError(
            f"Report folder not found: {REPORT_DIR}"
        )

    page_dir, _ = (
        find_executive_page()
    )

    visual_files = list(
        (
            page_dir
            / "visuals"
        ).rglob(
            "visual.json"
        )
    )

    visual_count = len(
        visual_files
    )

    bound_count = (
        count_bound_visuals(
            page_dir
        )
    )

    if visual_count != EXPECTED_VISUALS:
        raise RuntimeError(
            f"Expected current Executive Overview to have "
            f"{EXPECTED_VISUALS} visuals; found {visual_count}. "
            "No changes made."
        )

    if bound_count != EXPECTED_BOUND_VISUALS:
        raise RuntimeError(
            f"Expected current Executive Overview to have "
            f"{EXPECTED_BOUND_VISUALS} bound/query visuals; found {bound_count}. "
            "No changes made."
        )

    # Confirm all expected card/title pairs exist before making any change.
    missing = []

    for card_name, band_name in (
        CARD_BAND_PAIRS
    ):
        for name in (
            card_name,
            band_name,
        ):
            if not visual_path(
                page_dir,
                name,
            ).exists():
                missing.append(
                    name
                )

    if missing:
        raise RuntimeError(
            "Required visual(s) missing: "
            + ", ".join(
                sorted(
                    set(
                        missing
                    )
                )
            )
        )

    positions_before = (
        capture_positions(
            page_dir
        )
    )

    create_backup()

    patch_rows = []

    for card_name, band_name in (
        CARD_BAND_PAIRS
    ):
        patch_rows.append(
            align_card_body_below_band(
                page_dir,
                card_name,
                band_name,
            )
        )

    positions_after = (
        capture_positions(
            page_dir
        )
    )

    alignment_check, alignment_rows = (
        validate_pair_alignment(
            page_dir
        )
    )

    geometry_check, unexpected_geometry = (
        validate_non_target_geometry(
            positions_before,
            positions_after,
        )
    )

    bottom_check = (
        validate_bottom_preservation(
            patch_rows
        )
    )

    unsupported = (
        validate_no_unsupported_page_roots()
    )

    full_page = (
        validate_no_full_page_visuals()
    )

    json_failures = (
        validate_json_tree()
    )

    preserved_user_pairs = sum(
        row[
            "Action"
        ]
        == "PRESERVED_USER_LAYOUT"
        for row in patch_rows
    )

    patched_pairs = sum(
        row[
            "Action"
        ]
        == "ALIGNED_BELOW_TITLE_BAND"
        for row in patch_rows
    )

    validation_rows = [
        alignment_check,
        geometry_check,
        bottom_check,
        {
            "CheckID": "ALIGN-004",
            "TestName": "Visual count preserved",
            "Expected": EXPECTED_VISUALS,
            "Actual": len(
                list(
                    (
                        page_dir
                        / "visuals"
                    ).rglob(
                        "visual.json"
                    )
                )
            ),
            "Status": (
                "PASS"
                if len(
                    list(
                        (
                            page_dir
                            / "visuals"
                        ).rglob(
                            "visual.json"
                        )
                    )
                )
                == EXPECTED_VISUALS
                else "FAIL"
            ),
        },
        {
            "CheckID": "ALIGN-005",
            "TestName": "Bound/query visuals preserved",
            "Expected": EXPECTED_BOUND_VISUALS,
            "Actual": count_bound_visuals(
                page_dir
            ),
            "Status": (
                "PASS"
                if count_bound_visuals(
                    page_dir
                )
                == EXPECTED_BOUND_VISUALS
                else "FAIL"
            ),
        },
        {
            "CheckID": "ALIGN-006",
            "TestName": "Unsupported page-root properties",
            "Expected": 0,
            "Actual": len(
                unsupported
            ),
            "Status": (
                "PASS"
                if not unsupported
                else "FAIL"
            ),
        },
        {
            "CheckID": "ALIGN-007",
            "TestName": "Full-page selectable visuals",
            "Expected": 0,
            "Actual": len(
                full_page
            ),
            "Status": (
                "PASS"
                if not full_page
                else "FAIL"
            ),
        },
        {
            "CheckID": "ALIGN-008",
            "TestName": "PBIR JSON parse failures",
            "Expected": 0,
            "Actual": len(
                json_failures
            ),
            "Status": (
                "PASS"
                if not json_failures
                else "FAIL"
            ),
        },
    ]

    failures = sum(
        row[
            "Status"
        ] == "FAIL"
        for row in validation_rows
    )

    summary_df = pd.DataFrame(
        [
            (
                "PatchStatus",
                "PASS"
                if failures == 0
                else "FAIL",
            ),
            (
                "ScriptVersion",
                SCRIPT_VERSION,
            ),
            (
                "PatchedAtLocal",
                datetime.now()
                .isoformat(
                    timespec="seconds"
                ),
            ),
            (
                "Page",
                "01 Executive Overview",
            ),
            (
                "CardBandPairs",
                len(
                    CARD_BAND_PAIRS
                ),
            ),
            (
                "PreservedUserAdjustedPairs",
                preserved_user_pairs,
            ),
            (
                "RemainingPairsPatched",
                patched_pairs,
            ),
            (
                "BodyGapPixels",
                BODY_GAP,
            ),
            (
                "PatchRule",
                (
                    "If card already begins below its title band, preserve it. "
                    "Otherwise move only the card top below the band while "
                    "preserving the original card bottom."
                ),
            ),
            (
                "NonCardGeometryChanges",
                len(
                    unexpected_geometry
                ),
            ),
            (
                "ValidationFailures",
                failures,
            ),
            (
                "BackupFile",
                str(
                    BACKUP_FILE
                ),
            ),
        ],
        columns=[
            "Item",
            "Value",
        ],
    )

    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        REPORT_FILE,
        engine="openpyxl",
    ) as writer:
        summary_df.to_excel(
            writer,
            sheet_name="00_Summary",
            index=False,
        )

        pd.DataFrame(
            patch_rows
        ).to_excel(
            writer,
            sheet_name="01_Patches",
            index=False,
        )

        pd.DataFrame(
            validation_rows
        ).to_excel(
            writer,
            sheet_name="02_Validation",
            index=False,
        )

        pd.DataFrame(
            alignment_rows
        ).to_excel(
            writer,
            sheet_name="03_Final_Alignment",
            index=False,
        )

        if unexpected_geometry:
            pd.DataFrame(
                unexpected_geometry
            ).to_excel(
                writer,
                sheet_name="04_Unexpected_Geometry",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No visual geometry changed outside the approved "
                            "card-body patch targets."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="04_Unexpected_Geometry",
                index=False,
            )

        if unsupported:
            pd.DataFrame(
                unsupported
            ).to_excel(
                writer,
                sheet_name="05_Schema_Offenders",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No unsupported page-root properties."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="05_Schema_Offenders",
                index=False,
            )

        if full_page:
            pd.DataFrame(
                full_page
            ).to_excel(
                writer,
                sheet_name="06_FullPage_Offenders",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No full-page selectable visuals."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="06_FullPage_Offenders",
                index=False,
            )

        if json_failures:
            pd.DataFrame(
                json_failures
            ).to_excel(
                writer,
                sheet_name="07_JSON_Failures",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No JSON parse failures."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="07_JSON_Failures",
                index=False,
            )

    style_report(
        REPORT_FILE
    )

    print()
    print(
        "01 Executive card-body alignment patch completed."
    )
    print(
        f"Patch status                  : "
        f"{'PASS' if failures == 0 else 'FAIL'}"
    )
    print(
        f"Card/title pairs checked      : "
        f"{len(CARD_BAND_PAIRS)}"
    )
    print(
        f"User-adjusted pairs preserved : "
        f"{preserved_user_pairs}"
    )
    print(
        f"Remaining pairs patched       : "
        f"{patched_pairs}"
    )
    print(
        f"Non-card geometry changes     : "
        f"{len(unexpected_geometry)}"
    )
    print(
        f"Visuals preserved             : "
        f"{EXPECTED_VISUALS}"
    )
    print(
        f"Bound/query visuals preserved : "
        f"{EXPECTED_BOUND_VISUALS}"
    )
    print(
        f"Full-page selectable visuals  : "
        f"{len(full_page)}"
    )
    print(
        f"Unsupported page properties   : "
        f"{len(unsupported)}"
    )
    print(
        f"Validation failures           : "
        f"{failures}"
    )
    print(
        f"Backup                        : "
        f"{BACKUP_FILE}"
    )
    print(
        f"Review evidence               : "
        f"{REPORT_FILE}"
    )

    if failures:
        print()
        print(
            "Validation failed. Do not open/save the PBIP "
            "until the evidence workbook is reviewed."
        )
        return 2

    print()
    print(
        "Next: open TransparencyInCoverage.pbip and inspect page 01. "
        "The card rows you already fixed manually must remain unchanged, "
        "while every remaining card body should now start below its title band."
    )

    return 0


if __name__ == "__main__":
    try:
        sys.exit(
            main()
        )

    except Exception as exc:
        print()
        print(
            "EXECUTIVE CARD-BODY ALIGNMENT PATCH FAILED"
        )
        print(
            str(
                exc
            )
        )
        sys.exit(
            1
        )
