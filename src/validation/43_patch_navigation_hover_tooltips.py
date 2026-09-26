from __future__ import annotations

from datetime import datetime
from pathlib import Path
import hashlib
import json
import os
import subprocess
import sys
import zipfile

import pandas as pd


SCRIPT_VERSION = "1.0.0"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
REPORT_DIR = (
    PROJECT_ROOT
    / "powerbi"
    / "TransparencyInCoverage.Report"
)
PAGES_DIR = (
    REPORT_DIR
    / "definition"
    / "pages"
)

EXPECTED_PAGES = [
    "00 INDEX",
    "01 Executive Overview",
    "02 Network Comparison",
    "03 Denials",
    "04 Denial Reasons",
    "05 Appeals",
    "06 Resubmissions",
    "07 Enrollment",
    "08 Issuer & State Explorer",
    "09 Data Availability",
    "10 Data Quality & Methodology",
]

EXPECTED_INDEX_TILE_LINKS = 10
EXPECTED_HOME_LINKS = 10

TOOLTIP_HOME = (
    "Press Ctrl+Click to return to INDEX"
)

TOOLTIP_OPEN_TEMPLATE = (
    "Press Ctrl+Click to open {page}"
)


# =====================================================================
# Paths / environment
# =====================================================================

def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )

            with winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                key_path,
            ) as key:
                raw_value, _ = winreg.QueryValueEx(
                    key,
                    "Desktop",
                )

            return Path(
                os.path.expandvars(raw_value)
            ).resolve()

        except Exception:
            pass

    return (
        Path.home()
        / "Desktop"
    ).resolve()


REVIEW_DIR = (
    get_windows_desktop()
    / "Transparency_PUF_Review"
)

BACKUP_FILE = (
    REVIEW_DIR
    / "43_PRE_NAVIGATION_TOOLTIPS_BACKUP.zip"
)

REPORT_FILE = (
    REVIEW_DIR
    / "43_NAVIGATION_TOOLTIPS_PATCH_REPORT.xlsx"
)


def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False

    result = subprocess.run(
        [
            "tasklist",
            "/FI",
            "IMAGENAME eq PBIDesktop.exe",
        ],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(
            subprocess,
            "CREATE_NO_WINDOW",
            0,
        ),
    )

    return (
        "PBIDesktop.exe"
        in result.stdout
    )


def read_json(
    path: Path,
):
    return json.loads(
        path.read_text(
            encoding="utf-8-sig",
        )
    )


def write_json(
    path: Path,
    payload,
) -> None:
    path.write_text(
        json.dumps(
            payload,
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )


def sha256(
    path: Path,
) -> str:
    h = hashlib.sha256()

    with path.open(
        "rb"
    ) as f:
        for chunk in iter(
            lambda: f.read(
                1024 * 1024
            ),
            b"",
        ):
            h.update(
                chunk
            )

    return h.hexdigest()


# =====================================================================
# PBIR expression helpers
# =====================================================================

def expr_bool(
    value: bool,
) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": (
                    "true"
                    if value
                    else "false"
                )
            }
        }
    }


def expr_string(
    value: str,
) -> dict:
    escaped = value.replace(
        "'",
        "''",
    )

    return {
        "expr": {
            "Literal": {
                "Value": f"'{escaped}'"
            }
        }
    }


def decode_literal(
    obj,
):
    if not isinstance(
        obj,
        dict,
    ):
        return None

    value = (
        obj.get(
            "expr",
            {},
        )
        .get(
            "Literal",
            {},
        )
        .get(
            "Value"
        )
    )

    if not isinstance(
        value,
        str,
    ):
        return value

    if (
        len(
            value
        )
        >= 2
        and value.startswith(
            "'"
        )
        and value.endswith(
            "'"
        )
    ):
        return (
            value[1:-1]
            .replace(
                "''",
                "'",
            )
        )

    if value == "true":
        return True

    if value == "false":
        return False

    return value


# =====================================================================
# Discovery / backup
# =====================================================================

def find_pages() -> dict[
    str,
    tuple[Path, dict]
]:
    pages = {}

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        display_name = (
            payload.get(
                "displayName"
            )
        )

        if display_name:
            pages[
                display_name
            ] = (
                page_json.parent,
                payload,
            )

    return pages


def report_hashes() -> dict[
    str,
    str,
]:
    result = {}

    for path in REPORT_DIR.rglob(
        "*"
    ):
        if not path.is_file():
            continue

        result[
            str(
                path.relative_to(
                    REPORT_DIR
                )
            ).replace(
                "\\",
                "/",
            )
        ] = sha256(
            path
        )

    return result


def create_backup() -> None:
    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if BACKUP_FILE.exists():
        BACKUP_FILE.unlink()

    with zipfile.ZipFile(
        BACKUP_FILE,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zf:
        for path in sorted(
            p
            for p in REPORT_DIR.rglob(
                "*"
            )
            if p.is_file()
        ):
            archive_name = (
                Path(
                    REPORT_DIR.name
                )
                / path.relative_to(
                    REPORT_DIR
                )
            )

            zf.write(
                path,
                str(
                    archive_name
                ),
            )


# =====================================================================
# visualLink helpers
# =====================================================================

def get_visual_links(
    payload: dict,
) -> list[
    tuple[
        str,
        list,
    ]
]:
    visual = payload.get(
        "visual",
        {},
    )

    containers = []

    for branch_name in (
        "visualContainerObjects",
        "objects",
    ):
        branch = visual.get(
            branch_name,
            {},
        )

        links = branch.get(
            "visualLink"
        )

        if isinstance(
            links,
            list,
        ):
            containers.append(
                (
                    branch_name,
                    links,
                )
            )

    return containers


def extract_page_navigation(
    payload: dict,
) -> list[dict]:
    visual = payload.get(
        "visual",
        {},
    )

    if visual.get(
        "visualType"
    ) != "actionButton":
        return []

    result = []

    for (
        branch_name,
        links,
    ) in get_visual_links(
        payload
    ):
        for idx, link in enumerate(
            links
        ):
            props = link.get(
                "properties",
                {},
            )

            link_type = (
                decode_literal(
                    props.get(
                        "type"
                    )
                )
            )

            if (
                link_type
                != "PageNavigation"
            ):
                continue

            target = (
                decode_literal(
                    props.get(
                        "navigationSection"
                    )
                )
            )

            result.append(
                {
                    "branch": (
                        branch_name
                    ),
                    "index": idx,
                    "target": target,
                    "properties": props,
                }
            )

    return result


def patch_link_properties(
    props: dict,
    tooltip_text: str,
) -> None:
    # `tooltip` is the direct custom Action tooltip property.
    # `enabledTooltip` is also populated for Desktop compatibility.
    # Default tooltip is disabled so the user sees one concise custom instruction.
    props[
        "showDefaultTooltip"
    ] = expr_bool(
        False
    )

    props[
        "tooltip"
    ] = expr_string(
        tooltip_text
    )

    props[
        "enabledTooltip"
    ] = expr_string(
        tooltip_text
    )

    # Keep the action active; this does not alter navigation target or geometry.
    props[
        "show"
    ] = expr_bool(
        True
    )


# =====================================================================
# Patch
# =====================================================================

def patch_navigation_tooltips(
    pages: dict,
) -> list[dict]:
    index_internal_name = (
        pages[
            "00 INDEX"
        ][1][
            "name"
        ]
    )

    internal_to_display = {
        page_payload[
            "name"
        ]: display_name
        for display_name, (
            _,
            page_payload,
        ) in pages.items()
    }

    patch_rows = []

    for display_name, (
        page_dir,
        _,
    ) in pages.items():
        visuals_dir = (
            page_dir
            / "visuals"
        )

        if not visuals_dir.exists():
            continue

        for visual_json in sorted(
            visuals_dir.rglob(
                "visual.json"
            )
        ):
            payload = read_json(
                visual_json
            )

            nav_links = (
                extract_page_navigation(
                    payload
                )
            )

            if not nav_links:
                continue

            should_write = False

            for nav in nav_links:
                target = nav[
                    "target"
                ]

                is_index_tile = (
                    display_name
                    == "00 INDEX"
                    and target
                    in internal_to_display
                    and target
                    != index_internal_name
                )

                is_home_button = (
                    display_name
                    != "00 INDEX"
                    and target
                    == index_internal_name
                )

                if (
                    not is_index_tile
                    and not is_home_button
                ):
                    continue

                if is_index_tile:
                    target_display = (
                        internal_to_display[
                            target
                        ]
                    )

                    tooltip_text = (
                        TOOLTIP_OPEN_TEMPLATE.format(
                            page=target_display
                        )
                    )

                    role = (
                        "INDEX tile"
                    )

                else:
                    target_display = (
                        "00 INDEX"
                    )
                    tooltip_text = (
                        TOOLTIP_HOME
                    )
                    role = (
                        "Page INDEX button"
                    )

                before_props = json.dumps(
                    nav[
                        "properties"
                    ],
                    ensure_ascii=False,
                    sort_keys=True,
                )

                patch_link_properties(
                    nav[
                        "properties"
                    ],
                    tooltip_text,
                )

                after_props = json.dumps(
                    nav[
                        "properties"
                    ],
                    ensure_ascii=False,
                    sort_keys=True,
                )

                patch_rows.append(
                    {
                        "Page": (
                            display_name
                        ),
                        "Visual": (
                            payload.get(
                                "name"
                            )
                        ),
                        "Role": role,
                        "Target": (
                            target_display
                        ),
                        "Tooltip": (
                            tooltip_text
                        ),
                        "Branch": (
                            nav[
                                "branch"
                            ]
                        ),
                        "Changed": (
                            before_props
                            != after_props
                        ),
                        "RelativePath": str(
                            visual_json.relative_to(
                                REPORT_DIR
                            )
                        ).replace(
                            "\\",
                            "/",
                        ),
                    }
                )

                should_write = True

            if should_write:
                write_json(
                    visual_json,
                    payload,
                )

    return patch_rows


# =====================================================================
# Validation
# =====================================================================

def validate_json_tree() -> list[dict]:
    failures = []

    for path in (
        REPORT_DIR
        / "definition"
    ).rglob(
        "*.json"
    ):
        try:
            read_json(
                path
            )

        except Exception as exc:
            failures.append(
                {
                    "RelativePath": str(
                        path.relative_to(
                            PROJECT_ROOT
                        )
                    ),
                    "Error": str(
                        exc
                    ),
                }
            )

    return failures


def validate_patched_links(
    pages: dict,
) -> list[dict]:
    index_internal_name = (
        pages[
            "00 INDEX"
        ][1][
            "name"
        ]
    )

    internal_to_display = {
        page_payload[
            "name"
        ]: display_name
        for display_name, (
            _,
            page_payload,
        ) in pages.items()
    }

    rows = []

    for display_name, (
        page_dir,
        _,
    ) in pages.items():
        visuals_dir = (
            page_dir
            / "visuals"
        )

        for visual_json in sorted(
            visuals_dir.rglob(
                "visual.json"
            )
        ):
            payload = read_json(
                visual_json
            )

            for nav in extract_page_navigation(
                payload
            ):
                target = nav[
                    "target"
                ]

                if (
                    display_name
                    == "00 INDEX"
                    and target
                    in internal_to_display
                    and target
                    != index_internal_name
                ):
                    expected = (
                        TOOLTIP_OPEN_TEMPLATE.format(
                            page=(
                                internal_to_display[
                                    target
                                ]
                            )
                        )
                    )
                    role = (
                        "INDEX tile"
                    )

                elif (
                    display_name
                    != "00 INDEX"
                    and target
                    == index_internal_name
                ):
                    expected = (
                        TOOLTIP_HOME
                    )
                    role = (
                        "Page INDEX button"
                    )

                else:
                    continue

                props = nav[
                    "properties"
                ]

                actual_tooltip = (
                    decode_literal(
                        props.get(
                            "tooltip"
                        )
                    )
                )

                actual_enabled = (
                    decode_literal(
                        props.get(
                            "enabledTooltip"
                        )
                    )
                )

                actual_default = (
                    decode_literal(
                        props.get(
                            "showDefaultTooltip"
                        )
                    )
                )

                status = (
                    "PASS"
                    if (
                        actual_tooltip
                        == expected
                        and actual_enabled
                        == expected
                        and actual_default
                        is False
                    )
                    else "FAIL"
                )

                rows.append(
                    {
                        "Page": (
                            display_name
                        ),
                        "Visual": (
                            payload.get(
                                "name"
                            )
                        ),
                        "Role": role,
                        "ExpectedTooltip": (
                            expected
                        ),
                        "ActualTooltip": (
                            actual_tooltip
                        ),
                        "EnabledTooltip": (
                            actual_enabled
                        ),
                        "ShowDefaultTooltip": (
                            actual_default
                        ),
                        "Status": status,
                    }
                )

    return rows


def style_report(
    path: Path,
) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import (
        Alignment,
        Font,
        PatternFill,
    )
    from openpyxl.utils import (
        get_column_letter,
    )

    wb = load_workbook(
        path
    )

    fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )
    font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if (
            ws.max_row
            and ws.max_column
        ):
            ws.auto_filter.ref = (
                ws.dimensions
            )

        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            max_len = 0

            for row_idx in range(
                1,
                min(
                    ws.max_row,
                    160,
                )
                + 1,
            ):
                value = ws.cell(
                    row=row_idx,
                    column=col_idx,
                ).value

                if value is not None:
                    max_len = max(
                        max_len,
                        len(
                            str(
                                value
                            )
                        ),
                    )

            ws.column_dimensions[
                get_column_letter(
                    col_idx
                )
            ].width = min(
                max(
                    max_len + 2,
                    10,
                ),
                80,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    wb.save(
        path
    )


# =====================================================================
# Main
# =====================================================================

def main() -> int:
    print(
        "=" * 78
    )
    print(
        "Transparency in Coverage PUF — "
        "Navigation Hover Tooltip Patch"
    )
    print(
        "=" * 78
    )

    print(
        f"Script version : "
        f"{SCRIPT_VERSION}"
    )
    print(
        f"Report folder  : "
        f"{REPORT_DIR}"
    )
    print(
        f"Backup         : "
        f"{BACKUP_FILE}"
    )
    print(
        f"Review file    : "
        f"{REPORT_FILE}"
    )

    if power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is running. "
            "Save and close Power BI Desktop completely "
            "before applying this PBIR patch."
        )

    pages = find_pages()

    if set(
        pages
    ) != set(
        EXPECTED_PAGES
    ):
        raise RuntimeError(
            "Current report page inventory does not match "
            "the approved 11-page report. No changes made."
        )

    before_hashes = (
        report_hashes()
    )

    create_backup()

    patch_rows = (
        patch_navigation_tooltips(
            pages
        )
    )

    after_hashes = (
        report_hashes()
    )

    changed_files = sorted(
        rel_path
        for rel_path in (
            set(
                before_hashes
            )
            | set(
                after_hashes
            )
        )
        if (
            before_hashes.get(
                rel_path
            )
            != after_hashes.get(
                rel_path
            )
        )
    )

    json_failures = (
        validate_json_tree()
    )

    validation_rows = (
        validate_patched_links(
            pages
        )
    )

    index_tile_rows = [
        row
        for row in validation_rows
        if row[
            "Role"
        ]
        == "INDEX tile"
    ]

    home_rows = [
        row
        for row in validation_rows
        if row[
            "Role"
        ]
        == "Page INDEX button"
    ]

    link_failures = sum(
        row[
            "Status"
        ]
        == "FAIL"
        for row in (
            validation_rows
        )
    )

    expected_changed_paths = {
        row[
            "RelativePath"
        ]
        for row in patch_rows
    }

    unexpected_changed_files = [
        path
        for path in changed_files
        if path
        not in expected_changed_paths
    ]

    validation_summary = [
        {
            "Check": (
                "INDEX navigation tile tooltips"
            ),
            "Expected": (
                EXPECTED_INDEX_TILE_LINKS
            ),
            "Actual": len(
                index_tile_rows
            ),
            "Status": (
                "PASS"
                if (
                    len(
                        index_tile_rows
                    )
                    == EXPECTED_INDEX_TILE_LINKS
                    and all(
                        row[
                            "Status"
                        ]
                        == "PASS"
                        for row in (
                            index_tile_rows
                        )
                    )
                )
                else "FAIL"
            ),
        },
        {
            "Check": (
                "INDEX buttons on analytical pages"
            ),
            "Expected": (
                EXPECTED_HOME_LINKS
            ),
            "Actual": len(
                home_rows
            ),
            "Status": (
                "PASS"
                if (
                    len(
                        home_rows
                    )
                    == EXPECTED_HOME_LINKS
                    and all(
                        row[
                            "Status"
                        ]
                        == "PASS"
                        for row in (
                            home_rows
                        )
                    )
                )
                else "FAIL"
            ),
        },
        {
            "Check": (
                "Navigation tooltip validation failures"
            ),
            "Expected": 0,
            "Actual": (
                link_failures
            ),
            "Status": (
                "PASS"
                if link_failures
                == 0
                else "FAIL"
            ),
        },
        {
            "Check": (
                "Unexpected report files changed"
            ),
            "Expected": 0,
            "Actual": len(
                unexpected_changed_files
            ),
            "Status": (
                "PASS"
                if not unexpected_changed_files
                else "FAIL"
            ),
        },
        {
            "Check": (
                "JSON parse failures"
            ),
            "Expected": 0,
            "Actual": len(
                json_failures
            ),
            "Status": (
                "PASS"
                if not json_failures
                else "FAIL"
            ),
        },
    ]

    failures = sum(
        row[
            "Status"
        ]
        == "FAIL"
        for row in (
            validation_summary
        )
    )

    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        REPORT_FILE,
        engine="openpyxl",
    ) as writer:
        pd.DataFrame(
            [
                (
                    "PatchStatus",
                    (
                        "PASS"
                        if failures
                        == 0
                        else "FAIL"
                    ),
                ),
                (
                    "ScriptVersion",
                    SCRIPT_VERSION,
                ),
                (
                    "PatchedAtLocal",
                    datetime.now()
                    .isoformat(
                        timespec="seconds"
                    ),
                ),
                (
                    "IndexTileTooltips",
                    len(
                        index_tile_rows
                    ),
                ),
                (
                    "AnalyticalPageIndexTooltips",
                    len(
                        home_rows
                    ),
                ),
                (
                    "TotalNavigationTooltips",
                    len(
                        validation_rows
                    ),
                ),
                (
                    "ChangedVisualFiles",
                    len(
                        changed_files
                    ),
                ),
                (
                    "UnexpectedChangedFiles",
                    len(
                        unexpected_changed_files
                    ),
                ),
                (
                    "ValidationFailures",
                    failures,
                ),
                (
                    "BackupFile",
                    str(
                        BACKUP_FILE
                    ),
                ),
            ],
            columns=[
                "Item",
                "Value",
            ],
        ).to_excel(
            writer,
            sheet_name="00_Summary",
            index=False,
        )

        pd.DataFrame(
            validation_summary
        ).to_excel(
            writer,
            sheet_name="01_Validation",
            index=False,
        )

        pd.DataFrame(
            validation_rows
        ).to_excel(
            writer,
            sheet_name="02_Navigation_Tooltips",
            index=False,
        )

        pd.DataFrame(
            patch_rows
        ).to_excel(
            writer,
            sheet_name="03_Patched_Visuals",
            index=False,
        )

        if unexpected_changed_files:
            pd.DataFrame(
                {
                    "UnexpectedChangedFile": (
                        unexpected_changed_files
                    )
                }
            ).to_excel(
                writer,
                sheet_name="04_Unexpected_Changes",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No unexpected report files changed. "
                            "Geometry, colors, borders and data bindings were preserved."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="04_Unexpected_Changes",
                index=False,
            )

    style_report(
        REPORT_FILE
    )

    print()
    print(
        "Navigation hover-tooltip patch completed."
    )
    print(
        f"Patch status                  : "
        f"{'PASS' if failures == 0 else 'FAIL'}"
    )
    print(
        f"INDEX tile tooltips           : "
        f"{len(index_tile_rows)}"
    )
    print(
        f"Analytical INDEX tooltips     : "
        f"{len(home_rows)}"
    )
    print(
        f"Total navigation tooltips     : "
        f"{len(validation_rows)}"
    )
    print(
        f"Changed visual files          : "
        f"{len(changed_files)}"
    )
    print(
        f"Unexpected report changes     : "
        f"{len(unexpected_changed_files)}"
    )
    print(
        f"JSON parse failures           : "
        f"{len(json_failures)}"
    )
    print(
        f"Validation failures           : "
        f"{failures}"
    )
    print(
        f"Backup                        : "
        f"{BACKUP_FILE}"
    )
    print(
        f"Review evidence               : "
        f"{REPORT_FILE}"
    )

    if failures:
        print()
        print(
            "Validation failed. Review the evidence workbook "
            "before opening/saving the PBIP."
        )
        return 2

    print()
    print(
        "Next: open TransparencyInCoverage.pbip. "
        "Hover an INDEX tile and an INDEX button on an analytical page. "
        "The custom tooltip should instruct you to use Ctrl+Click."
    )

    return 0


if __name__ == "__main__":
    try:
        sys.exit(
            main()
        )

    except Exception as exc:
        print()
        print(
            "NAVIGATION TOOLTIP PATCH FAILED"
        )
        print(
            str(
                exc
            )
        )
        sys.exit(
            1
        )
