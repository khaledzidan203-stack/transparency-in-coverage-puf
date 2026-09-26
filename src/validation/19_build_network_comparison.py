from __future__ import annotations

from datetime import datetime
from pathlib import Path
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile

import pandas as pd


SCRIPT_VERSION = "1.0.0"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
POWERBI_ROOT = PROJECT_ROOT / "powerbi"
REPORT_DIR = POWERBI_ROOT / "TransparencyInCoverage.Report"
SEMANTIC_DIR = POWERBI_ROOT / "TransparencyInCoverage.SemanticModel"
PAGES_DIR = REPORT_DIR / "definition" / "pages"

MEASURES_TMDL = (
    SEMANTIC_DIR
    / "definition"
    / "tables"
    / "_Measures.tmdl"
)
DIM_STATE_TMDL = (
    SEMANTIC_DIR
    / "definition"
    / "tables"
    / "DimState.tmdl"
)
DIM_ISSUER_TMDL = (
    SEMANTIC_DIR
    / "definition"
    / "tables"
    / "DimIssuer.tmdl"
)

VISUAL_SCHEMA = (
    "https://developer.microsoft.com/json-schemas/fabric/item/report/"
    "definition/visualContainer/2.7.0/schema.json"
)

CANVAS_WIDTH = 1920
CANVAS_HEIGHT = 1080

COLORS = {
    "canvas": "#F5F7FA",
    "surface": "#FFFFFF",
    "surface_alt": "#F9FAFB",
    "border": "#E7ECF2",
    "ink": "#172033",
    "secondary": "#667085",
    "muted": "#98A2B3",
    "navy": "#17324D",
    "blue": "#4F6BED",
    "blue_dark": "#425CC7",
    "blue_soft": "#EEF3FF",
    "teal": "#16A6A1",
    "teal_soft": "#EAF8F6",
    "amber": "#C8872C",
    "red": "#C64B4B",
}

FONT_REGULAR = (
    "'Segoe UI', wf_segoe-ui_normal, helvetica, arial, sans-serif"
)
FONT_SEMIBOLD = (
    "'Segoe UI Semibold', wf_segoe-ui_semibold, "
    "helvetica, arial, sans-serif"
)

EXPECTED_PAGES = [
    "00 INDEX",
    "01 Executive Overview",
    "02 Network Comparison",
    "03 Denials",
    "04 Denial Reasons",
    "05 Appeals",
    "06 Resubmissions",
    "07 Enrollment",
    "08 Issuer & State Explorer",
    "09 Data Availability",
    "10 Data Quality & Methodology",
]

REQUIRED_MEASURES = {
    "Issuer Count",
    "Plan Count",
    "Issuer Claims Received - In Network",
    "Issuer Claims Received - Out of Network",
    "Issuer Claims Denied - In Network",
    "Issuer Claims Denied - Out of Network",
    "Issuer In-Network Denial Rate",
    "Issuer Out-of-Network Denial Rate",
    "Issuer Network Denial Rate Gap",
}

REQUIRED_COLUMNS = {
    "DimState": {
        "StateCode",
    },
    "DimIssuer": {
        "IssuerName",
    },
}

EXPECTED_VISUAL_NAMES = {
    "p02_context",
    "p02_page_title",
    "p02_index_surface",
    "p02_index_label",
    "p02_index_nav",
    "p02_state_slicer",
    "p02_issuer_slicer",
    "p02_kpi_in_rate",
    "p02_kpi_out_rate",
    "p02_kpi_gap",
    "p02_kpi_issuers",
    "p02_kpi_plans",
    "p02_state_gap_chart",
    "p02_flow_title",
    "p02_flow_in_received",
    "p02_flow_out_received",
    "p02_flow_in_denied",
    "p02_flow_out_denied",
    "p02_issuer_table",
    "p02_methodology_note",
}

EXPECTED_BOUND_VISUALS = 13

UNSUPPORTED_PAGE_ROOT_PROPERTIES = {
    "horizontalAlignment",
    "verticalAlignment",
}


# =====================================================================
# Windows / paths
# =====================================================================

def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )

            with winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                key_path,
            ) as key:
                raw_value, _ = winreg.QueryValueEx(
                    key,
                    "Desktop",
                )

            return Path(
                os.path.expandvars(raw_value)
            ).resolve()

        except Exception:
            pass

    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = (
    get_windows_desktop()
    / "Transparency_PUF_Review"
)

BACKUP_FILE = (
    REVIEW_DIR
    / "19_PRE_NETWORK_COMPARISON_BACKUP.zip"
)

REPORT_FILE = (
    REVIEW_DIR
    / "19_NETWORK_COMPARISON_BUILD_REPORT.xlsx"
)


def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False

    result = subprocess.run(
        [
            "tasklist",
            "/FI",
            "IMAGENAME eq PBIDesktop.exe",
        ],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(
            subprocess,
            "CREATE_NO_WINDOW",
            0,
        ),
    )

    return "PBIDesktop.exe" in result.stdout


def read_json(path: Path):
    return json.loads(
        path.read_text(
            encoding="utf-8-sig",
        )
    )


def write_json(
    path: Path,
    payload,
) -> None:
    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    path.write_text(
        json.dumps(
            payload,
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )


def sha256(path: Path) -> str:
    h = hashlib.sha256()

    with path.open("rb") as f:
        for chunk in iter(
            lambda: f.read(1024 * 1024),
            b"",
        ):
            h.update(chunk)

    return h.hexdigest()


# =====================================================================
# PBIR literals / formatting
# =====================================================================

def expr_bool(value: bool) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": (
                    "true"
                    if value
                    else "false"
                )
            }
        }
    }


def expr_string(value: str) -> dict:
    escaped = value.replace(
        "'",
        "''",
    )

    return {
        "expr": {
            "Literal": {
                "Value": f"'{escaped}'"
            }
        }
    }


def expr_double(
    value: float | int,
) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": f"{value}D"
            }
        }
    }


def expr_int(
    value: int,
) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": f"{value}L"
            }
        }
    }


def solid_color(
    color: str,
) -> dict:
    return {
        "solid": {
            "color": expr_string(
                color
            )
        }
    }


def page_background_objects() -> dict:
    """
    Native page background only.
    Do NOT create a page-sized textbox/shape visual.
    """
    return {
        "background": [
            {
                "properties": {
                    "color": solid_color(
                        COLORS["canvas"]
                    ),
                    "transparency": expr_double(
                        0
                    ),
                }
            }
        ]
    }


# =====================================================================
# TMDL inventory validation
# =====================================================================

def parse_tmdl_measures(
    path: Path,
) -> set[str]:
    text = path.read_text(
        encoding="utf-8-sig",
    )

    measures: set[str] = set()

    for line in text.splitlines():
        stripped = line.lstrip()

        if not stripped.startswith(
            "measure "
        ):
            continue

        lhs = stripped.split(
            "=",
            1,
        )[0].strip()

        raw = lhs[
            len("measure "):
        ].strip()

        if (
            raw.startswith("'")
            and raw.endswith("'")
        ):
            raw = (
                raw[1:-1]
                .replace(
                    "''",
                    "'",
                )
            )

        measures.add(
            raw
        )

    return measures


def parse_tmdl_columns(
    path: Path,
) -> set[str]:
    text = path.read_text(
        encoding="utf-8-sig",
    )

    columns: set[str] = set()

    for line in text.splitlines():
        stripped = line.lstrip()

        if not stripped.startswith(
            "column "
        ):
            continue

        raw = stripped[
            len("column "):
        ].strip()

        if raw.startswith("'"):
            match = re.match(
                r"'((?:''|[^'])*)'",
                raw,
            )

            if match:
                columns.add(
                    match.group(1)
                    .replace(
                        "''",
                        "'",
                    )
                )

        else:
            columns.add(
                raw.split(
                    "\t",
                    1,
                )[0].strip()
            )

    return columns


# =====================================================================
# Page discovery / backup
# =====================================================================

def find_pages() -> dict[
    str,
    tuple[Path, dict]
]:
    pages = {}

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        display_name = payload.get(
            "displayName"
        )

        if display_name:
            pages[
                display_name
            ] = (
                page_json.parent,
                payload,
            )

    return pages


def create_backup() -> None:
    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if BACKUP_FILE.exists():
        BACKUP_FILE.unlink()

    with zipfile.ZipFile(
        BACKUP_FILE,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zf:
        for path in sorted(
            p
            for p in REPORT_DIR.rglob("*")
            if p.is_file()
        ):
            archive_name = (
                Path(
                    REPORT_DIR.name
                )
                / path.relative_to(
                    REPORT_DIR
                )
            )

            zf.write(
                path,
                str(
                    archive_name
                ),
            )


def report_hashes() -> dict[
    str,
    str,
]:
    return {
        str(
            path.relative_to(
                REPORT_DIR
            )
        ): sha256(
            path
        )
        for path in REPORT_DIR.rglob(
            "*"
        )
        if path.is_file()
    }


# =====================================================================
# Page sanitation
# =====================================================================

def prepare_target_page(
    page_dir: Path,
) -> None:
    page_json = (
        page_dir
        / "page.json"
    )

    payload = read_json(
        page_json
    )

    for prop in (
        UNSUPPORTED_PAGE_ROOT_PROPERTIES
    ):
        payload.pop(
            prop,
            None,
        )

    payload["width"] = (
        CANVAS_WIDTH
    )
    payload["height"] = (
        CANVAS_HEIGHT
    )
    payload["displayOption"] = (
        "FitToPage"
    )

    objects = payload.setdefault(
        "objects",
        {},
    )

    objects.update(
        page_background_objects()
    )

    write_json(
        page_json,
        payload,
    )


# =====================================================================
# Textbox / container helpers
# =====================================================================

def text_run(
    value: str,
    *,
    size: float,
    color: str,
    semibold: bool = False,
) -> dict:
    return {
        "value": value,
        "textStyle": {
            "fontFamily": (
                FONT_SEMIBOLD
                if semibold
                else FONT_REGULAR
            ),
            "fontSize": (
                f"{size}pt"
            ),
            "color": color,
        },
    }


def paragraph(
    *runs: dict,
    align: str = "center",
) -> dict:
    return {
        "textRuns": list(
            runs
        ),
        "horizontalTextAlignment": (
            align
        ),
    }


def container_objects(
    *,
    background: str | None = None,
    border: str | None = None,
    radius: int = 8,
    show_title: bool = False,
    title: str = "",
    title_size: float = 11,
    title_color: str = COLORS["ink"],
    subtitle: str | None = None,
) -> dict:
    result = {
        "title": [
            {
                "properties": {
                    "show": expr_bool(
                        show_title
                    ),
                    **(
                        {
                            "text": expr_string(
                                title
                            ),
                            "alignment": expr_string(
                                "center"
                            ),
                            "fontColor": solid_color(
                                title_color
                            ),
                            "fontFamily": expr_string(
                                FONT_SEMIBOLD
                            ),
                            "fontSize": expr_double(
                                title_size
                            ),
                            "titleWrap": expr_bool(
                                True
                            ),
                        }
                        if show_title
                        else {}
                    ),
                }
            }
        ],
        "background": [
            {
                "properties": {
                    "show": expr_bool(
                        background is not None
                    ),
                    **(
                        {
                            "color": solid_color(
                                background
                            ),
                            "transparency": expr_double(
                                0
                            ),
                        }
                        if background
                        is not None
                        else {}
                    ),
                }
            }
        ],
        "border": [
            {
                "properties": {
                    "show": expr_bool(
                        border is not None
                    ),
                    **(
                        {
                            "color": solid_color(
                                border
                            ),
                            "width": expr_double(
                                1
                            ),
                            "radius": expr_double(
                                radius
                            ),
                        }
                        if border
                        is not None
                        else {}
                    ),
                }
            }
        ],
        "dropShadow": [
            {
                "properties": {
                    "show": expr_bool(
                        False
                    )
                }
            }
        ],
        "visualHeader": [
            {
                "properties": {
                    "show": expr_bool(
                        False
                    )
                }
            }
        ],
    }

    if subtitle is not None:
        result["subTitle"] = [
            {
                "properties": {
                    "show": expr_bool(
                        True
                    ),
                    "text": expr_string(
                        subtitle
                    ),
                    "alignment": expr_string(
                        "center"
                    ),
                    "fontColor": solid_color(
                        COLORS["secondary"]
                    ),
                    "fontFamily": expr_string(
                        FONT_REGULAR
                    ),
                    "fontSize": expr_double(
                        9.5
                    ),
                }
            }
        ]

    return result


def make_textbox(
    *,
    name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    paragraphs: list[dict],
    background: str | None = None,
    border: str | None = None,
    radius: int = 8,
) -> dict:
    return {
        "$schema": (
            VISUAL_SCHEMA
        ),
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
        },
        "visual": {
            "visualType": (
                "textbox"
            ),
            "objects": {
                "general": [
                    {
                        "properties": {
                            "paragraphs": (
                                paragraphs
                            )
                        }
                    }
                ]
            },
            "visualContainerObjects": (
                container_objects(
                    background=background,
                    border=border,
                    radius=radius,
                )
            ),
            "drillFilterOtherVisuals": (
                True
            ),
        },
    }


# =====================================================================
# Field binding helpers
# =====================================================================

def measure_field(
    measure_name: str,
) -> dict:
    return {
        "Measure": {
            "Expression": {
                "SourceRef": {
                    "Entity": (
                        "_Measures"
                    )
                }
            },
            "Property": (
                measure_name
            ),
        }
    }


def column_field(
    entity: str,
    column_name: str,
) -> dict:
    return {
        "Column": {
            "Expression": {
                "SourceRef": {
                    "Entity": entity
                }
            },
            "Property": (
                column_name
            ),
        }
    }


def measure_projection(
    measure_name: str,
) -> dict:
    return {
        "field": measure_field(
            measure_name
        ),
        "queryRef": (
            f"_Measures."
            f"{measure_name}"
        ),
        "nativeQueryRef": (
            measure_name
        ),
    }


def column_projection(
    entity: str,
    column_name: str,
    *,
    active: bool = False,
) -> dict:
    result = {
        "field": column_field(
            entity,
            column_name,
        ),
        "queryRef": (
            f"{entity}."
            f"{column_name}"
        ),
        "nativeQueryRef": (
            column_name
        ),
    }

    if active:
        result["active"] = (
            True
        )

    return result


# =====================================================================
# Navigation button
# =====================================================================

def make_index_surface(
    *,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
) -> dict:
    return make_textbox(
        name="p02_index_surface",
        x=x,
        y=y,
        width=width,
        height=height,
        z=z,
        paragraphs=[
            paragraph(
                text_run(
                    " ",
                    size=6,
                    color=COLORS["blue"],
                )
            )
        ],
        background=COLORS["blue"],
        border=COLORS["blue_dark"],
        radius=9,
    )


def make_index_label(
    *,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
) -> dict:
    return make_textbox(
        name="p02_index_label",
        x=x,
        y=y,
        width=width,
        height=height,
        z=z,
        paragraphs=[
            paragraph(
                text_run(
                    "INDEX",
                    size=11,
                    color="#FFFFFF",
                    semibold=True,
                ),
                align="center",
            )
        ],
        background=None,
        border=None,
        radius=0,
    )


def make_index_nav(
    *,
    target_page_name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
) -> dict:
    return {
        "$schema": (
            VISUAL_SCHEMA
        ),
        "name": (
            "p02_index_nav"
        ),
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": 0,
        },
        "visual": {
            "visualType": (
                "actionButton"
            ),
            "objects": {
                "icon": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "outline": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "fill": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": {
                "title": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "background": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "border": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "visualHeader": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "visualLink": [
                    {
                        "properties": {
                            "show": expr_bool(
                                True
                            ),
                            "type": expr_string(
                                "PageNavigation"
                            ),
                            "navigationSection": expr_string(
                                target_page_name
                            ),
                            "showDefaultTooltip": expr_bool(
                                False
                            ),
                            "enabledTooltip": expr_string(
                                "Return to INDEX"
                            ),
                        }
                    }
                ],
            },
            "drillFilterOtherVisuals": (
                True
            ),
        },
    }


# =====================================================================
# Slicer
# =====================================================================

def make_dropdown_slicer(
    *,
    name: str,
    entity: str,
    column_name: str,
    title: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    tab_order: int,
) -> dict:
    return {
        "$schema": (
            VISUAL_SCHEMA
        ),
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": (
                tab_order
            ),
        },
        "visual": {
            "visualType": (
                "slicer"
            ),
            "query": {
                "queryState": {
                    "Values": {
                        "projections": [
                            column_projection(
                                entity,
                                column_name,
                                active=True,
                            )
                        ]
                    }
                }
            },
            "objects": {
                "data": [
                    {
                        "properties": {
                            "mode": expr_string(
                                "Dropdown"
                            )
                        }
                    }
                ],
                "selection": [
                    {
                        "properties": {
                            "singleSelect": expr_bool(
                                True
                            )
                        }
                    }
                ],
                "header": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": (
                container_objects(
                    background=COLORS["surface"],
                    border=COLORS["border"],
                    radius=9,
                    show_title=True,
                    title=title,
                    title_size=9.5,
                )
            ),
            "drillFilterOtherVisuals": (
                True
            ),
        },
    }


# =====================================================================
# Card
# =====================================================================

def make_card(
    *,
    name: str,
    measure_name: str,
    title: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    tab_order: int,
    value_color: str = COLORS["ink"],
    background: str = COLORS["surface"],
    title_size: float = 10.5,
    value_size: float = 28,
) -> dict:
    return {
        "$schema": (
            VISUAL_SCHEMA
        ),
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": (
                tab_order
            ),
        },
        "visual": {
            "visualType": (
                "card"
            ),
            "query": {
                "queryState": {
                    "Values": {
                        "projections": [
                            measure_projection(
                                measure_name
                            )
                        ]
                    }
                },
                "sortDefinition": {
                    "sort": [
                        {
                            "field": (
                                measure_field(
                                    measure_name
                                )
                            ),
                            "direction": (
                                "Descending"
                            ),
                        }
                    ],
                    "isDefaultSort": (
                        True
                    ),
                },
            },
            "objects": {
                "labels": [
                    {
                        "properties": {
                            "color": solid_color(
                                value_color
                            ),
                            "fontSize": expr_double(
                                value_size
                            ),
                            "bold": expr_bool(
                                True
                            ),
                            "fontFamily": expr_string(
                                FONT_SEMIBOLD
                            ),
                        }
                    }
                ],
                "categoryLabels": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": (
                container_objects(
                    background=background,
                    border=COLORS["border"],
                    radius=9,
                    show_title=True,
                    title=title,
                    title_size=title_size,
                )
            ),
            "drillFilterOtherVisuals": (
                True
            ),
        },
    }


# =====================================================================
# State gap column chart
# =====================================================================

def make_state_gap_chart(
    *,
    name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    tab_order: int,
) -> dict:
    measure_name = (
        "Issuer Network Denial Rate Gap"
    )

    return {
        "$schema": (
            VISUAL_SCHEMA
        ),
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": (
                tab_order
            ),
        },
        "visual": {
            "visualType": (
                "columnChart"
            ),
            "query": {
                "queryState": {
                    "Category": {
                        "projections": [
                            column_projection(
                                "DimState",
                                "StateCode",
                                active=True,
                            )
                        ]
                    },
                    "Y": {
                        "projections": [
                            measure_projection(
                                measure_name
                            )
                        ]
                    },
                },
                "sortDefinition": {
                    "sort": [
                        {
                            "field": (
                                measure_field(
                                    measure_name
                                )
                            ),
                            "direction": (
                                "Descending"
                            ),
                        }
                    ],
                    "isDefaultSort": (
                        True
                    ),
                },
            },
            "objects": {
                "labels": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "legend": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "categoryAxis": [
                    {
                        "properties": {
                            "showAxisTitle": expr_bool(
                                False
                            ),
                            "labelColor": solid_color(
                                COLORS["secondary"]
                            ),
                            "fontFamily": expr_string(
                                FONT_REGULAR
                            ),
                            "fontSize": expr_double(
                                8.5
                            ),
                            "concatenateLabels": expr_bool(
                                False
                            ),
                            "maxMarginFactor": expr_int(
                                34
                            ),
                            "innerPadding": expr_int(
                                16
                            ),
                            "switchAxisPosition": expr_bool(
                                False
                            ),
                        }
                    }
                ],
                "valueAxis": [
                    {
                        "properties": {
                            "show": expr_bool(
                                True
                            ),
                            "showAxisTitle": expr_bool(
                                False
                            ),
                            "labelColor": solid_color(
                                COLORS["muted"]
                            ),
                            "fontFamily": expr_string(
                                FONT_REGULAR
                            ),
                            "fontSize": expr_double(
                                8
                            ),
                            "gridlineShow": expr_bool(
                                False
                            ),
                        }
                    }
                ],
                "general": [
                    {
                        "properties": {
                            "responsive": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "zoom": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "dataPoint": [
                    {
                        "properties": {
                            "fill": solid_color(
                                COLORS["blue"]
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": (
                container_objects(
                    background=COLORS["surface"],
                    border=COLORS["border"],
                    radius=9,
                    show_title=True,
                    title=(
                        "Where Is the Network Denial-Rate Gap Largest?"
                    ),
                    title_size=12.5,
                    subtitle=(
                        "State view • Out-of-network rate minus in-network rate"
                    ),
                )
            ),
            "drillFilterOtherVisuals": (
                True
            ),
        },
    }


# =====================================================================
# Issuer detail table
# =====================================================================

def make_issuer_table(
    *,
    name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    tab_order: int,
) -> dict:
    fields = [
        column_projection(
            "DimIssuer",
            "IssuerName",
        ),
        measure_projection(
            "Issuer In-Network Denial Rate"
        ),
        measure_projection(
            "Issuer Out-of-Network Denial Rate"
        ),
        measure_projection(
            "Issuer Network Denial Rate Gap"
        ),
        measure_projection(
            "Issuer Claims Received - In Network"
        ),
        measure_projection(
            "Issuer Claims Received - Out of Network"
        ),
    ]

    return {
        "$schema": (
            VISUAL_SCHEMA
        ),
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": (
                tab_order
            ),
        },
        "visual": {
            "visualType": (
                "tableEx"
            ),
            "query": {
                "queryState": {
                    "Values": {
                        "projections": (
                            fields
                        )
                    }
                },
                "sortDefinition": {
                    "sort": [
                        {
                            "field": (
                                measure_field(
                                    "Issuer Network Denial Rate Gap"
                                )
                            ),
                            "direction": (
                                "Descending"
                            ),
                        }
                    ],
                    "isDefaultSort": (
                        True
                    ),
                },
            },
            "objects": {
                "grid": [
                    {
                        "properties": {
                            "gridHorizontal": expr_bool(
                                True
                            ),
                            "gridHorizontalColor": solid_color(
                                COLORS["border"]
                            ),
                            "rowPadding": expr_double(
                                5
                            ),
                        }
                    }
                ],
                "columnHeaders": [
                    {
                        "properties": {
                            "alignment": expr_string(
                                "Center"
                            )
                        }
                    }
                ],
                "values": [
                    {
                        "properties": {
                            "backColorSecondary": solid_color(
                                COLORS["surface_alt"]
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": (
                container_objects(
                    background=COLORS["surface"],
                    border=COLORS["border"],
                    radius=9,
                    show_title=True,
                    title=(
                        "Issuer Detail — Network Rates & Volume"
                    ),
                    title_size=12.5,
                    subtitle=(
                        "Sorted by denial-rate gap; use State and Issuer filters above"
                    ),
                )
            ),
            "drillFilterOtherVisuals": (
                True
            ),
        },
    }


# =====================================================================
# Visual IO
# =====================================================================

def write_visual(
    page_dir: Path,
    visual: dict,
) -> Path:
    target = (
        page_dir
        / "visuals"
        / visual["name"]
        / "visual.json"
    )

    write_json(
        target,
        visual,
    )

    return target


# =====================================================================
# Build page
# =====================================================================

def build_network_page(
    page_dir: Path,
    index_page_name: str,
) -> None:
    visuals_dir = (
        page_dir
        / "visuals"
    )

    if visuals_dir.exists():
        shutil.rmtree(
            visuals_dir
        )

    visuals_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    # ---------------------------------------------------------------
    # Header: centered title, compact context, geometric-center INDEX.
    # ---------------------------------------------------------------

    write_visual(
        page_dir,
        make_textbox(
            name="p02_context",
            x=52,
            y=31,
            width=290,
            height=43,
            z=100,
            paragraphs=[
                paragraph(
                    text_run(
                        "2026 PUF  •  EXPERIENCE YEAR 2024",
                        size=9,
                        color=COLORS["blue"],
                        semibold=True,
                    )
                )
            ],
            background=COLORS["blue_soft"],
            border=None,
            radius=8,
        ),
    )

    write_visual(
        page_dir,
        make_textbox(
            name="p02_page_title",
            x=390,
            y=18,
            width=1140,
            height=84,
            z=110,
            paragraphs=[
                paragraph(
                    text_run(
                        "Network Comparison",
                        size=29,
                        color=COLORS["ink"],
                        semibold=True,
                    )
                ),
                paragraph(
                    text_run(
                        (
                            "Compare governed in-network and out-of-network "
                            "denial behavior, then locate where the gap concentrates"
                        ),
                        size=10.5,
                        color=COLORS["secondary"],
                    )
                ),
            ],
            background=None,
            border=None,
            radius=0,
        ),
    )

    index_x = 1648
    index_y = 25
    index_w = 220
    index_h = 56
    label_h = 22
    label_y = (
        index_y
        + (
            index_h
            - label_h
        )
        / 2
    )

    write_visual(
        page_dir,
        make_index_surface(
            x=index_x,
            y=index_y,
            width=index_w,
            height=index_h,
            z=130,
        ),
    )

    write_visual(
        page_dir,
        make_index_label(
            x=index_x,
            y=label_y,
            width=index_w,
            height=label_h,
            z=140,
        ),
    )

    write_visual(
        page_dir,
        make_index_nav(
            target_page_name=(
                index_page_name
            ),
            x=index_x,
            y=index_y,
            width=index_w,
            height=index_h,
            z=5000,
        ),
    )

    # ---------------------------------------------------------------
    # Filter band. Two slicers only: state -> issuer.
    # ---------------------------------------------------------------

    write_visual(
        page_dir,
        make_dropdown_slicer(
            name="p02_state_slicer",
            entity="DimState",
            column_name="StateCode",
            title="State",
            x=548,
            y=102,
            width=300,
            height=62,
            z=200,
            tab_order=1,
        ),
    )

    write_visual(
        page_dir,
        make_dropdown_slicer(
            name="p02_issuer_slicer",
            entity="DimIssuer",
            column_name="IssuerName",
            title="Issuer",
            x=864,
            y=102,
            width=508,
            height=62,
            z=200,
            tab_order=2,
        ),
    )

    # ---------------------------------------------------------------
    # KPI strip: comparison first, population context second.
    # ---------------------------------------------------------------

    margin = 52
    gap = 16
    total_width = (
        CANVAS_WIDTH
        - 2 * margin
    )
    card_width = (
        total_width
        - 4 * gap
    ) / 5

    kpis = [
        (
            "p02_kpi_in_rate",
            "Issuer In-Network Denial Rate",
            "In-Network Denial Rate",
            COLORS["navy"],
            COLORS["surface"],
        ),
        (
            "p02_kpi_out_rate",
            "Issuer Out-of-Network Denial Rate",
            "Out-of-Network Denial Rate",
            COLORS["blue"],
            COLORS["blue_soft"],
        ),
        (
            "p02_kpi_gap",
            "Issuer Network Denial Rate Gap",
            "Network Gap",
            COLORS["teal"],
            COLORS["teal_soft"],
        ),
        (
            "p02_kpi_issuers",
            "Issuer Count",
            "Issuers in Context",
            COLORS["ink"],
            COLORS["surface"],
        ),
        (
            "p02_kpi_plans",
            "Plan Count",
            "Plans in Context",
            COLORS["ink"],
            COLORS["surface"],
        ),
    ]

    for idx, (
        visual_name,
        measure_name,
        title,
        value_color,
        background,
    ) in enumerate(
        kpis
    ):
        write_visual(
            page_dir,
            make_card(
                name=visual_name,
                measure_name=(
                    measure_name
                ),
                title=title,
                x=(
                    margin
                    + idx
                    * (
                        card_width
                        + gap
                    )
                ),
                y=180,
                width=card_width,
                height=132,
                z=400 + idx,
                tab_order=10 + idx,
                value_color=value_color,
                background=background,
                title_size=10,
                value_size=28,
            ),
        )

    # ---------------------------------------------------------------
    # Main comparison: state gap chart.
    # ---------------------------------------------------------------

    write_visual(
        page_dir,
        make_state_gap_chart(
            name="p02_state_gap_chart",
            x=52,
            y=334,
            width=1148,
            height=336,
            z=700,
            tab_order=30,
        ),
    )

    # ---------------------------------------------------------------
    # Selected-context claims flow, 2x2 cards.
    # ---------------------------------------------------------------

    write_visual(
        page_dir,
        make_textbox(
            name="p02_flow_title",
            x=1216,
            y=334,
            width=652,
            height=44,
            z=250,
            paragraphs=[
                paragraph(
                    text_run(
                        "SELECTED CONTEXT — CLAIMS FLOW",
                        size=12,
                        color=COLORS["ink"],
                        semibold=True,
                    )
                ),
                paragraph(
                    text_run(
                        "Reported received and denied claims by network status",
                        size=9,
                        color=COLORS["secondary"],
                    )
                ),
            ],
            background=None,
            border=None,
            radius=0,
        ),
    )

    flow_x = 1216
    flow_y = 390
    flow_w = 652
    flow_gap = 16
    flow_card_w = (
        flow_w
        - flow_gap
    ) / 2
    flow_card_h = 132

    flow_cards = [
        (
            "p02_flow_in_received",
            "Issuer Claims Received - In Network",
            "In-Network Received",
            COLORS["navy"],
            COLORS["surface_alt"],
            flow_x,
            flow_y,
        ),
        (
            "p02_flow_out_received",
            "Issuer Claims Received - Out of Network",
            "Out-of-Network Received",
            COLORS["blue"],
            COLORS["blue_soft"],
            (
                flow_x
                + flow_card_w
                + flow_gap
            ),
            flow_y,
        ),
        (
            "p02_flow_in_denied",
            "Issuer Claims Denied - In Network",
            "In-Network Denied",
            COLORS["navy"],
            COLORS["surface_alt"],
            flow_x,
            (
                flow_y
                + flow_card_h
                + flow_gap
            ),
        ),
        (
            "p02_flow_out_denied",
            "Issuer Claims Denied - Out of Network",
            "Out-of-Network Denied",
            COLORS["blue"],
            COLORS["blue_soft"],
            (
                flow_x
                + flow_card_w
                + flow_gap
            ),
            (
                flow_y
                + flow_card_h
                + flow_gap
            ),
        ),
    ]

    for idx, (
        name,
        measure,
        title,
        value_color,
        background,
        x,
        y,
    ) in enumerate(
        flow_cards
    ):
        write_visual(
            page_dir,
            make_card(
                name=name,
                measure_name=measure,
                title=title,
                x=x,
                y=y,
                width=flow_card_w,
                height=flow_card_h,
                z=800 + idx,
                tab_order=40 + idx,
                value_color=value_color,
                background=background,
                title_size=9.5,
                value_size=23,
            ),
        )

    # ---------------------------------------------------------------
    # Detail gradient: issuer table at bottom.
    # Scroll is intentional here because this is the investigation layer.
    # ---------------------------------------------------------------

    write_visual(
        page_dir,
        make_issuer_table(
            name="p02_issuer_table",
            x=52,
            y=692,
            width=1816,
            height=322,
            z=900,
            tab_order=60,
        ),
    )

    # ---------------------------------------------------------------
    # Methodology strip.
    # ---------------------------------------------------------------

    write_visual(
        page_dir,
        make_textbox(
            name="p02_methodology_note",
            x=52,
            y=1026,
            width=1816,
            height=32,
            z=200,
            paragraphs=[
                paragraph(
                    text_run(
                        (
                            "Availability-aware ratio-of-totals  •  "
                            "Network gap = out-of-network denial rate − "
                            "in-network denial rate  •  Detail table is "
                            "intentionally scrollable for issuer investigation"
                        ),
                        size=8.8,
                        color=COLORS["muted"],
                        semibold=True,
                    )
                )
            ],
            background=None,
            border=None,
            radius=0,
        ),
    )


# =====================================================================
# Validation
# =====================================================================

def walk_json(
    obj,
):
    if isinstance(
        obj,
        dict,
    ):
        yield obj

        for value in (
            obj.values()
        ):
            yield from walk_json(
                value
            )

    elif isinstance(
        obj,
        list,
    ):
        for value in obj:
            yield from walk_json(
                value
            )


def validate_json_tree() -> list[dict]:
    failures = []

    for path in sorted(
        (
            REPORT_DIR
            / "definition"
        ).rglob(
            "*.json"
        )
    ):
        try:
            read_json(
                path
            )

        except Exception as exc:
            failures.append(
                {
                    "RelativePath": (
                        str(
                            path.relative_to(
                                PROJECT_ROOT
                            )
                        )
                    ),
                    "Error": (
                        str(
                            exc
                        )
                    ),
                }
            )

    return failures


def validate_no_unsupported_page_roots(
    pages: dict,
) -> list[dict]:
    offenders = []

    for display_name, (
        page_dir,
        _,
    ) in pages.items():
        payload = read_json(
            page_dir
            / "page.json"
        )

        for prop in sorted(
            UNSUPPORTED_PAGE_ROOT_PROPERTIES
        ):
            if prop in payload:
                offenders.append(
                    {
                        "Page": display_name,
                        "UnsupportedProperty": prop,
                    }
                )

    return offenders


def validate_no_full_page_visuals(
    pages: dict,
) -> list[dict]:
    offenders = []

    for display_name, (
        page_dir,
        _,
    ) in pages.items():
        visuals_dir = (
            page_dir
            / "visuals"
        )

        if not visuals_dir.exists():
            continue

        for path in visuals_dir.rglob(
            "visual.json"
        ):
            payload = read_json(
                path
            )

            pos = payload.get(
                "position",
                {},
            )

            width = float(
                pos.get(
                    "width",
                    0,
                )
                or 0
            )
            height = float(
                pos.get(
                    "height",
                    0,
                )
                or 0
            )

            if (
                width
                >= CANVAS_WIDTH * 0.94
                and height
                >= CANVAS_HEIGHT * 0.94
            ):
                offenders.append(
                    {
                        "Page": (
                            display_name
                        ),
                        "VisualName": (
                            payload.get(
                                "name"
                            )
                        ),
                        "VisualType": (
                            payload.get(
                                "visual",
                                {},
                            ).get(
                                "visualType"
                            )
                        ),
                        "Width": width,
                        "Height": height,
                    }
                )

    return offenders


def validate_binding_inventory(
    page_dir: Path,
) -> list[dict]:
    rows = []

    for path in sorted(
        (
            page_dir
            / "visuals"
        ).rglob(
            "visual.json"
        )
    ):
        payload = read_json(
            path
        )

        visual = payload.get(
            "visual",
            {},
        )

        if not visual.get(
            "query"
        ):
            continue

        fields = []

        for node in walk_json(
            visual.get(
                "query",
                {},
            )
        ):
            if "Measure" in node:
                m = node[
                    "Measure"
                ]

                prop = m.get(
                    "Property"
                )

                if prop:
                    fields.append(
                        (
                            "Measure",
                            "_Measures",
                            prop,
                        )
                    )

            if "Column" in node:
                c = node[
                    "Column"
                ]

                prop = c.get(
                    "Property"
                )

                entity = (
                    c.get(
                        "Expression",
                        {},
                    )
                    .get(
                        "SourceRef",
                        {},
                    )
                    .get(
                        "Entity"
                    )
                )

                if prop:
                    fields.append(
                        (
                            "Column",
                            entity,
                            prop,
                        )
                    )

        rows.append(
            {
                "VisualName": (
                    payload.get(
                        "name"
                    )
                ),
                "VisualType": (
                    visual.get(
                        "visualType"
                    )
                ),
                "Bindings": (
                    " | ".join(
                        f"{kind}:"
                        f"{entity}."
                        f"{field}"
                        for (
                            kind,
                            entity,
                            field,
                        )
                        in fields
                    )
                ),
            }
        )

    return rows


def validate_index_geometry(
    page_dir: Path,
) -> dict:
    surface = read_json(
        page_dir
        / "visuals"
        / "p02_index_surface"
        / "visual.json"
    )

    label = read_json(
        page_dir
        / "visuals"
        / "p02_index_label"
        / "visual.json"
    )

    surface_pos = (
        surface[
            "position"
        ]
    )
    label_pos = (
        label[
            "position"
        ]
    )

    expected_label_y = (
        surface_pos[
            "y"
        ]
        + (
            surface_pos[
                "height"
            ]
            - label_pos[
                "height"
            ]
        )
        / 2
    )

    label_text = (
        label[
            "visual"
        ][
            "objects"
        ][
            "general"
        ][0][
            "properties"
        ][
            "paragraphs"
        ][0][
            "textRuns"
        ][0][
            "value"
        ]
    )

    label_align = (
        label[
            "visual"
        ][
            "objects"
        ][
            "general"
        ][0][
            "properties"
        ][
            "paragraphs"
        ][0][
            "horizontalTextAlignment"
        ]
    )

    passed = (
        abs(
            float(
                label_pos[
                    "y"
                ]
            )
            - float(
                expected_label_y
            )
        )
        < 0.001
        and float(
            label_pos[
                "x"
            ]
        )
        == float(
            surface_pos[
                "x"
            ]
        )
        and float(
            label_pos[
                "width"
            ]
        )
        == float(
            surface_pos[
                "width"
            ]
        )
        and label_text
        == "INDEX"
        and label_align
        == "center"
    )

    return {
        "CheckID": (
            "NET-INDEX"
        ),
        "TestName": (
            "INDEX label geometric centering"
        ),
        "Expected": (
            "centered"
        ),
        "Actual": (
            "centered"
            if passed
            else "not-centered"
        ),
        "Status": (
            "PASS"
            if passed
            else "FAIL"
        ),
    }


def validate_page_title_center(
    page_dir: Path,
) -> dict:
    payload = read_json(
        page_dir
        / "visuals"
        / "p02_page_title"
        / "visual.json"
    )

    paragraphs = (
        payload[
            "visual"
        ][
            "objects"
        ][
            "general"
        ][0][
            "properties"
        ][
            "paragraphs"
        ]
    )

    passed = all(
        p.get(
            "horizontalTextAlignment"
        )
        == "center"
        for p in paragraphs
    )

    return {
        "CheckID": (
            "NET-TITLE"
        ),
        "TestName": (
            "Page title centered"
        ),
        "Expected": (
            "center"
        ),
        "Actual": (
            "center"
            if passed
            else "not-centered"
        ),
        "Status": (
            "PASS"
            if passed
            else "FAIL"
        ),
    }


def validate_slicers(
    page_dir: Path,
) -> dict:
    expected = {
        "p02_state_slicer": (
            "DimState",
            "StateCode",
        ),
        "p02_issuer_slicer": (
            "DimIssuer",
            "IssuerName",
        ),
    }

    failures = []

    for visual_name, (
        expected_entity,
        expected_column,
    ) in expected.items():
        payload = read_json(
            page_dir
            / "visuals"
            / visual_name
            / "visual.json"
        )

        query = (
            payload.get(
                "visual",
                {},
            )
            .get(
                "query",
                {},
            )
        )

        found = []

        for node in walk_json(
            query
        ):
            if "Column" in node:
                column = node[
                    "Column"
                ]

                entity = (
                    column.get(
                        "Expression",
                        {},
                    )
                    .get(
                        "SourceRef",
                        {},
                    )
                    .get(
                        "Entity"
                    )
                )

                prop = column.get(
                    "Property"
                )

                if (
                    entity
                    and prop
                ):
                    found.append(
                        (
                            entity,
                            prop,
                        )
                    )

        if (
            expected_entity,
            expected_column,
        ) not in found:
            failures.append(
                visual_name
            )

    return {
        "CheckID": (
            "NET-SLICERS"
        ),
        "TestName": (
            "State and Issuer slicer bindings"
        ),
        "Expected": 2,
        "Actual": (
            2
            - len(
                failures
            )
        ),
        "Status": (
            "PASS"
            if not failures
            else "FAIL"
        ),
    }


# =====================================================================
# Evidence workbook styling
# =====================================================================

def style_report(
    path: Path,
) -> None:
    from openpyxl import (
        load_workbook,
    )
    from openpyxl.styles import (
        Alignment,
        Font,
        PatternFill,
    )
    from openpyxl.utils import (
        get_column_letter,
    )

    wb = load_workbook(
        path
    )

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = (
            "A2"
        )

        if (
            ws.max_row
            and ws.max_column
        ):
            ws.auto_filter.ref = (
                ws.dimensions
            )

        for cell in ws[1]:
            cell.fill = (
                header_fill
            )
            cell.font = (
                header_font
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            max_len = 0

            for row_idx in range(
                1,
                min(
                    ws.max_row,
                    250,
                )
                + 1,
            ):
                value = ws.cell(
                    row=row_idx,
                    column=col_idx,
                ).value

                if value is not None:
                    max_len = max(
                        max_len,
                        len(
                            str(
                                value
                            )
                        ),
                    )

            ws.column_dimensions[
                get_column_letter(
                    col_idx
                )
            ].width = min(
                max(
                    max_len + 2,
                    10,
                ),
                70,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    wb.save(
        path
    )


# =====================================================================
# Main
# =====================================================================

def main() -> int:
    print("=" * 78)
    print(
        "Transparency in Coverage PUF — "
        "02 Network Comparison Build"
    )
    print("=" * 78)

    print(
        f"Script version : "
        f"{SCRIPT_VERSION}"
    )
    print(
        f"Report folder  : "
        f"{REPORT_DIR}"
    )
    print(
        f"Backup         : "
        f"{BACKUP_FILE}"
    )
    print(
        f"Review file    : "
        f"{REPORT_FILE}"
    )

    # ---------------------------------------------------------------
    # Safety gates
    # ---------------------------------------------------------------

    if power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is running. "
            "Save the PBIP and close Power BI Desktop completely "
            "before changing PBIR files."
        )

    if not REPORT_DIR.exists():
        raise FileNotFoundError(
            f"Report folder not found: "
            f"{REPORT_DIR}"
        )

    for required_file in (
        MEASURES_TMDL,
        DIM_STATE_TMDL,
        DIM_ISSUER_TMDL,
    ):
        if not required_file.exists():
            raise FileNotFoundError(
                f"Required semantic-model file not found: "
                f"{required_file}"
            )

    pages = find_pages()

    if set(
        pages
    ) != set(
        EXPECTED_PAGES
    ):
        raise RuntimeError(
            "Current report page inventory does not match "
            "the approved 11-page report shell. "
            "No changes made."
        )

    installed_measures = (
        parse_tmdl_measures(
            MEASURES_TMDL
        )
    )

    missing_measures = sorted(
        REQUIRED_MEASURES
        - installed_measures
    )

    if missing_measures:
        raise RuntimeError(
            "Required governed measures are missing: "
            + ", ".join(
                missing_measures
            )
        )

    installed_columns = {
        "DimState": (
            parse_tmdl_columns(
                DIM_STATE_TMDL
            )
        ),
        "DimIssuer": (
            parse_tmdl_columns(
                DIM_ISSUER_TMDL
            )
        ),
    }

    missing_columns = []

    for table_name, required in (
        REQUIRED_COLUMNS.items()
    ):
        missing = sorted(
            required
            - installed_columns[
                table_name
            ]
        )

        for column_name in missing:
            missing_columns.append(
                f"{table_name}."
                f"{column_name}"
            )

    if missing_columns:
        raise RuntimeError(
            "Required semantic columns are missing: "
            + ", ".join(
                missing_columns
            )
        )

    # Do not proceed if V3 page-schema problem has returned anywhere.
    unsupported_before = (
        validate_no_unsupported_page_roots(
            pages
        )
    )

    if unsupported_before:
        raise RuntimeError(
            "Unsupported page-root properties are present before build: "
            + "; ".join(
                f"{row['Page']}:{row['UnsupportedProperty']}"
                for row in unsupported_before
            )
            + ". Fix schema state before building page 02."
        )

    index_dir, index_page = (
        pages[
            "00 INDEX"
        ]
    )

    network_dir, network_page = (
        pages[
            "02 Network Comparison"
        ]
    )

    index_page_name = (
        index_page[
            "name"
        ]
    )

    if (
        network_page.get(
            "displayName"
        )
        != "02 Network Comparison"
    ):
        raise RuntimeError(
            "Target page identity mismatch. "
            "No changes made."
        )

    # ---------------------------------------------------------------
    # Backup + before hashes
    # ---------------------------------------------------------------

    create_backup()

    before_hashes = (
        report_hashes()
    )

    # ---------------------------------------------------------------
    # Destructive rebuild of page 02 only.
    # ---------------------------------------------------------------

    prepare_target_page(
        network_dir
    )

    build_network_page(
        network_dir,
        index_page_name,
    )

    # Refresh inventory after build.
    pages = find_pages()

    # ---------------------------------------------------------------
    # Offline validation
    # ---------------------------------------------------------------

    json_failures = (
        validate_json_tree()
    )

    unsupported_after = (
        validate_no_unsupported_page_roots(
            pages
        )
    )

    full_page_offenders = (
        validate_no_full_page_visuals(
            pages
        )
    )

    visual_files = sorted(
        (
            network_dir
            / "visuals"
        ).rglob(
            "visual.json"
        )
    )

    actual_visual_names = {
        read_json(
            path
        ).get(
            "name"
        )
        for path in visual_files
    }

    binding_rows = (
        validate_binding_inventory(
            network_dir
        )
    )

    page_json = read_json(
        network_dir
        / "page.json"
    )

    native_background = (
        page_json.get(
            "objects",
            {},
        )
        .get(
            "background",
            [],
        )
    )

    validation_rows = [
        {
            "CheckID": (
                "NET-001"
            ),
            "TestName": (
                "Approved page inventory"
            ),
            "Expected": len(
                EXPECTED_PAGES
            ),
            "Actual": len(
                pages
            ),
            "Status": (
                "PASS"
                if set(
                    pages
                )
                == set(
                    EXPECTED_PAGES
                )
                else "FAIL"
            ),
        },
        {
            "CheckID": (
                "NET-002"
            ),
            "TestName": (
                "Network page visual inventory"
            ),
            "Expected": len(
                EXPECTED_VISUAL_NAMES
            ),
            "Actual": len(
                actual_visual_names
            ),
            "Status": (
                "PASS"
                if actual_visual_names
                == EXPECTED_VISUAL_NAMES
                else "FAIL"
            ),
        },
        {
            "CheckID": (
                "NET-003"
            ),
            "TestName": (
                "Bound/query visuals"
            ),
            "Expected": (
                EXPECTED_BOUND_VISUALS
            ),
            "Actual": len(
                binding_rows
            ),
            "Status": (
                "PASS"
                if len(
                    binding_rows
                )
                == EXPECTED_BOUND_VISUALS
                else "FAIL"
            ),
        },
        {
            "CheckID": (
                "NET-004"
            ),
            "TestName": (
                "Native page background"
            ),
            "Expected": (
                "present"
            ),
            "Actual": (
                "present"
                if native_background
                else "missing"
            ),
            "Status": (
                "PASS"
                if native_background
                else "FAIL"
            ),
        },
        {
            "CheckID": (
                "NET-005"
            ),
            "TestName": (
                "Unsupported page-root properties"
            ),
            "Expected": 0,
            "Actual": len(
                unsupported_after
            ),
            "Status": (
                "PASS"
                if not
                unsupported_after
                else "FAIL"
            ),
        },
        {
            "CheckID": (
                "NET-006"
            ),
            "TestName": (
                "Full-page selectable visuals"
            ),
            "Expected": 0,
            "Actual": len(
                full_page_offenders
            ),
            "Status": (
                "PASS"
                if not
                full_page_offenders
                else "FAIL"
            ),
        },
        {
            "CheckID": (
                "NET-007"
            ),
            "TestName": (
                "PBIR JSON parse failures"
            ),
            "Expected": 0,
            "Actual": len(
                json_failures
            ),
            "Status": (
                "PASS"
                if not
                json_failures
                else "FAIL"
            ),
        },
        validate_page_title_center(
            network_dir
        ),
        validate_index_geometry(
            network_dir
        ),
        validate_slicers(
            network_dir
        ),
    ]

    failures = sum(
        row[
            "Status"
        ]
        == "FAIL"
        for row
        in validation_rows
    )

    # ---------------------------------------------------------------
    # Evidence
    # ---------------------------------------------------------------

    after_hashes = (
        report_hashes()
    )

    file_changes = []

    for rel_path in sorted(
        set(
            before_hashes
        )
        | set(
            after_hashes
        )
    ):
        before = before_hashes.get(
            rel_path
        )
        after = after_hashes.get(
            rel_path
        )

        if before == after:
            continue

        if (
            before is not None
            and after is not None
        ):
            change_state = (
                "MODIFIED"
            )

        elif after is not None:
            change_state = (
                "CREATED"
            )

        else:
            change_state = (
                "DELETED"
            )

        file_changes.append(
            {
                "RelativePath": (
                    rel_path
                ),
                "ChangeState": (
                    change_state
                ),
                "BeforeSHA256": (
                    before
                ),
                "AfterSHA256": (
                    after
                ),
            }
        )

    visual_inventory = []

    for path in visual_files:
        payload = read_json(
            path
        )

        visual_inventory.append(
            {
                "VisualName": (
                    payload.get(
                        "name"
                    )
                ),
                "VisualType": (
                    payload.get(
                        "visual",
                        {},
                    ).get(
                        "visualType"
                    )
                ),
                "X": (
                    payload.get(
                        "position",
                        {},
                    ).get(
                        "x"
                    )
                ),
                "Y": (
                    payload.get(
                        "position",
                        {},
                    ).get(
                        "y"
                    )
                ),
                "Width": (
                    payload.get(
                        "position",
                        {},
                    ).get(
                        "width"
                    )
                ),
                "Height": (
                    payload.get(
                        "position",
                        {},
                    ).get(
                        "height"
                    )
                ),
                "RelativePath": (
                    str(
                        path.relative_to(
                            PROJECT_ROOT
                        )
                    )
                ),
            }
        )

    summary_df = pd.DataFrame(
        [
            (
                "BuildStatus",
                (
                    "PASS"
                    if failures
                    == 0
                    else "FAIL"
                ),
            ),
            (
                "ScriptVersion",
                SCRIPT_VERSION,
            ),
            (
                "BuiltAtLocal",
                datetime.now()
                .isoformat(
                    timespec="seconds"
                ),
            ),
            (
                "PageRebuilt",
                "02 Network Comparison",
            ),
            (
                "PageIntent",
                (
                    "Compare in-network vs out-of-network denial behavior "
                    "and locate the gap by state and issuer."
                ),
            ),
            (
                "Visuals",
                len(
                    visual_files
                ),
            ),
            (
                "BoundVisuals",
                len(
                    binding_rows
                ),
            ),
            (
                "Slicers",
                2,
            ),
            (
                "NativePageBackground",
                "YES",
            ),
            (
                "FullPageSelectableVisuals",
                len(
                    full_page_offenders
                ),
            ),
            (
                "UnsupportedPageRootProperties",
                len(
                    unsupported_after
                ),
            ),
            (
                "ValidationChecks",
                len(
                    validation_rows
                ),
            ),
            (
                "ValidationFailures",
                failures,
            ),
            (
                "BackupFile",
                str(
                    BACKUP_FILE
                ),
            ),
            (
                "DesignContract",
                (
                    "Centered page/chart/card titles; geometric-center INDEX; "
                    "native page background; no page-sized visual background; "
                    "detail table intentionally scrollable."
                ),
            ),
        ],
        columns=[
            "Item",
            "Value",
        ],
    )

    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        REPORT_FILE,
        engine="openpyxl",
    ) as writer:
        summary_df.to_excel(
            writer,
            sheet_name="00_Summary",
            index=False,
        )

        pd.DataFrame(
            validation_rows
        ).to_excel(
            writer,
            sheet_name="01_Validation",
            index=False,
        )

        pd.DataFrame(
            visual_inventory
        ).to_excel(
            writer,
            sheet_name="02_Visuals",
            index=False,
        )

        pd.DataFrame(
            binding_rows
        ).to_excel(
            writer,
            sheet_name="03_Bindings",
            index=False,
        )

        pd.DataFrame(
            file_changes
        ).to_excel(
            writer,
            sheet_name="04_File_Changes",
            index=False,
        )

        if unsupported_after:
            pd.DataFrame(
                unsupported_after
            ).to_excel(
                writer,
                sheet_name="05_Schema_Offenders",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No unsupported page-root properties."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="05_Schema_Offenders",
                index=False,
            )

        if full_page_offenders:
            pd.DataFrame(
                full_page_offenders
            ).to_excel(
                writer,
                sheet_name="06_FullPage_Offenders",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No full-page selectable visuals."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="06_FullPage_Offenders",
                index=False,
            )

        if json_failures:
            pd.DataFrame(
                json_failures
            ).to_excel(
                writer,
                sheet_name="07_JSON_Failures",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No JSON parse failures."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="07_JSON_Failures",
                index=False,
            )

    style_report(
        REPORT_FILE
    )

    print()
    print(
        "02 Network Comparison build completed."
    )
    print(
        f"Build status                 : "
        f"{'PASS' if failures == 0 else 'FAIL'}"
    )
    print(
        f"Visuals                      : "
        f"{len(visual_files)}"
    )
    print(
        f"Bound/query visuals          : "
        f"{len(binding_rows)}"
    )
    print(
        "State/Issuer slicers         : 2"
    )
    print(
        f"Full-page selectable visuals : "
        f"{len(full_page_offenders)}"
    )
    print(
        f"Unsupported page properties  : "
        f"{len(unsupported_after)}"
    )
    print(
        "Page/chart/card titles       : center"
    )
    print(
        "INDEX label                  : geometric center"
    )
    print(
        f"Validation failures          : "
        f"{failures}"
    )
    print(
        f"Backup                       : "
        f"{BACKUP_FILE}"
    )
    print(
        f"Review evidence              : "
        f"{REPORT_FILE}"
    )

    if failures:
        print()
        print(
            "Offline validation failed. "
            "Do not open/save the PBIP until the evidence report is reviewed."
        )
        return 2

    print()
    print(
        "Next: open TransparencyInCoverage.pbip and inspect "
        "02 Network Comparison. Test State and Issuer slicers, "
        "INDEX navigation, the state-gap chart, and issuer table. "
        "Click blank canvas space to confirm no page-sized visual is selected."
    )

    return 0


if __name__ == "__main__":
    try:
        sys.exit(
            main()
        )

    except Exception as exc:
        print()
        print(
            "NETWORK COMPARISON BUILD FAILED"
        )
        print(
            str(
                exc
            )
        )
        sys.exit(
            1
        )
