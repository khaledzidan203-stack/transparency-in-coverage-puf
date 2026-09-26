from __future__ import annotations

from datetime import datetime
from pathlib import Path
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile

import pandas as pd


SCRIPT_VERSION = "1.0.0"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
POWERBI_ROOT = PROJECT_ROOT / "powerbi"
REPORT_DIR = POWERBI_ROOT / "TransparencyInCoverage.Report"
SEMANTIC_DIR = POWERBI_ROOT / "TransparencyInCoverage.SemanticModel"
PAGES_DIR = REPORT_DIR / "definition" / "pages"

MEASURES_TMDL = (
    SEMANTIC_DIR
    / "definition"
    / "tables"
    / "_Measures.tmdl"
)
DIM_STATE_TMDL = (
    SEMANTIC_DIR
    / "definition"
    / "tables"
    / "DimState.tmdl"
)
DIM_ISSUER_TMDL = (
    SEMANTIC_DIR
    / "definition"
    / "tables"
    / "DimIssuer.tmdl"
)
DIM_METRIC_TMDL = (
    SEMANTIC_DIR
    / "definition"
    / "tables"
    / "DimMetric.tmdl"
)

VISUAL_SCHEMA = (
    "https://developer.microsoft.com/json-schemas/fabric/item/report/"
    "definition/visualContainer/2.7.0/schema.json"
)

CANVAS_WIDTH = 1920
CANVAS_HEIGHT = 1080

# ---------------------------------------------------------------------
# Storytelling palette — colors encode analytical role / magnitude.
# They are not quality ratings.
# ---------------------------------------------------------------------

COLORS = {
    "canvas": "#F5F7FA",
    "surface": "#FFFFFF",
    "surface_alt": "#F9FAFB",
    "border": "#E7ECF2",
    "ink": "#172033",
    "secondary": "#667085",
    "muted": "#98A2B3",
    "navy": "#17324D",
    "slate": "#344054",
    "blue": "#4F6BED",
    "blue_dark": "#425CC7",
    "blue_soft": "#EEF3FF",
    "blue_pale": "#E5ECFF",
    "teal": "#16A6A1",
    "teal_dark": "#0E7774",
    "teal_soft": "#EAF8F6",
    "amber": "#C8872C",
    "amber_dark": "#9C641A",
    "amber_soft": "#FFF4E3",
    "red": "#C64B4B",
    "red_dark": "#9F2F2F",
    "red_soft": "#FCE8E8",
    "white": "#FFFFFF",
}

FONT_REGULAR = (
    "'Segoe UI', wf_segoe-ui_normal, helvetica, arial, sans-serif"
)
FONT_SEMIBOLD = (
    "'Segoe UI Semibold', wf_segoe-ui_semibold, "
    "helvetica, arial, sans-serif"
)

EXPECTED_PAGES = [
    "00 INDEX",
    "01 Executive Overview",
    "02 Network Comparison",
    "03 Denials",
    "04 Denial Reasons",
    "05 Appeals",
    "06 Resubmissions",
    "07 Enrollment",
    "08 Issuer & State Explorer",
    "09 Data Availability",
    "10 Data Quality & Methodology",
]

REQUIRED_MEASURES = {
    "Issuer Count",
    "Plan Count",
    "Issuer Comparable Claims Received - Total",
    "Issuer Comparable Claims Denied - Total",
    "Issuer Comparable Overall Denial Rate",
    "Issuer In-Network Denial Rate",
    "Issuer Out-of-Network Denial Rate",
    "Issuer Network Denial Rate Gap",
    "Internal Appeal Overturn Rate",
    "External Appeal Overturn Rate",
    "Internal Appeals Filed",
    "External Appeals Filed",
    "Comparable Denial Reason Composition %",
    "Open DQ Exception Count",
    "Known Source Exception Row Count",
}

REQUIRED_COLUMNS = {
    "DimState": {
        "StateCode",
    },
    "DimIssuer": {
        "IssuerName",
    },
    "DimMetric": {
        "MetricDisplayName",
    },
}

UNSUPPORTED_PAGE_ROOT_PROPERTIES = {
    "horizontalAlignment",
    "verticalAlignment",
}

EXPECTED_BOUND_VISUALS = 20

EXPECTED_VISUAL_NAMES = {
    # Header / navigation / filters.
    "p01_context",
    "p01_page_title",
    "p01_index_surface",
    "p01_index_label",
    "p01_index_nav",
    "p01_state_slicer",
    "p01_issuer_slicer",
    "p01_dq_global",
    "p01_band_dq_global",

    # Executive KPI strip.
    "p01_kpi_issuers",
    "p01_band_kpi_issuers",
    "p01_kpi_plans",
    "p01_band_kpi_plans",
    "p01_kpi_claims",
    "p01_band_kpi_claims",
    "p01_kpi_denial_rate",
    "p01_band_kpi_denial_rate",
    "p01_kpi_appeal_rate",
    "p01_band_kpi_appeal_rate",

    # Network section.
    "p01_band_network_section",
    "p01_network_in",
    "p01_band_network_in",
    "p01_network_out",
    "p01_band_network_out",
    "p01_network_gap",
    "p01_band_network_gap",

    # Comparable claims section.
    "p01_band_claims_section",
    "p01_claims_received",
    "p01_band_claims_received",
    "p01_claims_denied",
    "p01_band_claims_denied",
    "p01_claims_rate",
    "p01_band_claims_rate",

    # Denial reasons.
    "p01_band_denial_reasons",
    "p01_denial_reason_chart",

    # Appeals.
    "p01_band_appeals",
    "p01_appeal_internal_rate",
    "p01_band_appeal_internal_rate",
    "p01_appeal_external_rate",
    "p01_band_appeal_external_rate",
    "p01_appeal_internal_filed",
    "p01_band_appeal_internal_filed",
    "p01_appeal_external_filed",
    "p01_band_appeal_external_filed",

    # Governance.
    "p01_band_governance",
    "p01_known_source",
    "p01_band_known_source",
    "p01_governance_note",
}


# =====================================================================
# Windows / paths
# =====================================================================

def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )

            with winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                key_path,
            ) as key:
                raw_value, _ = winreg.QueryValueEx(
                    key,
                    "Desktop",
                )

            return Path(
                os.path.expandvars(raw_value)
            ).resolve()

        except Exception:
            pass

    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = (
    get_windows_desktop()
    / "Transparency_PUF_Review"
)

BACKUP_FILE = (
    REVIEW_DIR
    / "25_PRE_EXECUTIVE_STORYTELLING_BACKUP.zip"
)

REPORT_FILE = (
    REVIEW_DIR
    / "25_EXECUTIVE_STORYTELLING_BUILD_REPORT.xlsx"
)


def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False

    result = subprocess.run(
        [
            "tasklist",
            "/FI",
            "IMAGENAME eq PBIDesktop.exe",
        ],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(
            subprocess,
            "CREATE_NO_WINDOW",
            0,
        ),
    )

    return "PBIDesktop.exe" in result.stdout


def read_json(path: Path):
    return json.loads(
        path.read_text(
            encoding="utf-8-sig",
        )
    )


def write_json(
    path: Path,
    payload,
) -> None:
    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    path.write_text(
        json.dumps(
            payload,
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )


def sha256(path: Path) -> str:
    h = hashlib.sha256()

    with path.open("rb") as f:
        for chunk in iter(
            lambda: f.read(1024 * 1024),
            b"",
        ):
            h.update(chunk)

    return h.hexdigest()


# =====================================================================
# PBIR expression helpers
# =====================================================================

def expr_bool(value: bool) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": "true" if value else "false"
            }
        }
    }


def expr_string(value: str) -> dict:
    escaped = value.replace(
        "'",
        "''",
    )

    return {
        "expr": {
            "Literal": {
                "Value": f"'{escaped}'"
            }
        }
    }


def expr_double(
    value: float | int,
) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": f"{value}D"
            }
        }
    }


def expr_int(
    value: int,
) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": f"{value}L"
            }
        }
    }


def solid_color(
    color: str,
) -> dict:
    return {
        "solid": {
            "color": expr_string(
                color
            )
        }
    }


def measure_expr(
    measure_name: str,
) -> dict:
    return {
        "Measure": {
            "Expression": {
                "SourceRef": {
                    "Entity": "_Measures"
                }
            },
            "Property": measure_name,
        }
    }


def literal_expr(
    value: float | int,
) -> dict:
    return {
        "Literal": {
            "Value": f"{value}D"
        }
    }


# =====================================================================
# Semantic inventory
# =====================================================================

def parse_tmdl_measures(
    path: Path,
) -> set[str]:
    text = path.read_text(
        encoding="utf-8-sig",
    )

    measures: set[str] = set()

    for line in text.splitlines():
        stripped = line.lstrip()

        if not stripped.startswith(
            "measure "
        ):
            continue

        lhs = stripped.split(
            "=",
            1,
        )[0].strip()

        raw = lhs[
            len("measure "):
        ].strip()

        if (
            raw.startswith("'")
            and raw.endswith("'")
        ):
            raw = (
                raw[1:-1]
                .replace(
                    "''",
                    "'",
                )
            )

        measures.add(
            raw
        )

    return measures


def parse_tmdl_columns(
    path: Path,
) -> set[str]:
    text = path.read_text(
        encoding="utf-8-sig",
    )

    columns: set[str] = set()

    for line in text.splitlines():
        stripped = line.lstrip()

        if not stripped.startswith(
            "column "
        ):
            continue

        raw = stripped[
            len("column "):
        ].strip()

        if raw.startswith("'"):
            match = re.match(
                r"'((?:''|[^'])*)'",
                raw,
            )

            if match:
                columns.add(
                    match.group(1)
                    .replace(
                        "''",
                        "'",
                    )
                )

        else:
            columns.add(
                raw.split(
                    "\t",
                    1,
                )[0].strip()
            )

    return columns


# =====================================================================
# Report discovery / backup
# =====================================================================

def find_pages() -> dict[
    str,
    tuple[Path, dict]
]:
    pages = {}

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        display_name = payload.get(
            "displayName"
        )

        if display_name:
            pages[
                display_name
            ] = (
                page_json.parent,
                payload,
            )

    return pages


def create_backup() -> None:
    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if BACKUP_FILE.exists():
        BACKUP_FILE.unlink()

    with zipfile.ZipFile(
        BACKUP_FILE,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zf:
        for path in sorted(
            p
            for p in REPORT_DIR.rglob("*")
            if p.is_file()
        ):
            archive_name = (
                Path(REPORT_DIR.name)
                / path.relative_to(
                    REPORT_DIR
                )
            )

            zf.write(
                path,
                str(
                    archive_name
                ),
            )


def file_hashes(
    root: Path,
) -> dict[str, str]:
    hashes = {}

    for path in root.rglob("*"):
        if not path.is_file():
            continue

        hashes[
            str(
                path.relative_to(
                    REPORT_DIR
                )
            )
        ] = sha256(
            path
        )

    return hashes


def prepare_target_page(
    page_dir: Path,
) -> None:
    page_json = (
        page_dir
        / "page.json"
    )

    payload = read_json(
        page_json
    )

    for prop in (
        UNSUPPORTED_PAGE_ROOT_PROPERTIES
    ):
        payload.pop(
            prop,
            None,
        )

    payload[
        "width"
    ] = CANVAS_WIDTH

    payload[
        "height"
    ] = CANVAS_HEIGHT

    payload[
        "displayOption"
    ] = "FitToPage"

    objects = payload.setdefault(
        "objects",
        {},
    )

    objects[
        "background"
    ] = [
        {
            "properties": {
                "color": solid_color(
                    COLORS["canvas"]
                ),
                "transparency": expr_double(
                    0
                ),
            }
        }
    ]

    write_json(
        page_json,
        payload,
    )


# =====================================================================
# Text / container helpers
# =====================================================================

def text_run(
    value: str,
    *,
    size: float,
    color: str,
    semibold: bool = False,
) -> dict:
    return {
        "value": value,
        "textStyle": {
            "fontFamily": (
                FONT_SEMIBOLD
                if semibold
                else FONT_REGULAR
            ),
            "fontSize": (
                f"{size}pt"
            ),
            "color": color,
        },
    }


def paragraph(
    *runs: dict,
    align: str = "center",
) -> dict:
    return {
        "textRuns": list(
            runs
        ),
        "horizontalTextAlignment": (
            align
        ),
    }


def container_objects(
    *,
    background: str | None = None,
    border: str | None = None,
    radius: int = 8,
    show_title: bool = False,
    title: str = "",
    title_size: float = 10,
    title_color: str = COLORS["ink"],
) -> dict:
    return {
        "title": [
            {
                "properties": {
                    "show": expr_bool(
                        show_title
                    ),
                    **(
                        {
                            "text": expr_string(
                                title
                            ),
                            "alignment": expr_string(
                                "center"
                            ),
                            "fontColor": solid_color(
                                title_color
                            ),
                            "fontFamily": expr_string(
                                FONT_SEMIBOLD
                            ),
                            "fontSize": expr_double(
                                title_size
                            ),
                            "titleWrap": expr_bool(
                                True
                            ),
                        }
                        if show_title
                        else {}
                    ),
                }
            }
        ],
        "background": [
            {
                "properties": {
                    "show": expr_bool(
                        background is not None
                    ),
                    **(
                        {
                            "color": solid_color(
                                background
                            ),
                            "transparency": expr_double(
                                0
                            ),
                        }
                        if background is not None
                        else {}
                    ),
                }
            }
        ],
        "border": [
            {
                "properties": {
                    "show": expr_bool(
                        border is not None
                    ),
                    **(
                        {
                            "color": solid_color(
                                border
                            ),
                            "width": expr_double(
                                1
                            ),
                            "radius": expr_double(
                                radius
                            ),
                        }
                        if border is not None
                        else {}
                    ),
                }
            }
        ],
        "dropShadow": [
            {
                "properties": {
                    "show": expr_bool(
                        False
                    )
                }
            }
        ],
        "visualHeader": [
            {
                "properties": {
                    "show": expr_bool(
                        False
                    )
                }
            }
        ],
    }


def make_textbox(
    *,
    name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    paragraphs: list[dict],
    background: str | None = None,
    border: str | None = None,
    radius: int = 8,
) -> dict:
    return {
        "$schema": VISUAL_SCHEMA,
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
        },
        "visual": {
            "visualType": "textbox",
            "objects": {
                "general": [
                    {
                        "properties": {
                            "paragraphs": paragraphs
                        }
                    }
                ]
            },
            "visualContainerObjects": (
                container_objects(
                    background=background,
                    border=border,
                    radius=radius,
                )
            ),
            "drillFilterOtherVisuals": True,
        },
    }


def make_title_band(
    *,
    name: str,
    title: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    background: str,
    subtitle: str | None = None,
    title_color: str = COLORS["white"],
    subtitle_color: str = COLORS["white"],
) -> dict:
    paragraphs = [
        paragraph(
            text_run(
                title,
                size=10.0,
                color=title_color,
                semibold=True,
            )
        )
    ]

    if subtitle:
        paragraphs.append(
            paragraph(
                text_run(
                    subtitle,
                    size=7.7,
                    color=subtitle_color,
                )
            )
        )

    return make_textbox(
        name=name,
        x=x,
        y=y,
        width=width,
        height=height,
        z=z,
        paragraphs=paragraphs,
        background=background,
        border=background,
        radius=8,
    )


# =====================================================================
# Query binding helpers
# =====================================================================

def measure_field(
    measure_name: str,
) -> dict:
    return {
        "Measure": {
            "Expression": {
                "SourceRef": {
                    "Entity": "_Measures"
                }
            },
            "Property": measure_name,
        }
    }


def column_field(
    entity: str,
    column_name: str,
) -> dict:
    return {
        "Column": {
            "Expression": {
                "SourceRef": {
                    "Entity": entity
                }
            },
            "Property": column_name,
        }
    }


def measure_projection(
    measure_name: str,
) -> dict:
    return {
        "field": measure_field(
            measure_name
        ),
        "queryRef": (
            f"_Measures.{measure_name}"
        ),
        "nativeQueryRef": measure_name,
    }


def column_projection(
    entity: str,
    column_name: str,
    *,
    active: bool = False,
) -> dict:
    result = {
        "field": column_field(
            entity,
            column_name,
        ),
        "queryRef": (
            f"{entity}.{column_name}"
        ),
        "nativeQueryRef": column_name,
    }

    if active:
        result[
            "active"
        ] = True

    return result


# =====================================================================
# Navigation / slicer / card builders
# =====================================================================

def make_index_surface(
    *,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
) -> dict:
    return make_textbox(
        name="p01_index_surface",
        x=x,
        y=y,
        width=width,
        height=height,
        z=z,
        paragraphs=[
            paragraph(
                text_run(
                    " ",
                    size=6,
                    color=COLORS["blue"],
                )
            )
        ],
        background=COLORS["blue"],
        border=COLORS["blue_dark"],
        radius=9,
    )


def make_index_label(
    *,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
) -> dict:
    return make_textbox(
        name="p01_index_label",
        x=x,
        y=y,
        width=width,
        height=height,
        z=z,
        paragraphs=[
            paragraph(
                text_run(
                    "INDEX",
                    size=11,
                    color=COLORS["white"],
                    semibold=True,
                )
            )
        ],
        background=None,
        border=None,
        radius=0,
    )


def make_index_nav(
    *,
    target_page_name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
) -> dict:
    return {
        "$schema": VISUAL_SCHEMA,
        "name": "p01_index_nav",
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": 0,
        },
        "visual": {
            "visualType": "actionButton",
            "objects": {
                "icon": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "outline": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "fill": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": {
                "title": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "background": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "border": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "visualHeader": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "visualLink": [
                    {
                        "properties": {
                            "show": expr_bool(
                                True
                            ),
                            "type": expr_string(
                                "PageNavigation"
                            ),
                            "navigationSection": expr_string(
                                target_page_name
                            ),
                            "showDefaultTooltip": expr_bool(
                                False
                            ),
                            "enabledTooltip": expr_string(
                                "Return to INDEX"
                            ),
                        }
                    }
                ],
            },
            "drillFilterOtherVisuals": True,
        },
    }


def make_dropdown_slicer(
    *,
    name: str,
    entity: str,
    column_name: str,
    title: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    tab_order: int,
) -> dict:
    return {
        "$schema": VISUAL_SCHEMA,
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": tab_order,
        },
        "visual": {
            "visualType": "slicer",
            "query": {
                "queryState": {
                    "Values": {
                        "projections": [
                            column_projection(
                                entity,
                                column_name,
                                active=True,
                            )
                        ]
                    }
                }
            },
            "objects": {
                "data": [
                    {
                        "properties": {
                            "mode": expr_string(
                                "Dropdown"
                            )
                        }
                    }
                ],
                "selection": [
                    {
                        "properties": {
                            "singleSelect": expr_bool(
                                True
                            )
                        }
                    }
                ],
                "header": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": (
                container_objects(
                    background=COLORS["surface"],
                    border=COLORS["border"],
                    radius=9,
                    show_title=True,
                    title=title,
                    title_size=9.3,
                )
            ),
            "drillFilterOtherVisuals": True,
        },
    }


def make_value_card(
    *,
    name: str,
    measure_name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    tab_order: int,
    value_color: str,
    background: str = COLORS["surface"],
    value_size: float = 27,
) -> dict:
    return {
        "$schema": VISUAL_SCHEMA,
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": tab_order,
        },
        "visual": {
            "visualType": "card",
            "query": {
                "queryState": {
                    "Values": {
                        "projections": [
                            measure_projection(
                                measure_name
                            )
                        ]
                    }
                },
                "sortDefinition": {
                    "sort": [
                        {
                            "field": measure_field(
                                measure_name
                            ),
                            "direction": "Descending",
                        }
                    ],
                    "isDefaultSort": True,
                },
            },
            "objects": {
                "labels": [
                    {
                        "properties": {
                            "color": solid_color(
                                value_color
                            ),
                            "fontSize": expr_double(
                                value_size
                            ),
                            "bold": expr_bool(
                                True
                            ),
                            "fontFamily": expr_string(
                                FONT_SEMIBOLD
                            ),
                        }
                    }
                ],
                "categoryLabels": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": (
                container_objects(
                    background=background,
                    border=COLORS["border"],
                    radius=9,
                    show_title=False,
                )
            ),
            "drillFilterOtherVisuals": True,
        },
    }


# =====================================================================
# Denial-reason chart with magnitude storytelling
# =====================================================================

def denial_reason_color_expression() -> dict:
    # Share magnitude only:
    # <5% light blue, 5–15% blue, 15–25% deep blue, >=25% amber.
    measure_name = (
        "Comparable Denial Reason Composition %"
    )

    def comparison(
        kind: int,
        threshold: float,
    ) -> dict:
        return {
            "Comparison": {
                "ComparisonKind": kind,
                "Left": measure_expr(
                    measure_name
                ),
                "Right": literal_expr(
                    threshold
                ),
            }
        }

    return {
        "Conditional": {
            "Cases": [
                {
                    "Condition": comparison(
                        2,
                        0.25,
                    ),
                    "Value": {
                        "Literal": {
                            "Value": (
                                f"'{COLORS['amber']}'"
                            )
                        }
                    },
                },
                {
                    "Condition": comparison(
                        2,
                        0.15,
                    ),
                    "Value": {
                        "Literal": {
                            "Value": (
                                f"'{COLORS['blue_dark']}'"
                            )
                        }
                    },
                },
                {
                    "Condition": comparison(
                        2,
                        0.05,
                    ),
                    "Value": {
                        "Literal": {
                            "Value": (
                                f"'{COLORS['blue']}'"
                            )
                        }
                    },
                },
                {
                    "Condition": comparison(
                        2,
                        0,
                    ),
                    "Value": {
                        "Literal": {
                            "Value": (
                                f"'{COLORS['blue_pale']}'"
                            )
                        }
                    },
                },
            ]
        }
    }


def make_denial_reason_chart(
    *,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    tab_order: int,
) -> dict:
    measure_name = (
        "Comparable Denial Reason Composition %"
    )

    dynamic_fill = {
        "solid": {
            "color": {
                "expr": (
                    denial_reason_color_expression()
                )
            }
        }
    }

    return {
        "$schema": VISUAL_SCHEMA,
        "name": "p01_denial_reason_chart",
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": tab_order,
        },
        "visual": {
            "visualType": "barChart",
            "query": {
                "queryState": {
                    "Category": {
                        "projections": [
                            column_projection(
                                "DimMetric",
                                "MetricDisplayName",
                                active=True,
                            )
                        ]
                    },
                    "Y": {
                        "projections": [
                            measure_projection(
                                measure_name
                            )
                        ]
                    },
                },
                "sortDefinition": {
                    "sort": [
                        {
                            "field": measure_field(
                                measure_name
                            ),
                            "direction": "Descending",
                        }
                    ],
                    "isDefaultSort": True,
                },
            },
            "objects": {
                "legend": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "categoryAxis": [
                    {
                        "properties": {
                            "showAxisTitle": expr_bool(
                                False
                            ),
                            "labelColor": solid_color(
                                COLORS["secondary"]
                            ),
                            "fontFamily": expr_string(
                                FONT_REGULAR
                            ),
                            "fontSize": expr_double(
                                8.1
                            ),
                            "concatenateLabels": expr_bool(
                                False
                            ),
                            "innerPadding": expr_int(
                                3
                            ),
                            "maxMarginFactor": expr_int(
                                45
                            ),
                        }
                    }
                ],
                "valueAxis": [
                    {
                        "properties": {
                            "show": expr_bool(
                                True
                            ),
                            "showAxisTitle": expr_bool(
                                False
                            ),
                            "labelColor": solid_color(
                                COLORS["muted"]
                            ),
                            "fontFamily": expr_string(
                                FONT_REGULAR
                            ),
                            "fontSize": expr_double(
                                8
                            ),
                            "gridlineShow": expr_bool(
                                False
                            ),
                        }
                    }
                ],
                "dataPoint": [
                    {
                        "properties": {
                            "fill": solid_color(
                                COLORS["blue"]
                            )
                        }
                    },
                    {
                        "properties": {
                            "fill": dynamic_fill
                        },
                        "selector": {
                            "data": [
                                {
                                    "dataViewWildcard": {
                                        "matchingOption": 1
                                    }
                                }
                            ]
                        },
                    },
                ],
                "labels": [
                    {
                        "properties": {
                            "show": expr_bool(
                                True
                            ),
                            "color": solid_color(
                                COLORS["ink"]
                            ),
                            "fontSize": expr_double(
                                8.3
                            ),
                            "fontFamily": expr_string(
                                FONT_SEMIBOLD
                            ),
                            "labelOverflow": expr_bool(
                                False
                            ),
                        }
                    }
                ],
                "zoom": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": (
                container_objects(
                    background=COLORS["surface"],
                    border=COLORS["border"],
                    radius=9,
                    show_title=False,
                )
            ),
            "drillFilterOtherVisuals": True,
        },
    }


# =====================================================================
# Visual IO
# =====================================================================

def write_visual(
    page_dir: Path,
    visual: dict,
) -> Path:
    target = (
        page_dir
        / "visuals"
        / visual["name"]
        / "visual.json"
    )

    write_json(
        target,
        visual,
    )

    return target


# =====================================================================
# Page 01 comprehensive storytelling rebuild
# =====================================================================

def build_executive_storytelling(
    page_dir: Path,
    index_page_name: str,
) -> None:
    visuals_dir = (
        page_dir
        / "visuals"
    )

    if visuals_dir.exists():
        shutil.rmtree(
            visuals_dir
        )

    visuals_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    # ---------------------------------------------------------------
    # Header / filters
    # ---------------------------------------------------------------

    write_visual(
        page_dir,
        make_textbox(
            name="p01_context",
            x=52,
            y=28,
            width=290,
            height=44,
            z=100,
            paragraphs=[
                paragraph(
                    text_run(
                        "2026 PUF  •  EXPERIENCE YEAR 2024",
                        size=9,
                        color=COLORS["blue"],
                        semibold=True,
                    )
                )
            ],
            background=COLORS["blue_soft"],
            border=COLORS["blue_pale"],
            radius=8,
        ),
    )

    write_visual(
        page_dir,
        make_textbox(
            name="p01_page_title",
            x=390,
            y=16,
            width=1140,
            height=76,
            z=110,
            paragraphs=[
                paragraph(
                    text_run(
                        "Executive Overview",
                        size=28,
                        color=COLORS["navy"],
                        semibold=True,
                    )
                ),
                paragraph(
                    text_run(
                        (
                            "Market scale → claims exposure → network disparity → "
                            "denial drivers → appeals → governance"
                        ),
                        size=10,
                        color=COLORS["secondary"],
                    )
                ),
            ],
            background=COLORS["blue_soft"],
            border=COLORS["blue_pale"],
            radius=9,
        ),
    )

    index_x = 1648
    index_y = 24
    index_w = 220
    index_h = 56
    label_h = 22
    label_y = (
        index_y
        + (
            index_h
            - label_h
        )
        / 2
    )

    write_visual(
        page_dir,
        make_index_surface(
            x=index_x,
            y=index_y,
            width=index_w,
            height=index_h,
            z=130,
        ),
    )

    write_visual(
        page_dir,
        make_index_label(
            x=index_x,
            y=label_y,
            width=index_w,
            height=label_h,
            z=140,
        ),
    )

    write_visual(
        page_dir,
        make_index_nav(
            target_page_name=index_page_name,
            x=index_x,
            y=index_y,
            width=index_w,
            height=index_h,
            z=5000,
        ),
    )

    write_visual(
        page_dir,
        make_dropdown_slicer(
            name="p01_state_slicer",
            entity="DimState",
            column_name="StateCode",
            title="State",
            x=548,
            y=104,
            width=300,
            height=62,
            z=220,
            tab_order=1,
        ),
    )

    write_visual(
        page_dir,
        make_dropdown_slicer(
            name="p01_issuer_slicer",
            entity="DimIssuer",
            column_name="IssuerName",
            title="Issuer",
            x=864,
            y=104,
            width=508,
            height=62,
            z=220,
            tab_order=2,
        ),
    )

    write_visual(
        page_dir,
        make_value_card(
            name="p01_dq_global",
            measure_name="Open DQ Exception Count",
            x=1388,
            y=104,
            width=480,
            height=62,
            z=230,
            tab_order=3,
            value_color=COLORS["amber"],
            background=COLORS["surface"],
            value_size=18,
        ),
    )

    write_visual(
        page_dir,
        make_title_band(
            name="p01_band_dq_global",
            title="GLOBAL OPEN DQ EXCEPTIONS",
            x=1388,
            y=104,
            width=480,
            height=28,
            z=1600,
            background=COLORS["amber"],
        ),
    )

    # ---------------------------------------------------------------
    # Executive KPI strip
    # ---------------------------------------------------------------

    margin = 52
    gap = 16
    total_width = (
        CANVAS_WIDTH
        - 2 * margin
    )
    card_width = (
        total_width
        - 4 * gap
    ) / 5

    kpis = [
        (
            "p01_kpi_issuers",
            "p01_band_kpi_issuers",
            "Issuer Count",
            "ISSUERS",
            COLORS["ink"],
            COLORS["surface"],
            COLORS["slate"],
        ),
        (
            "p01_kpi_plans",
            "p01_band_kpi_plans",
            "Plan Count",
            "PLANS",
            COLORS["ink"],
            COLORS["surface"],
            COLORS["slate"],
        ),
        (
            "p01_kpi_claims",
            "p01_band_kpi_claims",
            "Issuer Comparable Claims Received - Total",
            "COMPARABLE CLAIMS",
            COLORS["navy"],
            COLORS["blue_soft"],
            COLORS["navy"],
        ),
        (
            "p01_kpi_denial_rate",
            "p01_band_kpi_denial_rate",
            "Issuer Comparable Overall Denial Rate",
            "COMPARABLE DENIAL RATE",
            COLORS["amber_dark"],
            COLORS["amber_soft"],
            COLORS["amber"],
        ),
        (
            "p01_kpi_appeal_rate",
            "p01_band_kpi_appeal_rate",
            "Internal Appeal Overturn Rate",
            "INTERNAL APPEAL OVERTURN",
            COLORS["teal_dark"],
            COLORS["teal_soft"],
            COLORS["teal_dark"],
        ),
    ]

    for idx, (
        card_name,
        band_name,
        measure_name,
        band_title,
        value_color,
        background,
        band_color,
    ) in enumerate(
        kpis
    ):
        x = (
            margin
            + idx
            * (
                card_width
                + gap
            )
        )

        write_visual(
            page_dir,
            make_value_card(
                name=card_name,
                measure_name=measure_name,
                x=x,
                y=182,
                width=card_width,
                height=132,
                z=400 + idx,
                tab_order=10 + idx,
                value_color=value_color,
                background=background,
                value_size=27,
            ),
        )

        write_visual(
            page_dir,
            make_title_band(
                name=band_name,
                title=band_title,
                x=x,
                y=182,
                width=card_width,
                height=34,
                z=1700 + idx,
                background=band_color,
            ),
        )

    # ---------------------------------------------------------------
    # Mid-story: network disparity + comparable claims
    # ---------------------------------------------------------------

    left_x = 52
    right_x = 968
    section_width = 900
    inner_gap = 16
    inner_card_width = (
        section_width
        - 2 * inner_gap
    ) / 3

    write_visual(
        page_dir,
        make_title_band(
            name="p01_band_network_section",
            title="NETWORK DISPARITY — WHERE DOES DENIAL PRESSURE CHANGE?",
            subtitle=(
                "Read in-network vs out-of-network first, then the gap"
            ),
            x=left_x,
            y=334,
            width=section_width,
            height=44,
            z=1800,
            background=COLORS["navy"],
        ),
    )

    network_specs = [
        (
            "p01_network_in",
            "p01_band_network_in",
            "Issuer In-Network Denial Rate",
            "IN-NETWORK",
            COLORS["navy"],
            COLORS["surface"],
            COLORS["navy"],
        ),
        (
            "p01_network_out",
            "p01_band_network_out",
            "Issuer Out-of-Network Denial Rate",
            "OUT-OF-NETWORK",
            COLORS["blue"],
            COLORS["blue_soft"],
            COLORS["blue"],
        ),
        (
            "p01_network_gap",
            "p01_band_network_gap",
            "Issuer Network Denial Rate Gap",
            "NETWORK GAP",
            COLORS["amber_dark"],
            COLORS["amber_soft"],
            COLORS["amber"],
        ),
    ]

    for idx, (
        card_name,
        band_name,
        measure_name,
        band_title,
        value_color,
        background,
        band_color,
    ) in enumerate(
        network_specs
    ):
        x = (
            left_x
            + idx
            * (
                inner_card_width
                + inner_gap
            )
        )

        write_visual(
            page_dir,
            make_value_card(
                name=card_name,
                measure_name=measure_name,
                x=x,
                y=390,
                width=inner_card_width,
                height=166,
                z=600 + idx,
                tab_order=30 + idx,
                value_color=value_color,
                background=background,
                value_size=29,
            ),
        )

        write_visual(
            page_dir,
            make_title_band(
                name=band_name,
                title=band_title,
                x=x,
                y=390,
                width=inner_card_width,
                height=30,
                z=1900 + idx,
                background=band_color,
            ),
        )

    write_visual(
        page_dir,
        make_title_band(
            name="p01_band_claims_section",
            title="COMPARABLE CLAIMS — HOW LARGE IS THE DENIAL EXPOSURE?",
            subtitle=(
                "Received and denied claims from the same eligible issuer population"
            ),
            x=right_x,
            y=334,
            width=section_width,
            height=44,
            z=1800,
            background=COLORS["slate"],
        ),
    )

    claims_specs = [
        (
            "p01_claims_received",
            "p01_band_claims_received",
            "Issuer Comparable Claims Received - Total",
            "CLAIMS RECEIVED",
            COLORS["navy"],
            COLORS["surface"],
            COLORS["navy"],
        ),
        (
            "p01_claims_denied",
            "p01_band_claims_denied",
            "Issuer Comparable Claims Denied - Total",
            "CLAIMS DENIED",
            COLORS["amber_dark"],
            COLORS["amber_soft"],
            COLORS["amber"],
        ),
        (
            "p01_claims_rate",
            "p01_band_claims_rate",
            "Issuer Comparable Overall Denial Rate",
            "OVERALL DENIAL RATE",
            COLORS["blue"],
            COLORS["blue_soft"],
            COLORS["blue"],
        ),
    ]

    for idx, (
        card_name,
        band_name,
        measure_name,
        band_title,
        value_color,
        background,
        band_color,
    ) in enumerate(
        claims_specs
    ):
        x = (
            right_x
            + idx
            * (
                inner_card_width
                + inner_gap
            )
        )

        write_visual(
            page_dir,
            make_value_card(
                name=card_name,
                measure_name=measure_name,
                x=x,
                y=390,
                width=inner_card_width,
                height=166,
                z=600 + idx,
                tab_order=40 + idx,
                value_color=value_color,
                background=background,
                value_size=28,
            ),
        )

        write_visual(
            page_dir,
            make_title_band(
                name=band_name,
                title=band_title,
                x=x,
                y=390,
                width=inner_card_width,
                height=30,
                z=1900 + idx,
                background=band_color,
            ),
        )

    # ---------------------------------------------------------------
    # Lower story: denial drivers + appeals outcome
    # ---------------------------------------------------------------

    write_visual(
        page_dir,
        make_title_band(
            name="p01_band_denial_reasons",
            title="DENIAL DRIVERS — WHAT ACCOUNTS FOR THE MIX?",
            subtitle=(
                "Color encodes share magnitude only: light blue → blue → deep blue → amber"
            ),
            x=52,
            y=580,
            width=1120,
            height=46,
            z=1800,
            background=COLORS["blue_dark"],
        ),
    )

    write_visual(
        page_dir,
        make_denial_reason_chart(
            x=52,
            y=626,
            width=1120,
            height=290,
            z=800,
            tab_order=60,
        ),
    )

    write_visual(
        page_dir,
        make_title_band(
            name="p01_band_appeals",
            title="POST-DENIAL JOURNEY — WHAT HAPPENS AFTER FILING?",
            subtitle=(
                "Filed volume and overturn rate provide the appeals context"
            ),
            x=1188,
            y=580,
            width=680,
            height=46,
            z=1800,
            background=COLORS["teal_dark"],
        ),
    )

    appeal_x = 1188
    appeal_y = 638
    appeal_w = 680
    appeal_gap = 16
    appeal_card_w = (
        appeal_w
        - appeal_gap
    ) / 2
    appeal_card_h = 130

    appeals = [
        (
            "p01_appeal_internal_rate",
            "p01_band_appeal_internal_rate",
            "Internal Appeal Overturn Rate",
            "INTERNAL OVERTURN",
            COLORS["teal_dark"],
            COLORS["teal_soft"],
            COLORS["teal_dark"],
            appeal_x,
            appeal_y,
        ),
        (
            "p01_appeal_external_rate",
            "p01_band_appeal_external_rate",
            "External Appeal Overturn Rate",
            "EXTERNAL OVERTURN",
            COLORS["teal"],
            COLORS["teal_soft"],
            COLORS["teal"],
            (
                appeal_x
                + appeal_card_w
                + appeal_gap
            ),
            appeal_y,
        ),
        (
            "p01_appeal_internal_filed",
            "p01_band_appeal_internal_filed",
            "Internal Appeals Filed",
            "INTERNAL FILED",
            COLORS["ink"],
            COLORS["surface"],
            COLORS["slate"],
            appeal_x,
            (
                appeal_y
                + appeal_card_h
                + appeal_gap
            ),
        ),
        (
            "p01_appeal_external_filed",
            "p01_band_appeal_external_filed",
            "External Appeals Filed",
            "EXTERNAL FILED",
            COLORS["ink"],
            COLORS["surface"],
            COLORS["slate"],
            (
                appeal_x
                + appeal_card_w
                + appeal_gap
            ),
            (
                appeal_y
                + appeal_card_h
                + appeal_gap
            ),
        ),
    ]

    for idx, (
        card_name,
        band_name,
        measure_name,
        band_title,
        value_color,
        background,
        band_color,
        x,
        y,
    ) in enumerate(
        appeals
    ):
        write_visual(
            page_dir,
            make_value_card(
                name=card_name,
                measure_name=measure_name,
                x=x,
                y=y,
                width=appeal_card_w,
                height=appeal_card_h,
                z=900 + idx,
                tab_order=70 + idx,
                value_color=value_color,
                background=background,
                value_size=23,
            ),
        )

        write_visual(
            page_dir,
            make_title_band(
                name=band_name,
                title=band_title,
                x=x,
                y=y,
                width=appeal_card_w,
                height=28,
                z=2000 + idx,
                background=band_color,
            ),
        )

    # ---------------------------------------------------------------
    # Governance footer — global context, explicitly not slicer-sensitive
    # ---------------------------------------------------------------

    write_visual(
        page_dir,
        make_title_band(
            name="p01_band_governance",
            title="GOVERNANCE CONTEXT — READ THE STORY WITH THE DATA QUALITY",
            x=52,
            y=932,
            width=1816,
            height=38,
            z=1800,
            background=COLORS["amber_dark"],
        ),
    )

    write_visual(
        page_dir,
        make_value_card(
            name="p01_known_source",
            measure_name=(
                "Known Source Exception Row Count"
            ),
            x=52,
            y=982,
            width=260,
            height=72,
            z=700,
            tab_order=90,
            value_color=COLORS["red_dark"],
            background=COLORS["red_soft"],
            value_size=18,
        ),
    )

    write_visual(
        page_dir,
        make_title_band(
            name="p01_band_known_source",
            title="GLOBAL KNOWN-SOURCE ROWS",
            x=52,
            y=982,
            width=260,
            height=26,
            z=2100,
            background=COLORS["red"],
        ),
    )

    write_visual(
        page_dir,
        make_textbox(
            name="p01_governance_note",
            x=328,
            y=992,
            width=1540,
            height=50,
            z=700,
            paragraphs=[
                paragraph(
                    text_run(
                        (
                            "State/Issuer filters drive business KPIs and charts  •  "
                            "Governance cards are global  •  Availability-aware governed KPIs  •  "
                            "Source anomalies remain visible  •  No clipping or imputation  •  "
                            "Colors encode magnitude / analytical role, not issuer quality"
                        ),
                        size=8.7,
                        color=COLORS["secondary"],
                        semibold=True,
                    )
                )
            ],
            background=COLORS["surface_alt"],
            border=COLORS["border"],
            radius=8,
        ),
    )


# =====================================================================
# Validation
# =====================================================================

def walk_json(
    obj,
):
    if isinstance(
        obj,
        dict,
    ):
        yield obj

        for value in obj.values():
            yield from walk_json(
                value
            )

    elif isinstance(
        obj,
        list,
    ):
        for value in obj:
            yield from walk_json(
                value
            )


def validate_json_tree() -> list[dict]:
    failures = []

    for path in sorted(
        (
            REPORT_DIR
            / "definition"
        ).rglob(
            "*.json"
        )
    ):
        try:
            read_json(
                path
            )

        except Exception as exc:
            failures.append(
                {
                    "RelativePath": str(
                        path.relative_to(
                            PROJECT_ROOT
                        )
                    ),
                    "Error": str(
                        exc
                    ),
                }
            )

    return failures


def validate_no_unsupported_page_roots(
    pages: dict,
) -> list[dict]:
    offenders = []

    for display_name, (
        page_dir,
        _,
    ) in pages.items():
        payload = read_json(
            page_dir
            / "page.json"
        )

        for prop in sorted(
            UNSUPPORTED_PAGE_ROOT_PROPERTIES
        ):
            if prop in payload:
                offenders.append(
                    {
                        "Page": display_name,
                        "UnsupportedProperty": prop,
                    }
                )

    return offenders


def validate_no_full_page_visuals(
    pages: dict,
) -> list[dict]:
    offenders = []

    for display_name, (
        page_dir,
        _,
    ) in pages.items():
        visuals_dir = (
            page_dir
            / "visuals"
        )

        if not visuals_dir.exists():
            continue

        for path in visuals_dir.rglob(
            "visual.json"
        ):
            payload = read_json(
                path
            )

            pos = payload.get(
                "position",
                {},
            )

            width = float(
                pos.get(
                    "width",
                    0,
                )
                or 0
            )
            height = float(
                pos.get(
                    "height",
                    0,
                )
                or 0
            )

            if (
                width
                >= CANVAS_WIDTH * 0.94
                and height
                >= CANVAS_HEIGHT * 0.94
            ):
                offenders.append(
                    {
                        "Page": display_name,
                        "VisualName": (
                            payload.get(
                                "name"
                            )
                        ),
                        "Width": width,
                        "Height": height,
                    }
                )

    return offenders


def count_bound_visuals(
    page_dir: Path,
) -> int:
    count = 0

    for visual_json in (
        page_dir
        / "visuals"
    ).rglob(
        "visual.json"
    ):
        payload = read_json(
            visual_json
        )

        if (
            payload.get(
                "visual",
                {},
            ).get(
                "query"
            )
        ):
            count += 1

    return count


def validate_denial_chart_story(
    page_dir: Path,
) -> dict:
    payload = read_json(
        page_dir
        / "visuals"
        / "p01_denial_reason_chart"
        / "visual.json"
    )

    data_points = (
        payload.get(
            "visual",
            {},
        )
        .get(
            "objects",
            {},
        )
        .get(
            "dataPoint",
            [],
        )
    )

    has_conditional = any(
        "Conditional" in node
        for node in walk_json(
            data_points
        )
    )

    has_wildcard = any(
        (
            item.get(
                "selector",
                {},
            )
            .get(
                "data",
                [{}],
            )[0]
            .get(
                "dataViewWildcard",
                {},
            )
            .get(
                "matchingOption"
            )
            == 1
        )
        for item in data_points
        if item.get(
            "selector"
        )
    )

    return {
        "CheckID": "EXEC-001",
        "TestName": (
            "Denial-reason magnitude colors"
        ),
        "Expected": (
            "Conditional + wildcard point formatting"
        ),
        "Actual": (
            f"conditional={has_conditional}; "
            f"wildcard={has_wildcard}"
        ),
        "Status": (
            "PASS"
            if (
                has_conditional
                and has_wildcard
            )
            else "FAIL"
        ),
    }


def validate_title_bands(
    page_dir: Path,
) -> dict:
    bands = []

    for path in (
        page_dir
        / "visuals"
    ).rglob(
        "visual.json"
    ):
        payload = read_json(
            path
        )

        if str(
            payload.get(
                "name",
                "",
            )
        ).startswith(
            "p01_band_"
        ):
            bands.append(
                payload.get(
                    "name"
                )
            )

    expected = 22

    return {
        "CheckID": "EXEC-002",
        "TestName": (
            "Colored title/story bands"
        ),
        "Expected": expected,
        "Actual": len(
            bands
        ),
        "Status": (
            "PASS"
            if len(
                bands
            )
            == expected
            else "FAIL"
        ),
    }


def validate_index_geometry(
    page_dir: Path,
) -> dict:
    surface = read_json(
        page_dir
        / "visuals"
        / "p01_index_surface"
        / "visual.json"
    )

    label = read_json(
        page_dir
        / "visuals"
        / "p01_index_label"
        / "visual.json"
    )

    s = surface[
        "position"
    ]
    l = label[
        "position"
    ]

    expected_y = (
        s["y"]
        + (
            s["height"]
            - l["height"]
        )
        / 2
    )

    passed = (
        abs(
            float(
                l["y"]
            )
            - float(
                expected_y
            )
        )
        < 0.001
        and float(
            l["x"]
        )
        == float(
            s["x"]
        )
        and float(
            l["width"]
        )
        == float(
            s["width"]
        )
    )

    return {
        "CheckID": "EXEC-003",
        "TestName": (
            "INDEX geometric centering"
        ),
        "Expected": "centered",
        "Actual": (
            "centered"
            if passed
            else "not-centered"
        ),
        "Status": (
            "PASS"
            if passed
            else "FAIL"
        ),
    }


def style_report(
    path: Path,
) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import (
        Alignment,
        Font,
        PatternFill,
    )
    from openpyxl.utils import (
        get_column_letter,
    )

    wb = load_workbook(
        path
    )

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if (
            ws.max_row
            and ws.max_column
        ):
            ws.auto_filter.ref = (
                ws.dimensions
            )

        for cell in ws[1]:
            cell.fill = (
                header_fill
            )
            cell.font = (
                header_font
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            max_len = 0

            for row_idx in range(
                1,
                min(
                    ws.max_row,
                    250,
                )
                + 1,
            ):
                value = ws.cell(
                    row=row_idx,
                    column=col_idx,
                ).value

                if value is not None:
                    max_len = max(
                        max_len,
                        len(
                            str(
                                value
                            )
                        ),
                    )

            ws.column_dimensions[
                get_column_letter(
                    col_idx
                )
            ].width = min(
                max(
                    max_len + 2,
                    10,
                ),
                75,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    wb.save(
        path
    )


# =====================================================================
# Main
# =====================================================================

def main() -> int:
    print("=" * 78)
    print(
        "Transparency in Coverage PUF — "
        "01 Executive Overview Storytelling Rebuild"
    )
    print("=" * 78)

    print(
        f"Script version : {SCRIPT_VERSION}"
    )
    print(
        f"Report folder  : {REPORT_DIR}"
    )
    print(
        f"Backup         : {BACKUP_FILE}"
    )
    print(
        f"Review file    : {REPORT_FILE}"
    )

    if power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is running. "
            "Save and close it completely before changing PBIR files."
        )

    if not REPORT_DIR.exists():
        raise FileNotFoundError(
            f"Report folder not found: {REPORT_DIR}"
        )

    for required_file in (
        MEASURES_TMDL,
        DIM_STATE_TMDL,
        DIM_ISSUER_TMDL,
        DIM_METRIC_TMDL,
    ):
        if not required_file.exists():
            raise FileNotFoundError(
                f"Required semantic-model file not found: "
                f"{required_file}"
            )

    pages = find_pages()

    if set(
        pages
    ) != set(
        EXPECTED_PAGES
    ):
        raise RuntimeError(
            "Current report page inventory does not match "
            "the approved 11-page shell. No changes made."
        )

    measures = parse_tmdl_measures(
        MEASURES_TMDL
    )

    missing_measures = sorted(
        REQUIRED_MEASURES
        - measures
    )

    if missing_measures:
        raise RuntimeError(
            "Required governed measures are missing: "
            + ", ".join(
                missing_measures
            )
        )

    installed_columns = {
        "DimState": parse_tmdl_columns(
            DIM_STATE_TMDL
        ),
        "DimIssuer": parse_tmdl_columns(
            DIM_ISSUER_TMDL
        ),
        "DimMetric": parse_tmdl_columns(
            DIM_METRIC_TMDL
        ),
    }

    missing_columns = []

    for table_name, required in (
        REQUIRED_COLUMNS.items()
    ):
        for column_name in sorted(
            required
            - installed_columns[
                table_name
            ]
        ):
            missing_columns.append(
                f"{table_name}.{column_name}"
            )

    if missing_columns:
        raise RuntimeError(
            "Required semantic columns are missing: "
            + ", ".join(
                missing_columns
            )
        )

    unsupported_before = (
        validate_no_unsupported_page_roots(
            pages
        )
    )

    if unsupported_before:
        raise RuntimeError(
            "Unsupported page-root properties are present before build: "
            + "; ".join(
                f"{row['Page']}:{row['UnsupportedProperty']}"
                for row in unsupported_before
            )
            + ". No changes made."
        )

    index_dir, index_page = (
        pages[
            "00 INDEX"
        ]
    )

    executive_dir, _ = (
        pages[
            "01 Executive Overview"
        ]
    )

    network_dir, _ = (
        pages[
            "02 Network Comparison"
        ]
    )

    index_page_name = (
        index_page[
            "name"
        ]
    )

    # Lock all non-target report files, especially the approved page 02.
    all_before = file_hashes(
        REPORT_DIR
    )

    target_prefix = str(
        executive_dir.relative_to(
            REPORT_DIR
        )
    ).replace(
        "\\",
        "/",
    )

    create_backup()

    prepare_target_page(
        executive_dir
    )

    build_executive_storytelling(
        executive_dir,
        index_page_name,
    )

    pages = find_pages()

    all_after = file_hashes(
        REPORT_DIR
    )

    non_target_changes = []

    for rel_path in sorted(
        set(
            all_before
        )
        | set(
            all_after
        )
    ):
        normalized = rel_path.replace(
            "\\",
            "/",
        )

        if normalized.startswith(
            target_prefix
        ):
            continue

        if (
            all_before.get(
                rel_path
            )
            != all_after.get(
                rel_path
            )
        ):
            non_target_changes.append(
                {
                    "RelativePath": rel_path,
                    "BeforeSHA256": all_before.get(
                        rel_path
                    ),
                    "AfterSHA256": all_after.get(
                        rel_path
                    ),
                }
            )

    json_failures = (
        validate_json_tree()
    )

    unsupported_after = (
        validate_no_unsupported_page_roots(
            pages
        )
    )

    full_page_offenders = (
        validate_no_full_page_visuals(
            pages
        )
    )

    visual_files = sorted(
        (
            executive_dir
            / "visuals"
        ).rglob(
            "visual.json"
        )
    )

    actual_names = {
        read_json(
            path
        ).get(
            "name"
        )
        for path in visual_files
    }

    bound_after = (
        count_bound_visuals(
            executive_dir
        )
    )

    validation_rows = [
        {
            "CheckID": "EXEC-000",
            "TestName": (
                "Executive visual inventory"
            ),
            "Expected": len(
                EXPECTED_VISUAL_NAMES
            ),
            "Actual": len(
                actual_names
            ),
            "Status": (
                "PASS"
                if actual_names
                == EXPECTED_VISUAL_NAMES
                else "FAIL"
            ),
        },
        {
            "CheckID": "EXEC-BOUND",
            "TestName": (
                "Executive bound/query visuals"
            ),
            "Expected": EXPECTED_BOUND_VISUALS,
            "Actual": bound_after,
            "Status": (
                "PASS"
                if bound_after
                == EXPECTED_BOUND_VISUALS
                else "FAIL"
            ),
        },
        validate_denial_chart_story(
            executive_dir
        ),
        validate_title_bands(
            executive_dir
        ),
        validate_index_geometry(
            executive_dir
        ),
        {
            "CheckID": "EXEC-LOCK",
            "TestName": (
                "All non-Executive report files preserved"
            ),
            "Expected": 0,
            "Actual": len(
                non_target_changes
            ),
            "Status": (
                "PASS"
                if not non_target_changes
                else "FAIL"
            ),
        },
        {
            "CheckID": "EXEC-SCHEMA",
            "TestName": (
                "Unsupported page-root properties"
            ),
            "Expected": 0,
            "Actual": len(
                unsupported_after
            ),
            "Status": (
                "PASS"
                if not unsupported_after
                else "FAIL"
            ),
        },
        {
            "CheckID": "EXEC-BG",
            "TestName": (
                "Full-page selectable visuals"
            ),
            "Expected": 0,
            "Actual": len(
                full_page_offenders
            ),
            "Status": (
                "PASS"
                if not full_page_offenders
                else "FAIL"
            ),
        },
        {
            "CheckID": "EXEC-JSON",
            "TestName": (
                "PBIR JSON parse failures"
            ),
            "Expected": 0,
            "Actual": len(
                json_failures
            ),
            "Status": (
                "PASS"
                if not json_failures
                else "FAIL"
            ),
        },
    ]

    failures = sum(
        row[
            "Status"
        ] == "FAIL"
        for row in validation_rows
    )

    summary_df = pd.DataFrame(
        [
            (
                "BuildStatus",
                (
                    "PASS"
                    if failures == 0
                    else "FAIL"
                ),
            ),
            (
                "ScriptVersion",
                SCRIPT_VERSION,
            ),
            (
                "BuiltAtLocal",
                datetime.now()
                .isoformat(
                    timespec="seconds"
                ),
            ),
            (
                "PageRebuilt",
                "01 Executive Overview",
            ),
            (
                "StoryFlow",
                (
                    "Market scale → claims exposure → network disparity → "
                    "denial drivers → appeals → governance"
                ),
            ),
            (
                "VisualsAfter",
                len(
                    visual_files
                ),
            ),
            (
                "BoundVisualsAfter",
                bound_after,
            ),
            (
                "ColoredTitleBands",
                22,
            ),
            (
                "DenialReasonChart",
                (
                    "Conditional magnitude colors: light blue / blue / "
                    "deep blue / amber"
                ),
            ),
            (
                "GlobalDQCard",
                (
                    "Open DQ Exception Count; explicitly global"
                ),
            ),
            (
                "BusinessFilters",
                (
                    "State + Issuer"
                ),
            ),
            (
                "NonTargetFilesChanged",
                len(
                    non_target_changes
                ),
            ),
            (
                "FullPageSelectableVisuals",
                len(
                    full_page_offenders
                ),
            ),
            (
                "UnsupportedPageProperties",
                len(
                    unsupported_after
                ),
            ),
            (
                "ValidationFailures",
                failures,
            ),
            (
                "BackupFile",
                str(
                    BACKUP_FILE
                ),
            ),
            (
                "InterpretationRule",
                (
                    "Colors communicate analytical role / magnitude only; "
                    "they are not issuer-quality ratings."
                ),
            ),
        ],
        columns=[
            "Item",
            "Value",
        ],
    )

    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        REPORT_FILE,
        engine="openpyxl",
    ) as writer:
        summary_df.to_excel(
            writer,
            sheet_name="00_Summary",
            index=False,
        )

        pd.DataFrame(
            validation_rows
        ).to_excel(
            writer,
            sheet_name="01_Validation",
            index=False,
        )

        if non_target_changes:
            pd.DataFrame(
                non_target_changes
            ).to_excel(
                writer,
                sheet_name="02_NonTarget_Changes",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No files outside 01 Executive Overview changed. "
                            "Approved page 02 and all other pages were preserved."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="02_NonTarget_Changes",
                index=False,
            )

        if full_page_offenders:
            pd.DataFrame(
                full_page_offenders
            ).to_excel(
                writer,
                sheet_name="03_FullPage_Offenders",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No full-page selectable visuals."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="03_FullPage_Offenders",
                index=False,
            )

        if unsupported_after:
            pd.DataFrame(
                unsupported_after
            ).to_excel(
                writer,
                sheet_name="04_Schema_Offenders",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No unsupported page-root properties."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="04_Schema_Offenders",
                index=False,
            )

        if json_failures:
            pd.DataFrame(
                json_failures
            ).to_excel(
                writer,
                sheet_name="05_JSON_Failures",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No JSON parse failures."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="05_JSON_Failures",
                index=False,
            )

    style_report(
        REPORT_FILE
    )

    print()
    print(
        "01 Executive Overview storytelling rebuild completed."
    )
    print(
        f"Build status                 : "
        f"{'PASS' if failures == 0 else 'FAIL'}"
    )
    print(
        f"Visuals after                : "
        f"{len(visual_files)}"
    )
    print(
        f"Bound/query visuals          : "
        f"{bound_after}"
    )
    print(
        "Colored title bands          : 22"
    )
    print(
        "Business filters             : State + Issuer"
    )
    print(
        "Denial-reason chart          : conditional magnitude colors"
    )
    print(
        f"Non-target files changed     : "
        f"{len(non_target_changes)}"
    )
    print(
        f"Full-page selectable visuals : "
        f"{len(full_page_offenders)}"
    )
    print(
        f"Unsupported page properties  : "
        f"{len(unsupported_after)}"
    )
    print(
        f"Validation failures          : "
        f"{failures}"
    )
    print(
        f"Backup                       : "
        f"{BACKUP_FILE}"
    )
    print(
        f"Review evidence              : "
        f"{REPORT_FILE}"
    )

    if failures:
        print()
        print(
            "Offline validation failed. "
            "Do not open/save the PBIP until the evidence report is reviewed."
        )
        return 2

    print()
    print(
        "Next: open TransparencyInCoverage.pbip and inspect "
        "01 Executive Overview. Verify title bands, State/Issuer filters, "
        "denial-reason story colors, appeals panel, governance footer, "
        "INDEX navigation, and blank-canvas behavior."
    )

    return 0


if __name__ == "__main__":
    try:
        sys.exit(
            main()
        )

    except Exception as exc:
        print()
        print(
            "EXECUTIVE STORYTELLING REBUILD FAILED"
        )
        print(
            str(
                exc
            )
        )
        sys.exit(
            1
        )
