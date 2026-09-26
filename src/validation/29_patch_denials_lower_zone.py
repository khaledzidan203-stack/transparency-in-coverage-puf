from __future__ import annotations

from datetime import datetime
from pathlib import Path
import copy
import json
import os
import subprocess
import sys
import zipfile

import pandas as pd


SCRIPT_VERSION = "1.0.0"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
REPORT_DIR = (
    PROJECT_ROOT
    / "powerbi"
    / "TransparencyInCoverage.Report"
)
PAGES_DIR = REPORT_DIR / "definition" / "pages"

EXPECTED_VISUALS = 39
EXPECTED_BOUND_VISUALS = 16

CANVAS_WIDTH = 1920
CANVAS_HEIGHT = 1080

# Only these lower-third visuals are allowed to move.
ALLOWED_GEOMETRY_TARGETS = {
    "p03_band_market",
    "p03_market_chart",
    "p03_band_offering",
    "p03_offering_chart",
    "p03_band_plan_type",
    "p03_plan_type_chart",
    "p03_band_table",
    "p03_rate_legend",
    "p03_issuer_table",
}

SEGMENT_BAND_HEIGHT = 24.0
SEGMENT_CHART_HEIGHT = 96.0
TABLE_BAND_HEIGHT = 32.0
LEGEND_HEIGHT = 20.0

SEGMENT_TOP_GAP = 10.0
TABLE_TOP_GAP = 12.0
LEGEND_GAP = 2.0
TABLE_GAP = 2.0
TABLE_BOTTOM_GAP = 6.0

MIN_TABLE_HEIGHT = 125.0

COLORS = {
    "border": "#E7ECF2",
    "surface_alt": "#F9FAFB",
    "ink": "#172033",
    "secondary": "#667085",
}

UNSUPPORTED_PAGE_ROOT_PROPERTIES = {
    "horizontalAlignment",
    "verticalAlignment",
}


# =====================================================================
# Paths / environment
# =====================================================================

def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )

            with winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                key_path,
            ) as key:
                raw_value, _ = winreg.QueryValueEx(
                    key,
                    "Desktop",
                )

            return Path(
                os.path.expandvars(raw_value)
            ).resolve()

        except Exception:
            pass

    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = (
    get_windows_desktop()
    / "Transparency_PUF_Review"
)

BACKUP_FILE = (
    REVIEW_DIR
    / "29_PRE_DENIALS_LOWER_ZONE_PATCH_BACKUP.zip"
)

REPORT_FILE = (
    REVIEW_DIR
    / "29_DENIALS_LOWER_ZONE_PATCH_REPORT.xlsx"
)


def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False

    result = subprocess.run(
        [
            "tasklist",
            "/FI",
            "IMAGENAME eq PBIDesktop.exe",
        ],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(
            subprocess,
            "CREATE_NO_WINDOW",
            0,
        ),
    )

    return "PBIDesktop.exe" in result.stdout


def read_json(path: Path):
    return json.loads(
        path.read_text(
            encoding="utf-8-sig",
        )
    )


def write_json(
    path: Path,
    payload,
) -> None:
    path.write_text(
        json.dumps(
            payload,
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )


def expr_double(
    value: float | int,
) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": f"{value}D"
            }
        }
    }


def expr_bool(value: bool) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": "true" if value else "false"
            }
        }
    }


# =====================================================================
# Page / visual helpers
# =====================================================================

def find_denials_page() -> tuple[Path, dict]:
    matches = []

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        if payload.get(
            "displayName"
        ) == "03 Denials":
            matches.append(
                (
                    page_json.parent,
                    payload,
                )
            )

    if len(matches) != 1:
        raise RuntimeError(
            "Expected exactly one '03 Denials' page; "
            f"found {len(matches)}."
        )

    return matches[0]


def visual_path(
    page_dir: Path,
    visual_name: str,
) -> Path:
    return (
        page_dir
        / "visuals"
        / visual_name
        / "visual.json"
    )


def get_visual(
    page_dir: Path,
    visual_name: str,
) -> dict:
    path = visual_path(
        page_dir,
        visual_name,
    )

    if not path.exists():
        raise FileNotFoundError(
            f"Visual not found: {visual_name}"
        )

    return read_json(
        path
    )


def save_visual(
    page_dir: Path,
    visual_name: str,
    payload: dict,
) -> None:
    write_json(
        visual_path(
            page_dir,
            visual_name,
        ),
        payload,
    )


def capture_positions(
    page_dir: Path,
) -> dict[str, dict]:
    result = {}

    for path in (
        page_dir
        / "visuals"
    ).rglob(
        "visual.json"
    ):
        payload = read_json(
            path
        )

        name = payload.get(
            "name"
        )

        if name:
            result[
                name
            ] = copy.deepcopy(
                payload.get(
                    "position",
                    {},
                )
            )

    return result


def count_bound_visuals(
    page_dir: Path,
) -> int:
    count = 0

    for path in (
        page_dir
        / "visuals"
    ).rglob(
        "visual.json"
    ):
        payload = read_json(
            path
        )

        if (
            payload.get(
                "visual",
                {},
            ).get(
                "query"
            )
        ):
            count += 1

    return count


def create_backup() -> None:
    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if BACKUP_FILE.exists():
        BACKUP_FILE.unlink()

    with zipfile.ZipFile(
        BACKUP_FILE,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zf:
        for path in sorted(
            p
            for p in REPORT_DIR.rglob("*")
            if p.is_file()
        ):
            archive_name = (
                Path(REPORT_DIR.name)
                / path.relative_to(
                    REPORT_DIR
                )
            )

            zf.write(
                path,
                str(
                    archive_name
                ),
            )


# =====================================================================
# Patch
# =====================================================================

def set_position(
    payload: dict,
    *,
    y: float | None = None,
    height: float | None = None,
) -> None:
    pos = payload[
        "position"
    ]

    if y is not None:
        pos[
            "y"
        ] = y

    if height is not None:
        pos[
            "height"
        ] = height


def patch_lower_zone(
    page_dir: Path,
) -> list[dict]:
    patch_rows = []

    section = get_visual(
        page_dir,
        "p03_band_segments",
    )

    methodology = get_visual(
        page_dir,
        "p03_methodology_note",
    )

    section_bottom = (
        float(
            section[
                "position"
            ][
                "y"
            ]
        )
        + float(
            section[
                "position"
            ][
                "height"
            ]
        )
    )

    methodology_y = float(
        methodology[
            "position"
        ][
            "y"
        ]
    )

    subband_y = (
        section_bottom
        + SEGMENT_TOP_GAP
    )

    chart_y = (
        subband_y
        + SEGMENT_BAND_HEIGHT
    )

    chart_bottom = (
        chart_y
        + SEGMENT_CHART_HEIGHT
    )

    table_band_y = (
        chart_bottom
        + TABLE_TOP_GAP
    )

    legend_y = (
        table_band_y
        + TABLE_BAND_HEIGHT
        + LEGEND_GAP
    )

    table_y = (
        legend_y
        + LEGEND_HEIGHT
        + TABLE_GAP
    )

    table_height = (
        methodology_y
        - TABLE_BOTTOM_GAP
        - table_y
    )

    if table_height < MIN_TABLE_HEIGHT:
        raise RuntimeError(
            "Not enough room to create a useful issuer investigation table. "
            f"Computed height={table_height:.2f}px; "
            f"minimum={MIN_TABLE_HEIGHT:.2f}px."
        )

    # Segment panels: preserve x/width/z, only compact vertically.
    pairs = [
        (
            "p03_band_market",
            "p03_market_chart",
        ),
        (
            "p03_band_offering",
            "p03_offering_chart",
        ),
        (
            "p03_band_plan_type",
            "p03_plan_type_chart",
        ),
    ]

    for band_name, chart_name in pairs:
        band = get_visual(
            page_dir,
            band_name,
        )
        chart = get_visual(
            page_dir,
            chart_name,
        )

        old_band = copy.deepcopy(
            band[
                "position"
            ]
        )
        old_chart = copy.deepcopy(
            chart[
                "position"
            ]
        )

        set_position(
            band,
            y=subband_y,
            height=SEGMENT_BAND_HEIGHT,
        )

        set_position(
            chart,
            y=chart_y,
            height=SEGMENT_CHART_HEIGHT,
        )

        save_visual(
            page_dir,
            band_name,
            band,
        )
        save_visual(
            page_dir,
            chart_name,
            chart,
        )

        patch_rows.append(
            {
                "Visual": band_name,
                "Action": "Compact segment title band",
                "Before": json.dumps(
                    old_band,
                    sort_keys=True,
                ),
                "After": json.dumps(
                    band[
                        "position"
                    ],
                    sort_keys=True,
                ),
            }
        )

        patch_rows.append(
            {
                "Visual": chart_name,
                "Action": "Compact segment chart to free investigation space",
                "Before": json.dumps(
                    old_chart,
                    sort_keys=True,
                ),
                "After": json.dumps(
                    chart[
                        "position"
                    ],
                    sort_keys=True,
                ),
            }
        )

    # Investigation title.
    table_band = get_visual(
        page_dir,
        "p03_band_table",
    )
    old_table_band = copy.deepcopy(
        table_band[
            "position"
        ]
    )

    set_position(
        table_band,
        y=table_band_y,
        height=TABLE_BAND_HEIGHT,
    )

    save_visual(
        page_dir,
        "p03_band_table",
        table_band,
    )

    patch_rows.append(
        {
            "Visual": "p03_band_table",
            "Action": "Move investigation header upward",
            "Before": json.dumps(
                old_table_band,
                sort_keys=True,
            ),
            "After": json.dumps(
                table_band[
                    "position"
                ],
                sort_keys=True,
            ),
        }
    )

    # Compact rate legend.
    legend = get_visual(
        page_dir,
        "p03_rate_legend",
    )
    old_legend = copy.deepcopy(
        legend[
            "position"
        ]
    )

    set_position(
        legend,
        y=legend_y,
        height=LEGEND_HEIGHT,
    )

    save_visual(
        page_dir,
        "p03_rate_legend",
        legend,
    )

    patch_rows.append(
        {
            "Visual": "p03_rate_legend",
            "Action": "Compact rate legend",
            "Before": json.dumps(
                old_legend,
                sort_keys=True,
            ),
            "After": json.dumps(
                legend[
                    "position"
                ],
                sort_keys=True,
            ),
        }
    )

    # Expand issuer table substantially.
    table = get_visual(
        page_dir,
        "p03_issuer_table",
    )
    old_table = copy.deepcopy(
        table[
            "position"
        ]
    )

    set_position(
        table,
        y=table_y,
        height=table_height,
    )

    table_visual = table[
        "visual"
    ]

    objects = table_visual.setdefault(
        "objects",
        {},
    )

    grid = objects.setdefault(
        "grid",
        [
            {
                "properties": {}
            }
        ],
    )

    if not grid:
        grid.append(
            {
                "properties": {}
            }
        )

    grid_props = grid[0].setdefault(
        "properties",
        {},
    )

    # Slightly denser rows so the table can expose more issuers at once.
    grid_props[
        "rowPadding"
    ] = expr_double(
        2
    )

    # Keep the table focused on investigation rather than decoration.
    headers = objects.setdefault(
        "columnHeaders",
        [
            {
                "properties": {}
            }
        ],
    )

    if not headers:
        headers.append(
            {
                "properties": {}
            }
        )

    header_props = headers[0].setdefault(
        "properties",
        {},
    )

    header_props[
        "fontSize"
    ] = expr_double(
        9
    )

    header_props[
        "bold"
    ] = expr_bool(
        True
    )

    save_visual(
        page_dir,
        "p03_issuer_table",
        table,
    )

    patch_rows.append(
        {
            "Visual": "p03_issuer_table",
            "Action": (
                "Expand issuer investigation table; reduce row padding "
                "to expose multiple issuer rows"
            ),
            "Before": json.dumps(
                old_table,
                sort_keys=True,
            ),
            "After": json.dumps(
                table[
                    "position"
                ],
                sort_keys=True,
            ),
        }
    )

    return patch_rows


# =====================================================================
# Validation
# =====================================================================

def validate_target_geometry(
    page_dir: Path,
) -> dict:
    table = get_visual(
        page_dir,
        "p03_issuer_table",
    )

    methodology = get_visual(
        page_dir,
        "p03_methodology_note",
    )

    table_pos = table[
        "position"
    ]

    table_bottom = (
        float(
            table_pos[
                "y"
            ]
        )
        + float(
            table_pos[
                "height"
            ]
        )
    )

    methodology_y = float(
        methodology[
            "position"
        ][
            "y"
        ]
    )

    passed = (
        float(
            table_pos[
                "height"
            ]
        )
        >= MIN_TABLE_HEIGHT
        and table_bottom
        <= methodology_y
        - TABLE_BOTTOM_GAP
        + 0.01
    )

    return {
        "CheckID": "LOWER-001",
        "TestName": (
            "Issuer investigation table has useful height"
        ),
        "Expected": (
            f"height >= {MIN_TABLE_HEIGHT:.0f}px and no methodology overlap"
        ),
        "Actual": (
            f"height={float(table_pos['height']):.2f}px; "
            f"bottom={table_bottom:.2f}; methodologyY={methodology_y:.2f}"
        ),
        "Status": (
            "PASS"
            if passed
            else "FAIL"
        ),
    }


def validate_non_target_geometry(
    before: dict[str, dict],
    after: dict[str, dict],
) -> tuple[dict, list[dict]]:
    changes = []

    for name in sorted(
        set(
            before
        )
        | set(
            after
        )
    ):
        if name in ALLOWED_GEOMETRY_TARGETS:
            continue

        if (
            before.get(
                name
            )
            != after.get(
                name
            )
        ):
            changes.append(
                {
                    "VisualName": name,
                    "Before": json.dumps(
                        before.get(
                            name
                        ),
                        sort_keys=True,
                    ),
                    "After": json.dumps(
                        after.get(
                            name
                        ),
                        sort_keys=True,
                    ),
                }
            )

    return (
        {
            "CheckID": "LOWER-002",
            "TestName": (
                "User-adjusted / non-target geometry preserved"
            ),
            "Expected": 0,
            "Actual": len(
                changes
            ),
            "Status": (
                "PASS"
                if not changes
                else "FAIL"
            ),
        },
        changes,
    )


def validate_segment_geometry(
    page_dir: Path,
) -> dict:
    failures = []

    for band_name, chart_name in [
        (
            "p03_band_market",
            "p03_market_chart",
        ),
        (
            "p03_band_offering",
            "p03_offering_chart",
        ),
        (
            "p03_band_plan_type",
            "p03_plan_type_chart",
        ),
    ]:
        band = get_visual(
            page_dir,
            band_name,
        )
        chart = get_visual(
            page_dir,
            chart_name,
        )

        band_bottom = (
            float(
                band[
                    "position"
                ][
                    "y"
                ]
            )
            + float(
                band[
                    "position"
                ][
                    "height"
                ]
            )
        )

        chart_y = float(
            chart[
                "position"
            ][
                "y"
            ]
        )

        if abs(
            chart_y
            - band_bottom
        ) > 0.01:
            failures.append(
                chart_name
            )

    return {
        "CheckID": "LOWER-003",
        "TestName": (
            "Segment charts begin directly below their title bands"
        ),
        "Expected": 0,
        "Actual": len(
            failures
        ),
        "Status": (
            "PASS"
            if not failures
            else "FAIL"
        ),
    }


def validate_no_unsupported_page_roots() -> list[dict]:
    offenders = []

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        for prop in sorted(
            UNSUPPORTED_PAGE_ROOT_PROPERTIES
        ):
            if prop in payload:
                offenders.append(
                    {
                        "Page": payload.get(
                            "displayName"
                        ),
                        "Property": prop,
                    }
                )

    return offenders


def validate_no_full_page_visuals() -> list[dict]:
    offenders = []

    for visual_json in PAGES_DIR.rglob(
        "visual.json"
    ):
        payload = read_json(
            visual_json
        )

        pos = payload.get(
            "position",
            {},
        )

        width = float(
            pos.get(
                "width",
                0,
            )
            or 0
        )
        height = float(
            pos.get(
                "height",
                0,
            )
            or 0
        )

        if (
            width
            >= CANVAS_WIDTH * 0.94
            and height
            >= CANVAS_HEIGHT * 0.94
        ):
            offenders.append(
                {
                    "VisualName": payload.get(
                        "name"
                    ),
                    "Width": width,
                    "Height": height,
                    "RelativePath": str(
                        visual_json.relative_to(
                            PROJECT_ROOT
                        )
                    ),
                }
            )

    return offenders


def validate_json_tree() -> list[dict]:
    failures = []

    for path in (
        REPORT_DIR
        / "definition"
    ).rglob(
        "*.json"
    ):
        try:
            read_json(
                path
            )

        except Exception as exc:
            failures.append(
                {
                    "RelativePath": str(
                        path.relative_to(
                            PROJECT_ROOT
                        )
                    ),
                    "Error": str(
                        exc
                    ),
                }
            )

    return failures


# =====================================================================
# Evidence workbook
# =====================================================================

def style_report(
    path: Path,
) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import (
        Alignment,
        Font,
        PatternFill,
    )
    from openpyxl.utils import (
        get_column_letter,
    )

    wb = load_workbook(
        path
    )

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if (
            ws.max_row
            and ws.max_column
        ):
            ws.auto_filter.ref = (
                ws.dimensions
            )

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            max_len = 0

            for row_idx in range(
                1,
                min(
                    ws.max_row,
                    220,
                )
                + 1,
            ):
                value = ws.cell(
                    row=row_idx,
                    column=col_idx,
                ).value

                if value is not None:
                    max_len = max(
                        max_len,
                        len(
                            str(
                                value
                            )
                        ),
                    )

            ws.column_dimensions[
                get_column_letter(
                    col_idx
                )
            ].width = min(
                max(
                    max_len + 2,
                    10,
                ),
                75,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    wb.save(
        path
    )


# =====================================================================
# Main
# =====================================================================

def main() -> int:
    print("=" * 78)
    print(
        "Transparency in Coverage PUF — "
        "03 Denials Lower-Zone Investigation Patch"
    )
    print("=" * 78)

    print(
        f"Script version : {SCRIPT_VERSION}"
    )
    print(
        f"Report folder  : {REPORT_DIR}"
    )
    print(
        f"Backup         : {BACKUP_FILE}"
    )
    print(
        f"Review file    : {REPORT_FILE}"
    )

    if power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is running. "
            "Save the PBIP and close Desktop completely before applying the patch."
        )

    page_dir, _ = (
        find_denials_page()
    )

    visual_count = len(
        list(
            (
                page_dir
                / "visuals"
            ).rglob(
                "visual.json"
            )
        )
    )

    bound_count = (
        count_bound_visuals(
            page_dir
        )
    )

    if visual_count != EXPECTED_VISUALS:
        raise RuntimeError(
            f"Expected {EXPECTED_VISUALS} Denials visuals; "
            f"found {visual_count}. No changes made."
        )

    if bound_count != EXPECTED_BOUND_VISUALS:
        raise RuntimeError(
            f"Expected {EXPECTED_BOUND_VISUALS} bound/query visuals; "
            f"found {bound_count}. No changes made."
        )

    # Verify all target visuals exist before backup / edits.
    missing = [
        name
        for name in ALLOWED_GEOMETRY_TARGETS
        if not visual_path(
            page_dir,
            name,
        ).exists()
    ]

    for required in [
        "p03_band_segments",
        "p03_methodology_note",
    ]:
        if not visual_path(
            page_dir,
            required,
        ).exists():
            missing.append(
                required
            )

    if missing:
        raise RuntimeError(
            "Required visual(s) missing: "
            + ", ".join(
                sorted(
                    missing
                )
            )
        )

    positions_before = (
        capture_positions(
            page_dir
        )
    )

    create_backup()

    patch_rows = (
        patch_lower_zone(
            page_dir
        )
    )

    positions_after = (
        capture_positions(
            page_dir
        )
    )

    target_check = (
        validate_target_geometry(
            page_dir
        )
    )

    geometry_check, unexpected = (
        validate_non_target_geometry(
            positions_before,
            positions_after,
        )
    )

    segment_check = (
        validate_segment_geometry(
            page_dir
        )
    )

    unsupported = (
        validate_no_unsupported_page_roots()
    )

    full_page = (
        validate_no_full_page_visuals()
    )

    json_failures = (
        validate_json_tree()
    )

    final_table = get_visual(
        page_dir,
        "p03_issuer_table",
    )

    final_table_height = float(
        final_table[
            "position"
        ][
            "height"
        ]
    )

    validation_rows = [
        target_check,
        geometry_check,
        segment_check,
        {
            "CheckID": "LOWER-004",
            "TestName": "Visual count preserved",
            "Expected": EXPECTED_VISUALS,
            "Actual": len(
                list(
                    (
                        page_dir
                        / "visuals"
                    ).rglob(
                        "visual.json"
                    )
                )
            ),
            "Status": "PASS",
        },
        {
            "CheckID": "LOWER-005",
            "TestName": "Bound/query visuals preserved",
            "Expected": EXPECTED_BOUND_VISUALS,
            "Actual": count_bound_visuals(
                page_dir
            ),
            "Status": (
                "PASS"
                if count_bound_visuals(
                    page_dir
                )
                == EXPECTED_BOUND_VISUALS
                else "FAIL"
            ),
        },
        {
            "CheckID": "LOWER-006",
            "TestName": "Unsupported page-root properties",
            "Expected": 0,
            "Actual": len(
                unsupported
            ),
            "Status": (
                "PASS"
                if not unsupported
                else "FAIL"
            ),
        },
        {
            "CheckID": "LOWER-007",
            "TestName": "Full-page selectable visuals",
            "Expected": 0,
            "Actual": len(
                full_page
            ),
            "Status": (
                "PASS"
                if not full_page
                else "FAIL"
            ),
        },
        {
            "CheckID": "LOWER-008",
            "TestName": "PBIR JSON parse failures",
            "Expected": 0,
            "Actual": len(
                json_failures
            ),
            "Status": (
                "PASS"
                if not json_failures
                else "FAIL"
            ),
        },
    ]

    failures = sum(
        row[
            "Status"
        ] == "FAIL"
        for row in validation_rows
    )

    summary_df = pd.DataFrame(
        [
            (
                "PatchStatus",
                "PASS"
                if failures == 0
                else "FAIL",
            ),
            (
                "ScriptVersion",
                SCRIPT_VERSION,
            ),
            (
                "PatchedAtLocal",
                datetime.now()
                .isoformat(
                    timespec="seconds"
                ),
            ),
            (
                "Page",
                "03 Denials",
            ),
            (
                "ProblemAddressed",
                (
                    "Issuer Investigation was too shallow to show issuer rows; "
                    "lower third rebalanced without touching the user's upper-page edits."
                ),
            ),
            (
                "SegmentChartHeight",
                SEGMENT_CHART_HEIGHT,
            ),
            (
                "IssuerTableHeightAfter",
                final_table_height,
            ),
            (
                "UnexpectedGeometryChanges",
                len(
                    unexpected
                ),
            ),
            (
                "VisualsPreserved",
                EXPECTED_VISUALS,
            ),
            (
                "BoundVisualsPreserved",
                EXPECTED_BOUND_VISUALS,
            ),
            (
                "ValidationFailures",
                failures,
            ),
            (
                "BackupFile",
                str(
                    BACKUP_FILE
                ),
            ),
        ],
        columns=[
            "Item",
            "Value",
        ],
    )

    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        REPORT_FILE,
        engine="openpyxl",
    ) as writer:
        summary_df.to_excel(
            writer,
            sheet_name="00_Summary",
            index=False,
        )

        pd.DataFrame(
            patch_rows
        ).to_excel(
            writer,
            sheet_name="01_Patches",
            index=False,
        )

        pd.DataFrame(
            validation_rows
        ).to_excel(
            writer,
            sheet_name="02_Validation",
            index=False,
        )

        if unexpected:
            pd.DataFrame(
                unexpected
            ).to_excel(
                writer,
                sheet_name="03_Unexpected_Geometry",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No geometry changed outside the approved lower-third targets. "
                            "The user's manual adjustments elsewhere were preserved."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="03_Unexpected_Geometry",
                index=False,
            )

        if unsupported:
            pd.DataFrame(
                unsupported
            ).to_excel(
                writer,
                sheet_name="04_Schema_Offenders",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No unsupported page-root properties."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="04_Schema_Offenders",
                index=False,
            )

        if json_failures:
            pd.DataFrame(
                json_failures
            ).to_excel(
                writer,
                sheet_name="05_JSON_Failures",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No JSON parse failures."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="05_JSON_Failures",
                index=False,
            )

    style_report(
        REPORT_FILE
    )

    print()
    print(
        "03 Denials lower-zone patch completed."
    )
    print(
        f"Patch status                  : "
        f"{'PASS' if failures == 0 else 'FAIL'}"
    )
    print(
        f"Issuer table height after     : "
        f"{final_table_height:.0f}px"
    )
    print(
        f"Segment chart height          : "
        f"{SEGMENT_CHART_HEIGHT:.0f}px"
    )
    print(
        f"Unexpected geometry changes   : "
        f"{len(unexpected)}"
    )
    print(
        f"Visuals preserved             : "
        f"{EXPECTED_VISUALS}"
    )
    print(
        f"Bound/query visuals preserved : "
        f"{EXPECTED_BOUND_VISUALS}"
    )
    print(
        f"Full-page selectable visuals  : "
        f"{len(full_page)}"
    )
    print(
        f"Unsupported page properties   : "
        f"{len(unsupported)}"
    )
    print(
        f"Validation failures           : "
        f"{failures}"
    )
    print(
        f"Backup                        : "
        f"{BACKUP_FILE}"
    )
    print(
        f"Review evidence               : "
        f"{REPORT_FILE}"
    )

    if failures:
        print()
        print(
            "Validation failed. Do not open/save the PBIP until reviewed."
        )
        return 2

    print()
    print(
        "Next: open TransparencyInCoverage.pbip and inspect only the lower third "
        "of 03 Denials. The three plan-segment panels should remain readable, "
        "while Issuer Investigation should now expose several issuer rows. "
        "All upper-page manual edits must remain unchanged."
    )

    return 0


if __name__ == "__main__":
    try:
        sys.exit(
            main()
        )

    except Exception as exc:
        print()
        print(
            "DENIALS LOWER-ZONE PATCH FAILED"
        )
        print(
            str(
                exc
            )
        )
        sys.exit(
            1
        )
