from __future__ import annotations

from datetime import datetime
from pathlib import Path
import json, os, shutil, subprocess, sys, zipfile
import pandas as pd

SCRIPT_VERSION = "1.0.0"
PROJECT_ROOT = Path(__file__).resolve().parents[2]
REPORT_DIR = PROJECT_ROOT / "powerbi" / "TransparencyInCoverage.Report"
PAGES_DIR = REPORT_DIR / "definition" / "pages"
CANVAS_WIDTH, CANVAS_HEIGHT = 1920, 1080

COLORS = {
    "canvas":"#F5F7FA","surface":"#FFFFFF","surface_alt":"#F9FAFB","border":"#E7ECF2",
    "ink":"#172033","secondary":"#667085","muted":"#98A2B3","navy":"#17324D",
    "slate":"#344054","blue":"#4F6BED","blue_dark":"#425CC7","blue_soft":"#EEF3FF",
    "blue_pale":"#E8EEFF","teal":"#16A6A1","teal_dark":"#0E7774","teal_soft":"#EAF8F6",
    "amber":"#C8872C","amber_dark":"#8B5E16","amber_soft":"#FFF4E3",
    "red":"#C64B4B","red_dark":"#9F2F2F","red_soft":"#FCE8E8","white":"#FFFFFF"
}
FONT_REGULAR = "'Segoe UI', wf_segoe-ui_normal, helvetica, arial, sans-serif"
FONT_SEMIBOLD = "'Segoe UI Semibold', wf_segoe-ui_semibold, helvetica, arial, sans-serif"
VISUAL_SCHEMA = "https://developer.microsoft.com/json-schemas/fabric/item/report/definition/visualContainer/2.7.0/schema.json"
EXPECTED_CURRENT_VISUALS = 21
EXPECTED_CURRENT_BOUND = 14
EXPECTED_AFTER_VISUALS = 34
EXPECTED_AFTER_BOUND = 14
UNSUPPORTED_ROOT = {"horizontalAlignment","verticalAlignment"}


def desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg
            k=r"Software\Microsoft\Windows\CurrentVersion\Explorer\User Shell Folders"
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER,k) as key:
                v,_=winreg.QueryValueEx(key,"Desktop")
            return Path(os.path.expandvars(v)).resolve()
        except Exception:
            pass
    return (Path.home()/"Desktop").resolve()

REVIEW_DIR = desktop()/"Transparency_PUF_Review"
BACKUP_FILE = REVIEW_DIR/"23_PRE_NETWORK_STORYTELLING_BACKUP.zip"
REPORT_FILE = REVIEW_DIR/"23_NETWORK_STORYTELLING_BUILD_REPORT.xlsx"


def pbi_running():
    if os.name != "nt": return False
    r=subprocess.run(["tasklist","/FI","IMAGENAME eq PBIDesktop.exe"],capture_output=True,text=True,check=False,
                     creationflags=getattr(subprocess,"CREATE_NO_WINDOW",0))
    return "PBIDesktop.exe" in r.stdout

def rj(p:Path): return json.loads(p.read_text(encoding="utf-8-sig"))
def wj(p:Path,o):
    p.parent.mkdir(parents=True,exist_ok=True)
    p.write_text(json.dumps(o,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")

def eb(v): return {"expr":{"Literal":{"Value":"true" if v else "false"}}}
def es(v): return {"expr":{"Literal":{"Value":f"'{v.replace("'","''")}'"}}}
def ed(v): return {"expr":{"Literal":{"Value":f"{v}D"}}}
def ei(v): return {"expr":{"Literal":{"Value":f"{v}L"}}}
def sc(c): return {"solid":{"color":es(c)}}
def mf(n): return {"Measure":{"Expression":{"SourceRef":{"Entity":"_Measures"}},"Property":n}}
def lit(v): return {"Literal":{"Value":f"{v}D"}}


def cond_color(measure, strong=True):
    # Descriptive magnitude bands only; not a quality score.
    pal = ([COLORS["teal_dark"],COLORS["navy"],COLORS["blue_dark"],COLORS["amber_dark"],COLORS["red_dark"]]
           if strong else [COLORS["teal_soft"],COLORS["blue_soft"],COLORS["blue_pale"],COLORS["amber_soft"],COLORS["red_soft"]])
    def cmp(kind,t):
        return {"Comparison":{"ComparisonKind":kind,"Left":mf(measure),"Right":lit(t)}}
    return {"Conditional":{"Cases":[
        {"Condition":cmp(2,0.50),"Value":{"Literal":{"Value":f"'{pal[4]}'"}}},
        {"Condition":cmp(2,0.25),"Value":{"Literal":{"Value":f"'{pal[3]}'"}}},
        {"Condition":cmp(2,0.10),"Value":{"Literal":{"Value":f"'{pal[2]}'"}}},
        {"Condition":cmp(2,0.00),"Value":{"Literal":{"Value":f"'{pal[1]}'"}}},
        {"Condition":cmp(4,0.00),"Value":{"Literal":{"Value":f"'{pal[0]}'"}}},
    ]}}


def find_network_page():
    m=[]
    for pj in PAGES_DIR.glob("*/page.json"):
        o=rj(pj)
        if o.get("displayName")=="02 Network Comparison": m.append((pj.parent,o))
    if len(m)!=1: raise RuntimeError(f"Expected exactly one Network Comparison page; found {len(m)}")
    return m[0]

def all_pages():
    out=[]
    for pj in PAGES_DIR.glob("*/page.json"):
        out.append((pj,rj(pj)))
    return out

def backup():
    REVIEW_DIR.mkdir(parents=True,exist_ok=True)
    if BACKUP_FILE.exists(): BACKUP_FILE.unlink()
    with zipfile.ZipFile(BACKUP_FILE,"w",zipfile.ZIP_DEFLATED) as z:
        for p in sorted(x for x in REPORT_DIR.rglob("*") if x.is_file()):
            z.write(p,str(Path(REPORT_DIR.name)/p.relative_to(REPORT_DIR)))

def count_bound(page):
    return sum(1 for p in (page/"visuals").rglob("visual.json") if rj(p).get("visual",{}).get("query"))

def vpath(page,name):
    p=page/"visuals"/name/"visual.json"
    if not p.exists(): raise FileNotFoundError(f"Visual not found: {name}")
    return p


def text_run(value,size,color,semi=False):
    return {"value":value,"textStyle":{"fontFamily":FONT_SEMIBOLD if semi else FONT_REGULAR,"fontSize":f"{size}pt","color":color}}
def paragraph(*runs): return {"textRuns":list(runs),"horizontalTextAlignment":"center"}
def container(background=None,border=None,radius=8):
    return {
        "title":[{"properties":{"show":eb(False)}}],
        "background":[{"properties":{"show":eb(background is not None),**({"color":sc(background),"transparency":ed(0)} if background else {})}}],
        "border":[{"properties":{"show":eb(border is not None),**({"color":sc(border),"width":ed(1),"radius":ed(radius)} if border else {})}}],
        "dropShadow":[{"properties":{"show":eb(False)}}],
        "visualHeader":[{"properties":{"show":eb(False)}}],
    }
def textbox(name,x,y,w,h,z,paras,bg=None,border=None,radius=8):
    return {"$schema":VISUAL_SCHEMA,"name":name,"position":{"x":x,"y":y,"z":z,"width":w,"height":h},
            "visual":{"visualType":"textbox","objects":{"general":[{"properties":{"paragraphs":paras}}]},
                      "visualContainerObjects":container(bg,border,radius),"drillFilterOtherVisuals":True}}
def band(name,title,x,y,w,h,z,bg,subtitle=None):
    ps=[paragraph(text_run(title,10.2,COLORS["white"],True))]
    if subtitle: ps.append(paragraph(text_run(subtitle,7.7,COLORS["white"],False)))
    return textbox(name,x,y,w,h,z,ps,bg,bg,8)
def write_visual(page,o):
    p=page/"visuals"/o["name"]/"visual.json"; wj(p,o); return p


def hide_native_title(obj):
    v=obj.setdefault("visual",{})
    vc=v.setdefault("visualContainerObjects",{})
    vc["title"]=[{"properties":{"show":eb(False)}}]
    vc.pop("subTitle",None)

def set_textbox_style(obj,bg,border,text_color=None):
    v=obj["visual"]
    vc=v.setdefault("visualContainerObjects",{})
    vc["background"]=[{"properties":{"show":eb(True),"color":sc(bg),"transparency":ed(0)}}]
    vc["border"]=[{"properties":{"show":eb(True),"color":sc(border),"width":ed(1),"radius":ed(9)}}]
    if text_color:
        ps=v["objects"]["general"][0]["properties"]["paragraphs"]
        for p in ps:
            for tr in p.get("textRuns",[]): tr.setdefault("textStyle",{})["color"]=text_color

def add_band_for_visual(page,visual_name,band_name,title,color,h=32,subtitle=None,z=1900):
    p=vpath(page,visual_name); o=rj(p); pos=o["position"]
    hide_native_title(o); wj(p,o)
    write_visual(page,band(band_name,title,pos["x"],pos["y"],pos["width"],h,z,color,subtitle))


def patch_chart(page):
    p=vpath(page,"p02_state_gap_chart"); o=rj(p); pos=o["position"]
    hide_native_title(o)
    # Reserve 52 px for title/story band.
    pos["y"] += 52; pos["height"] -= 52
    v=o["visual"]; objs=v.setdefault("objects",{})
    objs["dataPoint"]=[
        {"properties":{"fill":sc(COLORS["blue"])}},
        {"properties":{"fill":{"solid":{"color":{"expr":cond_color("Issuer Network Denial Rate Gap",True)}}}},
         "selector":{"data":[{"dataViewWildcard":{"matchingOption":1}}]}},
    ]
    wj(p,o)
    write_visual(page,band(
        "p02_band_state_gap","WHERE IS THE NETWORK GAP CONCENTRATED?",
        pos["x"],pos["y"]-52,pos["width"],52,1900,COLORS["navy"],
        "▼ negative  •  ● 0–10pp  •  ▲ 10–25pp  •  ▲ 25–50pp  •  ▲ ≥50pp"
    ))


def patch_flow_title(page):
    p=vpath(page,"p02_flow_title"); o=rj(p)
    set_textbox_style(o,COLORS["slate"],COLORS["slate"],COLORS["white"])
    # Slightly increase band height but keep section footprint intact.
    o["position"]["height"]=52
    wj(p,o)


def patch_table(page):
    p=vpath(page,"p02_issuer_table"); o=rj(p); pos=o["position"]
    hide_native_title(o)
    # Make room for title band + interpretation key.
    pos["y"] += 80; pos["height"] -= 80
    v=o["visual"]; objs=v.setdefault("objects",{})
    gap_ref="_Measures.Issuer Network Denial Rate Gap"
    in_rate="_Measures.Issuer In-Network Denial Rate"
    out_rate="_Measures.Issuer Out-of-Network Denial Rate"
    in_claim="_Measures.Issuer Claims Received - In Network"
    out_claim="_Measures.Issuer Claims Received - Out of Network"
    objs["values"]=[
        {"properties":{"backColorSecondary":sc(COLORS["surface_alt"]) }},
        {"properties":{
            "backColor":{"solid":{"color":{"expr":cond_color("Issuer Network Denial Rate Gap",False)}}},
            "fontColor":{"solid":{"color":{"expr":cond_color("Issuer Network Denial Rate Gap",True)}}},
         },"selector":{"data":[{"dataViewWildcard":{"matchingOption":1}}],"metadata":gap_ref}},
    ]
    objs["columnFormatting"]=[
        {"properties":{"fontColor":sc(COLORS["navy"]),"alignment":es("right")},"selector":{"metadata":in_rate}},
        {"properties":{"fontColor":sc(COLORS["blue"]),"alignment":es("right")},"selector":{"metadata":out_rate}},
        {"properties":{"alignment":es("right")},"selector":{"metadata":gap_ref}},
        {"properties":{"alignment":es("right"),"dataBars":{"positiveColor":sc("#DCE6F2"),"negativeColor":sc("#DCE6F2"),"axisColor":sc(COLORS["white"]),"reverseDirection":eb(False),"hideText":eb(False)}},"selector":{"metadata":in_claim}},
        {"properties":{"alignment":es("right"),"dataBars":{"positiveColor":sc("#D8E2FF"),"negativeColor":sc("#D8E2FF"),"axisColor":sc(COLORS["white"]),"reverseDirection":eb(False),"hideText":eb(False)}},"selector":{"metadata":out_claim}},
    ]
    wj(p,o)
    # Bands occupy the original top 80 px of table footprint.
    write_visual(page,band(
        "p02_band_table","ISSUER INVESTIGATION — FOLLOW THE COLOR",
        pos["x"],pos["y"]-80,pos["width"],46,1950,COLORS["navy"],
        "Sorted by gap • colors encode gap magnitude only, not issuer quality"
    ))
    write_visual(page,textbox(
        "p02_gap_legend",pos["x"],pos["y"]-32,pos["width"],30,1960,
        [paragraph(
            text_run("▼ NEGATIVE / LOWER GAP",8.6,COLORS["teal_dark"],True),
            text_run("     ● 0–10pp",8.6,COLORS["navy"],True),
            text_run("     ▲ 10–25pp",8.6,COLORS["blue_dark"],True),
            text_run("     ▲ 25–50pp",8.6,COLORS["amber_dark"],True),
            text_run("     ▲ ≥50pp",8.6,COLORS["red_dark"],True),
        )],COLORS["surface"],COLORS["border"],7
    ))


def patch_methodology(page):
    p=vpath(page,"p02_methodology_note"); o=rj(p)
    ps=o["visual"]["objects"]["general"][0]["properties"]["paragraphs"]
    ps[:]=[paragraph(text_run(
        "Color encodes network-gap magnitude only — not issuer quality  •  Gap = OON denial rate − IN denial rate  •  "
        "Availability-aware ratio-of-totals  •  Source-reported anomalies remain visible; no clipping or imputation",
        8.7,COLORS["muted"],True
    ))]
    wj(p,o)


def patch_page_title(page):
    p=vpath(page,"p02_page_title"); o=rj(p)
    set_textbox_style(o,COLORS["blue_soft"],COLORS["blue_pale"])
    wj(p,o)


def patch_cards(page):
    specs=[
        ("p02_dq_context","p02_band_dq","GLOBAL OPEN DQ EXCEPTIONS",COLORS["amber"],28),
        ("p02_kpi_in_rate","p02_band_kpi_in","IN-NETWORK DENIAL RATE",COLORS["navy"],34),
        ("p02_kpi_out_rate","p02_band_kpi_out","OUT-OF-NETWORK DENIAL RATE",COLORS["blue"],34),
        ("p02_kpi_gap","p02_band_kpi_gap","NETWORK GAP",COLORS["amber"],34),
        ("p02_kpi_issuers","p02_band_kpi_issuers","ISSUERS IN CONTEXT",COLORS["slate"],34),
        ("p02_kpi_plans","p02_band_kpi_plans","PLANS IN CONTEXT",COLORS["slate"],34),
        ("p02_flow_in_received","p02_band_flow_in_received","IN-NETWORK RECEIVED",COLORS["navy"],30),
        ("p02_flow_out_received","p02_band_flow_out_received","OUT-OF-NETWORK RECEIVED",COLORS["blue"],30),
        ("p02_flow_in_denied","p02_band_flow_in_denied","IN-NETWORK DENIED",COLORS["navy"],30),
        ("p02_flow_out_denied","p02_band_flow_out_denied","OUT-OF-NETWORK DENIED",COLORS["blue"],30),
    ]
    for vname,bname,title,color,h in specs:
        add_band_for_visual(page,vname,bname,title,color,h=h,z=1850)


def validate_json():
    bad=[]
    for p in (REPORT_DIR/"definition").rglob("*.json"):
        try: rj(p)
        except Exception as e: bad.append({"RelativePath":str(p.relative_to(PROJECT_ROOT)),"Error":str(e)})
    return bad

def validate_schema():
    bad=[]
    for pj,o in all_pages():
        for prop in UNSUPPORTED_ROOT:
            if prop in o: bad.append({"Page":o.get("displayName"),"Property":prop})
    return bad

def validate_full_page():
    bad=[]
    for p in PAGES_DIR.rglob("visual.json"):
        o=rj(p); pos=o.get("position",{}); w=float(pos.get("width",0) or 0); h=float(pos.get("height",0) or 0)
        if w>=CANVAS_WIDTH*.94 and h>=CANVAS_HEIGHT*.94:
            bad.append({"VisualName":o.get("name"),"Width":w,"Height":h,"RelativePath":str(p.relative_to(PROJECT_ROOT))})
    return bad

def walk(x):
    if isinstance(x,dict):
        yield x
        for v in x.values(): yield from walk(v)
    elif isinstance(x,list):
        for v in x: yield from walk(v)

def validation_story(page):
    rows=[]
    chart=rj(vpath(page,"p02_state_gap_chart"))
    dp=chart["visual"]["objects"].get("dataPoint",[])
    has_cond=any("Conditional" in n for n in walk(dp))
    wildcard=any(i.get("selector",{}).get("data",[{}])[0].get("dataViewWildcard",{}).get("matchingOption")==1 for i in dp if i.get("selector"))
    rows.append({"CheckID":"STORY-001","TestName":"State chart conditional colors","Expected":"Conditional + per-point wildcard","Actual":f"conditional={has_cond}; wildcard={wildcard}","Status":"PASS" if has_cond and wildcard else "FAIL"})
    table=rj(vpath(page,"p02_issuer_table")); objs=table["visual"]["objects"]
    has_gap=any("Conditional" in n for n in walk(objs.get("values",[])))
    bars=sum(1 for i in objs.get("columnFormatting",[]) if "dataBars" in i.get("properties",{}))
    rows.append({"CheckID":"STORY-002","TestName":"Table conditional colors + data bars","Expected":"Gap colors + 2 data bars","Actual":f"gap={has_gap}; dataBars={bars}","Status":"PASS" if has_gap and bars>=2 else "FAIL"})
    bands=[p for p in (page/"visuals").rglob("visual.json") if str(rj(p).get("name","")).startswith("p02_band_")]
    rows.append({"CheckID":"STORY-003","TestName":"Colored title bands","Expected":12,"Actual":len(bands),"Status":"PASS" if len(bands)==12 else "FAIL"})
    return rows


def style_report(path):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment,Font,PatternFill
    from openpyxl.utils import get_column_letter
    wb=load_workbook(path); hf=PatternFill("solid",fgColor="1F4E78"); hfont=Font(color="FFFFFF",bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes="A2"
        if ws.max_row and ws.max_column: ws.auto_filter.ref=ws.dimensions
        for c in ws[1]: c.fill=hf; c.font=hfont; c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        for ci in range(1,ws.max_column+1):
            m=0
            for ri in range(1,min(ws.max_row,200)+1):
                v=ws.cell(ri,ci).value
                if v is not None: m=max(m,len(str(v)))
            ws.column_dimensions[get_column_letter(ci)].width=min(max(m+2,10),70)
        for row in ws.iter_rows():
            for c in row: c.alignment=Alignment(vertical="top",wrap_text=True)
    wb.save(path)


def main():
    print("="*78)
    print("Transparency in Coverage PUF — 02 Network Comparison Storytelling Rebuild")
    print("="*78)
    print(f"Script version : {SCRIPT_VERSION}")
    print(f"Report folder  : {REPORT_DIR}")
    print(f"Backup         : {BACKUP_FILE}")
    print(f"Review file    : {REPORT_FILE}")

    if pbi_running(): raise RuntimeError("Power BI Desktop is running. Save and close it completely before changing PBIR files.")
    if not REPORT_DIR.exists(): raise FileNotFoundError(f"Report folder not found: {REPORT_DIR}")
    page,_=find_network_page()
    before_v=len(list((page/"visuals").rglob("visual.json"))); before_b=count_bound(page)
    if (before_v,before_b)!=(EXPECTED_CURRENT_VISUALS,EXPECTED_CURRENT_BOUND):
        raise RuntimeError(f"Expected approved post-22 baseline {EXPECTED_CURRENT_VISUALS}/{EXPECTED_CURRENT_BOUND}; found {before_v}/{before_b}. No changes made.")
    if validate_schema(): raise RuntimeError("Unsupported page-root properties already exist. No changes made.")

    backup()
    patch_page_title(page)
    patch_cards(page)
    patch_flow_title(page)
    patch_chart(page)
    patch_table(page)
    patch_methodology(page)

    after_v=len(list((page/"visuals").rglob("visual.json"))); after_b=count_bound(page)
    jbad=validate_json(); sbad=validate_schema(); fbad=validate_full_page(); story=validation_story(page)
    checks=story+[
        {"CheckID":"STORY-004","TestName":"Visual count","Expected":EXPECTED_AFTER_VISUALS,"Actual":after_v,"Status":"PASS" if after_v==EXPECTED_AFTER_VISUALS else "FAIL"},
        {"CheckID":"STORY-005","TestName":"Bound visual count","Expected":EXPECTED_AFTER_BOUND,"Actual":after_b,"Status":"PASS" if after_b==EXPECTED_AFTER_BOUND else "FAIL"},
        {"CheckID":"STORY-006","TestName":"Unsupported page-root properties","Expected":0,"Actual":len(sbad),"Status":"PASS" if not sbad else "FAIL"},
        {"CheckID":"STORY-007","TestName":"Full-page selectable visuals","Expected":0,"Actual":len(fbad),"Status":"PASS" if not fbad else "FAIL"},
        {"CheckID":"STORY-008","TestName":"PBIR JSON parse failures","Expected":0,"Actual":len(jbad),"Status":"PASS" if not jbad else "FAIL"},
    ]
    failures=sum(r["Status"]=="FAIL" for r in checks)
    summary=pd.DataFrame([
        ("BuildStatus","PASS" if failures==0 else "FAIL"),("ScriptVersion",SCRIPT_VERSION),("BuiltAtLocal",datetime.now().isoformat(timespec="seconds")),
        ("Page","02 Network Comparison"),("StoryFlow","Gap → concentration → volume → issuer investigation"),("VisualsAfter",after_v),("BoundVisualsAfter",after_b),
        ("TitleBands",12),("ChartColorLogic","<0 teal | 0–10pp navy/blue | 10–25pp deep blue | 25–50pp amber | >=50pp red"),
        ("TableStorytelling","Gap conditional cell colors + IN/OON claim data bars + colored rate columns"),("ValidationFailures",failures),("BackupFile",str(BACKUP_FILE)),
        ("InterpretationRule","Colors describe gap magnitude only; they are not an issuer-quality score."),
    ],columns=["Item","Value"])
    REVIEW_DIR.mkdir(parents=True,exist_ok=True)
    with pd.ExcelWriter(REPORT_FILE,engine="openpyxl") as w:
        summary.to_excel(w,sheet_name="00_Summary",index=False)
        pd.DataFrame(checks).to_excel(w,sheet_name="01_Validation",index=False)
        (pd.DataFrame(sbad) if sbad else pd.DataFrame([{"Info":"No unsupported page-root properties."}])).to_excel(w,sheet_name="02_Schema",index=False)
        (pd.DataFrame(fbad) if fbad else pd.DataFrame([{"Info":"No full-page selectable visuals."}])).to_excel(w,sheet_name="03_FullPage",index=False)
        (pd.DataFrame(jbad) if jbad else pd.DataFrame([{"Info":"No JSON parse failures."}])).to_excel(w,sheet_name="04_JSON",index=False)
    style_report(REPORT_FILE)

    print()
    print("02 Network Comparison storytelling rebuild completed.")
    print(f"Build status                 : {'PASS' if failures==0 else 'FAIL'}")
    print(f"Visuals after                : {after_v}")
    print(f"Bound/query visuals          : {after_b}")
    print("Colored title bands          : 12")
    print("State chart                  : conditional story colors")
    print("Issuer table                 : conditional gap colors + claim data bars")
    print(f"Full-page selectable visuals : {len(fbad)}")
    print(f"Unsupported page properties  : {len(sbad)}")
    print(f"Validation failures          : {failures}")
    print(f"Backup                       : {BACKUP_FILE}")
    print(f"Review evidence              : {REPORT_FILE}")
    if failures:
        print("\nOffline validation failed. Do not open/save the PBIP until reviewed.")
        return 2
    print("\nNext: open the PBIP and inspect page 02. Confirm title bands, state colors, table gap shading, claim data bars, INDEX navigation, and blank-canvas behavior.")
    return 0

if __name__=="__main__":
    try: sys.exit(main())
    except Exception as e:
        print("\nNETWORK STORYTELLING REBUILD FAILED")
        print(str(e))
        sys.exit(1)
