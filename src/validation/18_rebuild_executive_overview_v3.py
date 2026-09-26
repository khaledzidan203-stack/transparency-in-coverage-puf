from __future__ import annotations

from datetime import datetime
from pathlib import Path
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile

import pandas as pd


SCRIPT_VERSION = "3.0.0"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
POWERBI_ROOT = PROJECT_ROOT / "powerbi"
REPORT_DIR = POWERBI_ROOT / "TransparencyInCoverage.Report"
SEMANTIC_DIR = POWERBI_ROOT / "TransparencyInCoverage.SemanticModel"
PAGES_DIR = REPORT_DIR / "definition" / "pages"

MEASURES_TMDL = (
    SEMANTIC_DIR
    / "definition"
    / "tables"
    / "_Measures.tmdl"
)
DIM_METRIC_TMDL = (
    SEMANTIC_DIR
    / "definition"
    / "tables"
    / "DimMetric.tmdl"
)

VISUAL_SCHEMA = (
    "https://developer.microsoft.com/json-schemas/fabric/item/report/"
    "definition/visualContainer/2.7.0/schema.json"
)

CANVAS_WIDTH = 1920
CANVAS_HEIGHT = 1080

COLORS = {
    "canvas": "#F5F7FA",
    "surface": "#FFFFFF",
    "surface_alt": "#F9FAFB",
    "border": "#E7ECF2",
    "ink": "#172033",
    "secondary": "#667085",
    "muted": "#98A2B3",
    "navy": "#17324D",
    "blue": "#4F6BED",
    "blue_dark": "#425CC7",
    "blue_soft": "#EEF3FF",
    "teal": "#16A6A1",
    "teal_soft": "#EAF8F6",
    "amber": "#C8872C",
    "red": "#C64B4B",
}

FONT_REGULAR = (
    "'Segoe UI', wf_segoe-ui_normal, helvetica, arial, sans-serif"
)
FONT_SEMIBOLD = (
    "'Segoe UI Semibold', wf_segoe-ui_semibold, "
    "helvetica, arial, sans-serif"
)

EXPECTED_PAGES = [
    "00 INDEX",
    "01 Executive Overview",
    "02 Network Comparison",
    "03 Denials",
    "04 Denial Reasons",
    "05 Appeals",
    "06 Resubmissions",
    "07 Enrollment",
    "08 Issuer & State Explorer",
    "09 Data Availability",
    "10 Data Quality & Methodology",
]

REQUIRED_MEASURES = [
    "Issuer Count",
    "Plan Count",
    "Issuer Comparable Claims Received - Total",
    "Issuer Comparable Claims Denied - Total",
    "Issuer Comparable Overall Denial Rate",
    "Issuer In-Network Denial Rate",
    "Issuer Out-of-Network Denial Rate",
    "Issuer Network Denial Rate Gap",
    "Internal Appeal Overturn Rate",
    "External Appeal Overturn Rate",
    "Internal Appeals Filed",
    "External Appeals Filed",
    "Comparable Denial Reason Composition %",
    "Open DQ Exception Count",
    "Known Source Exception Row Count",
]

REQUIRED_DIM_METRIC_COLUMNS = {
    "MetricDisplayName",
}


# =====================================================================
# Windows / paths
# =====================================================================

def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )

            with winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                key_path,
            ) as key:
                raw_value, _ = winreg.QueryValueEx(
                    key,
                    "Desktop",
                )

            return Path(
                os.path.expandvars(raw_value)
            ).resolve()

        except Exception:
            pass

    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = (
    get_windows_desktop()
    / "Transparency_PUF_Review"
)

BACKUP_FILE = (
    REVIEW_DIR
    / "18_PRE_EXECUTIVE_REBUILD_V3_BACKUP.zip"
)

REPORT_FILE = (
    REVIEW_DIR
    / "18_EXECUTIVE_REBUILD_V3_REPORT.xlsx"
)


def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False

    result = subprocess.run(
        [
            "tasklist",
            "/FI",
            "IMAGENAME eq PBIDesktop.exe",
        ],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(
            subprocess,
            "CREATE_NO_WINDOW",
            0,
        ),
    )

    return "PBIDesktop.exe" in result.stdout


def read_json(path: Path):
    return json.loads(
        path.read_text(
            encoding="utf-8-sig",
        )
    )


def write_json(
    path: Path,
    payload,
) -> None:
    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    path.write_text(
        json.dumps(
            payload,
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )


def sha256(path: Path) -> str:
    h = hashlib.sha256()

    with path.open("rb") as f:
        for chunk in iter(
            lambda: f.read(1024 * 1024),
            b"",
        ):
            h.update(chunk)

    return h.hexdigest()


# =====================================================================
# PBIR literals / formatting
# =====================================================================

def expr_bool(value: bool) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": (
                    "true"
                    if value
                    else "false"
                )
            }
        }
    }


def expr_string(value: str) -> dict:
    escaped = value.replace(
        "'",
        "''",
    )

    return {
        "expr": {
            "Literal": {
                "Value": f"'{escaped}'"
            }
        }
    }


def expr_double(
    value: float | int,
) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": f"{value}D"
            }
        }
    }


def expr_int(
    value: int,
) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": f"{value}L"
            }
        }
    }


def solid_color(
    color: str,
) -> dict:
    return {
        "solid": {
            "color": expr_string(
                color
            )
        }
    }


def page_background_objects() -> dict:
    """
    IMPORTANT:
    The page background is a PAGE formatting object, not a full-page visual.
    This is the key fix for the old 'click blank area -> giant textbox selected'
    behavior.
    """
    return {
        "background": [
            {
                "properties": {
                    "color": solid_color(
                        COLORS["canvas"]
                    ),
                    "transparency": expr_double(
                        0
                    ),
                }
            }
        ]
    }


# =====================================================================
# TMDL validation
# =====================================================================

def parse_tmdl_measures(
    path: Path,
) -> set[str]:
    text = path.read_text(
        encoding="utf-8-sig",
    )

    measures: set[str] = set()

    for line in text.splitlines():
        stripped = line.lstrip()

        if not stripped.startswith(
            "measure "
        ):
            continue

        lhs = stripped.split(
            "=",
            1,
        )[0].strip()

        raw = lhs[
            len("measure "):
        ].strip()

        if (
            raw.startswith("'")
            and raw.endswith("'")
        ):
            raw = (
                raw[1:-1]
                .replace(
                    "''",
                    "'",
                )
            )

        measures.add(
            raw
        )

    return measures


def parse_tmdl_columns(
    path: Path,
) -> set[str]:
    text = path.read_text(
        encoding="utf-8-sig",
    )

    columns: set[str] = set()

    for line in text.splitlines():
        stripped = line.lstrip()

        if not stripped.startswith(
            "column "
        ):
            continue

        raw = stripped[
            len("column "):
        ].strip()

        if raw.startswith("'"):
            match = re.match(
                r"'((?:''|[^'])*)'",
                raw,
            )

            if match:
                columns.add(
                    match.group(1)
                    .replace(
                        "''",
                        "'",
                    )
                )

        else:
            columns.add(
                raw.split(
                    "\t",
                    1,
                )[0].strip()
            )

    return columns


# =====================================================================
# Pages / backup
# =====================================================================

def find_pages() -> dict[
    str,
    tuple[Path, dict]
]:
    pages = {}

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        display_name = payload.get(
            "displayName"
        )

        if display_name:
            pages[
                display_name
            ] = (
                page_json.parent,
                payload,
            )

    return pages


def create_backup() -> None:
    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if BACKUP_FILE.exists():
        BACKUP_FILE.unlink()

    with zipfile.ZipFile(
        BACKUP_FILE,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zf:
        for path in sorted(
            p
            for p in REPORT_DIR.rglob("*")
            if p.is_file()
        ):
            archive_name = (
                Path(
                    REPORT_DIR.name
                )
                / path.relative_to(
                    REPORT_DIR
                )
            )

            zf.write(
                path,
                str(
                    archive_name
                ),
            )


def report_hashes() -> dict[
    str,
    str,
]:
    return {
        str(
            path.relative_to(
                REPORT_DIR
            )
        ): sha256(
            path
        )
        for path in REPORT_DIR.rglob(
            "*"
        )
        if path.is_file()
    }


# =====================================================================
# Page background sanitation — global fix for ALL pages
# =====================================================================

def set_native_page_background(
    page_dir: Path,
) -> None:
    page_json = (
        page_dir
        / "page.json"
    )

    payload = read_json(
        page_json
    )

    # Current Power BI Desktop PBIR schema rejects horizontalAlignment and
    # verticalAlignment at the page root. Remove them if an earlier build
    # introduced them. Canvas placement is left to Desktop while FitToPage
    # controls the viewing behavior.
    payload.pop("horizontalAlignment", None)
    payload.pop("verticalAlignment", None)

    payload["width"] = (
        CANVAS_WIDTH
    )
    payload["height"] = (
        CANVAS_HEIGHT
    )
    payload["displayOption"] = (
        "FitToPage"
    )

    objects = payload.setdefault(
        "objects",
        {},
    )

    objects.update(
        page_background_objects()
    )

    write_json(
        page_json,
        payload,
    )


def remove_full_page_canvas_visuals(
    page_dir: Path,
) -> list[str]:
    removed = []

    visuals_dir = (
        page_dir
        / "visuals"
    )

    if not visuals_dir.exists():
        return removed

    for visual_json in list(
        visuals_dir.rglob(
            "visual.json"
        )
    ):
        payload = read_json(
            visual_json
        )

        name = str(
            payload.get(
                "name",
                "",
            )
        )

        pos = payload.get(
            "position",
            {},
        )

        width = float(
            pos.get(
                "width",
                0,
            )
            or 0
        )
        height = float(
            pos.get(
                "height",
                0,
            )
            or 0
        )

        visual_type = (
            payload.get(
                "visual",
                {},
            )
            .get(
                "visualType"
            )
        )

        is_named_canvas = (
            name == "idx_canvas"
            or name.endswith(
                "_canvas"
            )
        )

        is_full_page_textbox = (
            visual_type
            == "textbox"
            and width
            >= CANVAS_WIDTH * 0.94
            and height
            >= CANVAS_HEIGHT * 0.94
        )

        if (
            is_named_canvas
            or is_full_page_textbox
        ):
            visual_dir = (
                visual_json.parent
            )

            removed.append(
                name
            )

            shutil.rmtree(
                visual_dir
            )

    return removed


# =====================================================================
# Text / container builders
# =====================================================================

def text_run(
    value: str,
    *,
    size: float,
    color: str,
    semibold: bool = False,
) -> dict:
    return {
        "value": value,
        "textStyle": {
            "fontFamily": (
                FONT_SEMIBOLD
                if semibold
                else FONT_REGULAR
            ),
            "fontSize": (
                f"{size}pt"
            ),
            "color": color,
        },
    }


def paragraph(
    *runs: dict,
    align: str = "center",
) -> dict:
    return {
        "textRuns": list(
            runs
        ),
        "horizontalTextAlignment": (
            align
        ),
    }


def container_objects(
    *,
    background: str | None = None,
    border: str | None = None,
    radius: int = 8,
    show_title: bool = False,
    title: str = "",
    title_size: float = 11,
    title_color: str = COLORS["ink"],
    subtitle: str | None = None,
) -> dict:
    result = {
        "title": [
            {
                "properties": {
                    "show": expr_bool(
                        show_title
                    ),
                    **(
                        {
                            "text": expr_string(
                                title
                            ),
                            "alignment": expr_string(
                                "center"
                            ),
                            "fontColor": solid_color(
                                title_color
                            ),
                            "fontFamily": expr_string(
                                FONT_SEMIBOLD
                            ),
                            "fontSize": expr_double(
                                title_size
                            ),
                            "titleWrap": expr_bool(
                                True
                            ),
                        }
                        if show_title
                        else {}
                    ),
                }
            }
        ],
        "background": [
            {
                "properties": {
                    "show": expr_bool(
                        background is not None
                    ),
                    **(
                        {
                            "color": solid_color(
                                background
                            ),
                            "transparency": expr_double(
                                0
                            ),
                        }
                        if background
                        is not None
                        else {}
                    ),
                }
            }
        ],
        "border": [
            {
                "properties": {
                    "show": expr_bool(
                        border is not None
                    ),
                    **(
                        {
                            "color": solid_color(
                                border
                            ),
                            "width": expr_double(
                                1
                            ),
                            "radius": expr_double(
                                radius
                            ),
                        }
                        if border
                        is not None
                        else {}
                    ),
                }
            }
        ],
        "dropShadow": [
            {
                "properties": {
                    "show": expr_bool(
                        False
                    )
                }
            }
        ],
        "visualHeader": [
            {
                "properties": {
                    "show": expr_bool(
                        False
                    )
                }
            }
        ],
    }

    if subtitle is not None:
        result["subTitle"] = [
            {
                "properties": {
                    "show": expr_bool(
                        True
                    ),
                    "text": expr_string(
                        subtitle
                    ),
                    "alignment": expr_string(
                        "center"
                    ),
                    "fontColor": solid_color(
                        COLORS["secondary"]
                    ),
                    "fontFamily": expr_string(
                        FONT_REGULAR
                    ),
                    "fontSize": expr_double(
                        9.5
                    ),
                }
            }
        ]

    return result


def make_textbox(
    *,
    name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    paragraphs: list[dict],
    background: str | None = None,
    border: str | None = None,
    radius: int = 8,
) -> dict:
    return {
        "$schema": (
            VISUAL_SCHEMA
        ),
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
        },
        "visual": {
            "visualType": (
                "textbox"
            ),
            "objects": {
                "general": [
                    {
                        "properties": {
                            "paragraphs": (
                                paragraphs
                            )
                        }
                    }
                ]
            },
            "visualContainerObjects": (
                container_objects(
                    background=background,
                    border=border,
                    radius=radius,
                )
            ),
            "drillFilterOtherVisuals": (
                True
            ),
        },
    }


# =====================================================================
# Field binding helpers
# =====================================================================

def measure_field(
    measure_name: str,
) -> dict:
    return {
        "Measure": {
            "Expression": {
                "SourceRef": {
                    "Entity": (
                        "_Measures"
                    )
                }
            },
            "Property": (
                measure_name
            ),
        }
    }


def column_field(
    entity: str,
    column_name: str,
) -> dict:
    return {
        "Column": {
            "Expression": {
                "SourceRef": {
                    "Entity": entity
                }
            },
            "Property": (
                column_name
            ),
        }
    }


def measure_projection(
    measure_name: str,
) -> dict:
    return {
        "field": measure_field(
            measure_name
        ),
        "queryRef": (
            f"_Measures."
            f"{measure_name}"
        ),
        "nativeQueryRef": (
            measure_name
        ),
    }


def column_projection(
    entity: str,
    column_name: str,
    *,
    active: bool = False,
) -> dict:
    result = {
        "field": column_field(
            entity,
            column_name,
        ),
        "queryRef": (
            f"{entity}."
            f"{column_name}"
        ),
        "nativeQueryRef": (
            column_name
        ),
    }

    if active:
        result["active"] = (
            True
        )

    return result


# =====================================================================
# Bound visual builders
# =====================================================================

def make_card(
    *,
    name: str,
    measure_name: str,
    title: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    tab_order: int,
    value_color: str = COLORS["ink"],
    background: str = COLORS["surface"],
    title_size: float = 10.5,
    value_size: float = 29,
) -> dict:
    return {
        "$schema": (
            VISUAL_SCHEMA
        ),
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": (
                tab_order
            ),
        },
        "visual": {
            "visualType": (
                "card"
            ),
            "query": {
                "queryState": {
                    "Values": {
                        "projections": [
                            measure_projection(
                                measure_name
                            )
                        ]
                    }
                },
                "sortDefinition": {
                    "sort": [
                        {
                            "field": (
                                measure_field(
                                    measure_name
                                )
                            ),
                            "direction": (
                                "Descending"
                            ),
                        }
                    ],
                    "isDefaultSort": (
                        True
                    ),
                },
            },
            "objects": {
                "labels": [
                    {
                        "properties": {
                            "color": solid_color(
                                value_color
                            ),
                            "fontSize": expr_double(
                                value_size
                            ),
                            "bold": expr_bool(
                                True
                            ),
                            "fontFamily": expr_string(
                                FONT_SEMIBOLD
                            ),
                        }
                    }
                ],
                "categoryLabels": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": (
                container_objects(
                    background=background,
                    border=COLORS["border"],
                    radius=9,
                    show_title=True,
                    title=title,
                    title_size=title_size,
                )
            ),
            "drillFilterOtherVisuals": (
                True
            ),
        },
    }


def make_reason_bar_chart(
    *,
    name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    tab_order: int,
) -> dict:
    measure_name = (
        "Comparable Denial Reason Composition %"
    )

    return {
        "$schema": (
            VISUAL_SCHEMA
        ),
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": (
                tab_order
            ),
        },
        "visual": {
            "visualType": (
                "barChart"
            ),
            "query": {
                "queryState": {
                    "Category": {
                        "projections": [
                            column_projection(
                                "DimMetric",
                                "MetricDisplayName",
                                active=True,
                            )
                        ]
                    },
                    "Y": {
                        "projections": [
                            measure_projection(
                                measure_name
                            )
                        ]
                    },
                },
                "sortDefinition": {
                    "sort": [
                        {
                            "field": (
                                measure_field(
                                    measure_name
                                )
                            ),
                            "direction": (
                                "Descending"
                            ),
                        }
                    ],
                    "isDefaultSort": (
                        True
                    ),
                },
            },
            "objects": {
                "legend": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "categoryAxis": [
                    {
                        "properties": {
                            "showAxisTitle": expr_bool(
                                False
                            ),
                            "labelColor": solid_color(
                                COLORS["secondary"]
                            ),
                            "fontFamily": expr_string(
                                FONT_REGULAR
                            ),
                            "fontSize": expr_double(
                                8.5
                            ),
                            "concatenateLabels": expr_bool(
                                False
                            ),
                            "innerPadding": expr_int(
                                4
                            ),
                            "maxMarginFactor": expr_int(
                                42
                            ),
                        }
                    }
                ],
                "valueAxis": [
                    {
                        "properties": {
                            "show": expr_bool(
                                True
                            ),
                            "showAxisTitle": expr_bool(
                                False
                            ),
                            "labelColor": solid_color(
                                COLORS["muted"]
                            ),
                            "fontFamily": expr_string(
                                FONT_REGULAR
                            ),
                            "fontSize": expr_double(
                                8
                            ),
                            "gridlineShow": expr_bool(
                                False
                            ),
                        }
                    }
                ],
                "dataPoint": [
                    {
                        "properties": {
                            "fill": solid_color(
                                COLORS["blue"]
                            )
                        },
                        "selector": {
                            "data": [
                                {
                                    "dataViewWildcard": {
                                        "matchingOption": 1
                                    }
                                }
                            ]
                        },
                    }
                ],
                "labels": [
                    {
                        "properties": {
                            "show": expr_bool(
                                True
                            ),
                            "color": solid_color(
                                COLORS["ink"]
                            ),
                            "fontSize": expr_double(
                                8.5
                            ),
                            "fontFamily": expr_string(
                                FONT_SEMIBOLD
                            ),
                            "bold": expr_bool(
                                False
                            ),
                            "labelOverflow": expr_bool(
                                False
                            ),
                        }
                    }
                ],
                "zoom": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": (
                container_objects(
                    background=COLORS["surface"],
                    border=COLORS["border"],
                    radius=9,
                    show_title=True,
                    title=(
                        "Why Claims Were Denied"
                    ),
                    title_size=12.5,
                    subtitle=(
                        "Composition across fully comparable plans"
                    ),
                )
            ),
            "drillFilterOtherVisuals": (
                True
            ),
        },
    }


# =====================================================================
# INDEX button — no vertical-alignment guesswork
# Surface + precisely positioned label + transparent nav overlay.
# =====================================================================

def make_index_surface(
    *,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
) -> dict:
    return make_textbox(
        name="p01_index_surface",
        x=x,
        y=y,
        width=width,
        height=height,
        z=z,
        paragraphs=[
            paragraph(
                text_run(
                    " ",
                    size=6,
                    color=COLORS["blue"],
                )
            )
        ],
        background=COLORS["blue"],
        border=COLORS["blue_dark"],
        radius=9,
    )


def make_index_label(
    *,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
) -> dict:
    # Label height is deliberately smaller and geometrically centered
    # inside the button. This avoids Power BI textbox top-alignment.
    return make_textbox(
        name="p01_index_label",
        x=x,
        y=y,
        width=width,
        height=height,
        z=z,
        paragraphs=[
            paragraph(
                text_run(
                    "INDEX",
                    size=11,
                    color="#FFFFFF",
                    semibold=True,
                ),
                align="center",
            )
        ],
        background=None,
        border=None,
        radius=0,
    )


def make_index_nav(
    *,
    target_page_name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
) -> dict:
    return {
        "$schema": (
            VISUAL_SCHEMA
        ),
        "name": (
            "p01_index_nav"
        ),
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": 0,
        },
        "visual": {
            "visualType": (
                "actionButton"
            ),
            "objects": {
                "icon": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "outline": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "fill": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": {
                "title": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "background": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "border": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "visualHeader": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "visualLink": [
                    {
                        "properties": {
                            "show": expr_bool(
                                True
                            ),
                            "type": expr_string(
                                "PageNavigation"
                            ),
                            "navigationSection": expr_string(
                                target_page_name
                            ),
                            "showDefaultTooltip": expr_bool(
                                False
                            ),
                            "enabledTooltip": expr_string(
                                "Return to INDEX"
                            ),
                        }
                    }
                ],
            },
            "drillFilterOtherVisuals": (
                True
            ),
        },
    }


# =====================================================================
# Visual IO
# =====================================================================

def write_visual(
    page_dir: Path,
    visual: dict,
) -> Path:
    target = (
        page_dir
        / "visuals"
        / visual["name"]
        / "visual.json"
    )

    write_json(
        target,
        visual,
    )

    return target


# =====================================================================
# Executive Overview — V2 CLEAN REBUILD
# No page-sized textbox, no large decorative section surfaces.
# =====================================================================

def build_executive_page(
    page_dir: Path,
    index_page_name: str,
) -> None:
    visuals_dir = (
        page_dir
        / "visuals"
    )

    # Full destructive rebuild of visuals on this one page only.
    if visuals_dir.exists():
        shutil.rmtree(
            visuals_dir
        )

    visuals_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    # ---------------------------------------------------------------
    # Header
    # ---------------------------------------------------------------

    write_visual(
        page_dir,
        make_textbox(
            name="p01_context",
            x=52,
            y=31,
            width=290,
            height=43,
            z=100,
            paragraphs=[
                paragraph(
                    text_run(
                        "2026 PUF  •  EXPERIENCE YEAR 2024",
                        size=9,
                        color=COLORS["blue"],
                        semibold=True,
                    )
                )
            ],
            background=COLORS["blue_soft"],
            border=None,
            radius=8,
        ),
    )

    write_visual(
        page_dir,
        make_textbox(
            name="p01_page_title",
            x=390,
            y=18,
            width=1140,
            height=84,
            z=110,
            paragraphs=[
                paragraph(
                    text_run(
                        "Executive Overview",
                        size=29,
                        color=COLORS["ink"],
                        semibold=True,
                    )
                ),
                paragraph(
                    text_run(
                        (
                            "Market size → claims → network disparity → "
                            "denial reasons → appeals → governance"
                        ),
                        size=10.5,
                        color=COLORS["secondary"],
                    )
                ),
            ],
            background=None,
            border=None,
            radius=0,
        ),
    )

    index_x = 1648
    index_y = 25
    index_w = 220
    index_h = 56
    label_h = 22
    label_y = (
        index_y
        + (
            index_h
            - label_h
        )
        / 2
    )

    write_visual(
        page_dir,
        make_index_surface(
            x=index_x,
            y=index_y,
            width=index_w,
            height=index_h,
            z=130,
        ),
    )

    write_visual(
        page_dir,
        make_index_label(
            x=index_x,
            y=label_y,
            width=index_w,
            height=label_h,
            z=140,
        ),
    )

    write_visual(
        page_dir,
        make_index_nav(
            target_page_name=(
                index_page_name
            ),
            x=index_x,
            y=index_y,
            width=index_w,
            height=index_h,
            z=5000,
        ),
    )

    # ---------------------------------------------------------------
    # KPI strip
    # ---------------------------------------------------------------

    margin = 52
    gap = 16
    total_width = (
        CANVAS_WIDTH
        - 2 * margin
    )
    card_width = (
        total_width
        - 4 * gap
    ) / 5

    kpis = [
        (
            "p01_kpi_issuers",
            "Issuer Count",
            "Issuers",
            COLORS["ink"],
            COLORS["surface"],
        ),
        (
            "p01_kpi_plans",
            "Plan Count",
            "Plans",
            COLORS["ink"],
            COLORS["surface"],
        ),
        (
            "p01_kpi_claims",
            "Issuer Comparable Claims Received - Total",
            "Comparable Claims",
            COLORS["navy"],
            COLORS["blue_soft"],
        ),
        (
            "p01_kpi_denial_rate",
            "Issuer Comparable Overall Denial Rate",
            "Comparable Denial Rate",
            COLORS["blue"],
            COLORS["surface"],
        ),
        (
            "p01_kpi_appeal_rate",
            "Internal Appeal Overturn Rate",
            "Internal Appeal Overturn Rate",
            COLORS["teal"],
            COLORS["surface"],
        ),
    ]

    for idx, (
        visual_name,
        measure_name,
        title,
        value_color,
        background,
    ) in enumerate(
        kpis
    ):
        write_visual(
            page_dir,
            make_card(
                name=visual_name,
                measure_name=(
                    measure_name
                ),
                title=title,
                x=(
                    margin
                    + idx
                    * (
                        card_width
                        + gap
                    )
                ),
                y=120,
                width=card_width,
                height=134,
                z=400 + idx,
                tab_order=10 + idx,
                value_color=value_color,
                background=background,
                title_size=10.5,
                value_size=29,
            ),
        )

    # ---------------------------------------------------------------
    # Row 1 — Network + Comparable Claims
    # Small title textboxes only; no large decorative surfaces.
    # ---------------------------------------------------------------

    left_x = 52
    right_x = 968
    section_width = 900
    section_title_y = 274
    section_title_h = 40
    cards_y = 326
    cards_h = 238
    inner_gap = 16
    inner_card_width = (
        section_width
        - 2 * inner_gap
    ) / 3

    write_visual(
        page_dir,
        make_textbox(
            name="p01_network_title",
            x=left_x,
            y=section_title_y,
            width=section_width,
            height=section_title_h,
            z=200,
            paragraphs=[
                paragraph(
                    text_run(
                        "NETWORK DENIAL RATE",
                        size=12,
                        color=COLORS["ink"],
                        semibold=True,
                    )
                ),
                paragraph(
                    text_run(
                        "Issuer-level governed ratio of totals",
                        size=9,
                        color=COLORS["secondary"],
                    )
                ),
            ],
            background=None,
            border=None,
            radius=0,
        ),
    )

    network_cards = [
        (
            "p01_network_in",
            "Issuer In-Network Denial Rate",
            "In-Network",
            COLORS["navy"],
            COLORS["surface_alt"],
        ),
        (
            "p01_network_out",
            "Issuer Out-of-Network Denial Rate",
            "Out-of-Network",
            COLORS["blue"],
            COLORS["blue_soft"],
        ),
        (
            "p01_network_gap",
            "Issuer Network Denial Rate Gap",
            "Network Gap",
            COLORS["teal"],
            COLORS["teal_soft"],
        ),
    ]

    for idx, (
        name,
        measure,
        title,
        value_color,
        background,
    ) in enumerate(
        network_cards
    ):
        write_visual(
            page_dir,
            make_card(
                name=name,
                measure_name=measure,
                title=title,
                x=(
                    left_x
                    + idx
                    * (
                        inner_card_width
                        + inner_gap
                    )
                ),
                y=cards_y,
                width=inner_card_width,
                height=cards_h,
                z=500 + idx,
                tab_order=20 + idx,
                value_color=value_color,
                background=background,
                title_size=10.5,
                value_size=31,
            ),
        )

    write_visual(
        page_dir,
        make_textbox(
            name="p01_claims_title",
            x=right_x,
            y=section_title_y,
            width=section_width,
            height=section_title_h,
            z=200,
            paragraphs=[
                paragraph(
                    text_run(
                        "COMPARABLE CLAIMS",
                        size=12,
                        color=COLORS["ink"],
                        semibold=True,
                    )
                ),
                paragraph(
                    text_run(
                        (
                            "Same issuer population available for "
                            "received and denied claims"
                        ),
                        size=9,
                        color=COLORS["secondary"],
                    )
                ),
            ],
            background=None,
            border=None,
            radius=0,
        ),
    )

    claims_cards = [
        (
            "p01_claims_received",
            "Issuer Comparable Claims Received - Total",
            "Claims Received",
            COLORS["navy"],
            COLORS["surface_alt"],
        ),
        (
            "p01_claims_denied",
            "Issuer Comparable Claims Denied - Total",
            "Claims Denied",
            COLORS["blue"],
            COLORS["blue_soft"],
        ),
        (
            "p01_claims_rate",
            "Issuer Comparable Overall Denial Rate",
            "Overall Denial Rate",
            COLORS["teal"],
            COLORS["teal_soft"],
        ),
    ]

    for idx, (
        name,
        measure,
        title,
        value_color,
        background,
    ) in enumerate(
        claims_cards
    ):
        write_visual(
            page_dir,
            make_card(
                name=name,
                measure_name=measure,
                title=title,
                x=(
                    right_x
                    + idx
                    * (
                        inner_card_width
                        + inner_gap
                    )
                ),
                y=cards_y,
                width=inner_card_width,
                height=cards_h,
                z=500 + idx,
                tab_order=30 + idx,
                value_color=value_color,
                background=background,
                title_size=10.5,
                value_size=29,
            ),
        )

    # ---------------------------------------------------------------
    # Row 2 — denial reasons + appeals
    # ---------------------------------------------------------------

    write_visual(
        page_dir,
        make_reason_bar_chart(
            name=(
                "p01_denial_reason_chart"
            ),
            x=52,
            y=584,
            width=1120,
            height=360,
            z=700,
            tab_order=40,
        ),
    )

    write_visual(
        page_dir,
        make_textbox(
            name="p01_appeals_title",
            x=1188,
            y=584,
            width=680,
            height=44,
            z=200,
            paragraphs=[
                paragraph(
                    text_run(
                        "APPEALS OUTCOME",
                        size=12,
                        color=COLORS["ink"],
                        semibold=True,
                    )
                ),
                paragraph(
                    text_run(
                        "Filed volume and overturn rate",
                        size=9,
                        color=COLORS["secondary"],
                    )
                ),
            ],
            background=None,
            border=None,
            radius=0,
        ),
    )

    appeal_x = 1188
    appeal_y = 640
    appeal_w = 680
    appeal_gap = 16
    appeal_card_w = (
        appeal_w
        - appeal_gap
    ) / 2
    appeal_card_h = 144

    appeals = [
        (
            "p01_appeal_internal_rate",
            "Internal Appeal Overturn Rate",
            "Internal Overturn",
            COLORS["blue"],
            COLORS["blue_soft"],
            appeal_x,
            appeal_y,
        ),
        (
            "p01_appeal_external_rate",
            "External Appeal Overturn Rate",
            "External Overturn",
            COLORS["teal"],
            COLORS["teal_soft"],
            (
                appeal_x
                + appeal_card_w
                + appeal_gap
            ),
            appeal_y,
        ),
        (
            "p01_appeal_internal_filed",
            "Internal Appeals Filed",
            "Internal Filed",
            COLORS["ink"],
            COLORS["surface_alt"],
            appeal_x,
            (
                appeal_y
                + appeal_card_h
                + appeal_gap
            ),
        ),
        (
            "p01_appeal_external_filed",
            "External Appeals Filed",
            "External Filed",
            COLORS["ink"],
            COLORS["surface_alt"],
            (
                appeal_x
                + appeal_card_w
                + appeal_gap
            ),
            (
                appeal_y
                + appeal_card_h
                + appeal_gap
            ),
        ),
    ]

    for idx, (
        name,
        measure,
        title,
        value_color,
        background,
        x,
        y,
    ) in enumerate(
        appeals
    ):
        write_visual(
            page_dir,
            make_card(
                name=name,
                measure_name=measure,
                title=title,
                x=x,
                y=y,
                width=appeal_card_w,
                height=appeal_card_h,
                z=800 + idx,
                tab_order=50 + idx,
                value_color=value_color,
                background=background,
                title_size=10,
                value_size=24,
            ),
        )

    # ---------------------------------------------------------------
    # Governance strip
    # ---------------------------------------------------------------

    gov_y = 960
    gov_h = 88

    write_visual(
        page_dir,
        make_card(
            name="p01_dq_open",
            measure_name=(
                "Open DQ Exception Count"
            ),
            title=(
                "Open DQ Exceptions"
            ),
            x=52,
            y=gov_y,
            width=230,
            height=gov_h,
            z=600,
            tab_order=60,
            value_color=COLORS["amber"],
            background=COLORS["surface"],
            title_size=8.5,
            value_size=19,
        ),
    )

    write_visual(
        page_dir,
        make_card(
            name="p01_dq_known",
            measure_name=(
                "Known Source Exception Row Count"
            ),
            title=(
                "Known-Source Rows"
            ),
            x=298,
            y=gov_y,
            width=230,
            height=gov_h,
            z=600,
            tab_order=61,
            value_color=COLORS["red"],
            background=COLORS["surface"],
            title_size=8.5,
            value_size=19,
        ),
    )

    write_visual(
        page_dir,
        make_textbox(
            name="p01_governance_note",
            x=544,
            y=gov_y,
            width=1324,
            height=gov_h,
            z=600,
            paragraphs=[
                paragraph(
                    text_run(
                        (
                            "Governed context  •  Availability-aware KPIs  •  "
                            "Known-source anomaly preserved  •  "
                            "No clipping or imputation"
                        ),
                        size=9.5,
                        color=COLORS["secondary"],
                        semibold=True,
                    )
                )
            ],
            background=COLORS["surface_alt"],
            border=COLORS["border"],
            radius=9,
        ),
    )


# =====================================================================
# Other analytical pages — make INDEX label visually centered
# and remove legacy full-page canvas visuals.
# =====================================================================

def center_existing_page_header(
    page_dir: Path,
) -> bool:
    visual_files = list(
        (page_dir / "visuals")
        .rglob(
            "visual.json"
        )
    ) if (
        page_dir
        / "visuals"
    ).exists() else []

    changed = False

    for path in visual_files:
        payload = read_json(
            path
        )
        name = str(
            payload.get(
                "name",
                "",
            )
        )

        if not name.endswith(
            "_header"
        ):
            continue

        if (
            payload.get(
                "visual",
                {},
            ).get(
                "visualType"
            )
            != "textbox"
        ):
            continue

        paragraphs = (
            payload.get(
                "visual",
                {},
            )
            .get(
                "objects",
                {},
            )
            .get(
                "general",
                [{}],
            )[0]
            .get(
                "properties",
                {},
            )
            .get(
                "paragraphs",
                [],
            )
        )

        for p in paragraphs:
            p[
                "horizontalTextAlignment"
            ] = "center"

        write_json(
            path,
            payload,
        )

        changed = True

    return changed


def rebuild_existing_index_button_label(
    page_dir: Path,
) -> bool:
    """
    Existing shell pages created by the earlier script used one textbox for
    both surface and label. Power BI top-aligns the paragraph in edit mode.
    We split it into:
      - existing *_home_surface => blank blue surface
      - new *_home_label       => geometrically centered label
      - existing *_home_nav    => transparent navigation overlay
    """
    visuals_dir = (
        page_dir
        / "visuals"
    )

    if not visuals_dir.exists():
        return False

    surface_path = None
    nav_path = None

    for path in visuals_dir.rglob(
        "visual.json"
    ):
        payload = read_json(
            path
        )
        name = str(
            payload.get(
                "name",
                "",
            )
        )

        if name.endswith(
            "_home_surface"
        ):
            surface_path = (
                path
            )

        elif name.endswith(
            "_home_nav"
        ):
            nav_path = path

    if (
        surface_path is None
        or nav_path is None
    ):
        return False

    surface = read_json(
        surface_path
    )

    prefix = str(
        surface.get(
            "name"
        )
    ).replace(
        "_home_surface",
        "",
    )

    x = 1648
    y = 25
    width = 220
    height = 56
    label_h = 22
    label_y = (
        y
        + (
            height
            - label_h
        )
        / 2
    )

    surface[
        "position"
    ].update(
        {
            "x": x,
            "y": y,
            "width": width,
            "height": height,
        }
    )

    visual = surface[
        "visual"
    ]

    general = (
        visual
        .setdefault(
            "objects",
            {},
        )
        .setdefault(
            "general",
            [
                {
                    "properties": {}
                }
            ],
        )
    )

    if not general:
        general.append(
            {
                "properties": {}
            }
        )

    general[0].setdefault(
        "properties",
        {},
    )[
        "paragraphs"
    ] = [
        paragraph(
            text_run(
                " ",
                size=6,
                color=COLORS["blue"],
            )
        )
    ]

    vco = visual.setdefault(
        "visualContainerObjects",
        {},
    )

    vco["background"] = [
        {
            "properties": {
                "show": expr_bool(
                    True
                ),
                "color": solid_color(
                    COLORS["blue"]
                ),
                "transparency": expr_double(
                    0
                ),
            }
        }
    ]

    vco["border"] = [
        {
            "properties": {
                "show": expr_bool(
                    True
                ),
                "color": solid_color(
                    COLORS["blue_dark"]
                ),
                "width": expr_double(
                    1
                ),
                "radius": expr_double(
                    9
                ),
            }
        }
    ]

    write_json(
        surface_path,
        surface,
    )

    nav = read_json(
        nav_path
    )

    nav[
        "position"
    ].update(
        {
            "x": x,
            "y": y,
            "width": width,
            "height": height,
            "z": max(
                float(
                    nav.get(
                        "position",
                        {},
                    ).get(
                        "z",
                        0,
                    )
                    or 0
                ),
                5000,
            ),
        }
    )

    write_json(
        nav_path,
        nav,
    )

    label_name = (
        f"{prefix}_home_label"
    )

    label = make_textbox(
        name=label_name,
        x=x,
        y=label_y,
        width=width,
        height=label_h,
        z=4900,
        paragraphs=[
            paragraph(
                text_run(
                    "INDEX",
                    size=11,
                    color="#FFFFFF",
                    semibold=True,
                ),
                align="center",
            )
        ],
        background=None,
        border=None,
        radius=0,
    )

    write_visual(
        page_dir,
        label,
    )

    return True


# =====================================================================
# Validation
# =====================================================================

def walk_json(
    obj,
):
    if isinstance(
        obj,
        dict,
    ):
        yield obj

        for value in (
            obj.values()
        ):
            yield from walk_json(
                value
            )

    elif isinstance(
        obj,
        list,
    ):
        for value in obj:
            yield from walk_json(
                value
            )


def validate_binding_inventory(
    page_dir: Path,
) -> list[dict]:
    rows = []

    for path in sorted(
        (
            page_dir
            / "visuals"
        ).rglob(
            "visual.json"
        )
    ):
        payload = read_json(
            path
        )
        visual = payload.get(
            "visual",
            {},
        )

        if not visual.get(
            "query"
        ):
            continue

        fields = []

        for node in walk_json(
            visual.get(
                "query",
                {},
            )
        ):
            if "Measure" in node:
                m = node[
                    "Measure"
                ]

                prop = m.get(
                    "Property"
                )

                if prop:
                    fields.append(
                        (
                            "Measure",
                            "_Measures",
                            prop,
                        )
                    )

            if "Column" in node:
                c = node[
                    "Column"
                ]

                prop = c.get(
                    "Property"
                )

                entity = (
                    c.get(
                        "Expression",
                        {},
                    )
                    .get(
                        "SourceRef",
                        {},
                    )
                    .get(
                        "Entity"
                    )
                )

                if prop:
                    fields.append(
                        (
                            "Column",
                            entity,
                            prop,
                        )
                    )

        rows.append(
            {
                "VisualName": (
                    payload.get(
                        "name"
                    )
                ),
                "VisualType": (
                    visual.get(
                        "visualType"
                    )
                ),
                "Bindings": (
                    " | ".join(
                        f"{kind}:"
                        f"{entity}."
                        f"{field}"
                        for (
                            kind,
                            entity,
                            field,
                        )
                        in fields
                    )
                ),
            }
        )

    return rows


def validate_no_full_page_visuals(
    pages: dict,
) -> list[dict]:
    offenders = []

    for display_name, (
        page_dir,
        _,
    ) in pages.items():
        visuals_dir = (
            page_dir
            / "visuals"
        )

        if not visuals_dir.exists():
            continue

        for path in visuals_dir.rglob(
            "visual.json"
        ):
            payload = read_json(
                path
            )

            pos = payload.get(
                "position",
                {},
            )

            width = float(
                pos.get(
                    "width",
                    0,
                )
                or 0
            )
            height = float(
                pos.get(
                    "height",
                    0,
                )
                or 0
            )

            if (
                width
                >= CANVAS_WIDTH * 0.94
                and height
                >= CANVAS_HEIGHT * 0.94
            ):
                offenders.append(
                    {
                        "Page": (
                            display_name
                        ),
                        "VisualName": (
                            payload.get(
                                "name"
                            )
                        ),
                        "VisualType": (
                            payload.get(
                                "visual",
                                {},
                            ).get(
                                "visualType"
                            )
                        ),
                        "Width": width,
                        "Height": height,
                    }
                )

    return offenders


def validate_page_backgrounds(
    pages: dict,
) -> list[dict]:
    rows = []

    for display_name, (
        page_dir,
        _,
    ) in pages.items():
        payload = read_json(
            page_dir
            / "page.json"
        )

        background = (
            payload.get(
                "objects",
                {},
            )
            .get(
                "background",
                [],
            )
        )

        color = None
        transparency = None

        if background:
            props = (
                background[0]
                .get(
                    "properties",
                    {},
                )
            )

            color = (
                props.get(
                    "color",
                    {},
                )
                .get(
                    "solid",
                    {},
                )
                .get(
                    "color",
                    {},
                )
                .get(
                    "expr",
                    {},
                )
                .get(
                    "Literal",
                    {},
                )
                .get(
                    "Value"
                )
            )

            transparency = (
                props.get(
                    "transparency",
                    {},
                )
                .get(
                    "expr",
                    {},
                )
                .get(
                    "Literal",
                    {},
                )
                .get(
                    "Value"
                )
            )

        unsupported_root = [
            prop
            for prop in (
                "horizontalAlignment",
                "verticalAlignment",
            )
            if prop in payload
        ]

        rows.append(
            {
                "Page": display_name,
                "ExpectedColor": (
                    f"'{COLORS['canvas']}'"
                ),
                "ActualColor": color,
                "Transparency": transparency,
                "UnsupportedRootProperties": ", ".join(unsupported_root),
                "Status": (
                    "PASS"
                    if (
                        color
                        == f"'{COLORS['canvas']}'"
                        and transparency
                        == "0D"
                        and not unsupported_root
                    )
                    else "FAIL"
                ),
            }
        )

    return rows


def validate_json_tree() -> list[dict]:
    failures = []

    for path in sorted(
        (
            REPORT_DIR
            / "definition"
        ).rglob(
            "*.json"
        )
    ):
        try:
            read_json(
                path
            )

        except Exception as exc:
            failures.append(
                {
                    "RelativePath": (
                        str(
                            path.relative_to(
                                PROJECT_ROOT
                            )
                        )
                    ),
                    "Error": (
                        str(
                            exc
                        )
                    ),
                }
            )

    return failures


# =====================================================================
# Excel evidence formatting
# =====================================================================

def style_report(
    path: Path,
) -> None:
    from openpyxl import (
        load_workbook,
    )
    from openpyxl.styles import (
        Alignment,
        Font,
        PatternFill,
    )
    from openpyxl.utils import (
        get_column_letter,
    )

    wb = load_workbook(
        path
    )

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = (
            "A2"
        )

        if (
            ws.max_row
            and ws.max_column
        ):
            ws.auto_filter.ref = (
                ws.dimensions
            )

        for cell in ws[1]:
            cell.fill = (
                header_fill
            )
            cell.font = (
                header_font
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            max_len = 0

            for row_idx in range(
                1,
                min(
                    ws.max_row,
                    250,
                )
                + 1,
            ):
                value = ws.cell(
                    row=row_idx,
                    column=col_idx,
                ).value

                if value is not None:
                    max_len = max(
                        max_len,
                        len(
                            str(
                                value
                            )
                        ),
                    )

            ws.column_dimensions[
                get_column_letter(
                    col_idx
                )
            ].width = min(
                max(
                    max_len + 2,
                    10,
                ),
                70,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = (
                    Alignment(
                        vertical="top",
                        wrap_text=True,
                    )
                )

    wb.save(
        path
    )


# =====================================================================
# Main
# =====================================================================

def main() -> int:
    print("=" * 78)
    print(
        "Transparency in Coverage PUF — "
        "Executive Overview Clean Rebuild V3"
    )
    print("=" * 78)

    print(
        f"Script version : "
        f"{SCRIPT_VERSION}"
    )
    print(
        f"Report folder  : "
        f"{REPORT_DIR}"
    )
    print(
        f"Backup         : "
        f"{BACKUP_FILE}"
    )
    print(
        f"Review file    : "
        f"{REPORT_FILE}"
    )

    # ---------------------------------------------------------------
    # Hard safety gates
    # ---------------------------------------------------------------

    if power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is running. "
            "Save the PBIP and close Power BI Desktop completely "
            "before changing PBIR files."
        )

    if not REPORT_DIR.exists():
        raise FileNotFoundError(
            f"Report folder not found: "
            f"{REPORT_DIR}"
        )

    if not MEASURES_TMDL.exists():
        raise FileNotFoundError(
            f"Measure table not found: "
            f"{MEASURES_TMDL}"
        )

    if not DIM_METRIC_TMDL.exists():
        raise FileNotFoundError(
            f"DimMetric table not found: "
            f"{DIM_METRIC_TMDL}"
        )

    pages = find_pages()

    if set(
        pages
    ) != set(
        EXPECTED_PAGES
    ):
        raise RuntimeError(
            "Current report page inventory does not match "
            "the approved 11-page report shell. "
            "No changes made."
        )

    installed_measures = (
        parse_tmdl_measures(
            MEASURES_TMDL
        )
    )

    missing_measures = sorted(
        set(
            REQUIRED_MEASURES
        )
        - installed_measures
    )

    if missing_measures:
        raise RuntimeError(
            "Required governed measures are missing: "
            + ", ".join(
                missing_measures
            )
        )

    installed_metric_columns = (
        parse_tmdl_columns(
            DIM_METRIC_TMDL
        )
    )

    missing_metric_columns = sorted(
        REQUIRED_DIM_METRIC_COLUMNS
        - installed_metric_columns
    )

    if missing_metric_columns:
        raise RuntimeError(
            "Required DimMetric columns are missing: "
            + ", ".join(
                missing_metric_columns
            )
        )

    index_dir, index_page = (
        pages[
            "00 INDEX"
        ]
    )

    executive_dir, _ = (
        pages[
            "01 Executive Overview"
        ]
    )

    index_page_name = (
        index_page[
            "name"
        ]
    )

    # ---------------------------------------------------------------
    # Backup first
    # ---------------------------------------------------------------

    create_backup()
    before_hashes = (
        report_hashes()
    )

    # ---------------------------------------------------------------
    # GLOBAL FIX:
    # Native page backgrounds on every page; delete old full-page
    # canvas textboxes created by earlier iterations.
    # ---------------------------------------------------------------

    global_sanitation_rows = []

    for display_name, (
        page_dir,
        _,
    ) in pages.items():
        set_native_page_background(
            page_dir
        )

        removed = (
            remove_full_page_canvas_visuals(
                page_dir
            )
        )

        global_sanitation_rows.append(
            {
                "Page": display_name,
                "NativePageBackground": (
                    "YES"
                ),
                "RemovedLegacyCanvasVisuals": (
                    ", ".join(
                        removed
                    )
                    if removed
                    else ""
                ),
                "Status": (
                    "SANITIZED"
                ),
            }
        )

    # ---------------------------------------------------------------
    # Rebuild 01 from zero.
    # ---------------------------------------------------------------

    build_executive_page(
        executive_dir,
        index_page_name,
    )

    # ---------------------------------------------------------------
    # Future-page shell hardening:
    # centered titles + visibly centered INDEX labels on pages 02–10.
    # ---------------------------------------------------------------

    shell_rows = []

    for display_name in (
        EXPECTED_PAGES[2:]
    ):
        page_dir, _ = (
            pages[
                display_name
            ]
        )

        title_centered = (
            center_existing_page_header(
                page_dir
            )
        )

        index_centered = (
            rebuild_existing_index_button_label(
                page_dir
            )
        )

        shell_rows.append(
            {
                "Page": display_name,
                "TitleCentered": (
                    "YES"
                    if title_centered
                    else "NO_CHANGE"
                ),
                "IndexButtonRebuilt": (
                    "YES"
                    if index_centered
                    else "NO"
                ),
                "Status": (
                    "PATCHED"
                    if (
                        title_centered
                        or index_centered
                    )
                    else "UNCHANGED"
                ),
            }
        )

    # Refresh page objects after modifications.
    pages = find_pages()

    # ---------------------------------------------------------------
    # Validation
    # ---------------------------------------------------------------

    json_failures = (
        validate_json_tree()
    )

    page_background_rows = (
        validate_page_backgrounds(
            pages
        )
    )

    background_failures = sum(
        row[
            "Status"
        ] == "FAIL"
        for row
        in page_background_rows
    )

    full_page_offenders = (
        validate_no_full_page_visuals(
            pages
        )
    )

    binding_rows = (
        validate_binding_inventory(
            executive_dir
        )
    )

    exec_visual_files = sorted(
        (
            executive_dir
            / "visuals"
        ).rglob(
            "visual.json"
        )
    )

    expected_exec_names = {
        "p01_context",
        "p01_page_title",
        "p01_index_surface",
        "p01_index_label",
        "p01_index_nav",
        "p01_kpi_issuers",
        "p01_kpi_plans",
        "p01_kpi_claims",
        "p01_kpi_denial_rate",
        "p01_kpi_appeal_rate",
        "p01_network_title",
        "p01_network_in",
        "p01_network_out",
        "p01_network_gap",
        "p01_claims_title",
        "p01_claims_received",
        "p01_claims_denied",
        "p01_claims_rate",
        "p01_denial_reason_chart",
        "p01_appeals_title",
        "p01_appeal_internal_rate",
        "p01_appeal_external_rate",
        "p01_appeal_internal_filed",
        "p01_appeal_external_filed",
        "p01_dq_open",
        "p01_dq_known",
        "p01_governance_note",
    }

    actual_exec_names = set()

    for path in (
        exec_visual_files
    ):
        actual_exec_names.add(
            read_json(
                path
            ).get(
                "name"
            )
        )

    page_title_payload = read_json(
        executive_dir
        / "visuals"
        / "p01_page_title"
        / "visual.json"
    )

    title_paragraphs = (
        page_title_payload[
            "visual"
        ][
            "objects"
        ][
            "general"
        ][0][
            "properties"
        ][
            "paragraphs"
        ]
    )

    title_centered = all(
        p.get(
            "horizontalTextAlignment"
        )
        == "center"
        for p in title_paragraphs
    )

    index_surface = read_json(
        executive_dir
        / "visuals"
        / "p01_index_surface"
        / "visual.json"
    )

    index_label = read_json(
        executive_dir
        / "visuals"
        / "p01_index_label"
        / "visual.json"
    )

    surface_pos = (
        index_surface[
            "position"
        ]
    )

    label_pos = (
        index_label[
            "position"
        ]
    )

    expected_label_y = (
        surface_pos[
            "y"
        ]
        + (
            surface_pos[
                "height"
            ]
            - label_pos[
                "height"
            ]
        )
        / 2
    )

    label_text = (
        index_label[
            "visual"
        ][
            "objects"
        ][
            "general"
        ][0][
            "properties"
        ][
            "paragraphs"
        ][0][
            "textRuns"
        ][0][
            "value"
        ]
    )

    label_horizontal = (
        index_label[
            "visual"
        ][
            "objects"
        ][
            "general"
        ][0][
            "properties"
        ][
            "paragraphs"
        ][0][
            "horizontalTextAlignment"
        ]
    )

    index_geometrically_centered = (
        abs(
            float(
                label_pos[
                    "y"
                ]
            )
            - float(
                expected_label_y
            )
        )
        < 0.001
        and float(
            label_pos[
                "x"
            ]
        )
        == float(
            surface_pos[
                "x"
            ]
        )
        and float(
            label_pos[
                "width"
            ]
        )
        == float(
            surface_pos[
                "width"
            ]
        )
        and label_text
        == "INDEX"
        and label_horizontal
        == "center"
    )

    validation_rows = [
        {
            "CheckID": "V2-001",
            "TestName": "Approved page inventory",
            "Expected": len(
                EXPECTED_PAGES
            ),
            "Actual": len(
                pages
            ),
            "Status": (
                "PASS"
                if set(
                    pages
                )
                == set(
                    EXPECTED_PAGES
                )
                else "FAIL"
            ),
        },
        {
            "CheckID": "V2-002",
            "TestName": "Native page backgrounds",
            "Expected": 0,
            "Actual": (
                background_failures
            ),
            "Status": (
                "PASS"
                if background_failures
                == 0
                else "FAIL"
            ),
        },
        {
            "CheckID": "V2-003",
            "TestName": "Full-page selectable visuals",
            "Expected": 0,
            "Actual": len(
                full_page_offenders
            ),
            "Status": (
                "PASS"
                if not
                full_page_offenders
                else "FAIL"
            ),
        },
        {
            "CheckID": "V2-004",
            "TestName": "Executive visual inventory",
            "Expected": len(
                expected_exec_names
            ),
            "Actual": len(
                actual_exec_names
            ),
            "Status": (
                "PASS"
                if actual_exec_names
                == expected_exec_names
                else "FAIL"
            ),
        },
        {
            "CheckID": "V2-005",
            "TestName": "Executive bound visuals",
            "Expected": 18,
            "Actual": len(
                binding_rows
            ),
            "Status": (
                "PASS"
                if len(
                    binding_rows
                )
                == 18
                else "FAIL"
            ),
        },
        {
            "CheckID": "V2-006",
            "TestName": "Page title centered",
            "Expected": "center",
            "Actual": (
                "center"
                if title_centered
                else "not-centered"
            ),
            "Status": (
                "PASS"
                if title_centered
                else "FAIL"
            ),
        },
        {
            "CheckID": "V2-007",
            "TestName": "INDEX label geometric centering",
            "Expected": "centered",
            "Actual": (
                "centered"
                if index_geometrically_centered
                else "not-centered"
            ),
            "Status": (
                "PASS"
                if index_geometrically_centered
                else "FAIL"
            ),
        },
        {
            "CheckID": "V2-008",
            "TestName": "PBIR JSON parse failures",
            "Expected": 0,
            "Actual": len(
                json_failures
            ),
            "Status": (
                "PASS"
                if not
                json_failures
                else "FAIL"
            ),
        },
    ]

    failures = sum(
        row[
            "Status"
        ]
        == "FAIL"
        for row
        in validation_rows
    )

    # ---------------------------------------------------------------
    # Evidence
    # ---------------------------------------------------------------

    after_hashes = (
        report_hashes()
    )

    file_changes = []

    for rel_path in sorted(
        set(
            before_hashes
        )
        | set(
            after_hashes
        )
    ):
        before = before_hashes.get(
            rel_path
        )
        after = after_hashes.get(
            rel_path
        )

        if before == after:
            continue

        if (
            before is not None
            and after is not None
        ):
            change_state = (
                "MODIFIED"
            )

        elif (
            after is not None
        ):
            change_state = (
                "CREATED"
            )

        else:
            change_state = (
                "DELETED"
            )

        file_changes.append(
            {
                "RelativePath": (
                    rel_path
                ),
                "ChangeState": (
                    change_state
                ),
                "BeforeSHA256": (
                    before
                ),
                "AfterSHA256": (
                    after
                ),
            }
        )

    exec_inventory = []

    for path in (
        exec_visual_files
    ):
        payload = read_json(
            path
        )

        exec_inventory.append(
            {
                "VisualName": (
                    payload.get(
                        "name"
                    )
                ),
                "VisualType": (
                    payload.get(
                        "visual",
                        {},
                    ).get(
                        "visualType"
                    )
                ),
                "X": (
                    payload.get(
                        "position",
                        {},
                    ).get(
                        "x"
                    )
                ),
                "Y": (
                    payload.get(
                        "position",
                        {},
                    ).get(
                        "y"
                    )
                ),
                "Width": (
                    payload.get(
                        "position",
                        {},
                    ).get(
                        "width"
                    )
                ),
                "Height": (
                    payload.get(
                        "position",
                        {},
                    ).get(
                        "height"
                    )
                ),
                "RelativePath": (
                    str(
                        path.relative_to(
                            PROJECT_ROOT
                        )
                    )
                ),
            }
        )

    summary_df = pd.DataFrame(
        [
            (
                "BuildStatus",
                (
                    "PASS"
                    if failures
                    == 0
                    else "FAIL"
                ),
            ),
            (
                "ScriptVersion",
                SCRIPT_VERSION,
            ),
            (
                "BuiltAtLocal",
                datetime.now()
                .isoformat(
                    timespec="seconds"
                ),
            ),
            (
                "PageRebuilt",
                "01 Executive Overview",
            ),
            (
                "ExecutiveVisuals",
                len(
                    exec_visual_files
                ),
            ),
            (
                "ExecutiveBoundVisuals",
                len(
                    binding_rows
                ),
            ),
            (
                "FullPageSelectableVisuals",
                len(
                    full_page_offenders
                ),
            ),
            (
                "NativeBackgroundPages",
                len(
                    pages
                )
                - background_failures,
            ),
            (
                "PageTitleAlignment",
                "center",
            ),
            (
                "ChartCardTitleAlignment",
                "center",
            ),
            (
                "IndexButtonLabel",
                (
                    "Geometrically centered "
                    "inside blue button"
                ),
            ),
            (
                "ValidationChecks",
                len(
                    validation_rows
                ),
            ),
            (
                "ValidationFailures",
                failures,
            ),
            (
                "BackupFile",
                str(
                    BACKUP_FILE
                ),
            ),
            (
                "KeyDesignFix",
                (
                    "Page background uses supported page.json objects.background; "
                    "no full-page textbox/shape background and no unsupported "
                    "page-root alignment properties."
                ),
            ),
        ],
        columns=[
            "Item",
            "Value",
        ],
    )

    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        REPORT_FILE,
        engine="openpyxl",
    ) as writer:
        summary_df.to_excel(
            writer,
            sheet_name="00_Summary",
            index=False,
        )

        pd.DataFrame(
            validation_rows
        ).to_excel(
            writer,
            sheet_name="01_Validation",
            index=False,
        )

        pd.DataFrame(
            global_sanitation_rows
        ).to_excel(
            writer,
            sheet_name="02_Global_Background_Fix",
            index=False,
        )

        pd.DataFrame(
            shell_rows
        ).to_excel(
            writer,
            sheet_name="03_Future_Page_Shells",
            index=False,
        )

        pd.DataFrame(
            exec_inventory
        ).to_excel(
            writer,
            sheet_name="04_Executive_Visuals",
            index=False,
        )

        pd.DataFrame(
            binding_rows
        ).to_excel(
            writer,
            sheet_name="05_Bindings",
            index=False,
        )

        pd.DataFrame(
            page_background_rows
        ).to_excel(
            writer,
            sheet_name="06_Page_Backgrounds",
            index=False,
        )

        if full_page_offenders:
            pd.DataFrame(
                full_page_offenders
            ).to_excel(
                writer,
                sheet_name="07_FullPage_Offenders",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No full-page selectable "
                            "visual backgrounds remain."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="07_FullPage_Offenders",
                index=False,
            )

        pd.DataFrame(
            file_changes
        ).to_excel(
            writer,
            sheet_name="08_File_Changes",
            index=False,
        )

        if json_failures:
            pd.DataFrame(
                json_failures
            ).to_excel(
                writer,
                sheet_name="09_JSON_Failures",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No JSON parse failures"
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="09_JSON_Failures",
                index=False,
            )

    style_report(
        REPORT_FILE
    )

    print()
    print(
        "Executive Overview clean rebuild V3 completed."
    )
    print(
        f"Build status                 : "
        f"{'PASS' if failures == 0 else 'FAIL'}"
    )
    print(
        f"Executive visuals            : "
        f"{len(exec_visual_files)}"
    )
    print(
        f"Executive bound visuals      : "
        f"{len(binding_rows)}"
    )
    print(
        f"Full-page selectable visuals : "
        f"{len(full_page_offenders)}"
    )
    print(
        f"Native-background pages      : "
        f"{len(pages) - background_failures}/{len(pages)}"
    )
    print(
        "Page title alignment         : center"
    )
    print(
        "Chart/card title alignment   : center"
    )
    print(
        "INDEX label                  : geometric center"
    )
    print(
        f"Validation failures          : "
        f"{failures}"
    )
    print(
        f"Backup                       : "
        f"{BACKUP_FILE}"
    )
    print(
        f"Review evidence              : "
        f"{REPORT_FILE}"
    )

    if failures:
        print()
        print(
            "Offline validation failed. "
            "Do not open/save the PBIP until the evidence report is reviewed."
        )
        return 2

    print()
    print(
        "Next: open TransparencyInCoverage.pbip. "
        "Click a genuinely blank area on 01 Executive Overview: "
        "no page-sized textbox should be selectable or cover the charts. "
        "Then test INDEX navigation and inspect the rebuilt page."
    )

    return 0


if __name__ == "__main__":
    try:
        sys.exit(
            main()
        )

    except Exception as exc:
        print()
        print(
            "EXECUTIVE REBUILD V2 FAILED"
        )
        print(
            str(
                exc
            )
        )
        sys.exit(
            1
        )
