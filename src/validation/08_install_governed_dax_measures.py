from __future__ import annotations

from pathlib import Path
from datetime import datetime
import hashlib
import os
import re
import shutil
import subprocess
import sys
import textwrap
import uuid
import zipfile

import pandas as pd


SCRIPT_VERSION = "1.0.0"
MEASURE_TABLE = "_Measures"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
POWERBI_ROOT = PROJECT_ROOT / "powerbi"


def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, key_path) as key:
                raw_value, _ = winreg.QueryValueEx(key, "Desktop")
            return Path(os.path.expandvars(raw_value)).resolve()
        except Exception:
            pass
    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = get_windows_desktop() / "Transparency_PUF_Review"
REPORT_FILE = REVIEW_DIR / "08_DAX_INSTALL_REPORT.xlsx"
BACKUP_FILE = REVIEW_DIR / "08_PRE_DAX_MODEL_BACKUP.zip"


# ---------------------------------------------------------------------
# DAX building blocks
# ---------------------------------------------------------------------

def eligible_set(fact: str, entity_key: str, metric_key: int) -> str:
    return f"""
        CALCULATETABLE(
            VALUES({fact}[{entity_key}]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            {fact}[MetricKey] = {metric_key},
            {fact}[AvailabilityStatusKey] = 1
        )
    """.strip()


def intersect_sets(*sets: str) -> str:
    if not sets:
        raise ValueError("At least one set is required.")
    expr = sets[0]
    for item in sets[1:]:
        expr = f"INTERSECT(\n            {expr},\n            {item}\n        )"
    return expr


def ratio_measure(
    availability_fact: str,
    entity_key: str,
    fact_table: str,
    numerator_metric_key: int,
    denominator_metric_key: int,
    numerator_column: str,
    denominator_column: str,
) -> str:
    numerator_set = eligible_set(
        availability_fact, entity_key, numerator_metric_key
    )
    denominator_set = eligible_set(
        availability_fact, entity_key, denominator_metric_key
    )
    common = intersect_sets(numerator_set, denominator_set)

    return f"""
VAR EligibleEntities =
    {common}
VAR EligibleWithDenominator =
    FILTER(
        EligibleEntities,
        CALCULATE(
            SUM({fact_table}[{denominator_column}]),
            TREATAS(
                ROW("__EntityKey", [{entity_key}]),
                {fact_table}[{entity_key}]
            )
        ) > 0
    )
VAR Numerator =
    CALCULATE(
        SUM({fact_table}[{numerator_column}]),
        TREATAS(EligibleEntities, {fact_table}[{entity_key}])
    )
VAR Denominator =
    CALCULATE(
        SUM({fact_table}[{denominator_column}]),
        TREATAS(EligibleEntities, {fact_table}[{entity_key}])
    )
RETURN
    DIVIDE(Numerator, Denominator)
    """.strip()


# The generic helper above intentionally avoids being used for the final
# DAX because FILTER/ROW over a single-column table is easy to make
# unnecessarily complex. Final measures use simpler common-set logic.


MEASURES: list[dict[str, str]] = []


def add_measure(
    name: str,
    dax: str,
    fmt: str,
    folder: str,
    hidden: bool = False,
) -> None:
    MEASURES.append(
        {
            "name": name,
            "dax": textwrap.dedent(dax).strip(),
            "format": fmt,
            "folder": folder,
            "hidden": hidden,
        }
    )


# ---------------------------------------------------------------------
# 01 Population
# ---------------------------------------------------------------------

add_measure(
    "Issuer Count",
    "DISTINCTCOUNT(DimIssuer[IssuerID])",
    "#,##0",
    "01 Population",
)

add_measure(
    "Plan Count",
    "DISTINCTCOUNT(DimPlan[PlanID])",
    "#,##0",
    "01 Population",
)


# ---------------------------------------------------------------------
# 02 Claims — issuer primary
# ---------------------------------------------------------------------

add_measure(
    "Issuer Claims Received - In Network",
    """
    VAR Eligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 2,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    RETURN
        CALCULATE(
            SUM(FactIssuerTransparency[ClaimsReceivedInNetwork]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    """,
    "#,##0",
    "02 Claims",
)

add_measure(
    "Issuer Claims Received - Out of Network",
    """
    VAR Eligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 1,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    RETURN
        CALCULATE(
            SUM(FactIssuerTransparency[ClaimsReceivedOutOfNetwork]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    """,
    "#,##0",
    "02 Claims",
)

add_measure(
    "Issuer Comparable Claims Received - Total",
    """
    VAR InEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 2,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR OutEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 1,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR Eligible = INTERSECT(InEligible, OutEligible)
    RETURN
        CALCULATE(
            SUM(FactIssuerTransparency[ClaimsReceivedInNetwork])
                + SUM(FactIssuerTransparency[ClaimsReceivedOutOfNetwork]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    """,
    "#,##0",
    "02 Claims",
)


# ---------------------------------------------------------------------
# 03 Denials — issuer primary
# ---------------------------------------------------------------------

add_measure(
    "Issuer Claims Denied - In Network",
    """
    VAR Eligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 4,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    RETURN
        CALCULATE(
            SUM(FactIssuerTransparency[ClaimsDeniedInNetwork]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    """,
    "#,##0",
    "03 Denials",
)

add_measure(
    "Issuer Claims Denied - Out of Network",
    """
    VAR Eligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 3,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    RETURN
        CALCULATE(
            SUM(FactIssuerTransparency[ClaimsDeniedOutOfNetwork]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    """,
    "#,##0",
    "03 Denials",
)

add_measure(
    "Issuer Comparable Claims Denied - Total",
    """
    VAR InEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 4,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR OutEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 3,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR Eligible = INTERSECT(InEligible, OutEligible)
    RETURN
        CALCULATE(
            SUM(FactIssuerTransparency[ClaimsDeniedInNetwork])
                + SUM(FactIssuerTransparency[ClaimsDeniedOutOfNetwork]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    """,
    "#,##0",
    "03 Denials",
)

add_measure(
    "Issuer In-Network Denial Rate",
    """
    VAR DeniedEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 4,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR ReceivedEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 2,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR Eligible = INTERSECT(DeniedEligible, ReceivedEligible)
    VAR Numerator =
        CALCULATE(
            SUM(FactIssuerTransparency[ClaimsDeniedInNetwork]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    VAR Denominator =
        CALCULATE(
            SUM(FactIssuerTransparency[ClaimsReceivedInNetwork]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    RETURN DIVIDE(Numerator, Denominator)
    """,
    "0.00%",
    "03 Denials",
)

add_measure(
    "Issuer Out-of-Network Denial Rate",
    """
    VAR DeniedEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 3,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR ReceivedEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 1,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR Eligible = INTERSECT(DeniedEligible, ReceivedEligible)
    VAR Numerator =
        CALCULATE(
            SUM(FactIssuerTransparency[ClaimsDeniedOutOfNetwork]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    VAR Denominator =
        CALCULATE(
            SUM(FactIssuerTransparency[ClaimsReceivedOutOfNetwork]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    RETURN DIVIDE(Numerator, Denominator)
    """,
    "0.00%",
    "03 Denials",
)

add_measure(
    "Issuer Comparable Overall Denial Rate",
    """
    VAR RIn =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 2,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR ROut =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 1,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR DIn =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 4,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR DOut =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 3,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR Eligible =
        INTERSECT(INTERSECT(RIn, ROut), INTERSECT(DIn, DOut))
    VAR Numerator =
        CALCULATE(
            SUM(FactIssuerTransparency[ClaimsDeniedInNetwork])
                + SUM(FactIssuerTransparency[ClaimsDeniedOutOfNetwork]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    VAR Denominator =
        CALCULATE(
            SUM(FactIssuerTransparency[ClaimsReceivedInNetwork])
                + SUM(FactIssuerTransparency[ClaimsReceivedOutOfNetwork]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    RETURN DIVIDE(Numerator, Denominator)
    """,
    "0.00%",
    "03 Denials",
)


# ---------------------------------------------------------------------
# 04 Network Comparison
# ---------------------------------------------------------------------

add_measure(
    "Issuer Network Denial Rate Gap",
    "[Issuer Out-of-Network Denial Rate] - [Issuer In-Network Denial Rate]",
    "0.00%",
    "04 Network Comparison",
)


# ---------------------------------------------------------------------
# 05 Resubmissions
# ---------------------------------------------------------------------

add_measure(
    "Issuer Resubmission Events per 100 Denied - In Network",
    """
    VAR ResubEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 6,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR DeniedEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 4,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR Eligible = INTERSECT(ResubEligible, DeniedEligible)
    VAR Numerator =
        CALCULATE(
            SUM(FactIssuerTransparency[ClaimsResubmittedInNetwork]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    VAR Denominator =
        CALCULATE(
            SUM(FactIssuerTransparency[ClaimsDeniedInNetwork]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    RETURN DIVIDE(Numerator, Denominator) * 100
    """,
    "0.00",
    "05 Resubmissions",
)

add_measure(
    "Issuer Resubmission Events per 100 Denied - Out of Network",
    """
    VAR ResubEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 5,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR DeniedEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 3,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR Eligible = INTERSECT(ResubEligible, DeniedEligible)
    VAR Numerator =
        CALCULATE(
            SUM(FactIssuerTransparency[ClaimsResubmittedOutOfNetwork]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    VAR Denominator =
        CALCULATE(
            SUM(FactIssuerTransparency[ClaimsDeniedOutOfNetwork]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    RETURN DIVIDE(Numerator, Denominator) * 100
    """,
    "0.00",
    "05 Resubmissions",
)

add_measure(
    "Issuer Comparable Total Resubmission Events",
    """
    VAR InEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 6,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR OutEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 5,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR Eligible = INTERSECT(InEligible, OutEligible)
    RETURN
        CALCULATE(
            SUM(FactIssuerTransparency[ClaimsResubmittedInNetwork])
                + SUM(FactIssuerTransparency[ClaimsResubmittedOutOfNetwork]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    """,
    "#,##0",
    "05 Resubmissions",
)


# ---------------------------------------------------------------------
# 06 Appeals
# ---------------------------------------------------------------------

add_measure(
    "Internal Appeals Filed",
    """
    VAR Eligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 7,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    RETURN
        CALCULATE(
            SUM(FactIssuerTransparency[InternalAppealsFiled]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    """,
    "#,##0",
    "06 Appeals",
)

add_measure(
    "Internal Appeals Overturned",
    """
    VAR Eligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 8,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    RETURN
        CALCULATE(
            SUM(FactIssuerTransparency[InternalAppealsOverturned]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    """,
    "#,##0",
    "06 Appeals",
)

add_measure(
    "Internal Appeal Overturn Rate",
    """
    VAR FiledEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 7,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR OverturnedEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 8,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR Eligible = INTERSECT(FiledEligible, OverturnedEligible)
    VAR Numerator =
        CALCULATE(
            SUM(FactIssuerTransparency[InternalAppealsOverturned]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    VAR Denominator =
        CALCULATE(
            SUM(FactIssuerTransparency[InternalAppealsFiled]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    RETURN DIVIDE(Numerator, Denominator)
    """,
    "0.00%",
    "06 Appeals",
)

add_measure(
    "External Appeals Filed",
    """
    VAR Eligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 10,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    RETURN
        CALCULATE(
            SUM(FactIssuerTransparency[ExternalAppealsFiled]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    """,
    "#,##0",
    "06 Appeals",
)

add_measure(
    "External Appeals Overturned",
    """
    VAR Eligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 11,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    RETURN
        CALCULATE(
            SUM(FactIssuerTransparency[ExternalAppealsOverturned]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    """,
    "#,##0",
    "06 Appeals",
)

add_measure(
    "External Appeal Overturn Rate",
    """
    VAR FiledEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 10,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR OverturnedEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 11,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR Eligible = INTERSECT(FiledEligible, OverturnedEligible)
    VAR Numerator =
        CALCULATE(
            SUM(FactIssuerTransparency[ExternalAppealsOverturned]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    VAR Denominator =
        CALCULATE(
            SUM(FactIssuerTransparency[ExternalAppealsFiled]),
            TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
        )
    RETURN DIVIDE(Numerator, Denominator)
    """,
    "0.00%",
    "06 Appeals",
)


# ---------------------------------------------------------------------
# 07 Denial Reasons — selected reason from DimMetric
# ---------------------------------------------------------------------

add_measure(
    "Denial Reason Count",
    """
    VAR MetricKey = SELECTEDVALUE(DimMetric[MetricKey])
    VAR EligiblePlans =
        CALCULATETABLE(
            VALUES(FactPlanMetricAvailability[PlanKey]),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactPlanMetricAvailability[AvailabilityStatusKey] = 1
        )
    RETURN
        SWITCH(
            TRUE(),
            MetricKey = 19,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedReferralOrPriorAuthorization]),
                    TREATAS(EligiblePlans, FactPlanTransparency[PlanKey])
                ),
            MetricKey = 20,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedDueToOutOfNetworkProvider]),
                    TREATAS(EligiblePlans, FactPlanTransparency[PlanKey])
                ),
            MetricKey = 21,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedServicesExcluded]),
                    TREATAS(EligiblePlans, FactPlanTransparency[PlanKey])
                ),
            MetricKey = 22,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedNotMedicallyNecessaryExclBH]),
                    TREATAS(EligiblePlans, FactPlanTransparency[PlanKey])
                ),
            MetricKey = 23,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedNotMedicallyNecessaryBHOnly]),
                    TREATAS(EligiblePlans, FactPlanTransparency[PlanKey])
                ),
            MetricKey = 24,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedBenefitLimitReached]),
                    TREATAS(EligiblePlans, FactPlanTransparency[PlanKey])
                ),
            MetricKey = 25,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedMemberNotCovered]),
                    TREATAS(EligiblePlans, FactPlanTransparency[PlanKey])
                ),
            MetricKey = 26,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedInvestigationalExperimentalCosmetic]),
                    TREATAS(EligiblePlans, FactPlanTransparency[PlanKey])
                ),
            MetricKey = 27,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedAdministrativeReason]),
                    TREATAS(EligiblePlans, FactPlanTransparency[PlanKey])
                ),
            MetricKey = 28,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedOther]),
                    TREATAS(EligiblePlans, FactPlanTransparency[PlanKey])
                ),
            BLANK()
        )
    """,
    "#,##0",
    "07 Denial Reasons",
)

add_measure(
    "Comparable Denial Reason Composition %",
    """
    VAR MetricKey = SELECTEDVALUE(DimMetric[MetricKey])
    VAR FullyComparablePlans =
        SELECTCOLUMNS(
            FILTER(
                CALCULATETABLE(
                    SUMMARIZE(
                        FactPlanMetricAvailability,
                        FactPlanMetricAvailability[PlanKey],
                        "__AvailableReasonMetrics",
                            DISTINCTCOUNT(FactPlanMetricAvailability[MetricKey])
                    ),
                    REMOVEFILTERS(DimMetric),
                    REMOVEFILTERS(DimAvailabilityStatus),
                    FactPlanMetricAvailability[MetricKey] >= 19,
                    FactPlanMetricAvailability[MetricKey] <= 28,
                    FactPlanMetricAvailability[AvailabilityStatusKey] = 1
                ),
                [__AvailableReasonMetrics] = 10
            ),
            "PlanKey", FactPlanMetricAvailability[PlanKey]
        )
    VAR SelectedReason =
        SWITCH(
            TRUE(),
            MetricKey = 19,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedReferralOrPriorAuthorization]),
                    TREATAS(FullyComparablePlans, FactPlanTransparency[PlanKey])
                ),
            MetricKey = 20,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedDueToOutOfNetworkProvider]),
                    TREATAS(FullyComparablePlans, FactPlanTransparency[PlanKey])
                ),
            MetricKey = 21,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedServicesExcluded]),
                    TREATAS(FullyComparablePlans, FactPlanTransparency[PlanKey])
                ),
            MetricKey = 22,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedNotMedicallyNecessaryExclBH]),
                    TREATAS(FullyComparablePlans, FactPlanTransparency[PlanKey])
                ),
            MetricKey = 23,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedNotMedicallyNecessaryBHOnly]),
                    TREATAS(FullyComparablePlans, FactPlanTransparency[PlanKey])
                ),
            MetricKey = 24,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedBenefitLimitReached]),
                    TREATAS(FullyComparablePlans, FactPlanTransparency[PlanKey])
                ),
            MetricKey = 25,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedMemberNotCovered]),
                    TREATAS(FullyComparablePlans, FactPlanTransparency[PlanKey])
                ),
            MetricKey = 26,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedInvestigationalExperimentalCosmetic]),
                    TREATAS(FullyComparablePlans, FactPlanTransparency[PlanKey])
                ),
            MetricKey = 27,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedAdministrativeReason]),
                    TREATAS(FullyComparablePlans, FactPlanTransparency[PlanKey])
                ),
            MetricKey = 28,
                CALCULATE(
                    SUM(FactPlanTransparency[DeniedOther]),
                    TREATAS(FullyComparablePlans, FactPlanTransparency[PlanKey])
                ),
            BLANK()
        )
    VAR AllReasons =
        CALCULATE(
            SUM(FactPlanTransparency[DeniedReferralOrPriorAuthorization])
                + SUM(FactPlanTransparency[DeniedDueToOutOfNetworkProvider])
                + SUM(FactPlanTransparency[DeniedServicesExcluded])
                + SUM(FactPlanTransparency[DeniedNotMedicallyNecessaryExclBH])
                + SUM(FactPlanTransparency[DeniedNotMedicallyNecessaryBHOnly])
                + SUM(FactPlanTransparency[DeniedBenefitLimitReached])
                + SUM(FactPlanTransparency[DeniedMemberNotCovered])
                + SUM(FactPlanTransparency[DeniedInvestigationalExperimentalCosmetic])
                + SUM(FactPlanTransparency[DeniedAdministrativeReason])
                + SUM(FactPlanTransparency[DeniedOther]),
            TREATAS(FullyComparablePlans, FactPlanTransparency[PlanKey])
        )
    RETURN
        IF(
            MetricKey >= 19 && MetricKey <= 28,
            DIVIDE(SelectedReason, AllReasons),
            BLANK()
        )
    """,
    "0.00%",
    "07 Denial Reasons",
)


# ---------------------------------------------------------------------
# 08 Enrollment
# ---------------------------------------------------------------------

add_measure(
    "Reported Average Monthly Enrollment",
    """
    VAR Eligible =
        CALCULATETABLE(
            VALUES(FactPlanMetricAvailability[PlanKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactPlanMetricAvailability[MetricKey] = 29,
            FactPlanMetricAvailability[AvailabilityStatusKey] = 1
        )
    RETURN
        CALCULATE(
            SUM(FactPlanTransparency[AverageMonthlyEnrollment]),
            TREATAS(Eligible, FactPlanTransparency[PlanKey])
        )
    """,
    "#,##0.00",
    "08 Enrollment",
)

add_measure(
    "Reported Average Monthly Disenrollment",
    """
    VAR Eligible =
        CALCULATETABLE(
            VALUES(FactPlanMetricAvailability[PlanKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactPlanMetricAvailability[MetricKey] = 30,
            FactPlanMetricAvailability[AvailabilityStatusKey] = 1
        )
    RETURN
        CALCULATE(
            SUM(FactPlanTransparency[AverageMonthlyDisenrollment]),
            TREATAS(Eligible, FactPlanTransparency[PlanKey])
        )
    """,
    "#,##0.00",
    "08 Enrollment",
)

add_measure(
    "Comparable Disenrollment-to-Enrollment Ratio",
    """
    VAR EnrollmentEligible =
        CALCULATETABLE(
            VALUES(FactPlanMetricAvailability[PlanKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactPlanMetricAvailability[MetricKey] = 29,
            FactPlanMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR DisenrollmentEligible =
        CALCULATETABLE(
            VALUES(FactPlanMetricAvailability[PlanKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactPlanMetricAvailability[MetricKey] = 30,
            FactPlanMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR Eligible = INTERSECT(EnrollmentEligible, DisenrollmentEligible)
    VAR Numerator =
        CALCULATE(
            SUM(FactPlanTransparency[AverageMonthlyDisenrollment]),
            TREATAS(Eligible, FactPlanTransparency[PlanKey])
        )
    VAR Denominator =
        CALCULATE(
            SUM(FactPlanTransparency[AverageMonthlyEnrollment]),
            TREATAS(Eligible, FactPlanTransparency[PlanKey])
        )
    RETURN DIVIDE(Numerator, Denominator)
    """,
    "0.00%",
    "08 Enrollment",
)


# ---------------------------------------------------------------------
# 09 Availability — requires one selected metric
# ---------------------------------------------------------------------

add_measure(
    "Available Entity Count",
    """
    VAR EntityLevel = SELECTEDVALUE(DimMetric[EntityLevel])
    RETURN
        SWITCH(
            EntityLevel,
            "ISSUER",
                CALCULATE(
                    COUNTROWS(FactIssuerMetricAvailability),
                    REMOVEFILTERS(DimAvailabilityStatus),
                    FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
                ),
            "PLAN",
                CALCULATE(
                    COUNTROWS(FactPlanMetricAvailability),
                    REMOVEFILTERS(DimAvailabilityStatus),
                    FactPlanMetricAvailability[AvailabilityStatusKey] = 1
                ),
            BLANK()
        )
    """,
    "#,##0",
    "09 Availability",
)

add_measure(
    "Applicable Entity Count",
    """
    VAR EntityLevel = SELECTEDVALUE(DimMetric[EntityLevel])
    RETURN
        SWITCH(
            EntityLevel,
            "ISSUER",
                CALCULATE(
                    COUNTROWS(FactIssuerMetricAvailability),
                    REMOVEFILTERS(DimAvailabilityStatus),
                    FactIssuerMetricAvailability[AvailabilityStatusKey] IN {1, 2, 3, 7, 8}
                ),
            "PLAN",
                CALCULATE(
                    COUNTROWS(FactPlanMetricAvailability),
                    REMOVEFILTERS(DimAvailabilityStatus),
                    FactPlanMetricAvailability[AvailabilityStatusKey] IN {1, 2, 3, 7, 8}
                ),
            BLANK()
        )
    """,
    "#,##0",
    "09 Availability",
)

add_measure(
    "Data Availability Rate",
    "DIVIDE([Available Entity Count], [Applicable Entity Count])",
    "0.00%",
    "09 Availability",
)

add_measure(
    "Suppression Rate",
    """
    VAR EntityLevel = SELECTEDVALUE(DimMetric[EntityLevel])
    VAR Suppressed =
        SWITCH(
            EntityLevel,
            "ISSUER",
                CALCULATE(
                    COUNTROWS(FactIssuerMetricAvailability),
                    REMOVEFILTERS(DimAvailabilityStatus),
                    FactIssuerMetricAvailability[AvailabilityStatusKey] = 3
                ),
            "PLAN",
                CALCULATE(
                    COUNTROWS(FactPlanMetricAvailability),
                    REMOVEFILTERS(DimAvailabilityStatus),
                    FactPlanMetricAvailability[AvailabilityStatusKey] = 3
                ),
            BLANK()
        )
    RETURN DIVIDE(Suppressed, [Applicable Entity Count])
    """,
    "0.00%",
    "09 Availability",
)

add_measure(
    "Unavailable Rate",
    """
    VAR EntityLevel = SELECTEDVALUE(DimMetric[EntityLevel])
    VAR Unavailable =
        SWITCH(
            EntityLevel,
            "ISSUER",
                CALCULATE(
                    COUNTROWS(FactIssuerMetricAvailability),
                    REMOVEFILTERS(DimAvailabilityStatus),
                    FactIssuerMetricAvailability[AvailabilityStatusKey] = 2
                ),
            "PLAN",
                CALCULATE(
                    COUNTROWS(FactPlanMetricAvailability),
                    REMOVEFILTERS(DimAvailabilityStatus),
                    FactPlanMetricAvailability[AvailabilityStatusKey] = 2
                ),
            BLANK()
        )
    RETURN DIVIDE(Unavailable, [Applicable Entity Count])
    """,
    "0.00%",
    "09 Availability",
)

add_measure(
    "Structural Non-Applicability Count",
    """
    VAR EntityLevel = SELECTEDVALUE(DimMetric[EntityLevel])
    RETURN
        SWITCH(
            EntityLevel,
            "ISSUER",
                CALCULATE(
                    COUNTROWS(FactIssuerMetricAvailability),
                    REMOVEFILTERS(DimAvailabilityStatus),
                    FactIssuerMetricAvailability[AvailabilityStatusKey] IN {4, 5}
                ),
            "PLAN",
                CALCULATE(
                    COUNTROWS(FactPlanMetricAvailability),
                    REMOVEFILTERS(DimAvailabilityStatus),
                    FactPlanMetricAvailability[AvailabilityStatusKey] IN {4, 5}
                ),
            BLANK()
        )
    """,
    "#,##0",
    "09 Availability",
)


# ---------------------------------------------------------------------
# 10 Data Quality
# ---------------------------------------------------------------------

add_measure(
    "Open DQ Exception Count",
    """
    CALCULATE(
        COUNTROWS(DQExceptionRegister),
        DQExceptionRegister[ResolutionState] = "OPEN_REVIEW"
    )
    """,
    "#,##0",
    "10 Data Quality",
)

add_measure(
    "Known Source Exception Row Count",
    """
    CALCULATE(
        COUNTROWS(DQExceptionRegister),
        DQExceptionRegister[IsKnownSourceException] = TRUE()
    )
    """,
    "#,##0",
    "10 Data Quality",
)

add_measure(
    "Known Source Exception Entity Count",
    """
    CALCULATE(
        DISTINCTCOUNT(DQExceptionRegister[EntityKey]),
        DQExceptionRegister[IsKnownSourceException] = TRUE()
    )
    """,
    "#,##0",
    "10 Data Quality",
)

add_measure(
    "Current Context DQ Exception Count",
    """
    VAR IssuerIDs = VALUES(DimIssuer[IssuerID])
    VAR PlanIDs = VALUES(DimPlan[PlanID])
    VAR IssuerDQ =
        CALCULATE(
            COUNTROWS(DQExceptionRegister),
            DQExceptionRegister[Scope] = "ISSUER",
            TREATAS(IssuerIDs, DQExceptionRegister[EntityKey])
        )
    VAR PlanDQ =
        CALCULATE(
            COUNTROWS(DQExceptionRegister),
            DQExceptionRegister[Scope] = "PLAN",
            TREATAS(PlanIDs, DQExceptionRegister[EntityKey])
        )
    RETURN IssuerDQ + PlanDQ
    """,
    "#,##0",
    "10 Data Quality",
)

add_measure(
    "Current Context Has DQ Flag",
    'IF([Current Context DQ Exception Count] > 0, 1, 0)',
    "0",
    "10 Data Quality",
)

# ---------------------------------------------------------------------
# 11 Plan-level analytical counterparts
# ---------------------------------------------------------------------

add_measure(
    "Plan In-Network Denial Rate",
    """
    VAR DeniedEligible =
        CALCULATETABLE(
            VALUES(FactPlanMetricAvailability[PlanKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactPlanMetricAvailability[MetricKey] = 16,
            FactPlanMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR ReceivedEligible =
        CALCULATETABLE(
            VALUES(FactPlanMetricAvailability[PlanKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactPlanMetricAvailability[MetricKey] = 14,
            FactPlanMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR Eligible = INTERSECT(DeniedEligible, ReceivedEligible)
    VAR Numerator =
        CALCULATE(
            SUM(FactPlanTransparency[ClaimsDeniedInNetwork]),
            TREATAS(Eligible, FactPlanTransparency[PlanKey])
        )
    VAR Denominator =
        CALCULATE(
            SUM(FactPlanTransparency[ClaimsReceivedInNetwork]),
            TREATAS(Eligible, FactPlanTransparency[PlanKey])
        )
    RETURN DIVIDE(Numerator, Denominator)
    """,
    "0.00%",
    "11 Plan Analytics",
)

add_measure(
    "Plan Out-of-Network Denial Rate",
    """
    VAR DeniedEligible =
        CALCULATETABLE(
            VALUES(FactPlanMetricAvailability[PlanKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactPlanMetricAvailability[MetricKey] = 15,
            FactPlanMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR ReceivedEligible =
        CALCULATETABLE(
            VALUES(FactPlanMetricAvailability[PlanKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactPlanMetricAvailability[MetricKey] = 13,
            FactPlanMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR Eligible = INTERSECT(DeniedEligible, ReceivedEligible)
    VAR Numerator =
        CALCULATE(
            SUM(FactPlanTransparency[ClaimsDeniedOutOfNetwork]),
            TREATAS(Eligible, FactPlanTransparency[PlanKey])
        )
    VAR Denominator =
        CALCULATE(
            SUM(FactPlanTransparency[ClaimsReceivedOutOfNetwork]),
            TREATAS(Eligible, FactPlanTransparency[PlanKey])
        )
    RETURN DIVIDE(Numerator, Denominator)
    """,
    "0.00%",
    "11 Plan Analytics",
)

add_measure(
    "Plan Comparable Overall Denial Rate",
    """
    VAR RIn =
        CALCULATETABLE(
            VALUES(FactPlanMetricAvailability[PlanKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactPlanMetricAvailability[MetricKey] = 14,
            FactPlanMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR ROut =
        CALCULATETABLE(
            VALUES(FactPlanMetricAvailability[PlanKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactPlanMetricAvailability[MetricKey] = 13,
            FactPlanMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR DIn =
        CALCULATETABLE(
            VALUES(FactPlanMetricAvailability[PlanKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactPlanMetricAvailability[MetricKey] = 16,
            FactPlanMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR DOut =
        CALCULATETABLE(
            VALUES(FactPlanMetricAvailability[PlanKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactPlanMetricAvailability[MetricKey] = 15,
            FactPlanMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR Eligible =
        INTERSECT(INTERSECT(RIn, ROut), INTERSECT(DIn, DOut))
    VAR Numerator =
        CALCULATE(
            SUM(FactPlanTransparency[ClaimsDeniedInNetwork])
                + SUM(FactPlanTransparency[ClaimsDeniedOutOfNetwork]),
            TREATAS(Eligible, FactPlanTransparency[PlanKey])
        )
    VAR Denominator =
        CALCULATE(
            SUM(FactPlanTransparency[ClaimsReceivedInNetwork])
                + SUM(FactPlanTransparency[ClaimsReceivedOutOfNetwork]),
            TREATAS(Eligible, FactPlanTransparency[PlanKey])
        )
    RETURN DIVIDE(Numerator, Denominator)
    """,
    "0.00%",
    "11 Plan Analytics",
)

add_measure(
    "Plan Resubmission Events per 100 Denied - In Network",
    """
    VAR ResubEligible =
        CALCULATETABLE(
            VALUES(FactPlanMetricAvailability[PlanKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactPlanMetricAvailability[MetricKey] = 18,
            FactPlanMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR DeniedEligible =
        CALCULATETABLE(
            VALUES(FactPlanMetricAvailability[PlanKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactPlanMetricAvailability[MetricKey] = 16,
            FactPlanMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR Eligible = INTERSECT(ResubEligible, DeniedEligible)
    VAR Numerator =
        CALCULATE(
            SUM(FactPlanTransparency[ClaimsResubmittedInNetwork]),
            TREATAS(Eligible, FactPlanTransparency[PlanKey])
        )
    VAR Denominator =
        CALCULATE(
            SUM(FactPlanTransparency[ClaimsDeniedInNetwork]),
            TREATAS(Eligible, FactPlanTransparency[PlanKey])
        )
    RETURN DIVIDE(Numerator, Denominator) * 100
    """,
    "0.00",
    "11 Plan Analytics",
)

add_measure(
    "Plan Resubmission Events per 100 Denied - Out of Network",
    """
    VAR ResubEligible =
        CALCULATETABLE(
            VALUES(FactPlanMetricAvailability[PlanKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactPlanMetricAvailability[MetricKey] = 17,
            FactPlanMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR DeniedEligible =
        CALCULATETABLE(
            VALUES(FactPlanMetricAvailability[PlanKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactPlanMetricAvailability[MetricKey] = 15,
            FactPlanMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR Eligible = INTERSECT(ResubEligible, DeniedEligible)
    VAR Numerator =
        CALCULATE(
            SUM(FactPlanTransparency[ClaimsResubmittedOutOfNetwork]),
            TREATAS(Eligible, FactPlanTransparency[PlanKey])
        )
    VAR Denominator =
        CALCULATE(
            SUM(FactPlanTransparency[ClaimsDeniedOutOfNetwork]),
            TREATAS(Eligible, FactPlanTransparency[PlanKey])
        )
    RETURN DIVIDE(Numerator, Denominator) * 100
    """,
    "0.00",
    "11 Plan Analytics",
)


# ---------------------------------------------------------------------
# 12 Sensitivity
# ---------------------------------------------------------------------

add_measure(
    "Internal Appeal Overturn Rate - Excluding Known Source Exception",
    """
    VAR FiledEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 7,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR OverturnedEligible =
        CALCULATETABLE(
            VALUES(FactIssuerMetricAvailability[IssuerKey]),
            REMOVEFILTERS(DimMetric),
            REMOVEFILTERS(DimAvailabilityStatus),
            FactIssuerMetricAvailability[MetricKey] = 8,
            FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
        )
    VAR Eligible = INTERSECT(FiledEligible, OverturnedEligible)
    VAR ExceptionIssuerIDs =
        CALCULATETABLE(
            VALUES(DQExceptionRegister[EntityKey]),
            DQExceptionRegister[Scope] = "ISSUER",
            DQExceptionRegister[IsKnownSourceException] = TRUE()
        )
    VAR ExceptionIssuerKeys =
        CALCULATETABLE(
            VALUES(DimIssuer[IssuerKey]),
            TREATAS(ExceptionIssuerIDs, DimIssuer[IssuerID])
        )
    VAR FinalEligible = EXCEPT(Eligible, ExceptionIssuerKeys)
    VAR Numerator =
        CALCULATE(
            SUM(FactIssuerTransparency[InternalAppealsOverturned]),
            TREATAS(FinalEligible, FactIssuerTransparency[IssuerKey])
        )
    VAR Denominator =
        CALCULATE(
            SUM(FactIssuerTransparency[InternalAppealsFiled]),
            TREATAS(FinalEligible, FactIssuerTransparency[IssuerKey])
        )
    RETURN DIVIDE(Numerator, Denominator)
    """,
    "0.00%",
    "12 Sensitivity",
)


# ---------------------------------------------------------------------
# Utilities
# ---------------------------------------------------------------------

def find_semantic_model() -> Path:
    candidates = sorted(
        p for p in POWERBI_ROOT.glob("*.SemanticModel") if p.is_dir()
    )
    if len(candidates) != 1:
        raise RuntimeError(
            f"Expected exactly one *.SemanticModel folder; found {len(candidates)}."
        )
    return candidates[0]


def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False
    result = subprocess.run(
        ["tasklist", "/FI", "IMAGENAME eq PBIDesktop.exe"],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    return "PBIDesktop.exe" in result.stdout


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig")


def write_text(path: Path, text: str) -> None:
    path.write_text(text, encoding="utf-8")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def make_uuid(label: str) -> str:
    namespace = uuid.UUID("e144557e-9ae4-4eef-a222-8ff05dc21695")
    return str(uuid.uuid5(namespace, label))


def create_backup(semantic_model: Path) -> None:
    REVIEW_DIR.mkdir(parents=True, exist_ok=True)
    if BACKUP_FILE.exists():
        BACKUP_FILE.unlink()

    with zipfile.ZipFile(
        BACKUP_FILE,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zf:
        for path in sorted(p for p in semantic_model.rglob("*") if p.is_file()):
            arcname = Path(semantic_model.name) / path.relative_to(semantic_model)
            zf.write(path, arcname=str(arcname))


def tmdl_measure_block(measure: dict[str, str]) -> str:
    name = measure["name"].replace("'", "''")
    dax_lines = measure["dax"].splitlines()

    if len(dax_lines) == 1:
        first_line = f"\tmeasure '{name}' = {dax_lines[0]}"
        dax_body = []
    else:
        first_line = f"\tmeasure '{name}' ="
        dax_body = [f"\t\t\t{line}" if line else "" for line in dax_lines]

    lines = [first_line] + dax_body
    lines.append(f"\t\tformatString: {measure['format']}")
    lines.append(f"\t\tdisplayFolder: {measure['folder']}")
    if measure["hidden"]:
        lines.append("\t\tisHidden")
    lines.append(f"\t\tlineageTag: {make_uuid('measure:' + measure['name'])}")
    return "\n".join(lines)


def build_measure_table_tmdl() -> str:
    blocks = [
        f"table {MEASURE_TABLE}",
        f"\tlineageTag: {make_uuid('table:' + MEASURE_TABLE)}",
        "",
        "\tcolumn Placeholder",
        "\t\tdataType: int64",
        "\t\tisHidden",
        "\t\tformatString: 0",
        f"\t\tlineageTag: {make_uuid('column:' + MEASURE_TABLE + '.Placeholder')}",
        "\t\tsummarizeBy: none",
        "\t\tsourceColumn: Placeholder",
        "",
        "\t\tannotation SummarizationSetBy = User",
        "",
    ]

    for measure in MEASURES:
        blocks.append(tmdl_measure_block(measure))
        blocks.append("")

    blocks.extend(
        [
            f"\tpartition {MEASURE_TABLE} = m",
            "\t\tmode: import",
            "\t\tsource =",
            "\t\t\t\tlet",
            "\t\t\t\t    Source = #table(type table [Placeholder = Int64.Type], {})",
            "\t\t\t\tin",
            "\t\t\t\t    Source",
            "",
            "\tannotation PBI_NavigationStepName = Navigation",
            "",
            "\tannotation PBI_ResultType = Table",
            "",
        ]
    )

    return "\n".join(blocks)


def count_existing_measures(tables_dir: Path) -> int:
    count = 0
    for path in tables_dir.glob("*.tmdl"):
        text = read_text(path)
        count += len(re.findall(r"(?m)^\tmeasure\s+", text))
    return count


def parse_measure_names(text: str) -> list[str]:
    return [
        x.replace("''", "'")
        for x in re.findall(r"(?m)^\tmeasure\s+'(.+?)'\s*=", text)
    ]


def style_report(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row and ws.max_column:
            ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(1, ws.max_column + 1):
            max_len = 0
            for row_idx in range(1, min(ws.max_row, 200) + 1):
                value = ws.cell(row_idx, col_idx).value
                if value is not None:
                    max_len = max(max_len, len(str(value)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(
                max(max_len + 2, 10),
                60,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(path)


def main() -> int:
    print("=" * 78)
    print("Transparency in Coverage PUF — Governed DAX Installation")
    print("=" * 78)
    print(f"Script version : {SCRIPT_VERSION}")
    print(f"Power BI root  : {POWERBI_ROOT}")
    print(f"Measures       : {len(MEASURES)}")
    print(f"Review file    : {REPORT_FILE}")
    print(f"Backup         : {BACKUP_FILE}")

    if power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is running. Save the PBIP and close Power BI "
            "Desktop completely before installing DAX."
        )

    semantic_model = find_semantic_model()
    definition_dir = semantic_model / "definition"
    model_file = definition_dir / "model.tmdl"
    tables_dir = definition_dir / "tables"
    measure_table_file = tables_dir / f"{MEASURE_TABLE}.tmdl"
    relationships_file = definition_dir / "relationships.tmdl"

    for path in [model_file, relationships_file]:
        if not path.exists():
            raise FileNotFoundError(f"Required semantic-model file missing: {path}")

    existing_measures = count_existing_measures(tables_dir)
    if existing_measures != 0:
        raise RuntimeError(
            f"Expected zero existing measures before DAX installation; "
            f"found {existing_measures}. No changes were made."
        )

    if measure_table_file.exists():
        raise RuntimeError(
            f"{measure_table_file} already exists. No changes were made."
        )

    relationship_text = read_text(relationships_file)
    relationship_count = len(
        re.findall(r"(?m)^relationship\s+", relationship_text)
    )
    if relationship_count != 14:
        raise RuntimeError(
            f"Expected 14 governed relationships; found {relationship_count}. "
            "No changes were made."
        )

    create_backup(semantic_model)

    before_hashes = {
        str(p.relative_to(semantic_model)): sha256(p)
        for p in semantic_model.rglob("*")
        if p.is_file()
    }

    model_text = read_text(model_file)

    if re.search(rf"(?m)^ref table {re.escape(MEASURE_TABLE)}\s*$", model_text):
        raise RuntimeError(
            f"model.tmdl already references {MEASURE_TABLE}. No changes were made."
        )

    # Add the dedicated measure host directly before the culture reference.
    marker = "ref cultureInfo en-US"
    if marker not in model_text:
        raise RuntimeError(
            "Could not find the culture reference insertion marker in model.tmdl."
        )

    model_text = model_text.replace(
        marker,
        f"ref table {MEASURE_TABLE}\n\n{marker}",
        1,
    )
    write_text(model_file, model_text)

    measure_table_text = build_measure_table_tmdl()
    write_text(measure_table_file, measure_table_text)

    # Offline structural validation.
    installed_text = read_text(measure_table_file)
    installed_names = parse_measure_names(installed_text)
    expected_names = [m["name"] for m in MEASURES]

    checks = []

    def check(check_id, test, expected, actual, passed):
        checks.append(
            {
                "CheckID": check_id,
                "TestName": test,
                "Expected": expected,
                "Actual": actual,
                "Status": "PASS" if passed else "FAIL",
            }
        )

    check(
        "DAX-001",
        "Measure host file created",
        True,
        measure_table_file.exists(),
        measure_table_file.exists(),
    )
    check(
        "DAX-002",
        "Measure host referenced by model",
        True,
        f"ref table {MEASURE_TABLE}" in read_text(model_file),
        f"ref table {MEASURE_TABLE}" in read_text(model_file),
    )
    check(
        "DAX-003",
        "Installed measure count",
        len(expected_names),
        len(installed_names),
        len(installed_names) == len(expected_names),
    )
    check(
        "DAX-004",
        "Installed measure names",
        "Exact governed set",
        "Exact" if set(installed_names) == set(expected_names) else "Mismatch",
        set(installed_names) == set(expected_names),
    )
    check(
        "DAX-005",
        "Measure host relationships",
        0,
        len(
            re.findall(
                rf"(?mi)^(?:\s*fromColumn|\s*toColumn):\s*{re.escape(MEASURE_TABLE)}\.",
                relationship_text,
            )
        ),
        MEASURE_TABLE not in relationship_text,
    )
    check(
        "DAX-006",
        "Auto Date/Time remains disabled",
        False,
        "__PBI_TimeIntelligenceEnabled = 1" in read_text(model_file),
        "__PBI_TimeIntelligenceEnabled = 1" not in read_text(model_file),
    )

    validation_df = pd.DataFrame(checks)
    failures = int((validation_df["Status"] == "FAIL").sum())

    after_hashes = {
        str(p.relative_to(semantic_model)): sha256(p)
        for p in semantic_model.rglob("*")
        if p.is_file()
    }

    file_changes = []
    for rel_path in sorted(set(before_hashes) | set(after_hashes)):
        before = before_hashes.get(rel_path, "")
        after = after_hashes.get(rel_path, "")
        if before == after:
            continue
        if before and after:
            state = "MODIFIED"
        elif after:
            state = "CREATED"
        else:
            state = "DELETED"
        file_changes.append(
            {
                "RelativePath": rel_path,
                "ChangeState": state,
                "BeforeSHA256": before,
                "AfterSHA256": after,
            }
        )

    inventory_df = pd.DataFrame(
        [
            {
                "MeasureName": m["name"],
                "DisplayFolder": m["folder"],
                "FormatString": m["format"],
                "IsHidden": m["hidden"],
                "LineageTag": make_uuid("measure:" + m["name"]),
                "DAXExpression": m["dax"],
            }
            for m in MEASURES
        ]
    )

    folder_summary = (
        inventory_df.groupby("DisplayFolder")
        .size()
        .reset_index(name="MeasureCount")
        .sort_values("DisplayFolder")
    )

    summary_df = pd.DataFrame(
        [
            ("InstallStatus", "PASS" if failures == 0 else "FAIL"),
            ("ScriptVersion", SCRIPT_VERSION),
            ("InstalledAtLocal", datetime.now().isoformat(timespec="seconds")),
            ("SemanticModel", str(semantic_model)),
            ("MeasureHost", MEASURE_TABLE),
            ("MeasuresInstalled", len(MEASURES)),
            ("DisplayFolders", inventory_df["DisplayFolder"].nunique()),
            ("RelationshipsUnchanged", 14),
            ("ValidationChecks", len(validation_df)),
            ("ValidationFailures", failures),
            ("BackupFile", str(BACKUP_FILE)),
            ("EvidenceState", "PBIP SOURCE MODIFIED + OFFLINE DAX STRUCTURAL VALIDATION"),
        ],
        columns=["Item", "Value"],
    )

    REVIEW_DIR.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(REPORT_FILE, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="00_Summary", index=False)
        validation_df.to_excel(writer, sheet_name="01_Validation", index=False)
        folder_summary.to_excel(writer, sheet_name="02_Folders", index=False)
        inventory_df.to_excel(writer, sheet_name="03_Measure_Inventory", index=False)
        pd.DataFrame(file_changes).to_excel(
            writer, sheet_name="04_File_Changes", index=False
        )

    style_report(REPORT_FILE)

    print()
    print("Governed DAX installation completed.")
    print(f"Install status       : {'PASS' if failures == 0 else 'FAIL'}")
    print(f"Measure host         : {MEASURE_TABLE}")
    print(f"Measures installed   : {len(MEASURES)}")
    print(f"Display folders      : {inventory_df['DisplayFolder'].nunique()}")
    print(f"Validation failures  : {failures}")
    print(f"Backup               : {BACKUP_FILE}")
    print(f"Review evidence      : {REPORT_FILE}")

    if failures:
        print()
        print("Offline validation failed. Do not open/save the PBIP until reviewed.")
        return 2

    print()
    print("Next step: open the PBIP in Power BI Desktop.")
    print("If it opens without a semantic-model error, Save (Ctrl+S) and close it.")
    print("Do not build visuals yet.")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:
        print()
        print("DAX INSTALLATION FAILED")
        print(str(exc))
        sys.exit(1)
