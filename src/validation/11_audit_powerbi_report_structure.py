from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
import hashlib
import json
import os
import re
import sys

import pandas as pd


SCRIPT_VERSION = "1.0.0"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
POWERBI_ROOT = PROJECT_ROOT / "powerbi"


def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, key_path) as key:
                raw_value, _ = winreg.QueryValueEx(key, "Desktop")
            return Path(os.path.expandvars(raw_value)).resolve()
        except Exception:
            pass

    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = Path(os.environ.get("TRANSPARENCY_PUF_REVIEW_DIR", PROJECT_ROOT / ".local-review")).expanduser().resolve()
OUTPUT_FILE = REVIEW_DIR / "11_POWER_BI_REPORT_STRUCTURE_AUDIT.xlsx"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def load_json(path: Path) -> dict | list:
    with path.open("r", encoding="utf-8-sig") as f:
        return json.load(f)


def find_report_folder() -> Path:
    candidates = sorted(
        p for p in POWERBI_ROOT.glob("*.Report") if p.is_dir()
    )

    if len(candidates) != 1:
        raise RuntimeError(
            f"Expected exactly one *.Report folder under {POWERBI_ROOT}; "
            f"found {len(candidates)}: {[p.name for p in candidates]}"
        )

    return candidates[0]


def scalar(value):
    return isinstance(value, (str, int, float, bool)) or value is None


def walk_json(obj, prefix=""):
    """Yield (path, value) for all JSON nodes."""
    if isinstance(obj, dict):
        for key, value in obj.items():
            new_prefix = f"{prefix}.{key}" if prefix else str(key)
            yield new_prefix, value
            yield from walk_json(value, new_prefix)
    elif isinstance(obj, list):
        for idx, value in enumerate(obj):
            new_prefix = f"{prefix}[{idx}]"
            yield new_prefix, value
            yield from walk_json(value, new_prefix)


def first_scalar_by_keys(obj, keys: set[str]):
    """Find first scalar value whose terminal JSON key matches one of keys."""
    lowered = {k.lower() for k in keys}
    for path, value in walk_json(obj):
        terminal = re.sub(r"\[\d+\]$", "", path.split(".")[-1]).lower()
        if terminal in lowered and scalar(value):
            return value
    return None


def values_by_terminal_key(obj, key: str):
    out = []
    target = key.lower()

    for path, value in walk_json(obj):
        terminal = re.sub(r"\[\d+\]$", "", path.split(".")[-1]).lower()
        if terminal == target and scalar(value):
            out.append(value)

    return out


def extract_page_identity(page_json: dict, page_dir: Path) -> dict:
    page_name = (
        page_json.get("name")
        or page_json.get("id")
        or page_dir.name
    )

    display_name = (
        page_json.get("displayName")
        or page_json.get("displayname")
        or page_json.get("title")
        or page_name
    )

    width = first_scalar_by_keys(
        page_json,
        {"width", "pageWidth", "canvasWidth"},
    )
    height = first_scalar_by_keys(
        page_json,
        {"height", "pageHeight", "canvasHeight"},
    )

    return {
        "PageFolder": page_dir.name,
        "PageName": page_name,
        "DisplayName": display_name,
        "Width": width,
        "Height": height,
    }


def extract_visual_type(visual_json: dict) -> str:
    # Prefer canonical visualType keys.
    for key in ("visualType", "type"):
        values = values_by_terminal_key(visual_json, key)
        for value in values:
            if isinstance(value, str) and value.strip():
                low = value.lower()
                if low not in {
                    "visual",
                    "container",
                    "group",
                    "object",
                    "property",
                }:
                    return value.strip()

    return "UNKNOWN"


def extract_visual_identity(visual_json: dict, visual_dir: Path) -> tuple[str, str]:
    name = (
        visual_json.get("name")
        or visual_json.get("id")
        or visual_dir.name
    )

    title_candidates = []

    for key in ("title", "text", "displayName", "label"):
        for value in values_by_terminal_key(visual_json, key):
            if isinstance(value, str):
                cleaned = re.sub(r"\s+", " ", value).strip()
                if cleaned and len(cleaned) <= 250:
                    title_candidates.append(cleaned)

    title = title_candidates[0] if title_candidates else ""

    return str(name), title


def extract_position(visual_json: dict) -> dict:
    # PBIR has changed shape across versions. Search likely keys rather than
    # assuming one schema path.
    def find_num(keys):
        value = first_scalar_by_keys(visual_json, set(keys))
        if isinstance(value, (int, float)):
            return value

        try:
            if value is not None:
                return float(value)
        except Exception:
            pass

        return None

    return {
        "X": find_num({"x", "left"}),
        "Y": find_num({"y", "top"}),
        "Width": find_num({"width", "w"}),
        "Height": find_num({"height", "h"}),
        "Z": find_num({"z", "zIndex", "zindex"}),
    }


def detect_actions(visual_json: dict) -> list[str]:
    hits = []

    for path, value in walk_json(visual_json):
        path_low = path.lower()

        if any(
            token in path_low
            for token in (
                "action",
                "navigation",
                "bookmark",
                "pageNavigation".lower(),
                "destination",
            )
        ):
            if scalar(value) and value not in ("", None, False):
                hits.append(f"{path}={value}")

    # Keep report compact and deterministic.
    deduped = []
    seen = set()

    for hit in hits:
        if hit not in seen:
            seen.add(hit)
            deduped.append(hit)

    return deduped[:20]


def detect_home_candidate(
    visual_type: str,
    title: str,
    actions: list[str],
) -> str:
    blob = " ".join([visual_type, title, *actions]).lower()

    home_terms = (
        "home",
        "index",
        "homepage",
        "landing",
    )

    return "YES" if any(term in blob for term in home_terms) else "NO"


def audit_report(report_dir: Path):
    definition_dir = report_dir / "definition"
    pages_root = definition_dir / "pages"

    if not definition_dir.exists():
        raise RuntimeError(
            f"PBIR definition folder not found: {definition_dir}"
        )

    if not pages_root.exists():
        raise RuntimeError(
            f"PBIR pages folder not found: {pages_root}"
        )

    parse_errors = []
    page_rows = []
    visual_rows = []
    file_rows = []
    theme_rows = []

    all_json_files = sorted(definition_dir.rglob("*.json"))

    for path in sorted(
        p for p in report_dir.rglob("*") if p.is_file()
    ):
        file_rows.append(
            {
                "RelativePath": str(path.relative_to(PROJECT_ROOT)),
                "Extension": path.suffix.lower(),
                "Bytes": path.stat().st_size,
                "SHA256": sha256(path),
            }
        )

    # Capture potential report/theme metadata.
    for path in all_json_files:
        rel = str(path.relative_to(report_dir))

        try:
            data = load_json(path)
        except Exception as exc:
            parse_errors.append(
                {
                    "RelativePath": rel,
                    "Error": str(exc),
                }
            )
            continue

        rel_low = rel.lower()

        if "theme" in rel_low or path.name.lower() == "report.json":
            for key in (
                "name",
                "theme",
                "themeCollection",
                "baseTheme",
                "customTheme",
            ):
                values = values_by_terminal_key(data, key)

                for value in values[:20]:
                    theme_rows.append(
                        {
                            "RelativePath": rel,
                            "Key": key,
                            "Value": str(value)[:500],
                        }
                    )

    # A PBIR page is a directory directly or recursively under definition/pages
    # that contains page.json.
    page_json_files = sorted(pages_root.rglob("page.json"))

    if not page_json_files:
        raise RuntimeError(
            f"No page.json files found under {pages_root}"
        )

    page_order_map = {}

    pages_index = pages_root / "pages.json"
    if pages_index.exists():
        try:
            pages_index_json = load_json(pages_index)

            # Capture every scalar sequence-like reference. This works across
            # common PBIR pages.json shapes.
            order_counter = 1

            if isinstance(pages_index_json, dict):
                for candidate_key in (
                    "pageOrder",
                    "order",
                    "pages",
                    "sections",
                ):
                    candidate = pages_index_json.get(candidate_key)

                    if isinstance(candidate, list):
                        for item in candidate:
                            if isinstance(item, str):
                                page_order_map[item] = order_counter
                                order_counter += 1
                            elif isinstance(item, dict):
                                ident = (
                                    item.get("name")
                                    or item.get("id")
                                    or item.get("page")
                                )
                                if ident:
                                    page_order_map[str(ident)] = order_counter
                                    order_counter += 1
                        if page_order_map:
                            break
        except Exception as exc:
            parse_errors.append(
                {
                    "RelativePath": str(
                        pages_index.relative_to(report_dir)
                    ),
                    "Error": f"pages.json order parse: {exc}",
                }
            )

    for page_json_file in page_json_files:
        page_dir = page_json_file.parent

        try:
            page_json = load_json(page_json_file)
        except Exception as exc:
            parse_errors.append(
                {
                    "RelativePath": str(
                        page_json_file.relative_to(report_dir)
                    ),
                    "Error": str(exc),
                }
            )
            continue

        identity = extract_page_identity(
            page_json,
            page_dir,
        )

        page_name = str(identity["PageName"])
        display_name = str(identity["DisplayName"])

        order = (
            page_order_map.get(page_name)
            or page_order_map.get(page_dir.name)
        )

        visuals_dir = page_dir / "visuals"
        visual_json_files = (
            sorted(visuals_dir.rglob("visual.json"))
            if visuals_dir.exists()
            else []
        )

        page_visual_types = Counter()
        home_candidates = 0

        for visual_json_file in visual_json_files:
            visual_dir = visual_json_file.parent

            try:
                visual_json = load_json(visual_json_file)
            except Exception as exc:
                parse_errors.append(
                    {
                        "RelativePath": str(
                            visual_json_file.relative_to(report_dir)
                        ),
                        "Error": str(exc),
                    }
                )
                continue

            visual_type = extract_visual_type(visual_json)
            visual_name, title = extract_visual_identity(
                visual_json,
                visual_dir,
            )
            position = extract_position(visual_json)
            actions = detect_actions(visual_json)

            home_candidate = detect_home_candidate(
                visual_type,
                title,
                actions,
            )

            if home_candidate == "YES":
                home_candidates += 1

            page_visual_types[visual_type] += 1

            visual_rows.append(
                {
                    "PageOrder": order,
                    "PageName": page_name,
                    "PageDisplayName": display_name,
                    "VisualFolder": visual_dir.name,
                    "VisualName": visual_name,
                    "VisualType": visual_type,
                    "TitleOrLabel": title,
                    "X": position["X"],
                    "Y": position["Y"],
                    "Width": position["Width"],
                    "Height": position["Height"],
                    "Z": position["Z"],
                    "HasAction": "YES" if actions else "NO",
                    "HomeCandidate": home_candidate,
                    "ActionEvidence": " | ".join(actions),
                    "RelativePath": str(
                        visual_json_file.relative_to(PROJECT_ROOT)
                    ),
                }
            )

        page_rows.append(
            {
                **identity,
                "PageOrder": order,
                "VisualCount": len(visual_json_files),
                "HomeCandidateCount": home_candidates,
                "VisualTypeSummary": ", ".join(
                    f"{k}:{v}"
                    for k, v in sorted(
                        page_visual_types.items(),
                        key=lambda x: (-x[1], x[0]),
                    )
                ),
                "PageJsonPath": str(
                    page_json_file.relative_to(PROJECT_ROOT)
                ),
            }
        )

    pages_df = pd.DataFrame(page_rows)

    if "PageOrder" in pages_df.columns:
        pages_df = pages_df.sort_values(
            by=["PageOrder", "DisplayName"],
            na_position="last",
        )

    visuals_df = pd.DataFrame(visual_rows)

    if not visuals_df.empty:
        visuals_df = visuals_df.sort_values(
            by=["PageOrder", "PageDisplayName", "Y", "X"],
            na_position="last",
        )

    visual_type_rows = []

    if not visuals_df.empty:
        grouped = (
            visuals_df.groupby(
                ["PageDisplayName", "VisualType"],
                dropna=False,
            )
            .size()
            .reset_index(name="Count")
        )

        visual_type_rows = grouped.sort_values(
            ["PageDisplayName", "Count", "VisualType"],
            ascending=[True, False, True],
        )

    visual_types_df = pd.DataFrame(visual_type_rows)

    files_df = pd.DataFrame(file_rows)
    themes_df = pd.DataFrame(theme_rows)
    errors_df = pd.DataFrame(parse_errors)

    page_count = len(pages_df)
    visual_count = len(visuals_df)

    widths = (
        pages_df["Width"].dropna().astype(str).unique().tolist()
        if "Width" in pages_df.columns
        else []
    )
    heights = (
        pages_df["Height"].dropna().astype(str).unique().tolist()
        if "Height" in pages_df.columns
        else []
    )

    summary_df = pd.DataFrame(
        [
            ("AuditStatus", "PASS" if errors_df.empty else "PASS_WITH_PARSE_WARNINGS"),
            ("ScriptVersion", SCRIPT_VERSION),
            ("AuditedAtLocal", datetime.now().isoformat(timespec="seconds")),
            ("ProjectRoot", str(PROJECT_ROOT)),
            ("ReportFolder", str(report_dir)),
            ("Pages", page_count),
            ("Visuals", visual_count),
            ("JSONFiles", len(all_json_files)),
            ("ParseWarnings", len(errors_df)),
            ("ObservedPageWidths", ", ".join(widths)),
            ("ObservedPageHeights", ", ".join(heights)),
            (
                "EvidenceState",
                "READ-ONLY PBIR REPORT STRUCTURE AUDIT",
            ),
        ],
        columns=["Item", "Value"],
    )

    return {
        "00_Summary": summary_df,
        "01_Pages": pages_df,
        "02_Visuals": visuals_df,
        "03_Visual_Types": visual_types_df,
        "04_Theme_Metadata": themes_df,
        "05_File_Inventory": files_df,
        "06_Parse_Warnings": errors_df,
    }


def style_report(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            max_len = 0

            for row_idx in range(
                1,
                min(ws.max_row, 250) + 1,
            ):
                value = ws.cell(
                    row=row_idx,
                    column=col_idx,
                ).value

                if value is not None:
                    max_len = max(
                        max_len,
                        len(str(value)),
                    )

            ws.column_dimensions[
                get_column_letter(col_idx)
            ].width = min(
                max(max_len + 2, 10),
                60,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    wb.save(path)


def main() -> int:
    print("=" * 78)
    print("Transparency in Coverage PUF — Power BI Report Structure Audit")
    print("=" * 78)
    print(f"Script version : {SCRIPT_VERSION}")
    print(f"Project root   : {PROJECT_ROOT}")
    print(f"Power BI root  : {POWERBI_ROOT}")
    print(f"Review file    : {OUTPUT_FILE}")

    report_dir = find_report_folder()
    print(f"Report folder  : {report_dir}")

    audit = audit_report(report_dir)

    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        OUTPUT_FILE,
        engine="openpyxl",
    ) as writer:
        for sheet_name, df in audit.items():
            if df.empty:
                pd.DataFrame(
                    [{"Info": "No rows"}]
                ).to_excel(
                    writer,
                    sheet_name=sheet_name,
                    index=False,
                )
            else:
                df.to_excel(
                    writer,
                    sheet_name=sheet_name,
                    index=False,
                )

    style_report(OUTPUT_FILE)

    summary = audit["00_Summary"]
    values = dict(
        zip(
            summary["Item"],
            summary["Value"],
        )
    )

    print()
    print("PBIR report structure audit completed.")
    print(f"Audit status   : {values['AuditStatus']}")
    print(f"Pages          : {values['Pages']}")
    print(f"Visuals        : {values['Visuals']}")
    print(f"JSON files     : {values['JSONFiles']}")
    print(f"Parse warnings : {values['ParseWarnings']}")
    print(f"Review evidence: {OUTPUT_FILE}")
    print("PBIR files were read only.")

    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())

    except Exception as exc:
        print()
        print("REPORT STRUCTURE AUDIT FAILED")
        print(str(exc))
        sys.exit(1)
