from __future__ import annotations

from datetime import datetime
from pathlib import Path
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile

import pandas as pd


SCRIPT_VERSION = "1.0.0"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
POWERBI_ROOT = PROJECT_ROOT / "powerbi"
REPORT_DIR = POWERBI_ROOT / "TransparencyInCoverage.Report"
SEMANTIC_DIR = POWERBI_ROOT / "TransparencyInCoverage.SemanticModel"
PAGES_DIR = REPORT_DIR / "definition" / "pages"

MEASURES_TMDL = (
    SEMANTIC_DIR
    / "definition"
    / "tables"
    / "_Measures.tmdl"
)
DIM_METRIC_TMDL = (
    SEMANTIC_DIR
    / "definition"
    / "tables"
    / "DimMetric.tmdl"
)

VISUAL_SCHEMA = (
    "https://developer.microsoft.com/json-schemas/fabric/item/report/"
    "definition/visualContainer/2.7.0/schema.json"
)

CANVAS_WIDTH = 1920
CANVAS_HEIGHT = 1080

COLORS = {
    "canvas": "#F5F7FA",
    "surface": "#FFFFFF",
    "border": "#E7ECF2",
    "ink": "#172033",
    "secondary": "#667085",
    "muted": "#98A2B3",
    "navy": "#17324D",
    "blue": "#4F6BED",
    "blue_dark": "#425CC7",
    "blue_soft": "#EEF3FF",
    "teal": "#16A6A1",
    "teal_soft": "#EAF8F6",
    "amber": "#C8872C",
    "red": "#C64B4B",
    "neutral_soft": "#F9FAFB",
}

FONT_REGULAR = (
    "'Segoe UI', wf_segoe-ui_normal, helvetica, arial, sans-serif"
)
FONT_SEMIBOLD = (
    "'Segoe UI Semibold', wf_segoe-ui_semibold, "
    "helvetica, arial, sans-serif"
)

REQUIRED_MEASURES = [
    "Issuer Count",
    "Plan Count",
    "Issuer Comparable Claims Received - Total",
    "Issuer Comparable Claims Denied - Total",
    "Issuer Comparable Overall Denial Rate",
    "Issuer In-Network Denial Rate",
    "Issuer Out-of-Network Denial Rate",
    "Issuer Network Denial Rate Gap",
    "Internal Appeal Overturn Rate",
    "External Appeal Overturn Rate",
    "Internal Appeals Filed",
    "External Appeals Filed",
    "Comparable Denial Reason Composition %",
    "Open DQ Exception Count",
    "Known Source Exception Row Count",
]

REQUIRED_DIM_METRIC_COLUMNS = [
    "MetricDisplayName",
]

EXPECTED_PAGE_DISPLAY_NAMES = [
    "00 INDEX",
    "01 Executive Overview",
    "02 Network Comparison",
    "03 Denials",
    "04 Denial Reasons",
    "05 Appeals",
    "06 Resubmissions",
    "07 Enrollment",
    "08 Issuer & State Explorer",
    "09 Data Availability",
    "10 Data Quality & Methodology",
]


# ---------------------------------------------------------------------
# Environment / paths
# ---------------------------------------------------------------------

def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )
            with winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                key_path,
            ) as key:
                raw_value, _ = winreg.QueryValueEx(
                    key,
                    "Desktop",
                )

            return Path(
                os.path.expandvars(raw_value)
            ).resolve()

        except Exception:
            pass

    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = (
    get_windows_desktop()
    / "Transparency_PUF_Review"
)
BACKUP_FILE = (
    REVIEW_DIR
    / "14_PRE_EXECUTIVE_OVERVIEW_BACKUP.zip"
)
REPORT_FILE = (
    REVIEW_DIR
    / "14_EXECUTIVE_OVERVIEW_BUILD_REPORT.xlsx"
)


def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False

    result = subprocess.run(
        [
            "tasklist",
            "/FI",
            "IMAGENAME eq PBIDesktop.exe",
        ],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(
            subprocess,
            "CREATE_NO_WINDOW",
            0,
        ),
    )

    return "PBIDesktop.exe" in result.stdout


def read_json(path: Path):
    return json.loads(
        path.read_text(
            encoding="utf-8-sig",
        )
    )


def write_json(
    path: Path,
    payload,
) -> None:
    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    path.write_text(
        json.dumps(
            payload,
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )


def sha256(path: Path) -> str:
    h = hashlib.sha256()

    with path.open("rb") as f:
        for chunk in iter(
            lambda: f.read(1024 * 1024),
            b"",
        ):
            h.update(chunk)

    return h.hexdigest()


# ---------------------------------------------------------------------
# PBIR literal helpers
# ---------------------------------------------------------------------

def expr_bool(value: bool) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": (
                    "true"
                    if value
                    else "false"
                )
            }
        }
    }


def expr_string(value: str) -> dict:
    escaped = value.replace(
        "'",
        "''",
    )

    return {
        "expr": {
            "Literal": {
                "Value": f"'{escaped}'"
            }
        }
    }


def expr_double(
    value: float | int,
) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": f"{value}D"
            }
        }
    }


def expr_int(
    value: int,
) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": f"{value}L"
            }
        }
    }


def solid_color(
    value: str,
) -> dict:
    return {
        "solid": {
            "color": expr_string(value)
        }
    }


# ---------------------------------------------------------------------
# TMDL validation
# ---------------------------------------------------------------------

def parse_tmdl_measures(
    path: Path,
) -> set[str]:
    text = path.read_text(
        encoding="utf-8-sig",
    )

    measures = set()

    for line in text.splitlines():
        stripped = line.lstrip()

        if not stripped.startswith("measure "):
            continue

        lhs = stripped.split(
            "=",
            1,
        )[0].strip()

        raw_name = lhs[len("measure "):].strip()

        if (
            raw_name.startswith("'")
            and raw_name.endswith("'")
        ):
            raw_name = (
                raw_name[1:-1]
                .replace("''", "'")
            )

        measures.add(raw_name)

    return measures


def parse_tmdl_columns(
    path: Path,
) -> set[str]:
    text = path.read_text(
        encoding="utf-8-sig",
    )

    columns = set()

    for line in text.splitlines():
        stripped = line.lstrip()

        if not stripped.startswith("column "):
            continue

        raw_name = stripped[len("column "):].strip()

        if raw_name.startswith("'"):
            match = re.match(
                r"'((?:''|[^'])*)'",
                raw_name,
            )

            if match:
                columns.add(
                    match.group(1)
                    .replace("''", "'")
                )

        else:
            columns.add(
                raw_name.split(
                    "\t",
                    1,
                )[0].strip()
            )

    return columns


# ---------------------------------------------------------------------
# Page lookup / backup
# ---------------------------------------------------------------------

def find_pages() -> dict[str, tuple[Path, dict]]:
    pages = {}

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )
        display_name = payload.get(
            "displayName"
        )

        if display_name:
            pages[display_name] = (
                page_json.parent,
                payload,
            )

    return pages


def create_backup() -> None:
    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if BACKUP_FILE.exists():
        BACKUP_FILE.unlink()

    with zipfile.ZipFile(
        BACKUP_FILE,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zf:
        for path in sorted(
            p
            for p in REPORT_DIR.rglob("*")
            if p.is_file()
        ):
            archive_name = (
                Path(REPORT_DIR.name)
                / path.relative_to(
                    REPORT_DIR
                )
            )

            zf.write(
                path,
                str(archive_name),
            )


def report_hashes() -> dict[str, str]:
    return {
        str(
            path.relative_to(
                REPORT_DIR
            )
        ): sha256(path)
        for path in REPORT_DIR.rglob("*")
        if path.is_file()
    }


# ---------------------------------------------------------------------
# Textbox helpers
# ---------------------------------------------------------------------

def text_run(
    value: str,
    *,
    size: float,
    color: str,
    semibold: bool = False,
) -> dict:
    return {
        "value": value,
        "textStyle": {
            "fontFamily": (
                FONT_SEMIBOLD
                if semibold
                else FONT_REGULAR
            ),
            "fontSize": f"{size}pt",
            "color": color,
        },
    }


def paragraph(
    *runs: dict,
    align: str = "center",
) -> dict:
    return {
        "textRuns": list(runs),
        "horizontalTextAlignment": align,
    }


def container_objects(
    *,
    background: str | None = None,
    border: str | None = None,
    radius: int = 8,
    show_title: bool = False,
    title: str = "",
    title_size: float = 12,
    title_color: str = COLORS["ink"],
    title_alignment: str = "center",
    subtitle: str | None = None,
) -> dict:
    result = {
        "title": [
            {
                "properties": {
                    "show": expr_bool(
                        show_title
                    ),
                    **(
                        {
                            "text": expr_string(
                                title
                            ),
                            "alignment": expr_string(
                                title_alignment
                            ),
                            "fontColor": solid_color(
                                title_color
                            ),
                            "fontFamily": expr_string(
                                FONT_SEMIBOLD
                            ),
                            "fontSize": expr_double(
                                title_size
                            ),
                            "titleWrap": expr_bool(
                                True
                            ),
                        }
                        if show_title
                        else {}
                    ),
                }
            }
        ],
        "background": [
            {
                "properties": {
                    "show": expr_bool(
                        background is not None
                    ),
                    **(
                        {
                            "color": solid_color(
                                background
                            ),
                            "transparency": expr_double(
                                0
                            ),
                        }
                        if background is not None
                        else {}
                    ),
                }
            }
        ],
        "border": [
            {
                "properties": {
                    "show": expr_bool(
                        border is not None
                    ),
                    **(
                        {
                            "color": solid_color(
                                border
                            ),
                            "width": expr_double(
                                1
                            ),
                            "radius": expr_double(
                                radius
                            ),
                        }
                        if border is not None
                        else {}
                    ),
                }
            }
        ],
        "dropShadow": [
            {
                "properties": {
                    "show": expr_bool(
                        False
                    )
                }
            }
        ],
        "visualHeader": [
            {
                "properties": {
                    "show": expr_bool(
                        False
                    )
                }
            }
        ],
    }

    if subtitle is not None:
        result["subTitle"] = [
            {
                "properties": {
                    "show": expr_bool(
                        True
                    ),
                    "text": expr_string(
                        subtitle
                    ),
                    "fontColor": solid_color(
                        COLORS["secondary"]
                    ),
                    "fontFamily": expr_string(
                        FONT_REGULAR
                    ),
                    "fontSize": expr_double(
                        9.5
                    ),
                    "alignment": expr_string(
                        "center"
                    ),
                }
            }
        ]

    return result


def make_textbox(
    *,
    name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    paragraphs: list[dict],
    background: str | None = None,
    border: str | None = None,
    radius: int = 8,
    tab_order: int | None = None,
) -> dict:
    position = {
        "x": x,
        "y": y,
        "z": z,
        "width": width,
        "height": height,
    }

    if tab_order is not None:
        position["tabOrder"] = (
            tab_order
        )

    return {
        "$schema": VISUAL_SCHEMA,
        "name": name,
        "position": position,
        "visual": {
            "visualType": "textbox",
            "objects": {
                "general": [
                    {
                        "properties": {
                            "paragraphs": paragraphs
                        }
                    }
                ]
            },
            "visualContainerObjects": (
                container_objects(
                    background=background,
                    border=border,
                    radius=radius,
                )
            ),
            "drillFilterOtherVisuals": True,
        },
    }


# ---------------------------------------------------------------------
# Navigation
# ---------------------------------------------------------------------

def make_nav_button(
    *,
    name: str,
    target_page_name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    tab_order: int,
) -> dict:
    return {
        "$schema": VISUAL_SCHEMA,
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": tab_order,
        },
        "visual": {
            "visualType": "actionButton",
            "objects": {
                "icon": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "outline": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "fill": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": {
                "title": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "background": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "border": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "visualHeader": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "visualLink": [
                    {
                        "properties": {
                            "show": expr_bool(
                                True
                            ),
                            "type": expr_string(
                                "PageNavigation"
                            ),
                            "navigationSection": expr_string(
                                target_page_name
                            ),
                            "showDefaultTooltip": expr_bool(
                                False
                            ),
                            "enabledTooltip": expr_string(
                                "Return to INDEX"
                            ),
                        }
                    }
                ],
            },
            "drillFilterOtherVisuals": True,
        },
    }


def make_index_button_surface(
    *,
    name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
) -> dict:
    return make_textbox(
        name=name,
        x=x,
        y=y,
        width=width,
        height=height,
        z=z,
        paragraphs=[
            paragraph(
                text_run(
                    "INDEX",
                    size=11,
                    color="#FFFFFF",
                    semibold=True,
                ),
                align="center",
            )
        ],
        background=COLORS["blue"],
        border=COLORS["blue_dark"],
        radius=9,
    )


# ---------------------------------------------------------------------
# Data-field helpers
# ---------------------------------------------------------------------

def measure_field(
    measure_name: str,
) -> dict:
    return {
        "Measure": {
            "Expression": {
                "SourceRef": {
                    "Entity": "_Measures"
                }
            },
            "Property": measure_name,
        }
    }


def column_field(
    entity: str,
    column_name: str,
) -> dict:
    return {
        "Column": {
            "Expression": {
                "SourceRef": {
                    "Entity": entity
                }
            },
            "Property": column_name,
        }
    }


def measure_projection(
    measure_name: str,
) -> dict:
    return {
        "field": measure_field(
            measure_name
        ),
        "queryRef": (
            f"_Measures.{measure_name}"
        ),
        "nativeQueryRef": measure_name,
    }


def column_projection(
    entity: str,
    column_name: str,
    *,
    active: bool = False,
) -> dict:
    result = {
        "field": column_field(
            entity,
            column_name,
        ),
        "queryRef": (
            f"{entity}.{column_name}"
        ),
        "nativeQueryRef": column_name,
    }

    if active:
        result["active"] = True

    return result


# ---------------------------------------------------------------------
# Cards / charts
# ---------------------------------------------------------------------

def make_card(
    *,
    name: str,
    measure_name: str,
    title: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    tab_order: int,
    value_color: str = COLORS["ink"],
    background: str = COLORS["surface"],
    title_size: float = 11,
    value_size: float = 30,
) -> dict:
    return {
        "$schema": VISUAL_SCHEMA,
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": tab_order,
        },
        "visual": {
            "visualType": "card",
            "query": {
                "queryState": {
                    "Values": {
                        "projections": [
                            measure_projection(
                                measure_name
                            )
                        ]
                    }
                },
                "sortDefinition": {
                    "sort": [
                        {
                            "field": measure_field(
                                measure_name
                            ),
                            "direction": "Descending",
                        }
                    ],
                    "isDefaultSort": True,
                },
            },
            "objects": {
                "labels": [
                    {
                        "properties": {
                            "color": solid_color(
                                value_color
                            ),
                            "fontSize": expr_double(
                                value_size
                            ),
                            "bold": expr_bool(
                                True
                            ),
                            "fontFamily": expr_string(
                                FONT_SEMIBOLD
                            ),
                        }
                    }
                ],
                "categoryLabels": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": (
                container_objects(
                    background=background,
                    border=COLORS["border"],
                    radius=9,
                    show_title=True,
                    title=title,
                    title_size=title_size,
                    title_alignment="center",
                )
            ),
            "drillFilterOtherVisuals": True,
        },
    }


def make_reason_bar_chart(
    *,
    name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    tab_order: int,
) -> dict:
    measure_name = (
        "Comparable Denial Reason Composition %"
    )
    category_name = "MetricDisplayName"

    return {
        "$schema": VISUAL_SCHEMA,
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
            "tabOrder": tab_order,
        },
        "visual": {
            "visualType": "barChart",
            "query": {
                "queryState": {
                    "Category": {
                        "projections": [
                            column_projection(
                                "DimMetric",
                                category_name,
                                active=True,
                            )
                        ]
                    },
                    "Y": {
                        "projections": [
                            measure_projection(
                                measure_name
                            )
                        ]
                    },
                },
                "sortDefinition": {
                    "sort": [
                        {
                            "field": measure_field(
                                measure_name
                            ),
                            "direction": "Descending",
                        }
                    ],
                    "isDefaultSort": True,
                },
            },
            "objects": {
                "legend": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "categoryAxis": [
                    {
                        "properties": {
                            "showAxisTitle": expr_bool(
                                False
                            ),
                            "labelColor": solid_color(
                                COLORS["secondary"]
                            ),
                            "fontFamily": expr_string(
                                FONT_REGULAR
                            ),
                            "fontSize": expr_double(
                                9
                            ),
                            "concatenateLabels": expr_bool(
                                False
                            ),
                            "innerPadding": expr_int(
                                22
                            ),
                        }
                    }
                ],
                "valueAxis": [
                    {
                        "properties": {
                            "show": expr_bool(
                                True
                            ),
                            "showAxisTitle": expr_bool(
                                False
                            ),
                            "labelColor": solid_color(
                                COLORS["muted"]
                            ),
                            "fontFamily": expr_string(
                                FONT_REGULAR
                            ),
                            "fontSize": expr_double(
                                8
                            ),
                            "gridlineShow": expr_bool(
                                False
                            ),
                        }
                    }
                ],
                "dataPoint": [
                    {
                        "properties": {
                            "fill": solid_color(
                                COLORS["blue"]
                            )
                        },
                        "selector": {
                            "data": [
                                {
                                    "dataViewWildcard": {
                                        "matchingOption": 1
                                    }
                                }
                            ]
                        },
                    }
                ],
                "labels": [
                    {
                        "properties": {
                            "show": expr_bool(
                                True
                            ),
                            "color": solid_color(
                                COLORS["ink"]
                            ),
                            "fontSize": expr_double(
                                9
                            ),
                            "fontFamily": expr_string(
                                FONT_SEMIBOLD
                            ),
                            "bold": expr_bool(
                                False
                            ),
                            "labelOverflow": expr_bool(
                                False
                            ),
                        }
                    }
                ],
                "zoom": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": (
                container_objects(
                    background=COLORS["surface"],
                    border=COLORS["border"],
                    radius=9,
                    show_title=True,
                    title="Why Claims Were Denied",
                    title_size=12.5,
                    title_alignment="center",
                    subtitle=(
                        "Composition across fully comparable plans"
                    ),
                )
            ),
            "drillFilterOtherVisuals": True,
        },
    }


# ---------------------------------------------------------------------
# Visual IO
# ---------------------------------------------------------------------

def write_visual(
    page_dir: Path,
    visual: dict,
) -> Path:
    target = (
        page_dir
        / "visuals"
        / visual["name"]
        / "visual.json"
    )

    write_json(
        target,
        visual,
    )

    return target


def add_canvas(
    page_dir: Path,
) -> None:
    write_visual(
        page_dir,
        make_textbox(
            name="p01_canvas",
            x=0,
            y=0,
            width=CANVAS_WIDTH,
            height=CANVAS_HEIGHT,
            z=0,
            paragraphs=[
                paragraph(
                    text_run(
                        " ",
                        size=6,
                        color=COLORS["canvas"],
                    )
                )
            ],
            background=COLORS["canvas"],
            border=None,
            radius=0,
        ),
    )

    write_visual(
        page_dir,
        make_textbox(
            name="p01_accent",
            x=0,
            y=0,
            width=CANVAS_WIDTH,
            height=7,
            z=10,
            paragraphs=[
                paragraph(
                    text_run(
                        " ",
                        size=6,
                        color=COLORS["blue"],
                    )
                )
            ],
            background=COLORS["blue"],
            border=None,
            radius=0,
        ),
    )


def add_page_header(
    page_dir: Path,
    index_page_name: str,
) -> None:
    write_visual(
        page_dir,
        make_textbox(
            name="p01_context",
            x=64,
            y=46,
            width=290,
            height=52,
            z=100,
            paragraphs=[
                paragraph(
                    text_run(
                        "2026 PUF  •  EXPERIENCE YEAR 2024",
                        size=9,
                        color=COLORS["blue"],
                        semibold=True,
                    )
                )
            ],
            background=COLORS["blue_soft"],
            border=None,
            radius=8,
        ),
    )

    write_visual(
        page_dir,
        make_textbox(
            name="p01_header",
            x=390,
            y=28,
            width=1140,
            height=108,
            z=110,
            paragraphs=[
                paragraph(
                    text_run(
                        "Executive Overview",
                        size=28,
                        color=COLORS["ink"],
                        semibold=True,
                    )
                ),
                paragraph(
                    text_run(
                        (
                            "Market size → claims → network disparity → "
                            "denial reasons → appeals → governance"
                        ),
                        size=10.5,
                        color=COLORS["secondary"],
                    )
                ),
            ],
        ),
    )

    home_x = 1692
    home_y = 42
    home_w = 164
    home_h = 56

    write_visual(
        page_dir,
        make_index_button_surface(
            name="p01_home_surface",
            x=home_x,
            y=home_y,
            width=home_w,
            height=home_h,
            z=150,
        ),
    )

    write_visual(
        page_dir,
        make_nav_button(
            name="p01_home_nav",
            target_page_name=index_page_name,
            x=home_x,
            y=home_y,
            width=home_w,
            height=home_h,
            z=3200,
            tab_order=0,
        ),
    )


def add_kpi_row(
    page_dir: Path,
) -> None:
    y = 150
    h = 145
    gap = 16
    total_w = 1792
    card_w = (
        total_w - (4 * gap)
    ) / 5
    x0 = 64

    specs = [
        (
            "p01_kpi_issuers",
            "Issuer Count",
            "Issuers",
            COLORS["ink"],
            COLORS["surface"],
        ),
        (
            "p01_kpi_plans",
            "Plan Count",
            "Plans",
            COLORS["ink"],
            COLORS["surface"],
        ),
        (
            "p01_kpi_claims",
            "Issuer Comparable Claims Received - Total",
            "Comparable Claims",
            COLORS["navy"],
            COLORS["blue_soft"],
        ),
        (
            "p01_kpi_denial_rate",
            "Issuer Comparable Overall Denial Rate",
            "Comparable Denial Rate",
            COLORS["blue"],
            COLORS["surface"],
        ),
        (
            "p01_kpi_internal_appeal",
            "Internal Appeal Overturn Rate",
            "Internal Appeal Overturn Rate",
            COLORS["teal"],
            COLORS["surface"],
        ),
    ]

    for idx, (
        name,
        measure,
        title,
        value_color,
        background,
    ) in enumerate(specs):
        write_visual(
            page_dir,
            make_card(
                name=name,
                measure_name=measure,
                title=title,
                x=x0 + idx * (
                    card_w + gap
                ),
                y=y,
                width=card_w,
                height=h,
                z=500 + idx,
                tab_order=10 + idx,
                value_color=value_color,
                background=background,
                title_size=10.5,
                value_size=29,
            ),
        )


def add_network_section(
    page_dir: Path,
) -> None:
    x = 64
    y = 315
    w = 888
    h = 300

    write_visual(
        page_dir,
        make_textbox(
            name="p01_network_surface",
            x=x,
            y=y,
            width=w,
            height=h,
            z=300,
            paragraphs=[
                paragraph(
                    text_run(
                        "NETWORK DENIAL RATE",
                        size=12,
                        color=COLORS["ink"],
                        semibold=True,
                    )
                ),
                paragraph(
                    text_run(
                        "Issuer-level governed ratio of totals",
                        size=9.5,
                        color=COLORS["secondary"],
                    )
                ),
            ],
            background=COLORS["surface"],
            border=COLORS["border"],
            radius=9,
        ),
    )

    inner_y = y + 82
    inner_h = 185
    gap = 14
    inner_x = x + 24
    inner_total_w = w - 48
    inner_w = (
        inner_total_w - 2 * gap
    ) / 3

    specs = [
        (
            "p01_network_in",
            "Issuer In-Network Denial Rate",
            "In-Network",
            COLORS["navy"],
            COLORS["neutral_soft"],
        ),
        (
            "p01_network_out",
            "Issuer Out-of-Network Denial Rate",
            "Out-of-Network",
            COLORS["blue"],
            COLORS["blue_soft"],
        ),
        (
            "p01_network_gap",
            "Issuer Network Denial Rate Gap",
            "Network Gap",
            COLORS["teal"],
            COLORS["teal_soft"],
        ),
    ]

    for idx, (
        name,
        measure,
        title,
        value_color,
        background,
    ) in enumerate(specs):
        write_visual(
            page_dir,
            make_card(
                name=name,
                measure_name=measure,
                title=title,
                x=inner_x + idx * (
                    inner_w + gap
                ),
                y=inner_y,
                width=inner_w,
                height=inner_h,
                z=800 + idx,
                tab_order=20 + idx,
                value_color=value_color,
                background=background,
                title_size=10,
                value_size=31,
            ),
        )


def add_claims_section(
    page_dir: Path,
) -> None:
    x = 968
    y = 315
    w = 888
    h = 300

    write_visual(
        page_dir,
        make_textbox(
            name="p01_claims_surface",
            x=x,
            y=y,
            width=w,
            height=h,
            z=300,
            paragraphs=[
                paragraph(
                    text_run(
                        "COMPARABLE CLAIMS",
                        size=12,
                        color=COLORS["ink"],
                        semibold=True,
                    )
                ),
                paragraph(
                    text_run(
                        (
                            "Same issuer population available for "
                            "received and denied claims"
                        ),
                        size=9.5,
                        color=COLORS["secondary"],
                    )
                ),
            ],
            background=COLORS["surface"],
            border=COLORS["border"],
            radius=9,
        ),
    )

    inner_y = y + 82
    inner_h = 185
    gap = 14
    inner_x = x + 24
    inner_total_w = w - 48
    inner_w = (
        inner_total_w - 2 * gap
    ) / 3

    specs = [
        (
            "p01_claims_received",
            "Issuer Comparable Claims Received - Total",
            "Claims Received",
            COLORS["navy"],
            COLORS["neutral_soft"],
        ),
        (
            "p01_claims_denied",
            "Issuer Comparable Claims Denied - Total",
            "Claims Denied",
            COLORS["blue"],
            COLORS["blue_soft"],
        ),
        (
            "p01_claims_rate",
            "Issuer Comparable Overall Denial Rate",
            "Overall Denial Rate",
            COLORS["teal"],
            COLORS["teal_soft"],
        ),
    ]

    for idx, (
        name,
        measure,
        title,
        value_color,
        background,
    ) in enumerate(specs):
        write_visual(
            page_dir,
            make_card(
                name=name,
                measure_name=measure,
                title=title,
                x=inner_x + idx * (
                    inner_w + gap
                ),
                y=inner_y,
                width=inner_w,
                height=inner_h,
                z=800 + idx,
                tab_order=30 + idx,
                value_color=value_color,
                background=background,
                title_size=10,
                value_size=29,
            ),
        )


def add_reason_chart(
    page_dir: Path,
) -> None:
    write_visual(
        page_dir,
        make_reason_bar_chart(
            name="p01_denial_reason_chart",
            x=64,
            y=635,
            width=1098,
            height=325,
            z=700,
            tab_order=40,
        ),
    )


def add_appeals_section(
    page_dir: Path,
) -> None:
    x = 1178
    y = 635
    w = 678
    h = 325

    write_visual(
        page_dir,
        make_textbox(
            name="p01_appeals_surface",
            x=x,
            y=y,
            width=w,
            height=h,
            z=300,
            paragraphs=[
                paragraph(
                    text_run(
                        "APPEALS OUTCOME",
                        size=12,
                        color=COLORS["ink"],
                        semibold=True,
                    )
                ),
                paragraph(
                    text_run(
                        "Filed volume and overturn rate",
                        size=9.5,
                        color=COLORS["secondary"],
                    )
                ),
            ],
            background=COLORS["surface"],
            border=COLORS["border"],
            radius=9,
        ),
    )

    left = x + 22
    top = y + 76
    gap = 14
    card_w = (
        w - 44 - gap
    ) / 2
    card_h = 103

    specs = [
        (
            "p01_appeal_internal_rate",
            "Internal Appeal Overturn Rate",
            "Internal Overturn",
            COLORS["blue"],
            COLORS["blue_soft"],
            left,
            top,
        ),
        (
            "p01_appeal_external_rate",
            "External Appeal Overturn Rate",
            "External Overturn",
            COLORS["teal"],
            COLORS["teal_soft"],
            left + card_w + gap,
            top,
        ),
        (
            "p01_appeal_internal_filed",
            "Internal Appeals Filed",
            "Internal Filed",
            COLORS["ink"],
            COLORS["neutral_soft"],
            left,
            top + card_h + gap,
        ),
        (
            "p01_appeal_external_filed",
            "External Appeals Filed",
            "External Filed",
            COLORS["ink"],
            COLORS["neutral_soft"],
            left + card_w + gap,
            top + card_h + gap,
        ),
    ]

    for idx, (
        name,
        measure,
        title,
        value_color,
        background,
        card_x,
        card_y,
    ) in enumerate(specs):
        write_visual(
            page_dir,
            make_card(
                name=name,
                measure_name=measure,
                title=title,
                x=card_x,
                y=card_y,
                width=card_w,
                height=card_h,
                z=900 + idx,
                tab_order=50 + idx,
                value_color=value_color,
                background=background,
                title_size=9.5,
                value_size=22,
            ),
        )


def add_governance_strip(
    page_dir: Path,
) -> None:
    y = 980
    h = 72

    write_visual(
        page_dir,
        make_card(
            name="p01_dq_open",
            measure_name="Open DQ Exception Count",
            title="Open DQ Exceptions",
            x=64,
            y=y,
            width=260,
            height=h,
            z=600,
            tab_order=60,
            value_color=COLORS["amber"],
            background=COLORS["surface"],
            title_size=9,
            value_size=20,
        ),
    )

    write_visual(
        page_dir,
        make_card(
            name="p01_dq_known",
            measure_name="Known Source Exception Row Count",
            title="Known-Source Rows",
            x=340,
            y=y,
            width=260,
            height=h,
            z=600,
            tab_order=61,
            value_color=COLORS["red"],
            background=COLORS["surface"],
            title_size=9,
            value_size=20,
        ),
    )

    write_visual(
        page_dir,
        make_textbox(
            name="p01_governance_note",
            x=616,
            y=y,
            width=1240,
            height=h,
            z=600,
            paragraphs=[
                paragraph(
                    text_run(
                        (
                            "Governed context  •  Availability-aware KPIs  •  "
                            "Known-source anomaly preserved  •  No clipping or imputation"
                        ),
                        size=9.5,
                        color=COLORS["secondary"],
                        semibold=True,
                    )
                )
            ],
            background=COLORS["neutral_soft"],
            border=COLORS["border"],
            radius=9,
        ),
    )


# ---------------------------------------------------------------------
# Global shell styling requested by user:
# centered page titles + distinctive centered INDEX button.
# ---------------------------------------------------------------------

def set_textbox_centered(
    payload: dict,
) -> None:
    paragraphs = (
        payload.get("visual", {})
        .get("objects", {})
        .get("general", [{}])[0]
        .get("properties", {})
        .get("paragraphs", [])
    )

    for p in paragraphs:
        p["horizontalTextAlignment"] = (
            "center"
        )


def patch_home_surface(
    path: Path,
) -> None:
    payload = read_json(path)
    visual = payload.get(
        "visual",
        {},
    )

    if visual.get(
        "visualType"
    ) != "textbox":
        raise RuntimeError(
            f"Home surface is not textbox: {path}"
        )

    general = (
        visual.get("objects", {})
        .get("general", [{}])[0]
        .get("properties", {})
    )

    general["paragraphs"] = [
        paragraph(
            text_run(
                "INDEX",
                size=11,
                color="#FFFFFF",
                semibold=True,
            ),
            align="center",
        )
    ]

    vco = visual.setdefault(
        "visualContainerObjects",
        {},
    )

    vco["background"] = [
        {
            "properties": {
                "show": expr_bool(True),
                "color": solid_color(
                    COLORS["blue"]
                ),
                "transparency": expr_double(
                    0
                ),
            }
        }
    ]

    vco["border"] = [
        {
            "properties": {
                "show": expr_bool(True),
                "color": solid_color(
                    COLORS["blue_dark"]
                ),
                "width": expr_double(1),
                "radius": expr_double(9),
            }
        }
    ]

    payload["position"]["x"] = 1692
    payload["position"]["y"] = 42
    payload["position"]["width"] = 164
    payload["position"]["height"] = 56

    write_json(
        path,
        payload,
    )


def patch_header_surface(
    path: Path,
) -> None:
    payload = read_json(path)

    if (
        payload.get("visual", {})
        .get("visualType")
        != "textbox"
    ):
        raise RuntimeError(
            f"Header is not textbox: {path}"
        )

    set_textbox_centered(
        payload
    )

    payload["position"]["x"] = 390
    payload["position"]["y"] = 28
    payload["position"]["width"] = 1140
    payload["position"]["height"] = 108

    write_json(
        path,
        payload,
    )


def patch_home_nav(
    path: Path,
) -> None:
    payload = read_json(path)

    payload["position"]["x"] = 1692
    payload["position"]["y"] = 42
    payload["position"]["width"] = 164
    payload["position"]["height"] = 56

    write_json(
        path,
        payload,
    )


def patch_other_page_shells(
    pages: dict[str, tuple[Path, dict]],
) -> list[dict]:
    rows = []

    for display_name in EXPECTED_PAGE_DISPLAY_NAMES:
        if display_name in (
            "00 INDEX",
            "01 Executive Overview",
        ):
            continue

        page_dir, _ = pages[
            display_name
        ]

        visual_files = list(
            (page_dir / "visuals")
            .rglob("visual.json")
        )

        header_path = None
        home_surface_path = None
        home_nav_path = None

        for path in visual_files:
            payload = read_json(path)
            name = payload.get(
                "name",
                "",
            )

            if name.endswith(
                "_header"
            ):
                header_path = path

            elif name.endswith(
                "_home_surface"
            ):
                home_surface_path = path

            elif name.endswith(
                "_home_nav"
            ):
                home_nav_path = path

        if not all(
            [
                header_path,
                home_surface_path,
                home_nav_path,
            ]
        ):
            raise RuntimeError(
                f"Missing shell header/home visuals "
                f"on {display_name}."
            )

        patch_header_surface(
            header_path
        )
        patch_home_surface(
            home_surface_path
        )
        patch_home_nav(
            home_nav_path
        )

        rows.append(
            {
                "Page": display_name,
                "PageTitleCentered": "YES",
                "IndexButtonFill": COLORS["blue"],
                "IndexText": "INDEX",
                "IndexTextCentered": "YES",
                "Status": "PATCHED",
            }
        )

    return rows


# ---------------------------------------------------------------------
# Validation / report
# ---------------------------------------------------------------------

def validate_json_tree() -> list[dict]:
    failures = []

    for path in sorted(
        (REPORT_DIR / "definition")
        .rglob("*.json")
    ):
        try:
            read_json(path)

        except Exception as exc:
            failures.append(
                {
                    "RelativePath": str(
                        path.relative_to(
                            PROJECT_ROOT
                        )
                    ),
                    "Error": str(exc),
                }
            )

    return failures


def validate_visual_binding(
    visual_path: Path,
) -> dict:
    payload = read_json(
        visual_path
    )

    visual = payload.get(
        "visual",
        {},
    )

    fields = []

    def walk(obj):
        if isinstance(obj, dict):
            if "Measure" in obj:
                m = obj["Measure"]
                prop = m.get(
                    "Property"
                )

                if prop:
                    fields.append(
                        (
                            "Measure",
                            "_Measures",
                            prop,
                        )
                    )

            if "Column" in obj:
                c = obj["Column"]
                prop = c.get(
                    "Property"
                )
                source = (
                    c.get(
                        "Expression",
                        {},
                    )
                    .get(
                        "SourceRef",
                        {},
                    )
                    .get(
                        "Entity"
                    )
                )

                if prop:
                    fields.append(
                        (
                            "Column",
                            source,
                            prop,
                        )
                    )

            for value in obj.values():
                walk(value)

        elif isinstance(obj, list):
            for value in obj:
                walk(value)

    walk(
        visual.get(
            "query",
            {},
        )
    )

    return {
        "VisualName": payload.get(
            "name"
        ),
        "VisualType": visual.get(
            "visualType"
        ),
        "Bindings": " | ".join(
            f"{kind}:{entity}.{field}"
            for kind, entity, field
            in fields
        ),
    }


def style_report(
    path: Path,
) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import (
        Alignment,
        Font,
        PatternFill,
    )
    from openpyxl.utils import (
        get_column_letter,
    )

    wb = load_workbook(
        path
    )

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if (
            ws.max_row
            and ws.max_column
        ):
            ws.auto_filter.ref = (
                ws.dimensions
            )

        for cell in ws[1]:
            cell.fill = (
                header_fill
            )
            cell.font = (
                header_font
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            max_len = 0

            for row_idx in range(
                1,
                min(
                    ws.max_row,
                    250,
                )
                + 1,
            ):
                value = ws.cell(
                    row=row_idx,
                    column=col_idx,
                ).value

                if value is not None:
                    max_len = max(
                        max_len,
                        len(str(value)),
                    )

            ws.column_dimensions[
                get_column_letter(
                    col_idx
                )
            ].width = min(
                max(
                    max_len + 2,
                    10,
                ),
                65,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    wb.save(
        path
    )


# ---------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------

def main() -> int:
    print("=" * 78)
    print(
        "Transparency in Coverage PUF — "
        "01 Executive Overview Build"
    )
    print("=" * 78)
    print(
        f"Script version : "
        f"{SCRIPT_VERSION}"
    )
    print(
        f"Report folder  : "
        f"{REPORT_DIR}"
    )
    print(
        f"Backup         : "
        f"{BACKUP_FILE}"
    )
    print(
        f"Review file    : "
        f"{REPORT_FILE}"
    )

    # -----------------------------
    # Safety gates
    # -----------------------------

    if power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is running. "
            "Save the PBIP and close Power BI Desktop "
            "completely before changing PBIR files."
        )

    if not REPORT_DIR.exists():
        raise FileNotFoundError(
            f"Report folder not found: "
            f"{REPORT_DIR}"
        )

    if not MEASURES_TMDL.exists():
        raise FileNotFoundError(
            f"Measure table not found: "
            f"{MEASURES_TMDL}"
        )

    if not DIM_METRIC_TMDL.exists():
        raise FileNotFoundError(
            f"DimMetric TMDL not found: "
            f"{DIM_METRIC_TMDL}"
        )

    pages = find_pages()

    if set(
        EXPECTED_PAGE_DISPLAY_NAMES
    ) != set(
        pages
    ):
        raise RuntimeError(
            "Report page inventory does not match "
            "the approved 11-page shell. "
            "No changes made."
        )

    installed_measures = (
        parse_tmdl_measures(
            MEASURES_TMDL
        )
    )

    missing_measures = sorted(
        set(REQUIRED_MEASURES)
        - installed_measures
    )

    if missing_measures:
        raise RuntimeError(
            "Required measures are missing: "
            + ", ".join(
                missing_measures
            )
        )

    dim_metric_columns = (
        parse_tmdl_columns(
            DIM_METRIC_TMDL
        )
    )

    missing_columns = sorted(
        set(
            REQUIRED_DIM_METRIC_COLUMNS
        )
        - dim_metric_columns
    )

    if missing_columns:
        raise RuntimeError(
            "Required DimMetric columns are missing: "
            + ", ".join(
                missing_columns
            )
        )

    index_dir, index_payload = (
        pages["00 INDEX"]
    )
    executive_dir, _ = (
        pages["01 Executive Overview"]
    )

    index_page_name = (
        index_payload["name"]
    )

    existing_exec_visuals = sorted(
        (
            executive_dir
            / "visuals"
        ).rglob(
            "visual.json"
        )
    )

    unexpected_exec = []

    for path in existing_exec_visuals:
        payload = read_json(
            path
        )
        name = payload.get(
            "name",
            "",
        )

        if not name.startswith(
            "p01_"
        ):
            unexpected_exec.append(
                name
            )

    if unexpected_exec:
        raise RuntimeError(
            "01 Executive Overview contains visuals "
            "outside the approved shell namespace: "
            + ", ".join(
                unexpected_exec
            )
            + ". No changes made."
        )

    # -----------------------------
    # Backup / before state
    # -----------------------------

    create_backup()
    before_hashes = (
        report_hashes()
    )

    # -----------------------------
    # Apply global visual language
    # -----------------------------

    other_shell_patch_rows = (
        patch_other_page_shells(
            pages
        )
    )

    # -----------------------------
    # Rebuild 01 Executive Overview
    # -----------------------------

    visuals_dir = (
        executive_dir
        / "visuals"
    )

    if visuals_dir.exists():
        shutil.rmtree(
            visuals_dir
        )

    add_canvas(
        executive_dir
    )
    add_page_header(
        executive_dir,
        index_page_name,
    )
    add_kpi_row(
        executive_dir
    )
    add_network_section(
        executive_dir
    )
    add_claims_section(
        executive_dir
    )
    add_reason_chart(
        executive_dir
    )
    add_appeals_section(
        executive_dir
    )
    add_governance_strip(
        executive_dir
    )

    # -----------------------------
    # Validation
    # -----------------------------

    json_failures = (
        validate_json_tree()
    )

    exec_visual_files = sorted(
        (
            executive_dir
            / "visuals"
        ).rglob(
            "visual.json"
        )
    )

    binding_rows = []

    for path in exec_visual_files:
        payload = read_json(
            path
        )

        if (
            payload.get(
                "visual",
                {},
            ).get(
                "query"
            )
        ):
            binding_rows.append(
                validate_visual_binding(
                    path
                )
            )

    validation_rows = []

    validation_rows.append(
        {
            "CheckID": "EXE-001",
            "TestName": "01 page exists",
            "Expected": 1,
            "Actual": 1,
            "Status": "PASS",
        }
    )

    validation_rows.append(
        {
            "CheckID": "EXE-002",
            "TestName": "Required measures installed",
            "Expected": len(
                REQUIRED_MEASURES
            ),
            "Actual": len(
                REQUIRED_MEASURES
            )
            - len(
                missing_measures
            ),
            "Status": (
                "PASS"
                if not missing_measures
                else "FAIL"
            ),
        }
    )

    validation_rows.append(
        {
            "CheckID": "EXE-003",
            "TestName": "Required DimMetric columns",
            "Expected": len(
                REQUIRED_DIM_METRIC_COLUMNS
            ),
            "Actual": len(
                REQUIRED_DIM_METRIC_COLUMNS
            )
            - len(
                missing_columns
            ),
            "Status": (
                "PASS"
                if not missing_columns
                else "FAIL"
            ),
        }
    )

    validation_rows.append(
        {
            "CheckID": "EXE-004",
            "TestName": "PBIR JSON parse failures",
            "Expected": 0,
            "Actual": len(
                json_failures
            ),
            "Status": (
                "PASS"
                if not json_failures
                else "FAIL"
            ),
        }
    )

    expected_visual_names = {
        "p01_canvas",
        "p01_accent",
        "p01_context",
        "p01_header",
        "p01_home_surface",
        "p01_home_nav",
        "p01_kpi_issuers",
        "p01_kpi_plans",
        "p01_kpi_claims",
        "p01_kpi_denial_rate",
        "p01_kpi_internal_appeal",
        "p01_network_surface",
        "p01_network_in",
        "p01_network_out",
        "p01_network_gap",
        "p01_claims_surface",
        "p01_claims_received",
        "p01_claims_denied",
        "p01_claims_rate",
        "p01_denial_reason_chart",
        "p01_appeals_surface",
        "p01_appeal_internal_rate",
        "p01_appeal_external_rate",
        "p01_appeal_internal_filed",
        "p01_appeal_external_filed",
        "p01_dq_open",
        "p01_dq_known",
        "p01_governance_note",
    }

    actual_visual_names = set()

    for path in exec_visual_files:
        actual_visual_names.add(
            read_json(
                path
            ).get(
                "name"
            )
        )

    validation_rows.append(
        {
            "CheckID": "EXE-005",
            "TestName": "Executive visual inventory",
            "Expected": len(
                expected_visual_names
            ),
            "Actual": len(
                actual_visual_names
            ),
            "Status": (
                "PASS"
                if actual_visual_names
                == expected_visual_names
                else "FAIL"
            ),
        }
    )

    header_payload = read_json(
        executive_dir
        / "visuals"
        / "p01_header"
        / "visual.json"
    )

    header_paragraphs = (
        header_payload["visual"]
        ["objects"]
        ["general"][0]
        ["properties"]
        ["paragraphs"]
    )

    validation_rows.append(
        {
            "CheckID": "EXE-006",
            "TestName": "Page title centered",
            "Expected": "center",
            "Actual": ",".join(
                sorted(
                    {
                        p.get(
                            "horizontalTextAlignment",
                            "",
                        )
                        for p
                        in header_paragraphs
                    }
                )
            ),
            "Status": (
                "PASS"
                if all(
                    p.get(
                        "horizontalTextAlignment"
                    )
                    == "center"
                    for p
                    in header_paragraphs
                )
                else "FAIL"
            ),
        }
    )

    home_payload = read_json(
        executive_dir
        / "visuals"
        / "p01_home_surface"
        / "visual.json"
    )

    home_text = (
        home_payload["visual"]
        ["objects"]
        ["general"][0]
        ["properties"]
        ["paragraphs"][0]
        ["textRuns"][0]
        ["value"]
    )

    home_align = (
        home_payload["visual"]
        ["objects"]
        ["general"][0]
        ["properties"]
        ["paragraphs"][0]
        ["horizontalTextAlignment"]
    )

    validation_rows.append(
        {
            "CheckID": "EXE-007",
            "TestName": "INDEX button text",
            "Expected": "INDEX / center",
            "Actual": (
                f"{home_text} / "
                f"{home_align}"
            ),
            "Status": (
                "PASS"
                if (
                    home_text == "INDEX"
                    and home_align
                    == "center"
                )
                else "FAIL"
            ),
        }
    )

    failures = sum(
        row["Status"]
        == "FAIL"
        for row
        in validation_rows
    )

    # -----------------------------
    # Evidence
    # -----------------------------

    after_hashes = (
        report_hashes()
    )

    file_changes = []

    for rel_path in sorted(
        set(
            before_hashes
        )
        | set(
            after_hashes
        )
    ):
        before = before_hashes.get(
            rel_path
        )
        after = after_hashes.get(
            rel_path
        )

        if before == after:
            continue

        if (
            before is not None
            and after is not None
        ):
            state = "MODIFIED"

        elif after is not None:
            state = "CREATED"

        else:
            state = "DELETED"

        file_changes.append(
            {
                "RelativePath": rel_path,
                "ChangeState": state,
                "BeforeSHA256": before,
                "AfterSHA256": after,
            }
        )

    visual_inventory = []

    for path in exec_visual_files:
        payload = read_json(
            path
        )

        visual_inventory.append(
            {
                "VisualName": payload.get(
                    "name"
                ),
                "VisualType": (
                    payload.get(
                        "visual",
                        {},
                    ).get(
                        "visualType"
                    )
                ),
                "X": payload.get(
                    "position",
                    {},
                ).get(
                    "x"
                ),
                "Y": payload.get(
                    "position",
                    {},
                ).get(
                    "y"
                ),
                "Width": payload.get(
                    "position",
                    {},
                ).get(
                    "width"
                ),
                "Height": payload.get(
                    "position",
                    {},
                ).get(
                    "height"
                ),
                "RelativePath": str(
                    path.relative_to(
                        PROJECT_ROOT
                    )
                ),
            }
        )

    summary_df = pd.DataFrame(
        [
            (
                "BuildStatus",
                (
                    "PASS"
                    if failures == 0
                    else "FAIL"
                ),
            ),
            (
                "ScriptVersion",
                SCRIPT_VERSION,
            ),
            (
                "BuiltAtLocal",
                datetime.now()
                .isoformat(
                    timespec="seconds"
                ),
            ),
            (
                "Page",
                "01 Executive Overview",
            ),
            (
                "Visuals",
                len(
                    exec_visual_files
                ),
            ),
            (
                "BoundVisuals",
                len(
                    binding_rows
                ),
            ),
            (
                "ValidationChecks",
                len(
                    validation_rows
                ),
            ),
            (
                "ValidationFailures",
                failures,
            ),
            (
                "PageTitleAlignment",
                "center",
            ),
            (
                "ChartCardTitleAlignment",
                "center",
            ),
            (
                "IndexButton",
                (
                    "Distinct blue button; "
                    "INDEX centered"
                ),
            ),
            (
                "BackupFile",
                str(
                    BACKUP_FILE
                ),
            ),
            (
                "EvidenceState",
                (
                    "PBIR EXECUTIVE PAGE BUILD + "
                    "OFFLINE STRUCTURAL VALIDATION"
                ),
            ),
        ],
        columns=[
            "Item",
            "Value",
        ],
    )

    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        REPORT_FILE,
        engine="openpyxl",
    ) as writer:
        summary_df.to_excel(
            writer,
            sheet_name="00_Summary",
            index=False,
        )

        pd.DataFrame(
            validation_rows
        ).to_excel(
            writer,
            sheet_name="01_Validation",
            index=False,
        )

        pd.DataFrame(
            visual_inventory
        ).to_excel(
            writer,
            sheet_name="02_Visuals",
            index=False,
        )

        pd.DataFrame(
            binding_rows
        ).to_excel(
            writer,
            sheet_name="03_Bindings",
            index=False,
        )

        pd.DataFrame(
            other_shell_patch_rows
        ).to_excel(
            writer,
            sheet_name="04_Global_Shell_Style",
            index=False,
        )

        pd.DataFrame(
            file_changes
        ).to_excel(
            writer,
            sheet_name="05_File_Changes",
            index=False,
        )

        if json_failures:
            pd.DataFrame(
                json_failures
            ).to_excel(
                writer,
                sheet_name="06_JSON_Failures",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No JSON parse failures"
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="06_JSON_Failures",
                index=False,
            )

    style_report(
        REPORT_FILE
    )

    print()
    print(
        "01 Executive Overview build completed."
    )
    print(
        f"Build status          : "
        f"{'PASS' if failures == 0 else 'FAIL'}"
    )
    print(
        f"Executive visuals     : "
        f"{len(exec_visual_files)}"
    )
    print(
        f"Bound visuals         : "
        f"{len(binding_rows)}"
    )
    print(
        "Page title alignment  : center"
    )
    print(
        "Chart/card titles      : center"
    )
    print(
        "INDEX button           : blue + centered"
    )
    print(
        f"Validation failures   : "
        f"{failures}"
    )
    print(
        f"Backup                : "
        f"{BACKUP_FILE}"
    )
    print(
        f"Review evidence       : "
        f"{REPORT_FILE}"
    )

    if failures:
        print()
        print(
            "Offline validation failed. "
            "Do not open/save the PBIP until "
            "the evidence workbook is reviewed."
        )

        return 2

    print()
    print(
        "Next: open TransparencyInCoverage.pbip. "
        "Inspect 01 Executive Overview, test the INDEX button, "
        "and confirm that cards/charts render without a PBIR error. "
        "If correct, Save (Ctrl+S)."
    )

    return 0


if __name__ == "__main__":
    try:
        sys.exit(
            main()
        )

    except Exception as exc:
        print()
        print(
            "EXECUTIVE OVERVIEW BUILD FAILED"
        )
        print(
            str(exc)
        )
        sys.exit(
            1
        )
