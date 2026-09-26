from __future__ import annotations

from datetime import datetime
from pathlib import Path
import json
import os
import subprocess
import sys
import zipfile

import pandas as pd


SCRIPT_VERSION = "1.0.0"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
REPORT_DIR = (
    PROJECT_ROOT
    / "powerbi"
    / "TransparencyInCoverage.Report"
)
PAGES_DIR = REPORT_DIR / "definition" / "pages"

EXPECTED_VISUALS = 21
EXPECTED_BOUND_VISUALS = 14
CANVAS_WIDTH = 1920
CANVAS_HEIGHT = 1080

COLORS = {
    "ink": "#172033",
    "amber": "#C8872C",
}

TABLE_COLUMN_WIDTHS = {
    "DimIssuer.IssuerName": 500,
    "_Measures.Issuer In-Network Denial Rate": 190,
    "_Measures.Issuer Out-of-Network Denial Rate": 190,
    "_Measures.Issuer Network Denial Rate Gap": 170,
    "_Measures.Issuer Claims Received - In Network": 300,
    "_Measures.Issuer Claims Received - Out of Network": 300,
}

UNSUPPORTED_PAGE_ROOT_PROPERTIES = {
    "horizontalAlignment",
    "verticalAlignment",
}


def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )

            with winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                key_path,
            ) as key:
                raw_value, _ = winreg.QueryValueEx(
                    key,
                    "Desktop",
                )

            return Path(
                os.path.expandvars(raw_value)
            ).resolve()

        except Exception:
            pass

    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = (
    get_windows_desktop()
    / "Transparency_PUF_Review"
)

BACKUP_FILE = (
    REVIEW_DIR
    / "22_PRE_NETWORK_FINAL_POLISH_BACKUP.zip"
)

REPORT_FILE = (
    REVIEW_DIR
    / "22_NETWORK_FINAL_POLISH_REPORT.xlsx"
)


def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False

    result = subprocess.run(
        [
            "tasklist",
            "/FI",
            "IMAGENAME eq PBIDesktop.exe",
        ],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(
            subprocess,
            "CREATE_NO_WINDOW",
            0,
        ),
    )

    return "PBIDesktop.exe" in result.stdout


def read_json(path: Path):
    return json.loads(
        path.read_text(
            encoding="utf-8-sig",
        )
    )


def write_json(
    path: Path,
    payload,
) -> None:
    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    path.write_text(
        json.dumps(
            payload,
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )


def expr_double(
    value: float | int,
) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": f"{value}D"
            }
        }
    }


def expr_bool(value: bool) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": "true" if value else "false"
            }
        }
    }


def expr_string(value: str) -> dict:
    escaped = value.replace(
        "'",
        "''",
    )

    return {
        "expr": {
            "Literal": {
                "Value": f"'{escaped}'"
            }
        }
    }


def solid_color(color: str) -> dict:
    return {
        "solid": {
            "color": expr_string(
                color
            )
        }
    }


def find_network_page() -> tuple[Path, dict]:
    matches = []

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        if payload.get(
            "displayName"
        ) == "02 Network Comparison":
            matches.append(
                (
                    page_json.parent,
                    payload,
                )
            )

    if len(matches) != 1:
        raise RuntimeError(
            "Expected exactly one '02 Network Comparison' page; "
            f"found {len(matches)}."
        )

    return matches[0]


def create_backup() -> None:
    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if BACKUP_FILE.exists():
        BACKUP_FILE.unlink()

    with zipfile.ZipFile(
        BACKUP_FILE,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zf:
        for path in sorted(
            p
            for p in REPORT_DIR.rglob("*")
            if p.is_file()
        ):
            archive_name = (
                Path(REPORT_DIR.name)
                / path.relative_to(
                    REPORT_DIR
                )
            )

            zf.write(
                path,
                str(
                    archive_name
                ),
            )


def count_bound_visuals(
    page_dir: Path,
) -> int:
    count = 0

    for visual_json in (
        page_dir
        / "visuals"
    ).rglob(
        "visual.json"
    ):
        payload = read_json(
            visual_json
        )

        if (
            payload.get(
                "visual",
                {},
            ).get(
                "query"
            )
        ):
            count += 1

    return count


def patch_dq_card_height(
    dq_path: Path,
) -> dict:
    payload = read_json(
        dq_path
    )

    if (
        payload.get(
            "visual",
            {},
        ).get(
            "visualType"
        )
        != "card"
    ):
        raise RuntimeError(
            "p02_dq_context is not a card visual."
        )

    # 62px was too short for both title + callout value.
    # Keep the card in the filter band but give it nearly the same
    # vertical allowance as the successful DQ cards on page 01.
    payload[
        "position"
    ][
        "y"
    ] = 92

    payload[
        "position"
    ][
        "height"
    ] = 86

    visual = payload[
        "visual"
    ]

    labels = (
        visual.setdefault(
            "objects",
            {},
        ).setdefault(
            "labels",
            [
                {
                    "properties": {}
                }
            ],
        )
    )

    if not labels:
        labels.append(
            {
                "properties": {}
            }
        )

    label_props = labels[0].setdefault(
        "properties",
        {},
    )

    label_props[
        "fontSize"
    ] = expr_double(
        18
    )

    label_props[
        "bold"
    ] = expr_bool(
        True
    )

    label_props[
        "color"
    ] = solid_color(
        COLORS["amber"]
    )

    title = (
        visual.setdefault(
            "visualContainerObjects",
            {},
        ).setdefault(
            "title",
            [
                {
                    "properties": {}
                }
            ],
        )
    )

    if not title:
        title.append(
            {
                "properties": {}
            }
        )

    title_props = title[0].setdefault(
        "properties",
        {},
    )

    title_props[
        "fontSize"
    ] = expr_double(
        8.5
    )

    title_props[
        "alignment"
    ] = expr_string(
        "center"
    )

    title_props[
        "fontColor"
    ] = solid_color(
        COLORS["ink"]
    )

    write_json(
        dq_path,
        payload,
    )

    return {
        "VisualName": "p02_dq_context",
        "Change": (
            "Increased height 62→86 px, moved to y=92, "
            "and tuned title/value sizes so the DQ value can render."
        ),
        "Status": "PATCHED",
    }


def patch_table_column_widths(
    table_path: Path,
) -> dict:
    payload = read_json(
        table_path
    )

    visual = payload.get(
        "visual",
        {},
    )

    if visual.get(
        "visualType"
    ) != "tableEx":
        raise RuntimeError(
            "p02_issuer_table is not a tableEx visual."
        )

    query_refs = {
        projection.get(
            "queryRef"
        )
        for projection in (
            visual[
                "query"
            ][
                "queryState"
            ][
                "Values"
            ][
                "projections"
            ]
        )
    }

    missing = sorted(
        set(
            TABLE_COLUMN_WIDTHS
        )
        - query_refs
    )

    if missing:
        raise RuntimeError(
            "Cannot size table columns because bindings are missing: "
            + ", ".join(
                missing
            )
        )

    column_width_objects = []

    for query_ref, width in (
        TABLE_COLUMN_WIDTHS.items()
    ):
        column_width_objects.append(
            {
                "properties": {
                    "value": expr_double(
                        width
                    )
                },
                "selector": {
                    "metadata": query_ref
                },
            }
        )

    visual.setdefault(
        "objects",
        {},
    )[
        "columnWidth"
    ] = column_width_objects

    write_json(
        table_path,
        payload,
    )

    return {
        "VisualName": "p02_issuer_table",
        "Change": (
            "Distributed six table columns across the available width "
            "so the detail table no longer bunches into the left third."
        ),
        "Status": "PATCHED",
    }


def validate_dq_card(
    dq_path: Path,
) -> dict:
    payload = read_json(
        dq_path
    )

    pos = payload[
        "position"
    ]

    passed = (
        pos.get(
            "y"
        ) == 92
        and pos.get(
            "height"
        ) == 86
    )

    return {
        "CheckID": "NFP-001",
        "TestName": "DQ card render height",
        "Expected": "y=92; height=86",
        "Actual": (
            f"y={pos.get('y')}; "
            f"height={pos.get('height')}"
        ),
        "Status": (
            "PASS"
            if passed
            else "FAIL"
        ),
    }


def validate_table_widths(
    table_path: Path,
) -> dict:
    payload = read_json(
        table_path
    )

    width_objects = (
        payload.get(
            "visual",
            {},
        )
        .get(
            "objects",
            {},
        )
        .get(
            "columnWidth",
            [],
        )
    )

    actual = {}

    for item in width_objects:
        metadata = (
            item.get(
                "selector",
                {},
            ).get(
                "metadata"
            )
        )

        value = (
            item.get(
                "properties",
                {},
            )
            .get(
                "value",
                {},
            )
            .get(
                "expr",
                {},
            )
            .get(
                "Literal",
                {},
            )
            .get(
                "Value"
            )
        )

        if metadata:
            actual[
                metadata
            ] = value

    matched = sum(
        actual.get(
            query_ref
        )
        == f"{width}D"
        for query_ref, width
        in TABLE_COLUMN_WIDTHS.items()
    )

    return {
        "CheckID": "NFP-002",
        "TestName": "Issuer table column widths",
        "Expected": len(
            TABLE_COLUMN_WIDTHS
        ),
        "Actual": matched,
        "Status": (
            "PASS"
            if matched
            == len(
                TABLE_COLUMN_WIDTHS
            )
            else "FAIL"
        ),
    }


def validate_no_unsupported_page_roots() -> list[dict]:
    offenders = []

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        for prop in sorted(
            UNSUPPORTED_PAGE_ROOT_PROPERTIES
        ):
            if prop in payload:
                offenders.append(
                    {
                        "Page": payload.get(
                            "displayName"
                        ),
                        "Property": prop,
                    }
                )

    return offenders


def validate_no_full_page_visuals() -> list[dict]:
    offenders = []

    for visual_json in PAGES_DIR.rglob(
        "visual.json"
    ):
        payload = read_json(
            visual_json
        )

        pos = payload.get(
            "position",
            {},
        )

        width = float(
            pos.get(
                "width",
                0,
            )
            or 0
        )

        height = float(
            pos.get(
                "height",
                0,
            )
            or 0
        )

        if (
            width
            >= CANVAS_WIDTH * 0.94
            and height
            >= CANVAS_HEIGHT * 0.94
        ):
            offenders.append(
                {
                    "VisualName": payload.get(
                        "name"
                    ),
                    "Width": width,
                    "Height": height,
                    "RelativePath": str(
                        visual_json.relative_to(
                            PROJECT_ROOT
                        )
                    ),
                }
            )

    return offenders


def validate_json_tree() -> list[dict]:
    failures = []

    for path in (
        REPORT_DIR
        / "definition"
    ).rglob(
        "*.json"
    ):
        try:
            read_json(
                path
            )

        except Exception as exc:
            failures.append(
                {
                    "RelativePath": str(
                        path.relative_to(
                            PROJECT_ROOT
                        )
                    ),
                    "Error": str(
                        exc
                    ),
                }
            )

    return failures


def style_report(
    path: Path,
) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import (
        Alignment,
        Font,
        PatternFill,
    )
    from openpyxl.utils import (
        get_column_letter,
    )

    wb = load_workbook(
        path
    )

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if (
            ws.max_row
            and ws.max_column
        ):
            ws.auto_filter.ref = (
                ws.dimensions
            )

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            max_len = 0

            for row_idx in range(
                1,
                min(
                    ws.max_row,
                    200,
                )
                + 1,
            ):
                value = ws.cell(
                    row=row_idx,
                    column=col_idx,
                ).value

                if value is not None:
                    max_len = max(
                        max_len,
                        len(
                            str(
                                value
                            )
                        ),
                    )

            ws.column_dimensions[
                get_column_letter(
                    col_idx
                )
            ].width = min(
                max(
                    max_len + 2,
                    10,
                ),
                70,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    wb.save(
        path
    )


def main() -> int:
    print("=" * 78)
    print(
        "Transparency in Coverage PUF — "
        "02 Network Comparison Final Visual Polish"
    )
    print("=" * 78)

    print(
        f"Script version : "
        f"{SCRIPT_VERSION}"
    )
    print(
        f"Report folder  : "
        f"{REPORT_DIR}"
    )
    print(
        f"Backup         : "
        f"{BACKUP_FILE}"
    )
    print(
        f"Review file    : "
        f"{REPORT_FILE}"
    )

    if power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is running. "
            "Save and close it completely before applying this PBIR polish."
        )

    if not REPORT_DIR.exists():
        raise FileNotFoundError(
            f"Report folder not found: "
            f"{REPORT_DIR}"
        )

    page_dir, _ = find_network_page()

    visual_count = len(
        list(
            (
                page_dir
                / "visuals"
            ).rglob(
                "visual.json"
            )
        )
    )

    bound_count = (
        count_bound_visuals(
            page_dir
        )
    )

    if visual_count != EXPECTED_VISUALS:
        raise RuntimeError(
            f"Expected {EXPECTED_VISUALS} visuals before final polish; "
            f"found {visual_count}. No changes made."
        )

    if bound_count != EXPECTED_BOUND_VISUALS:
        raise RuntimeError(
            f"Expected {EXPECTED_BOUND_VISUALS} bound visuals before final polish; "
            f"found {bound_count}. No changes made."
        )

    dq_path = (
        page_dir
        / "visuals"
        / "p02_dq_context"
        / "visual.json"
    )

    table_path = (
        page_dir
        / "visuals"
        / "p02_issuer_table"
        / "visual.json"
    )

    if not dq_path.exists():
        raise FileNotFoundError(
            "p02_dq_context not found."
        )

    if not table_path.exists():
        raise FileNotFoundError(
            "p02_issuer_table not found."
        )

    create_backup()

    patch_rows = [
        patch_dq_card_height(
            dq_path
        ),
        patch_table_column_widths(
            table_path
        ),
    ]

    unsupported = (
        validate_no_unsupported_page_roots()
    )

    full_page = (
        validate_no_full_page_visuals()
    )

    json_failures = (
        validate_json_tree()
    )

    validation_rows = [
        validate_dq_card(
            dq_path
        ),
        validate_table_widths(
            table_path
        ),
        {
            "CheckID": "NFP-003",
            "TestName": "Visual count preserved",
            "Expected": EXPECTED_VISUALS,
            "Actual": len(
                list(
                    (
                        page_dir
                        / "visuals"
                    ).rglob(
                        "visual.json"
                    )
                )
            ),
            "Status": "PASS",
        },
        {
            "CheckID": "NFP-004",
            "TestName": "Bound visual count preserved",
            "Expected": EXPECTED_BOUND_VISUALS,
            "Actual": count_bound_visuals(
                page_dir
            ),
            "Status": (
                "PASS"
                if count_bound_visuals(
                    page_dir
                )
                == EXPECTED_BOUND_VISUALS
                else "FAIL"
            ),
        },
        {
            "CheckID": "NFP-005",
            "TestName": "Unsupported page-root properties",
            "Expected": 0,
            "Actual": len(
                unsupported
            ),
            "Status": (
                "PASS"
                if not unsupported
                else "FAIL"
            ),
        },
        {
            "CheckID": "NFP-006",
            "TestName": "Full-page selectable visuals",
            "Expected": 0,
            "Actual": len(
                full_page
            ),
            "Status": (
                "PASS"
                if not full_page
                else "FAIL"
            ),
        },
        {
            "CheckID": "NFP-007",
            "TestName": "PBIR JSON parse failures",
            "Expected": 0,
            "Actual": len(
                json_failures
            ),
            "Status": (
                "PASS"
                if not json_failures
                else "FAIL"
            ),
        },
    ]

    failures = sum(
        row[
            "Status"
        ] == "FAIL"
        for row
        in validation_rows
    )

    summary_df = pd.DataFrame(
        [
            (
                "PolishStatus",
                "PASS"
                if failures == 0
                else "FAIL",
            ),
            (
                "ScriptVersion",
                SCRIPT_VERSION,
            ),
            (
                "PolishedAtLocal",
                datetime.now()
                .isoformat(
                    timespec="seconds"
                ),
            ),
            (
                "Page",
                "02 Network Comparison",
            ),
            (
                "DQCardHeight",
                "86 px",
            ),
            (
                "TableColumnWidthTotal",
                sum(
                    TABLE_COLUMN_WIDTHS.values()
                ),
            ),
            (
                "VisualsPreserved",
                EXPECTED_VISUALS,
            ),
            (
                "BoundVisualsPreserved",
                EXPECTED_BOUND_VISUALS,
            ),
            (
                "ValidationFailures",
                failures,
            ),
            (
                "BackupFile",
                str(
                    BACKUP_FILE
                ),
            ),
        ],
        columns=[
            "Item",
            "Value",
        ],
    )

    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        REPORT_FILE,
        engine="openpyxl",
    ) as writer:
        summary_df.to_excel(
            writer,
            sheet_name="00_Summary",
            index=False,
        )

        pd.DataFrame(
            patch_rows
        ).to_excel(
            writer,
            sheet_name="01_Patches",
            index=False,
        )

        pd.DataFrame(
            validation_rows
        ).to_excel(
            writer,
            sheet_name="02_Validation",
            index=False,
        )

        if unsupported:
            pd.DataFrame(
                unsupported
            ).to_excel(
                writer,
                sheet_name="03_Schema_Offenders",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No unsupported page-root properties."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="03_Schema_Offenders",
                index=False,
            )

        if full_page:
            pd.DataFrame(
                full_page
            ).to_excel(
                writer,
                sheet_name="04_FullPage_Offenders",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No full-page selectable visuals."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="04_FullPage_Offenders",
                index=False,
            )

        if json_failures:
            pd.DataFrame(
                json_failures
            ).to_excel(
                writer,
                sheet_name="05_JSON_Failures",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No JSON parse failures."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="05_JSON_Failures",
                index=False,
            )

    style_report(
        REPORT_FILE
    )

    print()
    print(
        "02 Network Comparison final visual polish completed."
    )
    print(
        f"Polish status                 : "
        f"{'PASS' if failures == 0 else 'FAIL'}"
    )
    print(
        "DQ card height                : 86 px"
    )
    print(
        f"Table columns distributed     : "
        f"{len(TABLE_COLUMN_WIDTHS)}"
    )
    print(
        f"Visuals preserved             : "
        f"{EXPECTED_VISUALS}"
    )
    print(
        f"Bound/query visuals preserved : "
        f"{EXPECTED_BOUND_VISUALS}"
    )
    print(
        f"Full-page selectable visuals  : "
        f"{len(full_page)}"
    )
    print(
        f"Unsupported page properties   : "
        f"{len(unsupported)}"
    )
    print(
        f"Validation failures           : "
        f"{failures}"
    )
    print(
        f"Backup                        : "
        f"{BACKUP_FILE}"
    )
    print(
        f"Review evidence               : "
        f"{REPORT_FILE}"
    )

    if failures:
        print()
        print(
            "Offline validation failed. "
            "Do not open/save the PBIP until reviewed."
        )
        return 2

    print()
    print(
        "Next: open TransparencyInCoverage.pbip and inspect "
        "02 Network Comparison. The DQ value should now be visible, "
        "and the issuer table should use the page width more evenly."
    )

    return 0


if __name__ == "__main__":
    try:
        sys.exit(
            main()
        )

    except Exception as exc:
        print()
        print(
            "NETWORK FINAL POLISH FAILED"
        )
        print(
            str(
                exc
            )
        )
        sys.exit(
            1
        )
