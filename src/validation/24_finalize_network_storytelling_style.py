from __future__ import annotations

from datetime import datetime
from pathlib import Path
import copy
import json
import os
import subprocess
import sys
import zipfile

import pandas as pd


SCRIPT_VERSION = "1.0.0"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
REPORT_DIR = (
    PROJECT_ROOT
    / "powerbi"
    / "TransparencyInCoverage.Report"
)
PAGES_DIR = REPORT_DIR / "definition" / "pages"

EXPECTED_VISUALS = 34
EXPECTED_BOUND_VISUALS = 14
CANVAS_WIDTH = 1920
CANVAS_HEIGHT = 1080

COLORS = {
    "ink": "#172033",
    "navy": "#17324D",
    "blue": "#4F6BED",
    "blue_soft": "#EEF3FF",
    "blue_mid": "#D9E2FF",
    "blue_strong": "#C4D1FF",
    "amber": "#C8872C",
    "amber_dark": "#9C641A",
    "amber_soft": "#FFF4E3",
    "red": "#C64B4B",
    "red_dark": "#9F2F2F",
    "red_soft": "#FCE8E8",
    "teal": "#16A6A1",
    "teal_dark": "#0E7774",
    "teal_soft": "#EAF8F6",
    "white": "#FFFFFF",
}

GAP_MEASURE = "Issuer Network Denial Rate Gap"
IN_CLAIMS_MEASURE = "Issuer Claims Received - In Network"
OON_CLAIMS_MEASURE = "Issuer Claims Received - Out of Network"

GAP_REF = f"_Measures.{GAP_MEASURE}"
IN_CLAIMS_REF = f"_Measures.{IN_CLAIMS_MEASURE}"
OON_CLAIMS_REF = f"_Measures.{OON_CLAIMS_MEASURE}"

UNSUPPORTED_PAGE_ROOT_PROPERTIES = {
    "horizontalAlignment",
    "verticalAlignment",
}


def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )

            with winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                key_path,
            ) as key:
                raw_value, _ = winreg.QueryValueEx(
                    key,
                    "Desktop",
                )

            return Path(
                os.path.expandvars(raw_value)
            ).resolve()

        except Exception:
            pass

    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = (
    get_windows_desktop()
    / "Transparency_PUF_Review"
)

BACKUP_FILE = (
    REVIEW_DIR
    / "24_PRE_NETWORK_STORY_STYLE_LOCK_BACKUP.zip"
)

REPORT_FILE = (
    REVIEW_DIR
    / "24_NETWORK_STORY_STYLE_LOCK_REPORT.xlsx"
)


# =====================================================================
# Basic helpers
# =====================================================================

def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False

    result = subprocess.run(
        [
            "tasklist",
            "/FI",
            "IMAGENAME eq PBIDesktop.exe",
        ],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(
            subprocess,
            "CREATE_NO_WINDOW",
            0,
        ),
    )

    return "PBIDesktop.exe" in result.stdout


def read_json(path: Path):
    return json.loads(
        path.read_text(
            encoding="utf-8-sig",
        )
    )


def write_json(
    path: Path,
    payload,
) -> None:
    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    path.write_text(
        json.dumps(
            payload,
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )


def expr_bool(value: bool) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": "true" if value else "false"
            }
        }
    }


def expr_double(
    value: float | int,
) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": f"{value}D"
            }
        }
    }


def expr_string(value: str) -> dict:
    escaped = value.replace(
        "'",
        "''",
    )

    return {
        "expr": {
            "Literal": {
                "Value": f"'{escaped}'"
            }
        }
    }


def solid_color(
    color: str,
) -> dict:
    return {
        "solid": {
            "color": expr_string(
                color
            )
        }
    }


def measure_expr(
    measure_name: str,
) -> dict:
    return {
        "Measure": {
            "Expression": {
                "SourceRef": {
                    "Entity": "_Measures"
                }
            },
            "Property": measure_name,
        }
    }


def literal_expr(
    value: float | int,
) -> dict:
    return {
        "Literal": {
            "Value": f"{value}D"
        }
    }


def find_network_page() -> tuple[Path, dict]:
    matches = []

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        if payload.get(
            "displayName"
        ) == "02 Network Comparison":
            matches.append(
                (
                    page_json.parent,
                    payload,
                )
            )

    if len(matches) != 1:
        raise RuntimeError(
            "Expected exactly one '02 Network Comparison' page; "
            f"found {len(matches)}."
        )

    return matches[0]


def create_backup() -> None:
    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if BACKUP_FILE.exists():
        BACKUP_FILE.unlink()

    with zipfile.ZipFile(
        BACKUP_FILE,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zf:
        for path in sorted(
            p
            for p in REPORT_DIR.rglob("*")
            if p.is_file()
        ):
            archive_name = (
                Path(REPORT_DIR.name)
                / path.relative_to(
                    REPORT_DIR
                )
            )

            zf.write(
                path,
                str(
                    archive_name
                ),
            )


def count_bound_visuals(
    page_dir: Path,
) -> int:
    count = 0

    for visual_json in (
        page_dir
        / "visuals"
    ).rglob(
        "visual.json"
    ):
        payload = read_json(
            visual_json
        )

        if (
            payload.get(
                "visual",
                {},
            ).get(
                "query"
            )
        ):
            count += 1

    return count


def capture_positions(
    page_dir: Path,
) -> dict[str, dict]:
    positions = {}

    for visual_json in (
        page_dir
        / "visuals"
    ).rglob(
        "visual.json"
    ):
        payload = read_json(
            visual_json
        )

        name = payload.get(
            "name"
        )

        if name:
            positions[
                name
            ] = copy.deepcopy(
                payload.get(
                    "position",
                    {},
                )
            )

    return positions


# =====================================================================
# Conditional-formatting expressions
# =====================================================================

def comparison(
    measure_name: str,
    kind: int,
    threshold: float,
) -> dict:
    return {
        "Comparison": {
            "ComparisonKind": kind,
            "Left": measure_expr(
                measure_name
            ),
            "Right": literal_expr(
                threshold
            ),
        }
    }


def gap_conditional_color(
    *,
    soft: bool,
) -> dict:
    if soft:
        negative = COLORS["teal_soft"]
        small = COLORS["blue_soft"]
        moderate = "#E3E9FF"
        large = COLORS["amber_soft"]
        extreme = COLORS["red_soft"]
    else:
        negative = COLORS["teal_dark"]
        small = COLORS["navy"]
        moderate = COLORS["blue"]
        large = COLORS["amber_dark"]
        extreme = COLORS["red_dark"]

    return {
        "Conditional": {
            "Cases": [
                {
                    "Condition": comparison(
                        GAP_MEASURE,
                        2,
                        0.50,
                    ),
                    "Value": {
                        "Literal": {
                            "Value": f"'{extreme}'"
                        }
                    },
                },
                {
                    "Condition": comparison(
                        GAP_MEASURE,
                        2,
                        0.25,
                    ),
                    "Value": {
                        "Literal": {
                            "Value": f"'{large}'"
                        }
                    },
                },
                {
                    "Condition": comparison(
                        GAP_MEASURE,
                        2,
                        0.10,
                    ),
                    "Value": {
                        "Literal": {
                            "Value": f"'{moderate}'"
                        }
                    },
                },
                {
                    "Condition": comparison(
                        GAP_MEASURE,
                        2,
                        0,
                    ),
                    "Value": {
                        "Literal": {
                            "Value": f"'{small}'"
                        }
                    },
                },
                {
                    "Condition": comparison(
                        GAP_MEASURE,
                        4,
                        0,
                    ),
                    "Value": {
                        "Literal": {
                            "Value": f"'{negative}'"
                        }
                    },
                },
            ]
        }
    }


def linear_gradient_fill(
    measure_name: str,
    *,
    min_color: str,
    mid_color: str,
    max_color: str,
) -> dict:
    return {
        "FillRule": {
            "Input": measure_expr(
                measure_name
            ),
            "FillRule": {
                "linearGradient3": {
                    "min": {
                        "color": {
                            "Literal": {
                                "Value": f"'{min_color}'"
                            }
                        }
                    },
                    "mid": {
                        "color": {
                            "Literal": {
                                "Value": f"'{mid_color}'"
                            }
                        }
                    },
                    "max": {
                        "color": {
                            "Literal": {
                                "Value": f"'{max_color}'"
                            }
                        }
                    },
                    "nullColoringStrategy": {
                        "strategy": {
                            "Literal": {
                                "Value": "'noColor'"
                            }
                        }
                    },
                }
            },
        }
    }


# =====================================================================
# Patches — STYLE ONLY, positions are never written
# =====================================================================

def patch_gap_hero_color(
    gap_card_path: Path,
) -> dict:
    payload = read_json(
        gap_card_path
    )

    if (
        payload.get(
            "visual",
            {},
        ).get(
            "visualType"
        )
        != "card"
    ):
        raise RuntimeError(
            "p02_kpi_gap is not a card visual."
        )

    labels = (
        payload[
            "visual"
        ]
        .setdefault(
            "objects",
            {},
        )
        .setdefault(
            "labels",
            [
                {
                    "properties": {}
                }
            ],
        )
    )

    if not labels:
        labels.append(
            {
                "properties": {}
            }
        )

    props = labels[0].setdefault(
        "properties",
        {},
    )

    props[
        "color"
    ] = solid_color(
        COLORS["amber_dark"]
    )

    props[
        "bold"
    ] = expr_bool(
        True
    )

    write_json(
        gap_card_path,
        payload,
    )

    return {
        "VisualName": "p02_kpi_gap",
        "Change": (
            "Network Gap callout changed from teal to dark amber "
            "to match its attention-story title band."
        ),
        "GeometryChanged": "NO",
        "Status": "PATCHED",
    }


def patch_investigation_table(
    table_path: Path,
) -> dict:
    payload = read_json(
        table_path
    )

    visual = payload.get(
        "visual",
        {},
    )

    if visual.get(
        "visualType"
    ) != "tableEx":
        raise RuntimeError(
            "p02_issuer_table is not a tableEx visual."
        )

    objects = visual.setdefault(
        "objects",
        {},
    )

    # Preserve existing objects except for values entries targeting
    # the three storytelling metrics, which are rebuilt deterministically.
    existing_values = objects.get(
        "values",
        [],
    )

    preserved_values = []

    target_metadata = {
        GAP_REF,
        IN_CLAIMS_REF,
        OON_CLAIMS_REF,
    }

    for item in existing_values:
        metadata = (
            item.get(
                "selector",
                {},
            ).get(
                "metadata"
            )
        )

        if metadata in target_metadata:
            continue

        preserved_values.append(
            item
        )

    story_values = [
        # Gap: strongest attention signal.
        {
            "properties": {
                "backColor": {
                    "solid": {
                        "color": {
                            "expr": gap_conditional_color(
                                soft=True
                            )
                        }
                    }
                },
                "fontColor": {
                    "solid": {
                        "color": {
                            "expr": gap_conditional_color(
                                soft=False
                            )
                        }
                    }
                },
            },
            "selector": {
                "data": [
                    {
                        "dataViewWildcard": {
                            "matchingOption": 1
                        }
                    }
                ],
                "metadata": GAP_REF,
            },
        },
        # Claims: soft intensity shading. This is deliberately subtle;
        # the number remains readable while the eye sees relative volume.
        {
            "properties": {
                "backColor": {
                    "solid": {
                        "color": {
                            "expr": linear_gradient_fill(
                                IN_CLAIMS_MEASURE,
                                min_color=COLORS["white"],
                                mid_color="#EFF3F8",
                                max_color="#D7E3F1",
                            )
                        }
                    }
                }
            },
            "selector": {
                "data": [
                    {
                        "dataViewWildcard": {
                            "matchingOption": 1
                        }
                    }
                ],
                "metadata": IN_CLAIMS_REF,
            },
        },
        {
            "properties": {
                "backColor": {
                    "solid": {
                        "color": {
                            "expr": linear_gradient_fill(
                                OON_CLAIMS_MEASURE,
                                min_color=COLORS["white"],
                                mid_color=COLORS["blue_soft"],
                                max_color=COLORS["blue_strong"],
                            )
                        }
                    }
                }
            },
            "selector": {
                "data": [
                    {
                        "dataViewWildcard": {
                            "matchingOption": 1
                        }
                    }
                ],
                "metadata": OON_CLAIMS_REF,
            },
        },
    ]

    objects[
        "values"
    ] = (
        preserved_values
        + story_values
    )

    # Keep/strengthen dataBars too. If Desktop renders them in the
    # installed build they add a second cue; the gradient remains the
    # reliable fallback.
    column_formatting = objects.get(
        "columnFormatting",
        [],
    )

    new_column_formatting = []

    for item in column_formatting:
        metadata = (
            item.get(
                "selector",
                {},
            ).get(
                "metadata"
            )
        )

        if metadata in {
            IN_CLAIMS_REF,
            OON_CLAIMS_REF,
        }:
            continue

        new_column_formatting.append(
            item
        )

    new_column_formatting.extend(
        [
            {
                "properties": {
                    "alignment": expr_string(
                        "right"
                    ),
                    "dataBars": {
                        "positiveColor": solid_color(
                            "#B8C9DD"
                        ),
                        "negativeColor": solid_color(
                            "#B8C9DD"
                        ),
                        "axisColor": solid_color(
                            COLORS["white"]
                        ),
                        "reverseDirection": expr_bool(
                            False
                        ),
                        "hideText": expr_bool(
                            False
                        ),
                    },
                },
                "selector": {
                    "metadata": IN_CLAIMS_REF
                },
            },
            {
                "properties": {
                    "alignment": expr_string(
                        "right"
                    ),
                    "dataBars": {
                        "positiveColor": solid_color(
                            "#8FA8FF"
                        ),
                        "negativeColor": solid_color(
                            "#8FA8FF"
                        ),
                        "axisColor": solid_color(
                            COLORS["white"]
                        ),
                        "reverseDirection": expr_bool(
                            False
                        ),
                        "hideText": expr_bool(
                            False
                        ),
                    },
                },
                "selector": {
                    "metadata": OON_CLAIMS_REF
                },
            },
        ]
    )

    objects[
        "columnFormatting"
    ] = new_column_formatting

    write_json(
        table_path,
        payload,
    )

    return {
        "VisualName": "p02_issuer_table",
        "Change": (
            "Strengthened investigation story: Gap keeps conditional "
            "attention colors; IN/OON Claims receive soft volume-intensity "
            "background gradients plus data-bar cues where supported."
        ),
        "GeometryChanged": "NO",
        "Status": "PATCHED",
    }


def patch_legend_text(
    legend_path: Path,
) -> dict:
    payload = read_json(
        legend_path
    )

    if (
        payload.get(
            "visual",
            {},
        ).get(
            "visualType"
        )
        != "textbox"
    ):
        raise RuntimeError(
            "p02_gap_legend is not a textbox."
        )

    paragraphs = (
        payload[
            "visual"
        ][
            "objects"
        ][
            "general"
        ][0][
            "properties"
        ][
            "paragraphs"
        ]
    )

    # Preserve the actual geometry; only make the indicator wording clearer.
    runs = paragraphs[0][
        "textRuns"
    ]

    labels = [
        "▼ NEGATIVE / LOWER GAP",
        "     ● 0–10pp",
        "     ▲ 10–25pp",
        "     ▲▲ 25–50pp",
        "     ▲▲▲ ≥50pp",
    ]

    if len(
        runs
    ) >= len(
        labels
    ):
        for run, label in zip(
            runs,
            labels,
        ):
            run[
                "value"
            ] = label

    write_json(
        legend_path,
        payload,
    )

    return {
        "VisualName": "p02_gap_legend",
        "Change": (
            "Made the existing arrow legend more explicit: increasing "
            "arrow count communicates increasing gap magnitude."
        ),
        "GeometryChanged": "NO",
        "Status": "PATCHED",
    }


# =====================================================================
# Validation
# =====================================================================

def walk_json(obj):
    if isinstance(
        obj,
        dict,
    ):
        yield obj

        for value in obj.values():
            yield from walk_json(
                value
            )

    elif isinstance(
        obj,
        list,
    ):
        for value in obj:
            yield from walk_json(
                value
            )


def validate_layout_unchanged(
    before: dict[str, dict],
    after: dict[str, dict],
) -> tuple[dict, list[dict]]:
    changed = []

    names = sorted(
        set(
            before
        )
        | set(
            after
        )
    )

    for name in names:
        if (
            before.get(
                name
            )
            != after.get(
                name
            )
        ):
            changed.append(
                {
                    "VisualName": name,
                    "BeforePosition": json.dumps(
                        before.get(
                            name
                        ),
                        sort_keys=True,
                    ),
                    "AfterPosition": json.dumps(
                        after.get(
                            name
                        ),
                        sort_keys=True,
                    ),
                }
            )

    return (
        {
            "CheckID": "LOCK-001",
            "TestName": (
                "Manual visual geometry preserved"
            ),
            "Expected": 0,
            "Actual": len(
                changed
            ),
            "Status": (
                "PASS"
                if not changed
                else "FAIL"
            ),
        },
        changed,
    )


def validate_gap_card(
    gap_card_path: Path,
) -> dict:
    payload = read_json(
        gap_card_path
    )

    color = (
        payload[
            "visual"
        ][
            "objects"
        ][
            "labels"
        ][0][
            "properties"
        ][
            "color"
        ][
            "solid"
        ][
            "color"
        ][
            "expr"
        ][
            "Literal"
        ][
            "Value"
        ]
    )

    expected = (
        f"'{COLORS['amber_dark']}'"
    )

    return {
        "CheckID": "LOCK-002",
        "TestName": (
            "Network Gap hero color"
        ),
        "Expected": expected,
        "Actual": color,
        "Status": (
            "PASS"
            if color == expected
            else "FAIL"
        ),
    }


def validate_table_story(
    table_path: Path,
) -> dict:
    payload = read_json(
        table_path
    )

    objects = (
        payload[
            "visual"
        ].get(
            "objects",
            {},
        )
    )

    conditional_metadata = set()
    data_bar_metadata = set()

    for item in objects.get(
        "values",
        [],
    ):
        metadata = (
            item.get(
                "selector",
                {},
            ).get(
                "metadata"
            )
        )

        if metadata and any(
            "FillRule" in node
            or "Conditional" in node
            for node in walk_json(
                item
            )
        ):
            conditional_metadata.add(
                metadata
            )

    for item in objects.get(
        "columnFormatting",
        [],
    ):
        metadata = (
            item.get(
                "selector",
                {},
            ).get(
                "metadata"
            )
        )

        if (
            metadata
            and "dataBars"
            in item.get(
                "properties",
                {},
            )
        ):
            data_bar_metadata.add(
                metadata
            )

    expected_conditional = {
        GAP_REF,
        IN_CLAIMS_REF,
        OON_CLAIMS_REF,
    }

    expected_bars = {
        IN_CLAIMS_REF,
        OON_CLAIMS_REF,
    }

    passed = (
        expected_conditional.issubset(
            conditional_metadata
        )
        and expected_bars.issubset(
            data_bar_metadata
        )
    )

    return {
        "CheckID": "LOCK-003",
        "TestName": (
            "Table story cues"
        ),
        "Expected": (
            "Gap conditional + IN/OON claim gradients + IN/OON data bars"
        ),
        "Actual": (
            f"conditional={sorted(conditional_metadata)}; "
            f"dataBars={sorted(data_bar_metadata)}"
        ),
        "Status": (
            "PASS"
            if passed
            else "FAIL"
        ),
    }


def validate_no_unsupported_page_roots() -> list[dict]:
    offenders = []

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        for prop in sorted(
            UNSUPPORTED_PAGE_ROOT_PROPERTIES
        ):
            if prop in payload:
                offenders.append(
                    {
                        "Page": payload.get(
                            "displayName"
                        ),
                        "Property": prop,
                    }
                )

    return offenders


def validate_no_full_page_visuals() -> list[dict]:
    offenders = []

    for visual_json in PAGES_DIR.rglob(
        "visual.json"
    ):
        payload = read_json(
            visual_json
        )

        pos = payload.get(
            "position",
            {},
        )

        width = float(
            pos.get(
                "width",
                0,
            )
            or 0
        )

        height = float(
            pos.get(
                "height",
                0,
            )
            or 0
        )

        if (
            width
            >= CANVAS_WIDTH * 0.94
            and height
            >= CANVAS_HEIGHT * 0.94
        ):
            offenders.append(
                {
                    "VisualName": payload.get(
                        "name"
                    ),
                    "Width": width,
                    "Height": height,
                    "RelativePath": str(
                        visual_json.relative_to(
                            PROJECT_ROOT
                        )
                    ),
                }
            )

    return offenders


def validate_json_tree() -> list[dict]:
    failures = []

    for path in (
        REPORT_DIR
        / "definition"
    ).rglob(
        "*.json"
    ):
        try:
            read_json(
                path
            )

        except Exception as exc:
            failures.append(
                {
                    "RelativePath": str(
                        path.relative_to(
                            PROJECT_ROOT
                        )
                    ),
                    "Error": str(
                        exc
                    ),
                }
            )

    return failures


# =====================================================================
# Evidence workbook
# =====================================================================

def style_report(
    path: Path,
) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import (
        Alignment,
        Font,
        PatternFill,
    )
    from openpyxl.utils import (
        get_column_letter,
    )

    wb = load_workbook(
        path
    )

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if (
            ws.max_row
            and ws.max_column
        ):
            ws.auto_filter.ref = (
                ws.dimensions
            )

        for cell in ws[1]:
            cell.fill = (
                header_fill
            )
            cell.font = (
                header_font
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            max_len = 0

            for row_idx in range(
                1,
                min(
                    ws.max_row,
                    250,
                )
                + 1,
            ):
                value = ws.cell(
                    row=row_idx,
                    column=col_idx,
                ).value

                if value is not None:
                    max_len = max(
                        max_len,
                        len(
                            str(
                                value
                            )
                        ),
                    )

            ws.column_dimensions[
                get_column_letter(
                    col_idx
                )
            ].width = min(
                max(
                    max_len + 2,
                    10,
                ),
                75,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    wb.save(
        path
    )


# =====================================================================
# Main
# =====================================================================

def main() -> int:
    print("=" * 78)
    print(
        "Transparency in Coverage PUF — "
        "02 Network Story Style Lock"
    )
    print("=" * 78)

    print(
        f"Script version : "
        f"{SCRIPT_VERSION}"
    )
    print(
        f"Report folder  : "
        f"{REPORT_DIR}"
    )
    print(
        f"Backup         : "
        f"{BACKUP_FILE}"
    )
    print(
        f"Review file    : "
        f"{REPORT_FILE}"
    )

    if power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is running. "
            "Save and close it completely before applying this PBIR style patch."
        )

    if not REPORT_DIR.exists():
        raise FileNotFoundError(
            f"Report folder not found: "
            f"{REPORT_DIR}"
        )

    page_dir, _ = find_network_page()

    visual_count = len(
        list(
            (
                page_dir
                / "visuals"
            ).rglob(
                "visual.json"
            )
        )
    )

    bound_count = (
        count_bound_visuals(
            page_dir
        )
    )

    if visual_count != EXPECTED_VISUALS:
        raise RuntimeError(
            f"Expected current Storytelling page to have "
            f"{EXPECTED_VISUALS} visuals; found {visual_count}. "
            "No changes made."
        )

    if bound_count != EXPECTED_BOUND_VISUALS:
        raise RuntimeError(
            f"Expected current Storytelling page to have "
            f"{EXPECTED_BOUND_VISUALS} bound/query visuals; found {bound_count}. "
            "No changes made."
        )

    gap_card_path = (
        page_dir
        / "visuals"
        / "p02_kpi_gap"
        / "visual.json"
    )

    table_path = (
        page_dir
        / "visuals"
        / "p02_issuer_table"
        / "visual.json"
    )

    legend_path = (
        page_dir
        / "visuals"
        / "p02_gap_legend"
        / "visual.json"
    )

    for required in (
        gap_card_path,
        table_path,
        legend_path,
    ):
        if not required.exists():
            raise FileNotFoundError(
                f"Required visual file not found: {required}"
            )

    # The user's saved manual layout is now the immutable baseline.
    positions_before = (
        capture_positions(
            page_dir
        )
    )

    create_backup()

    patch_rows = [
        patch_gap_hero_color(
            gap_card_path
        ),
        patch_investigation_table(
            table_path
        ),
        patch_legend_text(
            legend_path
        ),
    ]

    positions_after = (
        capture_positions(
            page_dir
        )
    )

    layout_check, geometry_changes = (
        validate_layout_unchanged(
            positions_before,
            positions_after,
        )
    )

    unsupported = (
        validate_no_unsupported_page_roots()
    )

    full_page = (
        validate_no_full_page_visuals()
    )

    json_failures = (
        validate_json_tree()
    )

    validation_rows = [
        layout_check,
        validate_gap_card(
            gap_card_path
        ),
        validate_table_story(
            table_path
        ),
        {
            "CheckID": "LOCK-004",
            "TestName": "Visual count preserved",
            "Expected": EXPECTED_VISUALS,
            "Actual": len(
                list(
                    (
                        page_dir
                        / "visuals"
                    ).rglob(
                        "visual.json"
                    )
                )
            ),
            "Status": (
                "PASS"
                if len(
                    list(
                        (
                            page_dir
                            / "visuals"
                        ).rglob(
                            "visual.json"
                        )
                    )
                )
                == EXPECTED_VISUALS
                else "FAIL"
            ),
        },
        {
            "CheckID": "LOCK-005",
            "TestName": "Bound/query visual count preserved",
            "Expected": EXPECTED_BOUND_VISUALS,
            "Actual": count_bound_visuals(
                page_dir
            ),
            "Status": (
                "PASS"
                if count_bound_visuals(
                    page_dir
                )
                == EXPECTED_BOUND_VISUALS
                else "FAIL"
            ),
        },
        {
            "CheckID": "LOCK-006",
            "TestName": "Unsupported page-root properties",
            "Expected": 0,
            "Actual": len(
                unsupported
            ),
            "Status": (
                "PASS"
                if not unsupported
                else "FAIL"
            ),
        },
        {
            "CheckID": "LOCK-007",
            "TestName": "Full-page selectable visuals",
            "Expected": 0,
            "Actual": len(
                full_page
            ),
            "Status": (
                "PASS"
                if not full_page
                else "FAIL"
            ),
        },
        {
            "CheckID": "LOCK-008",
            "TestName": "PBIR JSON parse failures",
            "Expected": 0,
            "Actual": len(
                json_failures
            ),
            "Status": (
                "PASS"
                if not json_failures
                else "FAIL"
            ),
        },
    ]

    failures = sum(
        row[
            "Status"
        ] == "FAIL"
        for row
        in validation_rows
    )

    summary_df = pd.DataFrame(
        [
            (
                "PatchStatus",
                "PASS"
                if failures == 0
                else "FAIL",
            ),
            (
                "ScriptVersion",
                SCRIPT_VERSION,
            ),
            (
                "PatchedAtLocal",
                datetime.now()
                .isoformat(
                    timespec="seconds"
                ),
            ),
            (
                "Page",
                "02 Network Comparison",
            ),
            (
                "ManualLayoutPreserved",
                (
                    "YES"
                    if not geometry_changes
                    else "NO"
                ),
            ),
            (
                "NetworkGapCalloutColor",
                COLORS["amber_dark"],
            ),
            (
                "GapCellStory",
                (
                    "Conditional magnitude shading + font color"
                ),
            ),
            (
                "ClaimsStory",
                (
                    "Soft value-intensity gradients + data-bar cues"
                ),
            ),
            (
                "ArrowLegend",
                (
                    "▼ / ● / ▲ / ▲▲ / ▲▲▲"
                ),
            ),
            (
                "VisualsPreserved",
                EXPECTED_VISUALS,
            ),
            (
                "BoundVisualsPreserved",
                EXPECTED_BOUND_VISUALS,
            ),
            (
                "ValidationFailures",
                failures,
            ),
            (
                "BackupFile",
                str(
                    BACKUP_FILE
                ),
            ),
        ],
        columns=[
            "Item",
            "Value",
        ],
    )

    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        REPORT_FILE,
        engine="openpyxl",
    ) as writer:
        summary_df.to_excel(
            writer,
            sheet_name="00_Summary",
            index=False,
        )

        pd.DataFrame(
            patch_rows
        ).to_excel(
            writer,
            sheet_name="01_Patches",
            index=False,
        )

        pd.DataFrame(
            validation_rows
        ).to_excel(
            writer,
            sheet_name="02_Validation",
            index=False,
        )

        if geometry_changes:
            pd.DataFrame(
                geometry_changes
            ).to_excel(
                writer,
                sheet_name="03_Geometry_Changes",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No x/y/width/height/z/tabOrder values changed. "
                            "User's saved manual layout was preserved exactly."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="03_Geometry_Changes",
                index=False,
            )

        if unsupported:
            pd.DataFrame(
                unsupported
            ).to_excel(
                writer,
                sheet_name="04_Schema_Offenders",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No unsupported page-root properties."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="04_Schema_Offenders",
                index=False,
            )

        if full_page:
            pd.DataFrame(
                full_page
            ).to_excel(
                writer,
                sheet_name="05_FullPage_Offenders",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No full-page selectable visuals."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="05_FullPage_Offenders",
                index=False,
            )

        if json_failures:
            pd.DataFrame(
                json_failures
            ).to_excel(
                writer,
                sheet_name="06_JSON_Failures",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "No JSON parse failures."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="06_JSON_Failures",
                index=False,
            )

    style_report(
        REPORT_FILE
    )

    print()
    print(
        "02 Network Story Style Lock completed."
    )
    print(
        f"Patch status                  : "
        f"{'PASS' if failures == 0 else 'FAIL'}"
    )
    print(
        f"Manual layout preserved       : "
        f"{'YES' if not geometry_changes else 'NO'}"
    )
    print(
        "Network Gap callout           : amber story color"
    )
    print(
        "Issuer Gap cells              : conditional attention shading"
    )
    print(
        "IN/OON Claims                 : intensity shading + data-bar cues"
    )
    print(
        "Arrow legend                  : strengthened"
    )
    print(
        f"Visuals preserved             : "
        f"{EXPECTED_VISUALS}"
    )
    print(
        f"Bound/query visuals preserved : "
        f"{EXPECTED_BOUND_VISUALS}"
    )
    print(
        f"Full-page selectable visuals  : "
        f"{len(full_page)}"
    )
    print(
        f"Unsupported page properties   : "
        f"{len(unsupported)}"
    )
    print(
        f"Validation failures           : "
        f"{failures}"
    )
    print(
        f"Backup                        : "
        f"{BACKUP_FILE}"
    )
    print(
        f"Review evidence               : "
        f"{REPORT_FILE}"
    )

    if failures:
        print()
        print(
            "Validation failed. Do not open/save the PBIP "
            "until the evidence workbook is reviewed."
        )
        return 2

    print()
    print(
        "Next: open TransparencyInCoverage.pbip and inspect page 02. "
        "Your manual card sizes/positions must remain exactly as saved."
    )

    return 0


if __name__ == "__main__":
    try:
        sys.exit(
            main()
        )

    except Exception as exc:
        print()
        print(
            "NETWORK STORY STYLE LOCK FAILED"
        )
        print(
            str(
                exc
            )
        )
        sys.exit(
            1
        )
