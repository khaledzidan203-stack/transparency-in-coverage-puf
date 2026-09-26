from __future__ import annotations

from datetime import datetime
from pathlib import Path
import json
import os
import shutil
import subprocess
import sys
import zipfile

import pandas as pd


SCRIPT_VERSION = "1.0.0"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
SEMANTIC_DIR = (
    PROJECT_ROOT
    / "powerbi"
    / "TransparencyInCoverage.SemanticModel"
)
MEASURES_TMDL = (
    SEMANTIC_DIR
    / "definition"
    / "tables"
    / "_Measures.tmdl"
)

MEASURE_NAME = "Current Context DQ Exception Count"

OLD_ISSUER = (
    'DQExceptionRegister[Scope] = "ISSUER",'
)
OLD_PLAN = (
    'DQExceptionRegister[Scope] = "PLAN",'
)

NEW_ISSUER = (
    'KEEPFILTERS(DQExceptionRegister[Scope] = "ISSUER"),'
)
NEW_PLAN = (
    'KEEPFILTERS(DQExceptionRegister[Scope] = "PLAN"),'
)


def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )

            with winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                key_path,
            ) as key:
                raw_value, _ = winreg.QueryValueEx(
                    key,
                    "Desktop",
                )

            return Path(
                os.path.expandvars(raw_value)
            ).resolve()
        except Exception:
            pass

    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = (
    get_windows_desktop()
    / "Transparency_PUF_Review"
)

BACKUP_FILE = (
    REVIEW_DIR
    / "44_PRE_DQ_SCOPE_FILTER_FIX_BACKUP.zip"
)

REPORT_FILE = (
    REVIEW_DIR
    / "44_DQ_SCOPE_FILTER_FIX_REPORT.xlsx"
)


def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False

    result = subprocess.run(
        [
            "tasklist",
            "/FI",
            "IMAGENAME eq PBIDesktop.exe",
        ],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(
            subprocess,
            "CREATE_NO_WINDOW",
            0,
        ),
    )

    return "PBIDesktop.exe" in result.stdout


def extract_measure_block(
    text: str,
    measure_name: str,
) -> tuple[int, int, str]:
    start_marker = f"\tmeasure '{measure_name}' ="

    start = text.find(
        start_marker
    )

    if start < 0:
        raise RuntimeError(
            f"Measure not found: {measure_name}"
        )

    next_measure = text.find(
        "\n\tmeasure ",
        start + len(
            start_marker
        ),
    )

    next_column = text.find(
        "\n\tcolumn ",
        start + len(
            start_marker
        ),
    )

    candidates = [
        pos
        for pos in (
            next_measure,
            next_column,
        )
        if pos >= 0
    ]

    end = (
        min(
            candidates
        )
        if candidates
        else len(
            text
        )
    )

    return (
        start,
        end,
        text[
            start:end
        ],
    )


def create_backup() -> None:
    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if BACKUP_FILE.exists():
        BACKUP_FILE.unlink()

    with zipfile.ZipFile(
        BACKUP_FILE,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zf:
        for path in sorted(
            p
            for p in SEMANTIC_DIR.rglob("*")
            if p.is_file()
        ):
            zf.write(
                path,
                str(
                    Path(
                        SEMANTIC_DIR.name
                    )
                    / path.relative_to(
                        SEMANTIC_DIR
                    )
                ),
            )


def style_report(
    path: Path,
) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import (
        Alignment,
        Font,
        PatternFill,
    )
    from openpyxl.utils import (
        get_column_letter,
    )

    wb = load_workbook(
        path
    )

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if ws.max_row and ws.max_column:
            ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            max_len = 0

            for row_idx in range(
                1,
                min(
                    ws.max_row,
                    100,
                )
                + 1,
            ):
                value = ws.cell(
                    row=row_idx,
                    column=col_idx,
                ).value

                if value is not None:
                    max_len = max(
                        max_len,
                        len(
                            str(
                                value
                            )
                        ),
                    )

            ws.column_dimensions[
                get_column_letter(
                    col_idx
                )
            ].width = min(
                max(
                    max_len + 2,
                    12,
                ),
                90,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    wb.save(
        path
    )


def main() -> int:
    print("=" * 78)
    print(
        "Transparency in Coverage PUF — "
        "DQ Scope Filter Context Fix"
    )
    print("=" * 78)

    print(
        f"Script version : {SCRIPT_VERSION}"
    )
    print(
        f"Measure file   : {MEASURES_TMDL}"
    )
    print(
        f"Backup         : {BACKUP_FILE}"
    )
    print(
        f"Report         : {REPORT_FILE}"
    )

    if power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is running. "
            "Save and close Power BI Desktop completely before applying the DAX patch."
        )

    if not MEASURES_TMDL.exists():
        raise FileNotFoundError(
            f"Measure file not found: {MEASURES_TMDL}"
        )

    text = MEASURES_TMDL.read_text(
        encoding="utf-8-sig",
    )

    start, end, block = extract_measure_block(
        text,
        MEASURE_NAME,
    )

    already_fixed = (
        NEW_ISSUER in block
        and NEW_PLAN in block
    )

    direct_issuer_count = block.count(
        OLD_ISSUER
    )
    direct_plan_count = block.count(
        OLD_PLAN
    )

    if already_fixed:
        print()
        print(
            "Measure is already patched. No semantic-model change required."
        )
        return 0

    if (
        direct_issuer_count != 1
        or direct_plan_count != 1
    ):
        raise RuntimeError(
            "Unexpected Current Context DQ measure structure. "
            f"ISSUER direct-filter occurrences={direct_issuer_count}; "
            f"PLAN direct-filter occurrences={direct_plan_count}. "
            "No files changed."
        )

    create_backup()

    patched_block = (
        block.replace(
            OLD_ISSUER,
            NEW_ISSUER,
            1,
        )
        .replace(
            OLD_PLAN,
            NEW_PLAN,
            1,
        )
    )

    patched_text = (
        text[:start]
        + patched_block
        + text[end:]
    )

    MEASURES_TMDL.write_text(
        patched_text,
        encoding="utf-8",
    )

    verify_text = MEASURES_TMDL.read_text(
        encoding="utf-8-sig",
    )

    _, _, verify_block = extract_measure_block(
        verify_text,
        MEASURE_NAME,
    )

    validation = [
        {
            "CheckID": "DQ-SCOPE-001",
            "TestName": (
                "ISSUER Scope filter preserves visual/slicer context"
            ),
            "Expected": NEW_ISSUER,
            "Actual": (
                NEW_ISSUER
                if NEW_ISSUER
                in verify_block
                else "MISSING"
            ),
            "Status": (
                "PASS"
                if NEW_ISSUER
                in verify_block
                else "FAIL"
            ),
        },
        {
            "CheckID": "DQ-SCOPE-002",
            "TestName": (
                "PLAN Scope filter preserves visual/slicer context"
            ),
            "Expected": NEW_PLAN,
            "Actual": (
                NEW_PLAN
                if NEW_PLAN
                in verify_block
                else "MISSING"
            ),
            "Status": (
                "PASS"
                if NEW_PLAN
                in verify_block
                else "FAIL"
            ),
        },
        {
            "CheckID": "DQ-SCOPE-003",
            "TestName": (
                "Old direct Scope filters removed from target measure"
            ),
            "Expected": 0,
            "Actual": (
                verify_block.count(
                    OLD_ISSUER
                )
                + verify_block.count(
                    OLD_PLAN
                )
            ),
            "Status": (
                "PASS"
                if (
                    verify_block.count(
                        OLD_ISSUER
                    )
                    + verify_block.count(
                        OLD_PLAN
                    )
                    == 0
                )
                else "FAIL"
            ),
        },
    ]

    failures = sum(
        row[
            "Status"
        ]
        == "FAIL"
        for row in validation
    )

    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        REPORT_FILE,
        engine="openpyxl",
    ) as writer:
        pd.DataFrame(
            [
                (
                    "PatchStatus",
                    (
                        "PASS"
                        if failures == 0
                        else "FAIL"
                    ),
                ),
                (
                    "ScriptVersion",
                    SCRIPT_VERSION,
                ),
                (
                    "PatchedAtLocal",
                    datetime.now()
                    .isoformat(
                        timespec="seconds"
                    ),
                ),
                (
                    "MeasurePatched",
                    MEASURE_NAME,
                ),
                (
                    "RootCause",
                    (
                        "CALCULATE direct Scope filters replaced the chart's current "
                        "Scope category filter, causing both ISSUER and PLAN bars to "
                        "evaluate the full 38-row DQ register."
                    ),
                ),
                (
                    "Fix",
                    (
                        "Wrap ISSUER/PLAN Scope predicates in KEEPFILTERS so the measure "
                        "intersects with the current Scope filter instead of replacing it."
                    ),
                ),
                (
                    "ExpectedRuntimeTotal",
                    38,
                ),
                (
                    "ExpectedRuntimeIssuer",
                    5,
                ),
                (
                    "ExpectedRuntimePlan",
                    33,
                ),
                (
                    "ValidationFailures",
                    failures,
                ),
                (
                    "BackupFile",
                    str(
                        BACKUP_FILE
                    ),
                ),
            ],
            columns=[
                "Item",
                "Value",
            ],
        ).to_excel(
            writer,
            sheet_name="00_Summary",
            index=False,
        )

        pd.DataFrame(
            validation
        ).to_excel(
            writer,
            sheet_name="01_Validation",
            index=False,
        )

    style_report(
        REPORT_FILE
    )

    print()
    print(
        "DQ Scope filter-context patch completed."
    )
    print(
        f"Patch status        : "
        f"{'PASS' if failures == 0 else 'FAIL'}"
    )
    print(
        f"Measure patched     : {MEASURE_NAME}"
    )
    print(
        "Expected total      : 38"
    )
    print(
        "Expected ISSUER     : 5"
    )
    print(
        "Expected PLAN       : 33"
    )
    print(
        f"Validation failures : {failures}"
    )
    print(
        f"Backup              : {BACKUP_FILE}"
    )
    print(
        f"Review evidence     : {REPORT_FILE}"
    )

    if failures:
        print()
        print(
            "Offline validation failed. Do not open/save the PBIP until reviewed."
        )
        return 2

    print()
    print(
        "Next: open TransparencyInCoverage.pbip, Save (Ctrl+S), and inspect "
        "10 Data Quality & Methodology. The Scope chart should show ISSUER=5 and PLAN=33."
    )

    return 0


if __name__ == "__main__":
    try:
        sys.exit(
            main()
        )
    except Exception as exc:
        print()
        print(
            "DQ SCOPE FILTER FIX FAILED"
        )
        print(
            str(
                exc
            )
        )
        sys.exit(
            1
        )
