from __future__ import annotations

from datetime import datetime
from pathlib import Path
import hashlib
import json
import os
import subprocess
import sys
import zipfile

import pandas as pd


SCRIPT_VERSION = "1.0.0"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
REPORT_DIR = (
    PROJECT_ROOT
    / "powerbi"
    / "TransparencyInCoverage.Report"
)
PAGES_DIR = REPORT_DIR / "definition" / "pages"

BUILDER_FILE = (
    PROJECT_ROOT
    / "src"
    / "validation"
    / "33_build_appeals_storytelling.py"
)

EXPECTED_VISUALS = 34
EXPECTED_BOUND_VISUALS = 14
TARGET_PAGE = "05 Appeals"
TARGET_VISUAL = "p05_state_chart"

CANVAS_WIDTH = 1920
CANVAS_HEIGHT = 1080

UNSUPPORTED_PAGE_ROOT_PROPERTIES = {
    "horizontalAlignment",
    "verticalAlignment",
}


def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )

            with winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                key_path,
            ) as key:
                raw_value, _ = winreg.QueryValueEx(
                    key,
                    "Desktop",
                )

            return Path(
                os.path.expandvars(raw_value)
            ).resolve()

        except Exception:
            pass

    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = (
    get_windows_desktop()
    / "Transparency_PUF_Review"
)

BACKUP_FILE = (
    REVIEW_DIR
    / "34_PRE_APPEALS_VALIDATION_CHART_FIX_BACKUP.zip"
)

REPORT_FILE = (
    REVIEW_DIR
    / "34_APPEALS_VALIDATION_CHART_FIX_REPORT.xlsx"
)


def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False

    result = subprocess.run(
        [
            "tasklist",
            "/FI",
            "IMAGENAME eq PBIDesktop.exe",
        ],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(
            subprocess,
            "CREATE_NO_WINDOW",
            0,
        ),
    )

    return "PBIDesktop.exe" in result.stdout


def read_json(path: Path):
    return json.loads(
        path.read_text(
            encoding="utf-8-sig",
        )
    )


def write_json(path: Path, payload) -> None:
    path.write_text(
        json.dumps(
            payload,
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )


def sha256(path: Path) -> str:
    h = hashlib.sha256()

    with path.open("rb") as f:
        for chunk in iter(
            lambda: f.read(1024 * 1024),
            b"",
        ):
            h.update(chunk)

    return h.hexdigest()


def find_target_page() -> tuple[Path, dict]:
    matches = []

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        if payload.get(
            "displayName"
        ) == TARGET_PAGE:
            matches.append(
                (
                    page_json.parent,
                    payload,
                )
            )

    if len(matches) != 1:
        raise RuntimeError(
            f"Expected exactly one '{TARGET_PAGE}' page; "
            f"found {len(matches)}."
        )

    return matches[0]


def count_bound_visuals(page_dir: Path) -> int:
    count = 0

    for path in (
        page_dir
        / "visuals"
    ).rglob(
        "visual.json"
    ):
        payload = read_json(
            path
        )

        if (
            payload.get(
                "visual",
                {},
            ).get(
                "query"
            )
        ):
            count += 1

    return count


def report_hashes() -> dict[str, str]:
    result = {}

    for path in REPORT_DIR.rglob("*"):
        if path.is_file():
            result[
                str(
                    path.relative_to(
                        REPORT_DIR
                    )
                ).replace(
                    "\\",
                    "/",
                )
            ] = sha256(
                path
            )

    return result


def create_backup(
    target_page_dir: Path,
) -> None:
    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if BACKUP_FILE.exists():
        BACKUP_FILE.unlink()

    with zipfile.ZipFile(
        BACKUP_FILE,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zf:
        for path in sorted(
            p
            for p in target_page_dir.rglob("*")
            if p.is_file()
        ):
            zf.write(
                path,
                str(
                    Path("ReportPage")
                    / path.relative_to(
                        target_page_dir
                    )
                ),
            )

        if BUILDER_FILE.exists():
            zf.write(
                BUILDER_FILE,
                str(
                    Path("Source")
                    / BUILDER_FILE.name
                ),
            )


def patch_state_chart(
    page_dir: Path,
) -> dict:
    path = (
        page_dir
        / "visuals"
        / TARGET_VISUAL
        / "visual.json"
    )

    if not path.exists():
        raise FileNotFoundError(
            f"Target visual not found: {path}"
        )

    payload = read_json(
        path
    )

    visual = payload.get(
        "visual",
        {},
    )

    old_type = visual.get(
        "visualType"
    )

    if old_type not in {
        "columnChart",
        "clusteredColumnChart",
    }:
        raise RuntimeError(
            f"Unexpected state-chart visual type: {old_type}"
        )

    visual[
        "visualType"
    ] = "clusteredColumnChart"

    write_json(
        path,
        payload,
    )

    return {
        "Target": TARGET_VISUAL,
        "BeforeVisualType": old_type,
        "AfterVisualType": "clusteredColumnChart",
        "Status": (
            "PATCHED"
            if old_type != "clusteredColumnChart"
            else "ALREADY_CORRECT"
        ),
    }


def patch_builder() -> list[dict]:
    if not BUILDER_FILE.exists():
        return [
            {
                "Patch": "Builder source",
                "Status": "SKIPPED",
                "Details": (
                    f"Builder file not found at {BUILDER_FILE}"
                ),
            }
        ]

    text = BUILDER_FILE.read_text(
        encoding="utf-8-sig",
    )

    original = text

    if "EXPECTED_BOUND_VISUALS = 16" in text:
        text = text.replace(
            "EXPECTED_BOUND_VISUALS = 16",
            "EXPECTED_BOUND_VISUALS = 14",
            1,
        )
        bound_status = "PATCHED"
    elif "EXPECTED_BOUND_VISUALS = 14" in text:
        bound_status = "ALREADY_CORRECT"
    else:
        raise RuntimeError(
            "Could not locate EXPECTED_BOUND_VISUALS in builder source."
        )

    marker = "def make_state_chart(\n"
    marker_pos = text.find(
        marker
    )

    if marker_pos < 0:
        raise RuntimeError(
            "Could not locate make_state_chart() in builder source."
        )

    next_section = text.find(
        "\n# =====================================================================",
        marker_pos + len(
            marker
        ),
    )

    if next_section < 0:
        next_section = len(
            text
        )

    state_block = text[
        marker_pos:next_section
    ]

    if '"visualType": "columnChart"' in state_block:
        state_block = state_block.replace(
            '"visualType": "columnChart"',
            '"visualType": "clusteredColumnChart"',
            1,
        )
        chart_status = "PATCHED"
    elif '"visualType": "clusteredColumnChart"' in state_block:
        chart_status = "ALREADY_CORRECT"
    else:
        raise RuntimeError(
            "Could not locate the state chart visualType in builder source."
        )

    text = (
        text[
            :marker_pos
        ]
        + state_block
        + text[
            next_section:
        ]
    )

    if text != original:
        BUILDER_FILE.write_text(
            text,
            encoding="utf-8",
        )

    return [
        {
            "Patch": "EXPECTED_BOUND_VISUALS",
            "Status": bound_status,
            "Details": "16 -> 14",
        },
        {
            "Patch": "State chart visual type",
            "Status": chart_status,
            "Details": (
                "columnChart -> clusteredColumnChart"
            ),
        },
    ]


def validate_no_unsupported_page_roots() -> list[dict]:
    offenders = []

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        for prop in sorted(
            UNSUPPORTED_PAGE_ROOT_PROPERTIES
        ):
            if prop in payload:
                offenders.append(
                    {
                        "Page": payload.get(
                            "displayName"
                        ),
                        "Property": prop,
                    }
                )

    return offenders


def validate_no_full_page_visuals() -> list[dict]:
    offenders = []

    for visual_json in PAGES_DIR.rglob(
        "visual.json"
    ):
        payload = read_json(
            visual_json
        )

        pos = payload.get(
            "position",
            {},
        )

        width = float(
            pos.get(
                "width",
                0,
            )
            or 0
        )

        height = float(
            pos.get(
                "height",
                0,
            )
            or 0
        )

        if (
            width
            >= CANVAS_WIDTH * 0.94
            and height
            >= CANVAS_HEIGHT * 0.94
        ):
            offenders.append(
                {
                    "Visual": payload.get(
                        "name"
                    ),
                    "Width": width,
                    "Height": height,
                }
            )

    return offenders


def validate_json_tree() -> list[dict]:
    failures = []

    for path in (
        REPORT_DIR
        / "definition"
    ).rglob(
        "*.json"
    ):
        try:
            read_json(
                path
            )

        except Exception as exc:
            failures.append(
                {
                    "RelativePath": str(
                        path.relative_to(
                            PROJECT_ROOT
                        )
                    ),
                    "Error": str(
                        exc
                    ),
                }
            )

    return failures


def style_report(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import (
        Alignment,
        Font,
        PatternFill,
    )
    from openpyxl.utils import (
        get_column_letter,
    )

    wb = load_workbook(
        path
    )

    fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )
    font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if (
            ws.max_row
            and ws.max_column
        ):
            ws.auto_filter.ref = (
                ws.dimensions
            )

        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            max_len = 0

            for row_idx in range(
                1,
                min(
                    ws.max_row,
                    180,
                )
                + 1,
            ):
                value = ws.cell(
                    row=row_idx,
                    column=col_idx,
                ).value

                if value is not None:
                    max_len = max(
                        max_len,
                        len(
                            str(
                                value
                            )
                        ),
                    )

            ws.column_dimensions[
                get_column_letter(
                    col_idx
                )
            ].width = min(
                max(
                    max_len + 2,
                    10,
                ),
                75,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    wb.save(
        path
    )


def main() -> int:
    print("=" * 78)
    print(
        "Transparency in Coverage PUF — "
        "05 Appeals Validation + State Chart Fix"
    )
    print("=" * 78)

    print(
        f"Script version : {SCRIPT_VERSION}"
    )
    print(
        f"Report folder  : {REPORT_DIR}"
    )
    print(
        f"Builder file   : {BUILDER_FILE}"
    )
    print(
        f"Backup         : {BACKUP_FILE}"
    )
    print(
        f"Review file    : {REPORT_FILE}"
    )

    if power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is running. "
            "Save and close it completely before applying this patch."
        )

    page_dir, _ = (
        find_target_page()
    )

    visual_count = len(
        list(
            (
                page_dir
                / "visuals"
            ).rglob(
                "visual.json"
            )
        )
    )

    bound_count_before = (
        count_bound_visuals(
            page_dir
        )
    )

    if visual_count != EXPECTED_VISUALS:
        raise RuntimeError(
            f"Expected {EXPECTED_VISUALS} visuals on 05 Appeals; "
            f"found {visual_count}. No changes made."
        )

    if bound_count_before != EXPECTED_BOUND_VISUALS:
        raise RuntimeError(
            f"Expected the current page to contain {EXPECTED_BOUND_VISUALS} "
            f"query-bound visuals; found {bound_count_before}. "
            "This patch is intentionally scoped to the known 14-visual model."
        )

    before_hashes = (
        report_hashes()
    )

    target_prefix = str(
        page_dir.relative_to(
            REPORT_DIR
        )
    ).replace(
        "\\",
        "/",
    )

    create_backup(
        page_dir
    )

    chart_patch = (
        patch_state_chart(
            page_dir
        )
    )

    builder_patches = (
        patch_builder()
    )

    after_hashes = (
        report_hashes()
    )

    non_target_changes = []

    for rel_path in sorted(
        set(
            before_hashes
        )
        | set(
            after_hashes
        )
    ):
        if rel_path.startswith(
            target_prefix
        ):
            continue

        if (
            before_hashes.get(
                rel_path
            )
            != after_hashes.get(
                rel_path
            )
        ):
            non_target_changes.append(
                {
                    "RelativePath": rel_path,
                    "BeforeSHA256": before_hashes.get(
                        rel_path
                    ),
                    "AfterSHA256": after_hashes.get(
                        rel_path
                    ),
                }
            )

    chart_payload = read_json(
        page_dir
        / "visuals"
        / TARGET_VISUAL
        / "visual.json"
    )

    final_chart_type = (
        chart_payload.get(
            "visual",
            {},
        ).get(
            "visualType"
        )
    )

    final_visual_count = len(
        list(
            (
                page_dir
                / "visuals"
            ).rglob(
                "visual.json"
            )
        )
    )

    final_bound_count = (
        count_bound_visuals(
            page_dir
        )
    )

    unsupported = (
        validate_no_unsupported_page_roots()
    )
    full_page = (
        validate_no_full_page_visuals()
    )
    json_failures = (
        validate_json_tree()
    )

    validations = [
        {
            "Check": "Appeals visual count",
            "Expected": EXPECTED_VISUALS,
            "Actual": final_visual_count,
            "Status": (
                "PASS"
                if final_visual_count
                == EXPECTED_VISUALS
                else "FAIL"
            ),
        },
        {
            "Check": "Appeals query-bound visual count",
            "Expected": EXPECTED_BOUND_VISUALS,
            "Actual": final_bound_count,
            "Status": (
                "PASS"
                if final_bound_count
                == EXPECTED_BOUND_VISUALS
                else "FAIL"
            ),
        },
        {
            "Check": "State comparison chart type",
            "Expected": "clusteredColumnChart",
            "Actual": final_chart_type,
            "Status": (
                "PASS"
                if final_chart_type
                == "clusteredColumnChart"
                else "FAIL"
            ),
        },
        {
            "Check": "Non-target report files changed",
            "Expected": 0,
            "Actual": len(
                non_target_changes
            ),
            "Status": (
                "PASS"
                if not non_target_changes
                else "FAIL"
            ),
        },
        {
            "Check": "Unsupported page-root properties",
            "Expected": 0,
            "Actual": len(
                unsupported
            ),
            "Status": (
                "PASS"
                if not unsupported
                else "FAIL"
            ),
        },
        {
            "Check": "Full-page selectable visuals",
            "Expected": 0,
            "Actual": len(
                full_page
            ),
            "Status": (
                "PASS"
                if not full_page
                else "FAIL"
            ),
        },
        {
            "Check": "PBIR JSON parse failures",
            "Expected": 0,
            "Actual": len(
                json_failures
            ),
            "Status": (
                "PASS"
                if not json_failures
                else "FAIL"
            ),
        },
    ]

    failures = sum(
        row[
            "Status"
        ] == "FAIL"
        for row in validations
    )

    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        REPORT_FILE,
        engine="openpyxl",
    ) as writer:
        pd.DataFrame(
            [
                (
                    "PatchStatus",
                    (
                        "PASS"
                        if failures == 0
                        else "FAIL"
                    ),
                ),
                (
                    "ScriptVersion",
                    SCRIPT_VERSION,
                ),
                (
                    "PatchedAtLocal",
                    datetime.now()
                    .isoformat(
                        timespec="seconds"
                    ),
                ),
                (
                    "Page",
                    TARGET_PAGE,
                ),
                (
                    "RootCause",
                    (
                        "Builder expected 16 query-bound visuals, "
                        "but the designed page contains 14. "
                        "The state chart was also stacked rather than clustered."
                    ),
                ),
                (
                    "FinalBoundVisuals",
                    final_bound_count,
                ),
                (
                    "FinalStateChartType",
                    final_chart_type,
                ),
                (
                    "ValidationFailures",
                    failures,
                ),
                (
                    "BackupFile",
                    str(
                        BACKUP_FILE
                    ),
                ),
            ],
            columns=[
                "Item",
                "Value",
            ],
        ).to_excel(
            writer,
            sheet_name="00_Summary",
            index=False,
        )

        pd.DataFrame(
            [
                chart_patch
            ]
        ).to_excel(
            writer,
            sheet_name="01_Report_Patch",
            index=False,
        )

        pd.DataFrame(
            builder_patches
        ).to_excel(
            writer,
            sheet_name="02_Builder_Patch",
            index=False,
        )

        pd.DataFrame(
            validations
        ).to_excel(
            writer,
            sheet_name="03_Validation",
            index=False,
        )

        pd.DataFrame(
            non_target_changes
            or [
                {
                    "Info": (
                        "No report files outside 05 Appeals changed."
                    )
                }
            ]
        ).to_excel(
            writer,
            sheet_name="04_NonTarget_Changes",
            index=False,
        )

    style_report(
        REPORT_FILE
    )

    print()
    print(
        "05 Appeals validation + state chart fix completed."
    )
    print(
        f"Patch status                  : "
        f"{'PASS' if failures == 0 else 'FAIL'}"
    )
    print(
        f"Visuals                       : {final_visual_count}"
    )
    print(
        f"Bound/query visuals           : {final_bound_count}"
    )
    print(
        f"State chart type              : {final_chart_type}"
    )
    print(
        f"Non-target report changes     : {len(non_target_changes)}"
    )
    print(
        f"Unsupported page properties   : {len(unsupported)}"
    )
    print(
        f"Validation failures           : {failures}"
    )
    print(
        f"Backup                        : {BACKUP_FILE}"
    )
    print(
        f"Review evidence               : {REPORT_FILE}"
    )

    if failures:
        print()
        print(
            "Validation failed. Review the evidence workbook before saving PBIP."
        )
        return 2

    print()
    print(
        "Next: open TransparencyInCoverage.pbip. "
        "On 05 Appeals, the Internal and External state rates should now appear "
        "side-by-side rather than stacked. Everything else should remain unchanged."
    )

    return 0


if __name__ == "__main__":
    try:
        sys.exit(
            main()
        )

    except Exception as exc:
        print()
        print(
            "APPEALS VALIDATION + STATE CHART FIX FAILED"
        )
        print(
            str(
                exc
            )
        )
        sys.exit(
            1
        )
