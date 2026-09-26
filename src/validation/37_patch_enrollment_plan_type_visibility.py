from __future__ import annotations

from datetime import datetime
from pathlib import Path
import hashlib
import json
import os
import subprocess
import sys
import zipfile

import pandas as pd


SCRIPT_VERSION = "1.0.0"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
REPORT_DIR = PROJECT_ROOT / "powerbi" / "TransparencyInCoverage.Report"
PAGES_DIR = REPORT_DIR / "definition" / "pages"

BUILDER_FILE = (
    PROJECT_ROOT
    / "src"
    / "validation"
    / "36_build_enrollment_storytelling.py"
)

TARGET_PAGE = "07 Enrollment"
TARGET_VISUAL = "p07_plan_type_chart"

EXPECTED_VISUALS = 31
EXPECTED_BOUND_VISUALS = 13

TARGET_HEIGHT = 104.0
INVESTIGATION_TOP_Y = 714.0
MIN_GAP_PX = 2.0

CANVAS_WIDTH = 1920
CANVAS_HEIGHT = 1080

UNSUPPORTED_PAGE_ROOT_PROPERTIES = {
    "horizontalAlignment",
    "verticalAlignment",
}


def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )

            with winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                key_path,
            ) as key:
                raw_value, _ = winreg.QueryValueEx(
                    key,
                    "Desktop",
                )

            return Path(
                os.path.expandvars(raw_value)
            ).resolve()

        except Exception:
            pass

    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = (
    get_windows_desktop()
    / "Transparency_PUF_Review"
)

BACKUP_FILE = (
    REVIEW_DIR
    / "37_PRE_ENROLLMENT_PLAN_TYPE_VISIBILITY_BACKUP.zip"
)

REPORT_FILE = (
    REVIEW_DIR
    / "37_ENROLLMENT_PLAN_TYPE_VISIBILITY_REPORT.xlsx"
)


def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False

    result = subprocess.run(
        [
            "tasklist",
            "/FI",
            "IMAGENAME eq PBIDesktop.exe",
        ],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(
            subprocess,
            "CREATE_NO_WINDOW",
            0,
        ),
    )

    return "PBIDesktop.exe" in result.stdout


def read_json(path: Path):
    return json.loads(
        path.read_text(
            encoding="utf-8-sig",
        )
    )


def write_json(
    path: Path,
    payload,
) -> None:
    path.write_text(
        json.dumps(
            payload,
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )


def sha256(path: Path) -> str:
    h = hashlib.sha256()

    with path.open("rb") as f:
        for chunk in iter(
            lambda: f.read(1024 * 1024),
            b"",
        ):
            h.update(chunk)

    return h.hexdigest()


def find_target_page() -> tuple[Path, dict]:
    matches = []

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        if payload.get(
            "displayName"
        ) == TARGET_PAGE:
            matches.append(
                (
                    page_json.parent,
                    payload,
                )
            )

    if len(matches) != 1:
        raise RuntimeError(
            f"Expected exactly one '{TARGET_PAGE}' page; "
            f"found {len(matches)}."
        )

    return matches[0]


def target_visual_path(
    page_dir: Path,
) -> Path:
    return (
        page_dir
        / "visuals"
        / TARGET_VISUAL
        / "visual.json"
    )


def count_bound_visuals(
    page_dir: Path,
) -> int:
    count = 0

    for path in (
        page_dir
        / "visuals"
    ).rglob(
        "visual.json"
    ):
        payload = read_json(
            path
        )

        if (
            payload.get(
                "visual",
                {},
            ).get(
                "query"
            )
        ):
            count += 1

    return count


def report_hashes() -> dict[str, str]:
    result = {}

    for path in REPORT_DIR.rglob("*"):
        if path.is_file():
            result[
                str(
                    path.relative_to(
                        REPORT_DIR
                    )
                ).replace(
                    "\\",
                    "/",
                )
            ] = sha256(
                path
            )

    return result


def create_backup(
    page_dir: Path,
) -> None:
    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if BACKUP_FILE.exists():
        BACKUP_FILE.unlink()

    visual_path = target_visual_path(
        page_dir
    )

    with zipfile.ZipFile(
        BACKUP_FILE,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zf:
        zf.write(
            visual_path,
            str(
                Path("ReportVisual")
                / visual_path.name
            ),
        )

        if BUILDER_FILE.exists():
            zf.write(
                BUILDER_FILE,
                str(
                    Path("Source")
                    / BUILDER_FILE.name
                ),
            )


def patch_report_visual(
    page_dir: Path,
) -> dict:
    path = target_visual_path(
        page_dir
    )

    if not path.exists():
        raise FileNotFoundError(
            f"Target visual not found: {path}"
        )

    payload = read_json(
        path
    )

    position = payload.get(
        "position",
        {},
    )

    before = {
        "x": position.get("x"),
        "y": position.get("y"),
        "width": position.get("width"),
        "height": position.get("height"),
    }

    y = float(
        position.get(
            "y",
            0,
        )
        or 0
    )

    max_height = (
        INVESTIGATION_TOP_Y
        - MIN_GAP_PX
        - y
    )

    if TARGET_HEIGHT > max_height:
        raise RuntimeError(
            "Requested Plan Type chart height would overlap "
            "the Investigation section. "
            f"Target={TARGET_HEIGHT}px, max allowed={max_height:.2f}px."
        )

    position[
        "height"
    ] = TARGET_HEIGHT

    write_json(
        path,
        payload,
    )

    after = {
        "x": position.get("x"),
        "y": position.get("y"),
        "width": position.get("width"),
        "height": position.get("height"),
    }

    return {
        "Visual": TARGET_VISUAL,
        "Action": (
            "Increase height only so all Plan Type categories "
            "can render without the scrollbar"
        ),
        "Before": json.dumps(
            before,
            sort_keys=True,
        ),
        "After": json.dumps(
            after,
            sort_keys=True,
        ),
    }


def patch_builder_source() -> dict:
    if not BUILDER_FILE.exists():
        return {
            "Patch": "Builder source",
            "Status": "SKIPPED",
            "Details": (
                f"Builder file not found: {BUILDER_FILE}"
            ),
        }

    text = BUILDER_FILE.read_text(
        encoding="utf-8-sig",
    )

    original = text

    marker = 'name="p07_plan_type_chart",'

    marker_pos = text.find(
        marker
    )

    if marker_pos < 0:
        raise RuntimeError(
            "Could not locate p07_plan_type_chart "
            "inside 36_build_enrollment_storytelling.py."
        )

    block_end = text.find(
        "\n        ),",
        marker_pos,
    )

    if block_end < 0:
        raise RuntimeError(
            "Could not locate the end of the Plan Type chart call."
        )

    block = text[
        marker_pos:block_end
    ]

    if "height=84," in block:
        block = block.replace(
            "height=84,",
            f"height={int(TARGET_HEIGHT)},",
            1,
        )
        status = "PATCHED"

    elif f"height={int(TARGET_HEIGHT)}," in block:
        status = "ALREADY_CORRECT"

    else:
        raise RuntimeError(
            "Unexpected Plan Type chart height in builder source. "
            "No source change made."
        )

    text = (
        text[
            :marker_pos
        ]
        + block
        + text[
            block_end:
        ]
    )

    if text != original:
        BUILDER_FILE.write_text(
            text,
            encoding="utf-8",
        )

    return {
        "Patch": "Builder Plan Type chart height",
        "Status": status,
        "Details": (
            f"84 -> {int(TARGET_HEIGHT)} px"
        ),
    }


def validate_json_tree() -> list[dict]:
    failures = []

    for path in (
        REPORT_DIR
        / "definition"
    ).rglob(
        "*.json"
    ):
        try:
            read_json(
                path
            )

        except Exception as exc:
            failures.append(
                {
                    "RelativePath": str(
                        path.relative_to(
                            PROJECT_ROOT
                        )
                    ),
                    "Error": str(
                        exc
                    ),
                }
            )

    return failures


def validate_no_unsupported_page_roots() -> list[dict]:
    offenders = []

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        for prop in sorted(
            UNSUPPORTED_PAGE_ROOT_PROPERTIES
        ):
            if prop in payload:
                offenders.append(
                    {
                        "Page": payload.get(
                            "displayName"
                        ),
                        "Property": prop,
                    }
                )

    return offenders


def validate_no_full_page_visuals() -> list[dict]:
    offenders = []

    for visual_json in PAGES_DIR.rglob(
        "visual.json"
    ):
        payload = read_json(
            visual_json
        )

        pos = payload.get(
            "position",
            {},
        )

        width = float(
            pos.get(
                "width",
                0,
            )
            or 0
        )

        height = float(
            pos.get(
                "height",
                0,
            )
            or 0
        )

        if (
            width
            >= CANVAS_WIDTH * 0.94
            and height
            >= CANVAS_HEIGHT * 0.94
        ):
            offenders.append(
                {
                    "Visual": payload.get(
                        "name"
                    ),
                    "Width": width,
                    "Height": height,
                }
            )

    return offenders


def style_report(
    path: Path,
) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import (
        Alignment,
        Font,
        PatternFill,
    )
    from openpyxl.utils import (
        get_column_letter,
    )

    wb = load_workbook(
        path
    )

    fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )
    font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if (
            ws.max_row
            and ws.max_column
        ):
            ws.auto_filter.ref = (
                ws.dimensions
            )

        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            max_len = 0

            for row_idx in range(
                1,
                min(
                    ws.max_row,
                    180,
                )
                + 1,
            ):
                value = ws.cell(
                    row=row_idx,
                    column=col_idx,
                ).value

                if value is not None:
                    max_len = max(
                        max_len,
                        len(
                            str(
                                value
                            )
                        ),
                    )

            ws.column_dimensions[
                get_column_letter(
                    col_idx
                )
            ].width = min(
                max(
                    max_len + 2,
                    10,
                ),
                75,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    wb.save(
        path
    )


def main() -> int:
    print("=" * 78)
    print(
        "Transparency in Coverage PUF — "
        "07 Enrollment Plan Type Visibility Patch"
    )
    print("=" * 78)

    print(
        f"Script version : {SCRIPT_VERSION}"
    )
    print(
        f"Report folder  : {REPORT_DIR}"
    )
    print(
        f"Builder file   : {BUILDER_FILE}"
    )
    print(
        f"Backup         : {BACKUP_FILE}"
    )
    print(
        f"Review file    : {REPORT_FILE}"
    )

    if power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is running. "
            "Save and close it completely before applying this patch."
        )

    page_dir, _ = (
        find_target_page()
    )

    visual_count = len(
        list(
            (
                page_dir
                / "visuals"
            ).rglob(
                "visual.json"
            )
        )
    )

    bound_count = (
        count_bound_visuals(
            page_dir
        )
    )

    if visual_count != EXPECTED_VISUALS:
        raise RuntimeError(
            f"Expected {EXPECTED_VISUALS} visuals on 07 Enrollment; "
            f"found {visual_count}. No changes made."
        )

    if bound_count != EXPECTED_BOUND_VISUALS:
        raise RuntimeError(
            f"Expected {EXPECTED_BOUND_VISUALS} query-bound visuals; "
            f"found {bound_count}. No changes made."
        )

    target_path = (
        target_visual_path(
            page_dir
        )
    )

    before_target_hash = (
        sha256(
            target_path
        )
    )

    before_hashes = (
        report_hashes()
    )

    target_rel = str(
        target_path.relative_to(
            REPORT_DIR
        )
    ).replace(
        "\\",
        "/",
    )

    create_backup(
        page_dir
    )

    report_patch = (
        patch_report_visual(
            page_dir
        )
    )

    builder_patch = (
        patch_builder_source()
    )

    after_hashes = (
        report_hashes()
    )

    changed_report_files = []

    for rel_path in sorted(
        set(
            before_hashes
        )
        | set(
            after_hashes
        )
    ):
        if (
            before_hashes.get(
                rel_path
            )
            != after_hashes.get(
                rel_path
            )
        ):
            changed_report_files.append(
                rel_path
            )

    target_only = (
        changed_report_files
        == [
            target_rel
        ]
    )

    final_payload = read_json(
        target_path
    )

    position = final_payload[
        "position"
    ]

    final_y = float(
        position[
            "y"
        ]
    )
    final_height = float(
        position[
            "height"
        ]
    )
    final_bottom = (
        final_y
        + final_height
    )
    final_gap = (
        INVESTIGATION_TOP_Y
        - final_bottom
    )

    json_failures = (
        validate_json_tree()
    )
    unsupported = (
        validate_no_unsupported_page_roots()
    )
    full_page = (
        validate_no_full_page_visuals()
    )

    validations = [
        {
            "Check": "Only Plan Type report visual changed",
            "Expected": target_rel,
            "Actual": (
                ", ".join(
                    changed_report_files
                )
                if changed_report_files
                else "NONE"
            ),
            "Status": (
                "PASS"
                if target_only
                else "FAIL"
            ),
        },
        {
            "Check": "Plan Type chart height",
            "Expected": TARGET_HEIGHT,
            "Actual": final_height,
            "Status": (
                "PASS"
                if final_height
                == TARGET_HEIGHT
                else "FAIL"
            ),
        },
        {
            "Check": "Gap before Investigation section",
            "Expected": f">= {MIN_GAP_PX}px",
            "Actual": final_gap,
            "Status": (
                "PASS"
                if final_gap
                >= MIN_GAP_PX
                else "FAIL"
            ),
        },
        {
            "Check": "Enrollment visual count preserved",
            "Expected": EXPECTED_VISUALS,
            "Actual": visual_count,
            "Status": "PASS",
        },
        {
            "Check": "Enrollment bound visuals preserved",
            "Expected": EXPECTED_BOUND_VISUALS,
            "Actual": bound_count,
            "Status": "PASS",
        },
        {
            "Check": "Unsupported page-root properties",
            "Expected": 0,
            "Actual": len(
                unsupported
            ),
            "Status": (
                "PASS"
                if not unsupported
                else "FAIL"
            ),
        },
        {
            "Check": "Full-page selectable visuals",
            "Expected": 0,
            "Actual": len(
                full_page
            ),
            "Status": (
                "PASS"
                if not full_page
                else "FAIL"
            ),
        },
        {
            "Check": "PBIR JSON parse failures",
            "Expected": 0,
            "Actual": len(
                json_failures
            ),
            "Status": (
                "PASS"
                if not json_failures
                else "FAIL"
            ),
        },
    ]

    failures = sum(
        row[
            "Status"
        ] == "FAIL"
        for row in validations
    )

    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        REPORT_FILE,
        engine="openpyxl",
    ) as writer:
        pd.DataFrame(
            [
                (
                    "PatchStatus",
                    (
                        "PASS"
                        if failures == 0
                        else "FAIL"
                    ),
                ),
                (
                    "ScriptVersion",
                    SCRIPT_VERSION,
                ),
                (
                    "PatchedAtLocal",
                    datetime.now()
                    .isoformat(
                        timespec="seconds"
                    ),
                ),
                (
                    "Page",
                    TARGET_PAGE,
                ),
                (
                    "TargetVisual",
                    TARGET_VISUAL,
                ),
                (
                    "TargetHeight",
                    TARGET_HEIGHT,
                ),
                (
                    "FinalGapBeforeInvestigation",
                    final_gap,
                ),
                (
                    "ValidationFailures",
                    failures,
                ),
                (
                    "BackupFile",
                    str(
                        BACKUP_FILE
                    ),
                ),
            ],
            columns=[
                "Item",
                "Value",
            ],
        ).to_excel(
            writer,
            sheet_name="00_Summary",
            index=False,
        )

        pd.DataFrame(
            [
                report_patch
            ]
        ).to_excel(
            writer,
            sheet_name="01_Report_Patch",
            index=False,
        )

        pd.DataFrame(
            [
                builder_patch
            ]
        ).to_excel(
            writer,
            sheet_name="02_Builder_Patch",
            index=False,
        )

        pd.DataFrame(
            validations
        ).to_excel(
            writer,
            sheet_name="03_Validation",
            index=False,
        )

        pd.DataFrame(
            [
                {
                    "TargetBeforeSHA256": before_target_hash,
                    "TargetAfterSHA256": sha256(
                        target_path
                    ),
                    "ChangedReportFiles": (
                        ", ".join(
                            changed_report_files
                        )
                    ),
                }
            ]
        ).to_excel(
            writer,
            sheet_name="04_File_Evidence",
            index=False,
        )

    style_report(
        REPORT_FILE
    )

    print()
    print(
        "07 Enrollment Plan Type visibility patch completed."
    )
    print(
        f"Patch status                  : "
        f"{'PASS' if failures == 0 else 'FAIL'}"
    )
    print(
        f"Plan Type chart height        : {final_height:.0f}px"
    )
    print(
        f"Gap before Investigation      : {final_gap:.0f}px"
    )
    print(
        f"Report files changed          : {len(changed_report_files)}"
    )
    print(
        f"Non-target report changes     : "
        f"{0 if target_only else max(len(changed_report_files) - 1, 0)}"
    )
    print(
        f"Validation failures           : {failures}"
    )
    print(
        f"Backup                        : {BACKUP_FILE}"
    )
    print(
        f"Review evidence               : {REPORT_FILE}"
    )

    if failures:
        print()
        print(
            "Validation failed. Review the evidence workbook before saving PBIP."
        )
        return 2

    print()
    print(
        "Next: open TransparencyInCoverage.pbip and inspect only the Plan Type panel "
        "on 07 Enrollment. The panel should now expose all Plan Type categories "
        "without changing any other report visual."
    )

    return 0


if __name__ == "__main__":
    try:
        sys.exit(
            main()
        )

    except Exception as exc:
        print()
        print(
            "ENROLLMENT PLAN TYPE VISIBILITY PATCH FAILED"
        )
        print(
            str(
                exc
            )
        )
        sys.exit(
            1
        )
