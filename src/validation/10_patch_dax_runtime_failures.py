from __future__ import annotations

from pathlib import Path
from datetime import datetime
import hashlib
import os
import re
import subprocess
import sys
import zipfile

import pandas as pd


SCRIPT_VERSION = "1.0.0"
PROJECT_ROOT = Path(__file__).resolve().parents[2]
POWERBI_ROOT = PROJECT_ROOT / "powerbi"
MEASURE_TABLE_FILE = (
    POWERBI_ROOT
    / "TransparencyInCoverage.SemanticModel"
    / "definition"
    / "tables"
    / "_Measures.tmdl"
)


def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, key_path) as key:
                raw_value, _ = winreg.QueryValueEx(key, "Desktop")
            return Path(os.path.expandvars(raw_value)).resolve()
        except Exception:
            pass

    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = get_windows_desktop() / "Transparency_PUF_Review"
BACKUP_FILE = REVIEW_DIR / "10_PRE_DAX_RUNTIME_FIX_BACKUP.zip"
REPORT_FILE = REVIEW_DIR / "10_DAX_RUNTIME_FIX_REPORT.xlsx"


PATCHES = {
    "Issuer Comparable Claims Received - Total": r"""
VAR RIn =
    CALCULATETABLE(
        VALUES(FactIssuerMetricAvailability[IssuerKey]),
        REMOVEFILTERS(DimMetric),
        REMOVEFILTERS(DimAvailabilityStatus),
        FactIssuerMetricAvailability[MetricKey] = 2,
        FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
    )
VAR ROut =
    CALCULATETABLE(
        VALUES(FactIssuerMetricAvailability[IssuerKey]),
        REMOVEFILTERS(DimMetric),
        REMOVEFILTERS(DimAvailabilityStatus),
        FactIssuerMetricAvailability[MetricKey] = 1,
        FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
    )
VAR DIn =
    CALCULATETABLE(
        VALUES(FactIssuerMetricAvailability[IssuerKey]),
        REMOVEFILTERS(DimMetric),
        REMOVEFILTERS(DimAvailabilityStatus),
        FactIssuerMetricAvailability[MetricKey] = 4,
        FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
    )
VAR DOut =
    CALCULATETABLE(
        VALUES(FactIssuerMetricAvailability[IssuerKey]),
        REMOVEFILTERS(DimMetric),
        REMOVEFILTERS(DimAvailabilityStatus),
        FactIssuerMetricAvailability[MetricKey] = 3,
        FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
    )
VAR Eligible =
    INTERSECT(INTERSECT(RIn, ROut), INTERSECT(DIn, DOut))
RETURN
    CALCULATE(
        SUM(FactIssuerTransparency[ClaimsReceivedInNetwork])
            + SUM(FactIssuerTransparency[ClaimsReceivedOutOfNetwork]),
        TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
    )
""".strip(),

    "Issuer Comparable Claims Denied - Total": r"""
VAR RIn =
    CALCULATETABLE(
        VALUES(FactIssuerMetricAvailability[IssuerKey]),
        REMOVEFILTERS(DimMetric),
        REMOVEFILTERS(DimAvailabilityStatus),
        FactIssuerMetricAvailability[MetricKey] = 2,
        FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
    )
VAR ROut =
    CALCULATETABLE(
        VALUES(FactIssuerMetricAvailability[IssuerKey]),
        REMOVEFILTERS(DimMetric),
        REMOVEFILTERS(DimAvailabilityStatus),
        FactIssuerMetricAvailability[MetricKey] = 1,
        FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
    )
VAR DIn =
    CALCULATETABLE(
        VALUES(FactIssuerMetricAvailability[IssuerKey]),
        REMOVEFILTERS(DimMetric),
        REMOVEFILTERS(DimAvailabilityStatus),
        FactIssuerMetricAvailability[MetricKey] = 4,
        FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
    )
VAR DOut =
    CALCULATETABLE(
        VALUES(FactIssuerMetricAvailability[IssuerKey]),
        REMOVEFILTERS(DimMetric),
        REMOVEFILTERS(DimAvailabilityStatus),
        FactIssuerMetricAvailability[MetricKey] = 3,
        FactIssuerMetricAvailability[AvailabilityStatusKey] = 1
    )
VAR Eligible =
    INTERSECT(INTERSECT(RIn, ROut), INTERSECT(DIn, DOut))
RETURN
    CALCULATE(
        SUM(FactIssuerTransparency[ClaimsDeniedInNetwork])
            + SUM(FactIssuerTransparency[ClaimsDeniedOutOfNetwork]),
        TREATAS(Eligible, FactIssuerTransparency[IssuerKey])
    )
""".strip(),

    "Unavailable Rate": r"""
VAR EntityLevel = SELECTEDVALUE(DimMetric[EntityLevel])
VAR Unavailable =
    SWITCH(
        EntityLevel,
        "ISSUER",
            CALCULATE(
                COUNTROWS(FactIssuerMetricAvailability),
                REMOVEFILTERS(DimAvailabilityStatus),
                FactIssuerMetricAvailability[AvailabilityStatusKey] = 2
            ),
        "PLAN",
            CALCULATE(
                COUNTROWS(FactPlanMetricAvailability),
                REMOVEFILTERS(DimAvailabilityStatus),
                FactPlanMetricAvailability[AvailabilityStatusKey] = 2
            ),
        BLANK()
    )
RETURN
    IF(
        ISBLANK(EntityLevel),
        BLANK(),
        DIVIDE(COALESCE(Unavailable, 0), [Applicable Entity Count])
    )
""".strip(),

    "Suppression Rate": r"""
VAR EntityLevel = SELECTEDVALUE(DimMetric[EntityLevel])
VAR Suppressed =
    SWITCH(
        EntityLevel,
        "ISSUER",
            CALCULATE(
                COUNTROWS(FactIssuerMetricAvailability),
                REMOVEFILTERS(DimAvailabilityStatus),
                FactIssuerMetricAvailability[AvailabilityStatusKey] = 3
            ),
        "PLAN",
            CALCULATE(
                COUNTROWS(FactPlanMetricAvailability),
                REMOVEFILTERS(DimAvailabilityStatus),
                FactPlanMetricAvailability[AvailabilityStatusKey] = 3
            ),
        BLANK()
    )
RETURN
    IF(
        ISBLANK(EntityLevel),
        BLANK(),
        DIVIDE(COALESCE(Suppressed, 0), [Applicable Entity Count])
    )
""".strip(),

    "Current Context DQ Exception Count": r"""
VAR VisiblePlanIDs =
    VALUES(DimPlan[PlanID])
VAR VisibleIssuerKeys =
    VALUES(DimPlan[IssuerKey])
VAR VisibleIssuerIDs =
    CALCULATETABLE(
        VALUES(DimIssuer[IssuerID]),
        TREATAS(VisibleIssuerKeys, DimIssuer[IssuerKey])
    )
VAR IssuerDQ =
    CALCULATE(
        COUNTROWS(DQExceptionRegister),
        DQExceptionRegister[Scope] = "ISSUER",
        TREATAS(VisibleIssuerIDs, DQExceptionRegister[EntityKey])
    )
VAR PlanDQ =
    CALCULATE(
        COUNTROWS(DQExceptionRegister),
        DQExceptionRegister[Scope] = "PLAN",
        TREATAS(VisiblePlanIDs, DQExceptionRegister[EntityKey])
    )
RETURN
    COALESCE(IssuerDQ, 0) + COALESCE(PlanDQ, 0)
""".strip(),
}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False

    result = subprocess.run(
        ["tasklist", "/FI", "IMAGENAME eq PBIDesktop.exe"],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    return "PBIDesktop.exe" in result.stdout


def create_backup() -> None:
    semantic_model = (
        POWERBI_ROOT / "TransparencyInCoverage.SemanticModel"
    )

    REVIEW_DIR.mkdir(parents=True, exist_ok=True)

    if BACKUP_FILE.exists():
        BACKUP_FILE.unlink()

    with zipfile.ZipFile(
        BACKUP_FILE,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zf:
        for path in sorted(
            p for p in semantic_model.rglob("*") if p.is_file()
        ):
            archive_name = (
                Path(semantic_model.name)
                / path.relative_to(semantic_model)
            )
            zf.write(path, str(archive_name))


def locate_measure_block(
    lines: list[str],
    measure_name: str,
) -> tuple[int, int]:
    target = f"\tmeasure '{measure_name}' ="

    start = None

    for idx, line in enumerate(lines):
        if line.startswith(target):
            start = idx
            break

    if start is None:
        raise RuntimeError(
            f"Measure not found in _Measures.tmdl: {measure_name}"
        )

    end = len(lines)

    for idx in range(start + 1, len(lines)):
        line = lines[idx]

        if line.startswith("\tmeasure "):
            end = idx
            break

        if line.startswith("\tpartition "):
            end = idx
            break

    return start, end


def extract_metadata(
    block_lines: list[str],
) -> tuple[str, str, str | None]:
    format_string = None
    display_folder = None
    lineage_tag = None

    for line in block_lines:
        stripped = line.strip()

        if stripped.startswith("formatString:"):
            format_string = stripped.split(":", 1)[1].strip()

        elif stripped.startswith("displayFolder:"):
            display_folder = stripped.split(":", 1)[1].strip()

        elif stripped.startswith("lineageTag:"):
            lineage_tag = stripped.split(":", 1)[1].strip()

    if not format_string:
        raise RuntimeError(
            "Could not preserve formatString for measure block."
        )

    if not display_folder:
        raise RuntimeError(
            "Could not preserve displayFolder for measure block."
        )

    return format_string, display_folder, lineage_tag


def build_measure_block(
    measure_name: str,
    dax: str,
    format_string: str,
    display_folder: str,
    lineage_tag: str | None,
) -> list[str]:
    out = [f"\tmeasure '{measure_name}' ="]

    for dax_line in dax.splitlines():
        if dax_line:
            out.append(f"\t\t\t{dax_line}")
        else:
            out.append("")

    out.append(f"\t\tformatString: {format_string}")
    out.append(f"\t\tdisplayFolder: {display_folder}")

    if lineage_tag:
        out.append(f"\t\tlineageTag: {lineage_tag}")

    out.append("")
    return out


def patch_measure(
    lines: list[str],
    measure_name: str,
    dax: str,
) -> tuple[list[str], dict]:
    start, end = locate_measure_block(lines, measure_name)

    old_block = lines[start:end]
    format_string, display_folder, lineage_tag = extract_metadata(
        old_block
    )

    new_block = build_measure_block(
        measure_name=measure_name,
        dax=dax,
        format_string=format_string,
        display_folder=display_folder,
        lineage_tag=lineage_tag,
    )

    new_lines = lines[:start] + new_block + lines[end:]

    return new_lines, {
        "MeasureName": measure_name,
        "DisplayFolder": display_folder,
        "FormatString": format_string,
        "LineageTagPreserved": bool(lineage_tag),
        "PatchStatus": "APPLIED",
    }


def style_report(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)

    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if ws.max_row and ws.max_column:
            ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(1, ws.max_column + 1):
            max_len = 0

            for row_idx in range(
                1,
                min(ws.max_row, 200) + 1,
            ):
                value = ws.cell(row_idx, col_idx).value

                if value is not None:
                    max_len = max(max_len, len(str(value)))

            ws.column_dimensions[
                get_column_letter(col_idx)
            ].width = min(max(max_len + 2, 10), 60)

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    wb.save(path)


def main() -> int:
    print("=" * 78)
    print("Transparency in Coverage PUF — Runtime DAX Failure Patch")
    print("=" * 78)
    print(f"Script version : {SCRIPT_VERSION}")
    print(f"Measure file   : {MEASURE_TABLE_FILE}")
    print(f"Backup         : {BACKUP_FILE}")
    print(f"Report         : {REPORT_FILE}")

    if power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is running. Save the PBIP and close "
            "Power BI Desktop completely before applying this patch."
        )

    if not MEASURE_TABLE_FILE.exists():
        raise FileNotFoundError(
            f"Measure table file not found: {MEASURE_TABLE_FILE}"
        )

    create_backup()

    before_hash = sha256(MEASURE_TABLE_FILE)

    text = MEASURE_TABLE_FILE.read_text(encoding="utf-8-sig")
    lines = text.splitlines()

    patch_rows = []

    for measure_name, dax in PATCHES.items():
        lines, evidence = patch_measure(
            lines,
            measure_name,
            dax,
        )
        patch_rows.append(evidence)

    final_text = "\n".join(lines).rstrip() + "\n"
    MEASURE_TABLE_FILE.write_text(
        final_text,
        encoding="utf-8",
    )

    after_hash = sha256(MEASURE_TABLE_FILE)

    validation_rows = []

    for measure_name, expected_dax in PATCHES.items():
        start, end = locate_measure_block(
            final_text.splitlines(),
            measure_name,
        )

        block = "\n".join(
            final_text.splitlines()[start:end]
        )

        first_line = expected_dax.splitlines()[0]
        passed = first_line in block

        validation_rows.append(
            {
                "MeasureName": measure_name,
                "ExpectedMarker": first_line,
                "Status": "PASS" if passed else "FAIL",
            }
        )

    failures = sum(
        row["Status"] == "FAIL"
        for row in validation_rows
    )

    summary = pd.DataFrame(
        [
            ("PatchStatus", "PASS" if failures == 0 else "FAIL"),
            ("ScriptVersion", SCRIPT_VERSION),
            ("PatchedAtLocal", datetime.now().isoformat(timespec="seconds")),
            ("MeasuresPatched", len(PATCHES)),
            ("ValidationFailures", failures),
            ("BeforeSHA256", before_hash),
            ("AfterSHA256", after_hash),
            ("BackupFile", str(BACKUP_FILE)),
            (
                "Purpose",
                "Correct three runtime reconciliation failures and "
                "harden comparable/availability semantics.",
            ),
        ],
        columns=["Item", "Value"],
    )

    REVIEW_DIR.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(
        REPORT_FILE,
        engine="openpyxl",
    ) as writer:
        summary.to_excel(
            writer,
            sheet_name="00_Summary",
            index=False,
        )

        pd.DataFrame(patch_rows).to_excel(
            writer,
            sheet_name="01_Patched_Measures",
            index=False,
        )

        pd.DataFrame(validation_rows).to_excel(
            writer,
            sheet_name="02_Validation",
            index=False,
        )

        pd.DataFrame(
            [
                {
                    "Failure": "Issuer Comparable Claims Received - Total",
                    "RootCause": (
                        "Measure used only the two received-claim availability "
                        "sets instead of the exact four-metric comparable "
                        "population used by the governed overall denial KPI."
                    ),
                    "Correction": (
                        "Use common issuer set across received IN/OON and "
                        "denied IN/OON."
                    ),
                },
                {
                    "Failure": "Unavailable Rate",
                    "RootCause": (
                        "COUNTROWS on an empty filtered availability table "
                        "returned BLANK, so DIVIDE returned BLANK instead of "
                        "a governed zero."
                    ),
                    "Correction": (
                        "Use COALESCE(Unavailable, 0) while retaining BLANK "
                        "when no metric/entity level is selected."
                    ),
                },
                {
                    "Failure": "Current Context DQ Exception Count - Plan",
                    "RootCause": (
                        "A Plan filter does not propagate upstream from "
                        "DimPlan to DimIssuer in the single-direction model; "
                        "VALUES(DimIssuer[IssuerID]) therefore remained broad."
                    ),
                    "Correction": (
                        "Derive visible issuer keys from filtered DimPlan, "
                        "then TREATAS those keys onto DimIssuer before matching "
                        "the disconnected DQ register."
                    ),
                },
            ]
        ).to_excel(
            writer,
            sheet_name="03_Root_Cause",
            index=False,
        )

    style_report(REPORT_FILE)

    print()
    print("Runtime DAX patch completed.")
    print(f"Patch status        : {'PASS' if failures == 0 else 'FAIL'}")
    print(f"Measures patched    : {len(PATCHES)}")
    print(f"Validation failures : {failures}")
    print(f"Backup              : {BACKUP_FILE}")
    print(f"Review evidence     : {REPORT_FILE}")

    if failures:
        print()
        print(
            "Patch validation failed. Do not open/save the PBIP "
            "until the report is reviewed."
        )
        return 2

    print()
    print(
        "Next: open the PBIP, Save (Ctrl+S), keep Power BI open, "
        "then rerun 09_runtime_dax_reconciliation.py."
    )

    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())

    except Exception as exc:
        print()
        print("DAX PATCH FAILED")
        print(str(exc))
        sys.exit(1)
