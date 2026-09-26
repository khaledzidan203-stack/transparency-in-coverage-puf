from __future__ import annotations

from datetime import datetime
from pathlib import Path
import hashlib
import json
import os
import shutil
import subprocess
import sys
import zipfile

import pandas as pd


SCRIPT_VERSION = "1.0.0"

PROJECT_ROOT = Path(__file__).resolve().parents[2]
POWERBI_ROOT = PROJECT_ROOT / "powerbi"
REPORT_DIR = POWERBI_ROOT / "TransparencyInCoverage.Report"
PAGES_DIR = REPORT_DIR / "definition" / "pages"

VISUAL_SCHEMA = (
    "https://developer.microsoft.com/json-schemas/fabric/item/report/"
    "definition/visualContainer/2.7.0/schema.json"
)

CANVAS_WIDTH = 1920
CANVAS_HEIGHT = 1080

COLORS = {
    "canvas": "#F5F7FA",
    "surface": "#FFFFFF",
    "surface_alt": "#F9FAFB",
    "border": "#E7ECF2",
    "ink": "#172033",
    "secondary": "#667085",
    "muted": "#98A2B3",
    "navy": "#17324D",
    "slate": "#344054",
    "blue": "#4F6BED",
    "blue_dark": "#425CC7",
    "blue_soft": "#EEF3FF",
    "blue_pale": "#E5ECFF",
    "teal": "#16A6A1",
    "teal_dark": "#0E7774",
    "teal_soft": "#EAF8F6",
    "amber": "#C8872C",
    "amber_dark": "#9C641A",
    "amber_soft": "#FFF4E3",
    "red": "#C64B4B",
    "red_soft": "#FCE8E8",
    "white": "#FFFFFF",
}

FONT_REGULAR = (
    "'Segoe UI', wf_segoe-ui_normal, helvetica, arial, sans-serif"
)
FONT_SEMIBOLD = (
    "'Segoe UI Semibold', wf_segoe-ui_semibold, "
    "helvetica, arial, sans-serif"
)

EXPECTED_PAGES = [
    "00 INDEX",
    "01 Executive Overview",
    "02 Network Comparison",
    "03 Denials",
    "04 Denial Reasons",
    "05 Appeals",
    "06 Resubmissions",
    "07 Enrollment",
    "08 Issuer & State Explorer",
    "09 Data Availability",
    "10 Data Quality & Methodology",
]

UNSUPPORTED_PAGE_ROOT_PROPERTIES = {
    "horizontalAlignment",
    "verticalAlignment",
}

TILES = [
    {
        "page": "01 Executive Overview",
        "number": "01",
        "section": "OVERVIEW",
        "title": "Executive Overview",
        "subtitle": "Market size, claims exposure and governance context.",
        "band": COLORS["blue"],
        "body": COLORS["blue_soft"],
        "title_color": COLORS["navy"],
    },
    {
        "page": "02 Network Comparison",
        "number": "02",
        "section": "CLAIMS & DENIALS",
        "title": "Network Comparison",
        "subtitle": "In-network vs out-of-network denial behavior.",
        "band": COLORS["navy"],
        "body": COLORS["surface"],
        "title_color": COLORS["navy"],
    },
    {
        "page": "03 Denials",
        "number": "03",
        "section": "CLAIMS & DENIALS",
        "title": "Denials",
        "subtitle": "Volume, comparable rates and state concentration.",
        "band": COLORS["amber"],
        "body": COLORS["amber_soft"],
        "title_color": COLORS["amber_dark"],
    },
    {
        "page": "04 Denial Reasons",
        "number": "04",
        "section": "CLAIMS & DENIALS",
        "title": "Denial Reasons",
        "subtitle": "Reason mix and issuer investigation.",
        "band": COLORS["blue_dark"],
        "body": COLORS["blue_soft"],
        "title_color": COLORS["blue_dark"],
    },
    {
        "page": "05 Appeals",
        "number": "05",
        "section": "POST-DENIAL",
        "title": "Appeals",
        "subtitle": "Filed volume, overturn rates and sensitivity.",
        "band": COLORS["teal_dark"],
        "body": COLORS["teal_soft"],
        "title_color": COLORS["teal_dark"],
    },
    {
        "page": "06 Resubmissions",
        "number": "06",
        "section": "POST-DENIAL",
        "title": "Resubmissions",
        "subtitle": "Events per 100 denied across network settings.",
        "band": COLORS["blue"],
        "body": COLORS["blue_soft"],
        "title_color": COLORS["blue_dark"],
    },
    {
        "page": "07 Enrollment",
        "number": "07",
        "section": "MARKET & POPULATION",
        "title": "Enrollment",
        "subtitle": "Reported enrollment, coverage and comparable ratio.",
        "band": COLORS["teal"],
        "body": COLORS["teal_soft"],
        "title_color": COLORS["teal_dark"],
    },
    {
        "page": "08 Issuer & State Explorer",
        "number": "08",
        "section": "MARKET & POPULATION",
        "title": "Issuer & State Explorer",
        "subtitle": "Explore state and issuer patterns.",
        "band": COLORS["slate"],
        "body": COLORS["surface"],
        "title_color": COLORS["navy"],
    },
    {
        "page": "09 Data Availability",
        "number": "09",
        "section": "GOVERNANCE",
        "title": "Data Availability",
        "subtitle": "Coverage, suppression and structural applicability.",
        "band": COLORS["teal_dark"],
        "body": COLORS["teal_soft"],
        "title_color": COLORS["teal_dark"],
    },
    {
        "page": "10 Data Quality & Methodology",
        "number": "10",
        "section": "GOVERNANCE",
        "title": "Data Quality & Methodology",
        "subtitle": "Exceptions, known-source issues and analytical rules.",
        "band": COLORS["amber"],
        "body": COLORS["amber_soft"],
        "title_color": COLORS["amber_dark"],
    },
]


# =====================================================================
# Paths / environment
# =====================================================================

def get_windows_desktop() -> Path:
    if os.name == "nt":
        try:
            import winreg

            key_path = (
                r"Software\Microsoft\Windows\CurrentVersion"
                r"\Explorer\User Shell Folders"
            )

            with winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                key_path,
            ) as key:
                raw_value, _ = winreg.QueryValueEx(
                    key,
                    "Desktop",
                )

            return Path(
                os.path.expandvars(raw_value)
            ).resolve()

        except Exception:
            pass

    return (Path.home() / "Desktop").resolve()


REVIEW_DIR = (
    get_windows_desktop()
    / "Transparency_PUF_Review"
)

BACKUP_FILE = (
    REVIEW_DIR
    / "42_PRE_INDEX_STORYTELLING_REBUILD_BACKUP.zip"
)

REPORT_FILE = (
    REVIEW_DIR
    / "42_INDEX_STORYTELLING_REBUILD_REPORT.xlsx"
)


def power_bi_desktop_running() -> bool:
    if os.name != "nt":
        return False

    result = subprocess.run(
        [
            "tasklist",
            "/FI",
            "IMAGENAME eq PBIDesktop.exe",
        ],
        capture_output=True,
        text=True,
        check=False,
        creationflags=getattr(
            subprocess,
            "CREATE_NO_WINDOW",
            0,
        ),
    )

    return "PBIDesktop.exe" in result.stdout


def read_json(path: Path):
    return json.loads(
        path.read_text(
            encoding="utf-8-sig",
        )
    )


def write_json(
    path: Path,
    payload,
) -> None:
    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    path.write_text(
        json.dumps(
            payload,
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )


def sha256(path: Path) -> str:
    h = hashlib.sha256()

    with path.open("rb") as f:
        for chunk in iter(
            lambda: f.read(1024 * 1024),
            b"",
        ):
            h.update(chunk)

    return h.hexdigest()


def file_hashes(
    root: Path,
) -> dict[str, str]:
    result = {}

    for path in root.rglob("*"):
        if path.is_file():
            result[
                str(
                    path.relative_to(
                        REPORT_DIR
                    )
                ).replace(
                    "\\",
                    "/",
                )
            ] = sha256(
                path
            )

    return result


# =====================================================================
# PBIR expressions
# =====================================================================

def expr_bool(value: bool) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": "true" if value else "false"
            }
        }
    }


def expr_string(value: str) -> dict:
    escaped = value.replace(
        "'",
        "''",
    )

    return {
        "expr": {
            "Literal": {
                "Value": f"'{escaped}'"
            }
        }
    }


def expr_double(
    value: float | int,
) -> dict:
    return {
        "expr": {
            "Literal": {
                "Value": f"{value}D"
            }
        }
    }


def solid_color(
    color: str,
) -> dict:
    return {
        "solid": {
            "color": expr_string(
                color
            )
        }
    }


# =====================================================================
# Report discovery / backup
# =====================================================================

def find_pages() -> dict[
    str,
    tuple[Path, dict]
]:
    pages = {}

    for page_json in PAGES_DIR.glob(
        "*/page.json"
    ):
        payload = read_json(
            page_json
        )

        display_name = payload.get(
            "displayName"
        )

        if display_name:
            pages[
                display_name
            ] = (
                page_json.parent,
                payload,
            )

    return pages


def create_backup() -> None:
    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if BACKUP_FILE.exists():
        BACKUP_FILE.unlink()

    with zipfile.ZipFile(
        BACKUP_FILE,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zf:
        for path in sorted(
            p
            for p in REPORT_DIR.rglob("*")
            if p.is_file()
        ):
            archive_name = (
                Path(REPORT_DIR.name)
                / path.relative_to(
                    REPORT_DIR
                )
            )

            zf.write(
                path,
                str(
                    archive_name
                ),
            )


def prepare_index_page(
    page_dir: Path,
) -> None:
    page_json = (
        page_dir
        / "page.json"
    )

    payload = read_json(
        page_json
    )

    for prop in (
        UNSUPPORTED_PAGE_ROOT_PROPERTIES
    ):
        payload.pop(
            prop,
            None,
        )

    payload[
        "width"
    ] = CANVAS_WIDTH

    payload[
        "height"
    ] = CANVAS_HEIGHT

    payload[
        "displayOption"
    ] = "FitToPage"

    objects = payload.setdefault(
        "objects",
        {},
    )

    objects[
        "background"
    ] = [
        {
            "properties": {
                "color": solid_color(
                    COLORS["canvas"]
                ),
                "transparency": expr_double(
                    0
                ),
            }
        }
    ]

    write_json(
        page_json,
        payload,
    )


# =====================================================================
# Text / visual builders
# =====================================================================

def text_run(
    value: str,
    *,
    size: float,
    color: str,
    semibold: bool = False,
) -> dict:
    return {
        "value": value,
        "textStyle": {
            "fontFamily": (
                FONT_SEMIBOLD
                if semibold
                else FONT_REGULAR
            ),
            "fontSize": (
                f"{size}pt"
            ),
            "color": color,
        },
    }


def paragraph(
    *runs: dict,
    align: str = "center",
) -> dict:
    return {
        "textRuns": list(
            runs
        ),
        "horizontalTextAlignment": (
            align
        ),
    }


def container_objects(
    *,
    background: str | None = None,
    border: str | None = None,
    radius: int = 8,
) -> dict:
    return {
        "title": [
            {
                "properties": {
                    "show": expr_bool(
                        False
                    )
                }
            }
        ],
        "background": [
            {
                "properties": {
                    "show": expr_bool(
                        background is not None
                    ),
                    **(
                        {
                            "color": solid_color(
                                background
                            ),
                            "transparency": expr_double(
                                0
                            ),
                        }
                        if background is not None
                        else {}
                    ),
                }
            }
        ],
        "border": [
            {
                "properties": {
                    "show": expr_bool(
                        border is not None
                    ),
                    **(
                        {
                            "color": solid_color(
                                border
                            ),
                            "width": expr_double(
                                1
                            ),
                            "radius": expr_double(
                                radius
                            ),
                        }
                        if border is not None
                        else {}
                    ),
                }
            }
        ],
        "dropShadow": [
            {
                "properties": {
                    "show": expr_bool(
                        False
                    )
                }
            }
        ],
        "visualHeader": [
            {
                "properties": {
                    "show": expr_bool(
                        False
                    )
                }
            }
        ],
    }


def make_textbox(
    *,
    name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
    paragraphs: list[dict],
    background: str | None = None,
    border: str | None = None,
    radius: int = 8,
) -> dict:
    return {
        "$schema": VISUAL_SCHEMA,
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
        },
        "visual": {
            "visualType": "textbox",
            "objects": {
                "general": [
                    {
                        "properties": {
                            "paragraphs": paragraphs
                        }
                    }
                ]
            },
            "visualContainerObjects": (
                container_objects(
                    background=background,
                    border=border,
                    radius=radius,
                )
            ),
            "drillFilterOtherVisuals": True,
        },
    }


def make_section_band(
    *,
    name: str,
    title: str,
    x: float,
    y: float,
    width: float,
    background: str,
    z: int,
) -> dict:
    return make_textbox(
        name=name,
        x=x,
        y=y,
        width=width,
        height=32,
        z=z,
        paragraphs=[
            paragraph(
                text_run(
                    title,
                    size=11,
                    color=COLORS["white"],
                    semibold=True,
                )
            )
        ],
        background=background,
        border=background,
        radius=8,
    )


def make_tile_band(
    *,
    name: str,
    number: str,
    section: str,
    x: float,
    y: float,
    width: float,
    background: str,
    z: int,
) -> dict:
    return make_textbox(
        name=name,
        x=x,
        y=y,
        width=width,
        height=28,
        z=z,
        paragraphs=[
            paragraph(
                text_run(
                    f"{number}  {section}",
                    size=8.2,
                    color=COLORS["white"],
                    semibold=True,
                )
            )
        ],
        background=background,
        border=background,
        radius=8,
    )


def make_tile_body(
    *,
    name: str,
    title: str,
    subtitle: str,
    x: float,
    y: float,
    width: float,
    height: float,
    background: str,
    title_color: str,
    z: int,
) -> dict:
    return make_textbox(
        name=name,
        x=x,
        y=y,
        width=width,
        height=height,
        z=z,
        paragraphs=[
            paragraph(
                text_run(
                    " ",
                    size=3,
                    color=title_color,
                )
            ),
            paragraph(
                text_run(
                    title,
                    size=17,
                    color=title_color,
                    semibold=True,
                )
            ),
            paragraph(
                text_run(
                    subtitle,
                    size=8.8,
                    color=COLORS["secondary"],
                )
            ),
        ],
        background=background,
        border=COLORS["border"],
        radius=9,
    )


def make_nav_overlay(
    *,
    name: str,
    target_page_name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z: int,
) -> dict:
    return {
        "$schema": VISUAL_SCHEMA,
        "name": name,
        "position": {
            "x": x,
            "y": y,
            "z": z,
            "width": width,
            "height": height,
        },
        "visual": {
            "visualType": "actionButton",
            "objects": {
                "icon": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "outline": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "fill": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
            },
            "visualContainerObjects": {
                "title": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "background": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "border": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "visualHeader": [
                    {
                        "properties": {
                            "show": expr_bool(
                                False
                            )
                        }
                    }
                ],
                "visualLink": [
                    {
                        "properties": {
                            "show": expr_bool(
                                True
                            ),
                            "type": expr_string(
                                "PageNavigation"
                            ),
                            "navigationSection": expr_string(
                                target_page_name
                            ),
                            "showDefaultTooltip": expr_bool(
                                False
                            ),
                            "enabledTooltip": expr_string(
                                "Open "
                                + target_page_name
                            ),
                        }
                    }
                ],
            },
            "drillFilterOtherVisuals": True,
        },
    }


def write_visual(
    page_dir: Path,
    visual: dict,
) -> Path:
    target = (
        page_dir
        / "visuals"
        / visual["name"]
        / "visual.json"
    )

    write_json(
        target,
        visual,
    )

    return target


def add_tile(
    *,
    page_dir: Path,
    tile: dict,
    target_internal_name: str,
    x: float,
    y: float,
    width: float,
    body_height: float,
    idx: int,
) -> None:
    band_name = (
        f"idx_tile_{idx:02d}_band"
    )
    body_name = (
        f"idx_tile_{idx:02d}_body"
    )
    nav_name = (
        f"idx_tile_{idx:02d}_nav"
    )

    write_visual(
        page_dir,
        make_tile_band(
            name=band_name,
            number=tile[
                "number"
            ],
            section=tile[
                "section"
            ],
            x=x,
            y=y,
            width=width,
            background=tile[
                "band"
            ],
            z=1000 + idx,
        ),
    )

    write_visual(
        page_dir,
        make_tile_body(
            name=body_name,
            title=tile[
                "title"
            ],
            subtitle=tile[
                "subtitle"
            ],
            x=x,
            y=y + 30,
            width=width,
            height=body_height,
            background=tile[
                "body"
            ],
            title_color=tile[
                "title_color"
            ],
            z=800 + idx,
        ),
    )

    write_visual(
        page_dir,
        make_nav_overlay(
            name=nav_name,
            target_page_name=(
                target_internal_name
            ),
            x=x,
            y=y,
            width=width,
            height=30 + body_height,
            z=5000 + idx,
        ),
    )


# =====================================================================
# Page build
# =====================================================================

def build_index(
    *,
    page_dir: Path,
    page_map: dict[
        str,
        tuple[Path, dict]
    ],
) -> None:
    visuals_dir = (
        page_dir
        / "visuals"
    )

    if visuals_dir.exists():
        shutil.rmtree(
            visuals_dir
        )

    visuals_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    # Header.
    write_visual(
        page_dir,
        make_textbox(
            name="idx_identity",
            x=52,
            y=28,
            width=300,
            height=44,
            z=100,
            paragraphs=[
                paragraph(
                    text_run(
                        "HEALTHCARE TRANSPARENCY ANALYTICS",
                        size=8.8,
                        color=COLORS["blue"],
                        semibold=True,
                    )
                )
            ],
            background=COLORS["blue_soft"],
            border=COLORS["blue_pale"],
            radius=8,
        ),
    )

    write_visual(
        page_dir,
        make_textbox(
            name="idx_title",
            x=390,
            y=16,
            width=1140,
            height=76,
            z=110,
            paragraphs=[
                paragraph(
                    text_run(
                        "Transparency in Coverage",
                        size=28,
                        color=COLORS["navy"],
                        semibold=True,
                    )
                ),
                paragraph(
                    text_run(
                        "Claims → denials → appeals → availability → governance",
                        size=10,
                        color=COLORS["secondary"],
                    )
                ),
            ],
            background=COLORS["blue_soft"],
            border=COLORS["blue_pale"],
            radius=9,
        ),
    )

    write_visual(
        page_dir,
        make_textbox(
            name="idx_period",
            x=1598,
            y=28,
            width=270,
            height=44,
            z=100,
            paragraphs=[
                paragraph(
                    text_run(
                        "2026 PUF  •  EXPERIENCE YEAR 2024",
                        size=8.8,
                        color=COLORS["navy"],
                        semibold=True,
                    )
                )
            ],
            background=COLORS["surface"],
            border=COLORS["border"],
            radius=8,
        ),
    )

    # Overview.
    write_visual(
        page_dir,
        make_section_band(
            name="idx_section_overview",
            title="OVERVIEW",
            x=52,
            y=126,
            width=1816,
            background=COLORS["blue"],
            z=300,
        ),
    )

    tile_by_page = {
        tile[
            "page"
        ]: tile
        for tile in TILES
    }

    add_tile(
        page_dir=page_dir,
        tile=tile_by_page[
            "01 Executive Overview"
        ],
        target_internal_name=(
            page_map[
                "01 Executive Overview"
            ][1][
                "name"
            ]
        ),
        x=52,
        y=170,
        width=1816,
        body_height=92,
        idx=1,
    )

    # Claims & Denials.
    write_visual(
        page_dir,
        make_section_band(
            name="idx_section_claims",
            title="CLAIMS & DENIALS",
            x=52,
            y=318,
            width=1816,
            background=COLORS["navy"],
            z=300,
        ),
    )

    claims_pages = [
        "02 Network Comparison",
        "03 Denials",
        "04 Denial Reasons",
    ]

    claims_gap = 18
    claims_w = (
        1816
        - 2 * claims_gap
    ) / 3

    for i, page_name in enumerate(
        claims_pages
    ):
        add_tile(
            page_dir=page_dir,
            tile=tile_by_page[
                page_name
            ],
            target_internal_name=(
                page_map[
                    page_name
                ][1][
                    "name"
                ]
            ),
            x=52
            + i
            * (
                claims_w
                + claims_gap
            ),
            y=362,
            width=claims_w,
            body_height=92,
            idx=2 + i,
        )

    # Post-denial + Market/Population.
    left_x = 52
    left_w = 888
    split_gap = 20
    right_x = (
        left_x
        + left_w
        + split_gap
    )
    right_w = (
        1816
        - left_w
        - split_gap
    )

    write_visual(
        page_dir,
        make_section_band(
            name="idx_section_post",
            title="POST-DENIAL JOURNEY",
            x=left_x,
            y=510,
            width=left_w,
            background=COLORS["teal_dark"],
            z=300,
        ),
    )

    write_visual(
        page_dir,
        make_section_band(
            name="idx_section_market",
            title="MARKET & POPULATION",
            x=right_x,
            y=510,
            width=right_w,
            background=COLORS["blue_dark"],
            z=300,
        ),
    )

    two_gap = 16
    left_tile_w = (
        left_w
        - two_gap
    ) / 2

    for i, page_name in enumerate(
        [
            "05 Appeals",
            "06 Resubmissions",
        ]
    ):
        add_tile(
            page_dir=page_dir,
            tile=tile_by_page[
                page_name
            ],
            target_internal_name=(
                page_map[
                    page_name
                ][1][
                    "name"
                ]
            ),
            x=left_x
            + i
            * (
                left_tile_w
                + two_gap
            ),
            y=554,
            width=left_tile_w,
            body_height=92,
            idx=5 + i,
        )

    right_tile_w = (
        right_w
        - two_gap
    ) / 2

    for i, page_name in enumerate(
        [
            "07 Enrollment",
            "08 Issuer & State Explorer",
        ]
    ):
        add_tile(
            page_dir=page_dir,
            tile=tile_by_page[
                page_name
            ],
            target_internal_name=(
                page_map[
                    page_name
                ][1][
                    "name"
                ]
            ),
            x=right_x
            + i
            * (
                right_tile_w
                + two_gap
            ),
            y=554,
            width=right_tile_w,
            body_height=92,
            idx=7 + i,
        )

    # Governance.
    write_visual(
        page_dir,
        make_section_band(
            name="idx_section_governance",
            title="GOVERNANCE",
            x=52,
            y=702,
            width=1816,
            background=COLORS["amber_dark"],
            z=300,
        ),
    )

    governance_gap = 18
    governance_w = (
        1816
        - governance_gap
    ) / 2

    for i, page_name in enumerate(
        [
            "09 Data Availability",
            "10 Data Quality & Methodology",
        ]
    ):
        add_tile(
            page_dir=page_dir,
            tile=tile_by_page[
                page_name
            ],
            target_internal_name=(
                page_map[
                    page_name
                ][1][
                    "name"
                ]
            ),
            x=52
            + i
            * (
                governance_w
                + governance_gap
            ),
            y=746,
            width=governance_w,
            body_height=92,
            idx=9 + i,
        )

    # Minimal footer only.
    write_visual(
        page_dir,
        make_textbox(
            name="idx_footer",
            x=52,
            y=1014,
            width=1816,
            height=24,
            z=100,
            paragraphs=[
                paragraph(
                    text_run(
                        (
                            "Governed semantic model • Availability-aware KPIs • "
                            "Source anomalies preserved • No clipping or imputation"
                        ),
                        size=8.2,
                        color=COLORS["muted"],
                        semibold=True,
                    )
                )
            ],
            background=None,
            border=None,
            radius=0,
        ),
    )


# =====================================================================
# Validation
# =====================================================================

def validate_json_tree() -> list[dict]:
    failures = []

    for path in sorted(
        (
            REPORT_DIR
            / "definition"
        ).rglob(
            "*.json"
        )
    ):
        try:
            read_json(
                path
            )

        except Exception as exc:
            failures.append(
                {
                    "RelativePath": str(
                        path.relative_to(
                            PROJECT_ROOT
                        )
                    ),
                    "Error": str(
                        exc
                    ),
                }
            )

    return failures


def validate_no_unsupported_page_roots(
    pages: dict,
) -> list[dict]:
    offenders = []

    for display_name, (
        page_dir,
        _,
    ) in pages.items():
        payload = read_json(
            page_dir
            / "page.json"
        )

        for prop in sorted(
            UNSUPPORTED_PAGE_ROOT_PROPERTIES
        ):
            if prop in payload:
                offenders.append(
                    {
                        "Page": display_name,
                        "UnsupportedProperty": prop,
                    }
                )

    return offenders


def validate_no_full_page_visuals(
    pages: dict,
) -> list[dict]:
    offenders = []

    for display_name, (
        page_dir,
        _,
    ) in pages.items():
        visuals_dir = (
            page_dir
            / "visuals"
        )

        if not visuals_dir.exists():
            continue

        for path in visuals_dir.rglob(
            "visual.json"
        ):
            payload = read_json(
                path
            )

            pos = payload.get(
                "position",
                {},
            )

            width = float(
                pos.get(
                    "width",
                    0,
                )
                or 0
            )

            height = float(
                pos.get(
                    "height",
                    0,
                )
                or 0
            )

            if (
                width
                >= CANVAS_WIDTH * 0.94
                and height
                >= CANVAS_HEIGHT * 0.94
            ):
                offenders.append(
                    {
                        "Page": display_name,
                        "VisualName": payload.get(
                            "name"
                        ),
                        "Width": width,
                        "Height": height,
                    }
                )

    return offenders


def validate_navigation(
    index_dir: Path,
    pages: dict,
) -> list[dict]:
    results = []

    for idx, tile in enumerate(
        TILES,
        start=1,
    ):
        path = (
            index_dir
            / "visuals"
            / f"idx_tile_{idx:02d}_nav"
            / "visual.json"
        )

        actual = None

        try:
            payload = read_json(
                path
            )

            actual = (
                payload[
                    "visual"
                ][
                    "visualContainerObjects"
                ][
                    "visualLink"
                ][0][
                    "properties"
                ][
                    "navigationSection"
                ][
                    "expr"
                ][
                    "Literal"
                ][
                    "Value"
                ]
            )

            if (
                isinstance(
                    actual,
                    str,
                )
                and actual.startswith(
                    "'"
                )
                and actual.endswith(
                    "'"
                )
            ):
                actual = (
                    actual[1:-1]
                    .replace(
                        "''",
                        "'",
                    )
                )
        except Exception:
            actual = None

        expected = (
            pages[
                tile[
                    "page"
                ]
            ][1][
                "name"
            ]
        )

        results.append(
            {
                "Tile": tile[
                    "page"
                ],
                "ExpectedTarget": expected,
                "ActualTarget": actual,
                "Status": (
                    "PASS"
                    if actual == expected
                    else "FAIL"
                ),
            }
        )

    return results


def style_report(
    path: Path,
) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import (
        Alignment,
        Font,
        PatternFill,
    )
    from openpyxl.utils import (
        get_column_letter,
    )

    wb = load_workbook(
        path
    )

    fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )
    font = Font(
        color="FFFFFF",
        bold=True,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if (
            ws.max_row
            and ws.max_column
        ):
            ws.auto_filter.ref = (
                ws.dimensions
            )

        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            max_len = 0

            for row_idx in range(
                1,
                min(
                    ws.max_row,
                    180,
                )
                + 1,
            ):
                value = ws.cell(
                    row=row_idx,
                    column=col_idx,
                ).value

                if value is not None:
                    max_len = max(
                        max_len,
                        len(
                            str(
                                value
                            )
                        ),
                    )

            ws.column_dimensions[
                get_column_letter(
                    col_idx
                )
            ].width = min(
                max(
                    max_len + 2,
                    10,
                ),
                75,
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    wb.save(
        path
    )


# =====================================================================
# Main
# =====================================================================

def main() -> int:
    print("=" * 78)
    print(
        "Transparency in Coverage PUF — "
        "00 INDEX Storytelling Rebuild"
    )
    print("=" * 78)

    print(
        f"Script version : {SCRIPT_VERSION}"
    )
    print(
        f"Report folder  : {REPORT_DIR}"
    )
    print(
        f"Backup         : {BACKUP_FILE}"
    )
    print(
        f"Review file    : {REPORT_FILE}"
    )

    if power_bi_desktop_running():
        raise RuntimeError(
            "Power BI Desktop is running. "
            "Save the PBIP and close Power BI Desktop completely "
            "before changing PBIR files."
        )

    pages = find_pages()

    if set(
        pages
    ) != set(
        EXPECTED_PAGES
    ):
        raise RuntimeError(
            "Current report page inventory does not match "
            "the approved 11-page report. No changes made."
        )

    unsupported_before = (
        validate_no_unsupported_page_roots(
            pages
        )
    )

    if unsupported_before:
        raise RuntimeError(
            "Unsupported page-root properties are present before build: "
            + "; ".join(
                f"{row['Page']}:{row['UnsupportedProperty']}"
                for row in unsupported_before
            )
            + ". No changes made."
        )

    index_dir, _ = (
        pages[
            "00 INDEX"
        ]
    )

    before_hashes = (
        file_hashes(
            REPORT_DIR
        )
    )

    target_prefix = str(
        index_dir.relative_to(
            REPORT_DIR
        )
    ).replace(
        "\\",
        "/",
    )

    create_backup()

    prepare_index_page(
        index_dir
    )

    build_index(
        page_dir=index_dir,
        page_map=pages,
    )

    pages = find_pages()

    after_hashes = (
        file_hashes(
            REPORT_DIR
        )
    )

    non_target_changes = []

    for rel_path in sorted(
        set(
            before_hashes
        )
        | set(
            after_hashes
        )
    ):
        if rel_path.startswith(
            target_prefix
        ):
            continue

        if (
            before_hashes.get(
                rel_path
            )
            != after_hashes.get(
                rel_path
            )
        ):
            non_target_changes.append(
                rel_path
            )

    json_failures = (
        validate_json_tree()
    )

    unsupported_after = (
        validate_no_unsupported_page_roots(
            pages
        )
    )

    full_page = (
        validate_no_full_page_visuals(
            pages
        )
    )

    nav_results = (
        validate_navigation(
            index_dir,
            pages,
        )
    )

    nav_failures = sum(
        row[
            "Status"
        ] == "FAIL"
        for row in nav_results
    )

    visual_count = len(
        list(
            (
                index_dir
                / "visuals"
            ).rglob(
                "visual.json"
            )
        )
    )

    validation_rows = [
        {
            "Check": "Navigation tiles",
            "Expected": 10,
            "Actual": len(
                nav_results
            ),
            "Status": (
                "PASS"
                if (
                    len(
                        nav_results
                    )
                    == 10
                    and nav_failures == 0
                )
                else "FAIL"
            ),
        },
        {
            "Check": "Non-target report changes",
            "Expected": 0,
            "Actual": len(
                non_target_changes
            ),
            "Status": (
                "PASS"
                if not non_target_changes
                else "FAIL"
            ),
        },
        {
            "Check": "Unsupported page properties",
            "Expected": 0,
            "Actual": len(
                unsupported_after
            ),
            "Status": (
                "PASS"
                if not unsupported_after
                else "FAIL"
            ),
        },
        {
            "Check": "Full-page selectable visuals",
            "Expected": 0,
            "Actual": len(
                full_page
            ),
            "Status": (
                "PASS"
                if not full_page
                else "FAIL"
            ),
        },
        {
            "Check": "JSON parse failures",
            "Expected": 0,
            "Actual": len(
                json_failures
            ),
            "Status": (
                "PASS"
                if not json_failures
                else "FAIL"
            ),
        },
    ]

    failures = sum(
        row[
            "Status"
        ] == "FAIL"
        for row in validation_rows
    )

    REVIEW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        REPORT_FILE,
        engine="openpyxl",
    ) as writer:
        pd.DataFrame(
            [
                (
                    "BuildStatus",
                    (
                        "PASS"
                        if failures == 0
                        else "FAIL"
                    ),
                ),
                (
                    "ScriptVersion",
                    SCRIPT_VERSION,
                ),
                (
                    "BuiltAtLocal",
                    datetime.now()
                    .isoformat(
                        timespec="seconds"
                    ),
                ),
                (
                    "PageBuilt",
                    "00 INDEX",
                ),
                (
                    "VisualsAfter",
                    visual_count,
                ),
                (
                    "NavigationTiles",
                    10,
                ),
                (
                    "StoryGroups",
                    (
                        "Overview | Claims & Denials | Post-Denial Journey | "
                        "Market & Population | Governance"
                    ),
                ),
                (
                    "DesignPattern",
                    (
                        "Colored section bands + colored tile headers + "
                        "soft card bodies + centered copy"
                    ),
                ),
                (
                    "NonTargetFilesChanged",
                    len(
                        non_target_changes
                    ),
                ),
                (
                    "ValidationFailures",
                    failures,
                ),
                (
                    "BackupFile",
                    str(
                        BACKUP_FILE
                    ),
                ),
            ],
            columns=[
                "Item",
                "Value",
            ],
        ).to_excel(
            writer,
            sheet_name="00_Summary",
            index=False,
        )

        pd.DataFrame(
            validation_rows
        ).to_excel(
            writer,
            sheet_name="01_Validation",
            index=False,
        )

        pd.DataFrame(
            nav_results
        ).to_excel(
            writer,
            sheet_name="02_Navigation",
            index=False,
        )

        if non_target_changes:
            pd.DataFrame(
                {
                    "RelativePath": (
                        non_target_changes
                    )
                }
            ).to_excel(
                writer,
                sheet_name="03_NonTarget_Changes",
                index=False,
            )
        else:
            pd.DataFrame(
                [
                    {
                        "Info": (
                            "Only 00 INDEX changed. "
                            "Pages 01-10 were preserved."
                        )
                    }
                ]
            ).to_excel(
                writer,
                sheet_name="03_NonTarget_Changes",
                index=False,
            )

    style_report(
        REPORT_FILE
    )

    print()
    print(
        "00 INDEX storytelling rebuild completed."
    )
    print(
        f"Build status                 : "
        f"{'PASS' if failures == 0 else 'FAIL'}"
    )
    print(
        f"INDEX visuals                : {visual_count}"
    )
    print(
        "Navigation tiles             : 10"
    )
    print(
        "Story groups                 : 5"
    )
    print(
        "Tile design                  : colored header + soft body"
    )
    print(
        f"Non-target files changed     : "
        f"{len(non_target_changes)}"
    )
    print(
        f"Full-page selectable visuals : "
        f"{len(full_page)}"
    )
    print(
        f"Unsupported page properties  : "
        f"{len(unsupported_after)}"
    )
    print(
        f"Validation failures          : {failures}"
    )
    print(
        f"Backup                       : {BACKUP_FILE}"
    )
    print(
        f"Review evidence              : {REPORT_FILE}"
    )

    if failures:
        print()
        print(
            "Offline validation failed. "
            "Do not open/save the PBIP until the evidence report is reviewed."
        )
        return 2

    print()
    print(
        "Next: open TransparencyInCoverage.pbip and inspect 00 INDEX. "
        "Test one tile in every story group before Save (Ctrl+S)."
    )

    return 0


if __name__ == "__main__":
    try:
        sys.exit(
            main()
        )

    except Exception as exc:
        print()
        print(
            "INDEX STORYTELLING REBUILD FAILED"
        )
        print(
            str(
                exc
            )
        )
        sys.exit(
            1
        )
